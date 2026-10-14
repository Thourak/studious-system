# Fabric notebook source

# METADATA ********************

# META {
# META   "kernel_info": {
# META     "name": "synapse_pyspark"
# META   }
# META }

# MARKDOWN ********************

# **Nom de l'objet**
# # T4_NB_HRIS_LIB. C'est le nom affiché dans Microsoft Fabric et celui que les autres notebooks utilisent pour le charger. Le dossier du dépôt garde son nom historique NB_HRIS_LIB.Notebook afin de ne déplacer aucun fichier. Le préfixe T4 désigne la version des objets de ce dossier, comme pour tous les autres objets qui l'accompagnent.
# # **Type d'objet**
# # Notebook, exécuté dans Microsoft Fabric avec le moteur Synapse PySpark. Il n'est jamais lancé seul par le pipeline. Chaque notebook du flux le charge au début de son exécution par la commande magique %run, qui exécute ses cellules dans le même interpréteur que le notebook appelant.
# # **Chemin dans le dépôt**
# # TEST_MOTUL/TST_4/QUOTIDIEN/NB_HRIS_LIB.Notebook/notebook-content.py
# # **Description fonctionnelle**
# # Ce notebook est la bibliothèque commune du hub d'intégration des ressources humaines de Motul. Ce hub extrait les salariés depuis TalentSoft, applique les règles de transformation de l'existant, compare le résultat à l'état validé précédent et prépare la publication vers l'outil de paie ADP. Le notebook contient l'unique copie du code métier, organisé en vingt et un modules Python regroupés sous le nom hris. Auparavant, chaque notebook embarquait sa propre copie complète de ce code ; la copie unique supprime ce dédoublement et garantit que tous les notebooks exécutent exactement le même code. Le notebook ne lit et n'écrit aucune donnée et ne prépare aucune configuration d'environnement : il rend seulement des fonctions disponibles.
# # **Dépendances**
# # Le notebook n'appelle aucun autre notebook et ne lit aucune table. Pour vérifier l'empreinte de chaque module et l'enregistrer dans l'interpréteur, il n'utilise que la bibliothèque standard de Python. Les modules s'appuient ensuite, au moment où les notebooks métier les appellent, sur Spark et sur le format de table Delta fournis par Fabric, ainsi que sur la bibliothèque requests pour les appels réseau de l'acquisition, qui lit aussi les référentiels TalentSoft alimentant les tables de correspondance. Onze notebooks dépendent de celui-ci par la commande %run T4_NB_HRIS_LIB, répartis dans trois dossiers. Le dossier QUOTIDIEN, celui de ce notebook, contient T4_NB_HRIS_RUN_CONTROL, T4_NB_HRIS_TS_ACQUIRE, T4_NB_HRIS_REMUNERATION, T4_NB_HRIS_TRANSFORM, T4_NB_HRIS_COMPARE, T4_NB_HRIS_PROMOTE et T4_NB_HRIS_PUBLISH_ADP. Le dossier ONESHOT contient T4_NB_HRIS_SETUP, T4_NB_HRIS_REFERENCE_ADMIN et T4_NB_HRIS_BOOTSTRAP. Le dossier TESTS contient T4_NB_HRIS_TESTS. La commande %run désigne ce notebook par son seul nom. Son fonctionnement d'un dossier à l'autre n'est pas décrit par la documentation de Fabric : le premier lancement du pipeline T4_PL_HRIS_BOOTSTRAP le vérifie dès son activité Setup. Tout renommage de ce notebook impose de corriger cette commande dans chacun d'eux.
# # **Fonctionnement et logique de traitement**
# # La première cellule de code importe les quatre bibliothèques standard nécessaires. La cellule suivante définit la fonction register_hris_module. Celle-ci vérifie l'empreinte cryptographique du texte d'un module, calculée avec l'algorithme de hachage sécurisé SHA-256, puis crée le module dans l'interpréteur et l'exécute. En cas d'erreur, elle affiche le nom du module et la trace complète avant de relancer l'erreur. Chacune des vingt et une cellules suivantes enregistre un module, dans l'ordre de leurs dépendances : constantes et assainissement des messages, émulation des règles du langage de requête de l'existant, modèle de données, acquisition TalentSoft, transformation, comparaison, contrôle du run, publication, étapes du pipeline, lecture des journaux historiques, journal des cellules, tables de correspondance (partie TalentSoft lue par l'API, partie ADP saisie par le métier), flux de rémunération TalentSoft par API, puis premier chargement guidé (choix du mode et action suivante). La dernière cellule rassemble les empreintes chargées et affiche une ligne de synthèse. Pour modifier un module, il faut remplacer son texte puis recalculer son empreinte. Une empreinte qui ne correspond pas bloque le chargement, ce qui empêche d'exécuter un code modifié par erreur.
# # **Paramètres**
# # Ce notebook n'a aucun paramètre. Il s'exécute dans le contexte du notebook qui le charge et partage ses variables. Aucun nom de workspace, de Lakehouse ou de secret n'y est lu.
# # **Sorties produites**
# # Le notebook ne produit ni table ni fichier. Il enregistre les modules hris dans l'interpréteur du notebook appelant. Il définit aussi la fonction register_hris_module, que le notebook de tests réutilise pour charger ses propres modules, et le dictionnaire hris_library_digests, qui associe chaque module à son empreinte et prouve la version exécutée.
# # **Limitations connues et points d'attention**
# # Le chargement par %run se fait par le nom affiché du notebook dans le même workspace. Si le préfixe de version change, les commandes %run de tous les notebooks doivent être mises à jour. Faute de quoi, les notebooks échouent dès leur deuxième cellule, ce qui est préférable au chargement silencieux d'une autre version. Le texte des modules n'est plus produit par l'outil build_notebooks.py du dossier d'implémentation externe : ce notebook fait désormais référence et cet outil ne doit plus servir à régénérer ce dossier. Le simulateur de l'outil de paie n'est volontairement pas présent ici : il n'existe que dans le notebook de tests. Le temps de chargement dans Fabric n'a pas été mesuré.
# # **Responsable et contact**
# # Thibault Delalande, développeur du hub. Le responsable définitif de la maintenance reste à confirmer par l'équipe en charge des ressources humaines.

# CELL ********************

# ============= IMPORTS =============
# Standard libraries
import hashlib
import sys
import traceback
import types

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

def register_hris_module(name: str, digest: str, source: str) -> str:
    """
    Vérifie l'empreinte d'un module, l'enregistre dans l'interpréteur puis l'exécute.

    Args:
        name: Nom complet du module (« hris.common » ; les modules de test n'ont pas de préfixe).
        digest: Empreinte SHA-256 attendue du texte du module.
        source: Texte Python du module.

    Returns:
        L'empreinte vérifiée, conservée sur le module (attribut __hris_digest__).
    """
    try:
        if hashlib.sha256(source.encode("utf-8")).hexdigest() != digest:
            raise RuntimeError("empreinte invalide : " + name)
        package = sys.modules.get("hris") or types.ModuleType("hris")
        package.__path__ = []
        sys.modules["hris"] = package
        module = types.ModuleType(name)
        module.__file__ = name.replace(".", "/") + ".py"
        module.__hris_digest__ = digest
        sys.modules[name] = module
        exec(compile(source, module.__file__, "exec"), module.__dict__)
        if name.startswith("hris."):
            setattr(package, name.split(".", 1)[1], module)
        return digest
    except Exception as error:
        print(f"[T4_NB_HRIS_LIB] ❌ chargement du module {name} impossible : {type(error).__name__}")
        traceback.print_exc()
        raise

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.common =============
register_hris_module('hris.common', 'da64004aee4e83cbbedd33fdfbbed057845a34037de058aa2c8d68b3df1051db', r'''"""hris.common — constantes, erreurs, assainissement et utilitaires partagés.

Python pur (aucune dépendance Spark) : importable en local pour les tests et
embarqué tel quel dans les notebooks Fabric (voir tools/build_notebooks.py).
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
from typing import Any, Iterable, Optional

try:  # Python >= 3.9
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover
    ZoneInfo = None  # type: ignore

PROCESS_CODE = "TS_ADP_EMPLOYEE"

# Source TalentSoft : le tenant de PRODUCTION est la source réelle dans TOUS les
# environnements (consigne du mandat). Aucune bascule vers un tenant de test,
# aucun repli : toute autre URL est refusée (voir ts_acquire.assert_prd_source).
TS_PRD_BASE_URL = "https://motul.talent-soft.com"
TS_SECRET_CLIENT_ID = "client-id-ts"
TS_SECRET_CLIENT_SECRET = "client-secret-ts"

RUN_MODES = ("NORMAL", "INIT_REFERENCE", "RETRY_PUBLICATION")  # SHADOW : non retenu (D-CUT-03 À CONFIRMER)

# Statuts de ctl.run (05_TABLES_LOG_LAKEHOUSE.md)
RUN_STARTED = "STARTED"
RUN_ACQUIRED = "ACQUIRED"
RUN_TRANSFORMED = "TRANSFORMED"
RUN_COMPARED = "COMPARED"
RUN_BLOCKED = "BLOCKED"
RUN_VALIDATED = "VALIDATED"
RUN_PROMOTED = "PROMOTED"
RUN_PUBLISHED = "PUBLISHED"
RUN_PUBLISHED_PARTIAL = "PUBLISHED_PARTIAL"
RUN_NOT_PUBLISHED_NON_PRD = "NOT_PUBLISHED_NON_PRD"
RUN_SUCCEEDED = "SUCCEEDED"
RUN_FAILED = "FAILED"
RUN_SKIPPED_CONCURRENT = "SKIPPED_CONCURRENT"

SEVERITY_BLOCKING = "BLOCKING"
SEVERITY_WARNING = "WARNING"
SEVERITY_INFO = "INFO"


class HrisError(Exception):
    """Erreur fonctionnelle portant un code stable. Le message ne doit jamais
    contenir de donnée personnelle ni de secret (il est journalisé)."""

    def __init__(self, code: str, message: str = ""):
        super().__init__(f"{code}: {message}" if message else code)
        self.code = code
        self.safe_message = message


def utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


def to_utc_naive(value: dt.datetime) -> dt.datetime:
    """Horodatage UTC sans fuseau, convention des colonnes TIMESTAMP (session Spark en UTC)."""
    if value.tzinfo is not None:
        value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
    return value


def parse_trigger_time(value: str) -> dt.datetime:
    """@pipeline().TriggerTime (ISO 8601, UTC) -> datetime UTC aware."""
    if not value:
        raise HrisError("PARAM_INVALID", "p_trigger_time vide")
    text = value.strip().replace("Z", "+00:00")
    # fractions de seconde à 7 chiffres possibles dans les expressions Fabric
    text = re.sub(r"(\.\d{6})\d+", r"\1", text)
    try:
        parsed = dt.datetime.fromisoformat(text)
    except ValueError as exc:
        raise HrisError("PARAM_INVALID", "p_trigger_time illisible") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=dt.timezone.utc)
    return parsed.astimezone(dt.timezone.utc)


def business_date_of(trigger_utc: dt.datetime, timezone: str) -> dt.date:
    """Date métier = date locale du déclenchement (A08 D-RUN-02, À CONFIRMER : Europe/Paris)."""
    if ZoneInfo is None:  # pragma: no cover
        raise HrisError("TZ_UNAVAILABLE")
    return trigger_utc.astimezone(ZoneInfo(timezone)).date()


# ---------------------------------------------------------------------------
# Assainissement : aucun secret, aucune donnée personnelle dans les journaux
# ---------------------------------------------------------------------------
_SANITIZE_PATTERNS = [
    (re.compile(r"(?i)bearer\s+[A-Za-z0-9\-._~+/]+=*"), "Bearer ***"),
    (re.compile(r"(?i)(client_secret|client_id|access_token|password|token)=([^&\s]+)"), r"\1=***"),
    (re.compile(r"https?://[^\s'\"]+"), "<url>"),
    (re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"), "<email>"),
    (re.compile(r"\b[A-Z]{2}\d{2}[A-Z0-9 ]{10,30}\b"), "<iban>"),
    (re.compile(r"\d{4,}"), "<n>"),
]


def sanitize(text: Any, max_len: int = 500, mask_numbers: bool = True) -> str:
    """Rend un texte publiable dans ctl/log : masque URL, e-mails, IBAN, jetons et, par défaut,
    toute séquence de 4 chiffres ou plus (matricules, NIR, montants). `mask_numbers=False` est
    réservé aux messages GÉNÉRÉS par le code à partir de compteurs (aucune valeur métier)."""
    if text is None:
        return ""
    out = str(text)
    patterns = _SANITIZE_PATTERNS if mask_numbers else _SANITIZE_PATTERNS[:-1]
    for pattern, repl in patterns:
        out = pattern.sub(repl, out)
    return out[:max_len]


def safe_error(exc: BaseException) -> str:
    """Résumé d'exception publiable : classe + code fonctionnel, message assaini."""
    if isinstance(exc, HrisError):
        return f"{exc.code}: {sanitize(exc.safe_message, 300)}"
    return f"{type(exc).__name__}: {sanitize(str(exc), 300)}"


def sha256_hex(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def canonical_json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str)


def json_exit(payload: dict) -> str:
    """Valeur de sortie d'un notebook : JSON court, sans donnée personnelle."""
    return json.dumps(payload, ensure_ascii=False, default=str, separators=(",", ":"))


def require(params: dict, names: Iterable[str]) -> None:
    missing = [n for n in names if params.get(n) in (None, "")]
    if missing:
        raise HrisError("PARAM_MISSING", ",".join(missing))


def as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("1", "true", "yes", "oui")


class SecretValue:
    """Enveloppe d'une valeur secrète : jamais affichée par repr/str."""

    __slots__ = ("_v",)

    def __init__(self, value: Optional[str]):
        self._v = value

    def reveal(self) -> Optional[str]:
        return self._v

    def __repr__(self) -> str:
        return "SecretValue(***)"

    __str__ = __repr__
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.parity =============
register_hris_module('hris.parity', '12d5d086c85cf5a0c61aadd96e0bd3edbf1a17fa2ecbedc7346161bbb6a7854b', r'''"""hris.parity — émulation Python des sémantiques T-SQL du SQL de production (A09).

Implémentation de RÉFÉRENCE (Python pur) des règles A02 ; la version Spark
(hris.transform) doit produire exactement les mêmes résultats sur les jeux de tests
(test de parité exécuté en local et dans Fabric).

Points d'émulation non exacts, mesurés en non-régression (A08 D-PAR-01, D-PAR-02) :
  * ISDATE : formats numériques usuels uniquement (pas de mois en lettres) ;
  * ISNUMERIC : grammaire approchée (signe, symbole monétaire, virgules, exposant).
"""
from __future__ import annotations

import datetime as dt
import decimal
import json
import re
from typing import Any, Optional

# ---------------------------------------------------------------------------
# Lecture JSON (R-SQL-002) : équivalent de JSON_VALUE(..., '$.<clé>') en mode lax
# ---------------------------------------------------------------------------
JSON_VALUE_MAX_LEN = 4000  # JSON_VALUE renvoie NULL au-delà de 4000 caractères (mode lax)


def json_value_repr(value: Any) -> Optional[str]:
    """Texte que JSON_VALUE renverrait pour la valeur Python sérialisée par json.dump :
    None -> NULL ; booléen -> 'true'/'false' ; nombre -> représentation JSON ;
    chaîne -> la chaîne (NULL si > 4000) ; objet ou tableau -> NULL."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return json.dumps(value)
    if isinstance(value, str):
        return value if len(value) <= JSON_VALUE_MAX_LEN else None
    return None  # dict / list : non scalaire


def nvarchar(value: Optional[str], length: int = 50) -> Optional[str]:
    """CAST(x AS NVARCHAR(n)) : troncature silencieuse à n unités UTF-16 (R-SQL-002, D-PAR-03)."""
    if value is None:
        return None
    if all(ord(ch) <= 0xFFFF for ch in value):
        return value[:length]
    units = value.encode("utf-16-le", "surrogatepass")[: 2 * length]
    return units.decode("utf-16-le", "surrogatepass")


# ---------------------------------------------------------------------------
# Comparaisons de chaînes : SQL Server ignore les espaces de fin (= , <>, LEN)
# ---------------------------------------------------------------------------
def sql_rtrim(value: Optional[str]) -> Optional[str]:
    return None if value is None else value.rstrip(" ")


def sql_is_empty(value: Optional[str]) -> bool:
    """`x = ''` : vrai pour '' et '   ' ; NULL -> inconnu (faux dans un CASE WHEN)."""
    return value is not None and value.rstrip(" ") == ""


def sql_is_not_empty(value: Optional[str]) -> bool:
    """`x <> ''` : NULL -> inconnu (faux)."""
    return value is not None and value.rstrip(" ") != ""


def sql_len(value: Optional[str]) -> Optional[int]:
    return None if value is None else len(value.rstrip(" "))


def sql_eq(a: Optional[str], b: Optional[str]) -> bool:
    """Égalité de jointure : NULL ne correspond à rien ; espaces de fin ignorés."""
    return a is not None and b is not None and a.rstrip(" ") == b.rstrip(" ")


# ---------------------------------------------------------------------------
# ISDATE (R-SQL-033/034, R-SQL-043) — émulation us_english / DATEFORMAT mdy / datetime
# ---------------------------------------------------------------------------
_WS = " \t\r\n"
# Motifs acceptés (D-PAR-01). Les MÊMES expressions régulières gardent l'analyse Spark (hris.transform) :
# (garde, extraction de la partie date, ordre y/m/d)
_HH = r"(?:[01]\d|2[0-3])"
_HH1 = r"(?:[01]?\d|2[0-3])"
_MS = r"[0-5]\d"
DATE_PATTERNS = [
    ("ymd_dash", r"^\d{4}-\d{1,2}-\d{1,2}$", r"^(\d{4})-(\d{1,2})-(\d{1,2})$", "ymd"),
    ("ymd_compact", r"^\d{8}$", r"^(\d{4})(\d{2})(\d{2})$", "ymd"),
    ("ymd_iso_time", rf"^\d{{4}}-\d{{2}}-\d{{2}}T{_HH}:{_MS}(?::{_MS}(?:\.\d{{1,3}})?)?Z?$",
     r"^(\d{4})-(\d{2})-(\d{2})T", "ymd"),
    ("ymd_space_time", rf"^\d{{4}}-\d{{1,2}}-\d{{1,2}} {_HH1}:{_MS}(?::{_MS}(?:\.\d{{1,3}})?)?$",
     r"^(\d{4})-(\d{1,2})-(\d{1,2}) ", "ymd"),
    ("mdy_slash", r"^\d{1,2}/\d{1,2}/\d{4}$", r"^(\d{1,2})/(\d{1,2})/(\d{4})$", "mdy"),
]
_COMPILED = [(n, re.compile(g), re.compile(x), o) for n, g, x, o in DATE_PATTERNS]
DATETIME_MIN_YEAR = 1753


def parse_tsql_date(value: Optional[str]) -> Optional[dt.date]:
    """Date reconnue (partie date) ou None. Les espaces de début et de fin sont tolérés."""
    if value is None:
        return None
    text = value.strip(_WS)
    for _, guard, extract, order in _COMPILED:
        if not guard.match(text):
            continue
        g = extract.match(text).groups()
        y, mo, d = (int(g[0]), int(g[1]), int(g[2])) if order == "ymd" else (int(g[2]), int(g[0]), int(g[1]))
        try:
            return dt.date(y, mo, d)
        except ValueError:
            return None
    return None


def isdate(value: Optional[str]) -> int:
    """ISDATE(x) : 1 si convertible en datetime (année >= 1753), 0 sinon ; NULL -> 0."""
    parsed = parse_tsql_date(value)
    return 1 if parsed is not None and parsed.year >= DATETIME_MIN_YEAR else 0


SQL_EMPTY_DATE = dt.date(1900, 1, 1)


def convert_date(value: Optional[str]) -> Optional[dt.date]:
    """CONVERT(DATE, x) : NULL -> NULL ; '' (espaces compris) -> 1900-01-01 (R-SQL-052) ;
    valeur invalide -> erreur de conversion (le run échoue en parité)."""
    if value is None:
        return None
    if value.strip(_WS) == "":
        return SQL_EMPTY_DATE
    parsed = parse_tsql_date(value)
    if parsed is None:
        raise ValueError("CONVERSION_DATE")
    return parsed


# ---------------------------------------------------------------------------
# ISNUMERIC et CAST NUMERIC(12,2) (R-SQL-035/036/044/054)
# ---------------------------------------------------------------------------
_ISNUMERIC = re.compile(
    r"^[+-]?[$€£¥]?[ ]*[+-]?"
    r"(?P<body>\d[\d,]*\.?\d*|\.\d*|,[\d,]*\.?\d*)?"
    r"(?P<exp>[eEdD][+-]?\d+)?$"
)
_STRICT_DECIMAL = re.compile(r"^[ \t]*([+-]?)(\d+\.?\d*|\.\d+)[ \t]*$")
_NUMERIC_12_2_LIMIT = decimal.Decimal("10000000000")  # 10^(12-2)


def isnumeric(value: Optional[str]) -> int:
    """ISNUMERIC(x) approché (A08 D-PAR-02). NULL et '' -> 0."""
    if value is None:
        return 0
    text = value.strip(" \t\r\n")
    if text == "":
        return 0
    m = _ISNUMERIC.match(text)
    if not m:
        return 0
    body = m.group("body") or ""
    if m.group("exp") and not re.search(r"\d", body):
        return 0  # exposant sans mantisse (ex. 'e5')
    return 1


def cast_numeric_12_2(value: Optional[str]) -> Optional[decimal.Decimal]:
    """CAST(x AS NUMERIC(12,2)) strict : arrondi à 2 décimales (au plus proche, à distance
    de zéro en cas d'égalité). Toute valeur non convertible ou hors capacité lève
    ValueError (erreur de conversion T-SQL => échec du run en parité)."""
    if value is None:
        return None
    m = _STRICT_DECIMAL.match(value)
    if not m:
        raise ValueError("CONVERSION_NUMERIC")
    number = decimal.Decimal(m.group(1) + m.group(2))
    rounded = number.quantize(decimal.Decimal("0.01"), rounding=decimal.ROUND_HALF_UP)
    if abs(rounded) >= _NUMERIC_12_2_LIMIT:
        raise ValueError("NUMERIC_OVERFLOW")
    return rounded


def monthly_salary(value: Optional[str]) -> Optional[decimal.Decimal]:
    """CAST(mt_salaire_base AS NUMERIC(12,2)) / 12 -> DECIMAL(23,13) (R-SQL-054)."""
    base = cast_numeric_12_2(value)
    if base is None:
        return None
    ctx = decimal.Context(prec=38, rounding=decimal.ROUND_HALF_UP)
    return ctx.divide(base, decimal.Decimal(12)).quantize(
        decimal.Decimal("1e-13"), rounding=decimal.ROUND_HALF_UP, context=ctx)


def pct_flag(value: Optional[str], upper: int) -> int:
    """Indicateur de rejet « pourcentage / montant » : ISNUMERIC = 0 -> 0 ; sinon
    CAST(...) BETWEEN 0 AND upper -> 0, hors bornes -> 1 ; échec de CAST -> ValueError."""
    if isnumeric(value) == 0:
        return 0
    number = cast_numeric_12_2(value)
    return 0 if decimal.Decimal(0) <= number <= decimal.Decimal(upper) else 1


def utf8_truncate(text: Optional[str], max_bytes: int) -> Optional[str]:
    """CAST(... AS VARCHAR(n)) avec collation UTF-8 : n OCTETS (résumé R-SQL-066).
    Ne coupe jamais un caractère multi-octets (comportement exact SQL Server À CONFIRMER)."""
    if text is None:
        return None
    raw = text.encode("utf-8")
    if len(raw) <= max_bytes:
        return text
    return raw[:max_bytes].decode("utf-8", "ignore")
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.model =============
register_hris_module('hris.model', 'b386f3e30b4d9b1d4327a303d061c7cc2078aaa6e6acede45dba5742f1a948f8', r'''"""hris.model — contrat de données : colonnes, correspondances, règles de rejet, DDL.

Sources : A02 (R-SQL-004, § 4, § 5, R-SQL-056), A03 (17 tables), A07, 05 (DDL des tables
de contrôle et de log, qui fait foi). Python pur.

Extensions par rapport au DDL documenté (signalées dans docs/ECARTS_ET_DECISIONS.md) :
  * ctl.run.target_run_id (mode RETRY_PUBLICATION) ;
  * tables PROPOSÉES de publication par opération (pub.adp_operation, pub.adp_attempt,
    pub.adp_reconciliation, log.adp_manual_decision) : leur sémantique dépend du contrat
    final ADP, absent (docs/CONTRAT_ADP_POINTS_D_USAGE.md) ;
  * log.cell_execution : journal des cellules de code exécutées (hris.cell_log), sans donnée
    personnelle (refactorisation TST_4).
"""
from __future__ import annotations

from typing import Dict, List, Tuple

# ---------------------------------------------------------------------------
# R-SQL-004 : 69 colonnes source (clé JSON -> colonne), toutes NVARCHAR(50)
# ---------------------------------------------------------------------------
SOURCE_COLUMNS: List[Tuple[str, str]] = [
    ("employeeNumber", "id_unique"),
    ("LegacyID", "id_payroll"),
    ("Matricule_IT", "cd_matricule_it"),
    ("Identity_startDate", "dt_debut_identite"),
    ("LastName", "lb_nom"),
    ("FirstName", "lb_prenom"),
    ("BirthName", "lb_nom_naissance"),
    ("Title", "cd_emp_titre_ts"),
    ("Gender", "cd_emp_sexe_ts"),
    ("Nationality", "cd_emp_nationalite_ts"),
    ("hubSocialSecurityNumber", "id_numero_securite_sociale"),
    ("MaritalStatus_startDate", "dt_debut_situation_matrimoniale"),
    ("MaritalStatus", "cd_emp_situation_matrimoniale_ts"),
    ("CivilRegistration_startDate", "dt_debut_civil_registration"),
    ("Date_de_naissance", "dt_naissance"),
    ("Ville_de_naissance", "lb_naissance_ville"),
    ("code_postal_lieu_de_naissance", "lb_naissance_cp"),
    ("region_ou_etat_lieu_de_naissance", "lb_naissance_departement"),
    ("Pays_de_naissance", "cd_naissance_pays_ts"),
    ("LegalTownOrCity", "lb_fiscal_ville"),
    ("LegalCountry", "lb_fiscal_country"),
    ("LegalPostalCode", "lb_fiscal_cp"),
    ("LegalStreet", "lb_fiscal_street"),
    ("LegalStreetNumber", "lb_fiscal_streetnumber"),
    ("LegalAdditionalAddressInformation", "lb_fiscal_additional_address_information"),
    ("LegalStreetNumberComplement", "lb_fiscal_streetnumber_complement"),
    ("address_postalCode", "lb_adresse_cp"),
    ("adresse_city", "lb_adresse_ville"),
    ("adresse_countryCode", "cd_adresse_pays_ts"),
    ("ContactInformation_startDate", "dt_debut_contatct_info"),
    ("BusinessEmail", "lb_email_professionnel"),
    ("PersonalEmail", "lb_email_personnel"),
    ("BusinessMobile|BusinessPhone", "lb_telephone_professionnel"),  # COALESCE (R-SQL-003)
    ("PersonalMobile", "lb_telephone_personnel"),
    ("referenceAdmissionDate", "dt_entree_societe"),
    ("GroupStartDate", "dt_entree_groupe"),
    ("RecalculatedSeniorityDate", "dt_anciennete"),
    ("GroupEndDate", "dt_sortie"),
    ("CompanyStartDate", "dt_company_start"),
    ("HiringReason", "cd_ppc_raison_debut_contrat_ts"),
    ("TerminationReason", "cd_ppc_raison_fin_contrat_ts"),
    ("date_debut_contract", "dt_ppc_debut_contrat"),
    ("date_fin_contract", "dt_ppc_fin_contrat"),  # REPLACE(...,'None','') (R-SQL-003)
    ("NatureOfContract", "cd_ppc_contrat_ts"),
    ("FixedTermContractReason", "cd_recours_cdd"),
    ("OccupationalCategory", "cd_occupational_category"),
    ("FirstPeriodStartDate", "dt_ppc_debut_periode_essaie"),
    ("FirstPeriodEndDate", "dt_ppc_fin_periode_essaie"),
    ("Situation_starDate", "dt_situation_start"),
    ("Coefficient", "cd_ppc_convention_collective_lvl2_ts"),
    ("LegalStructure", "cd_geo_entite_ts"),
    ("GeographicOrganisationStructure", "cd_ppc_location_country_ts"),
    ("PPC_FTE", "pc_ppc_fte"),
    ("OrganisationalStructure", "cd_organisationalstructure_ts"),
    ("numero_ordre", "id_numero_ordre"),
    ("CostCenter", "cd_ppc_cost_center_ts"),
    ("pourcentage_imputation", "pc_imputation"),
    ("AccountHolder", "lb_account_holder"),
    ("IBANOrABA", "lb_iban"),
    ("BicOrSwift", "lb_bic"),
    ("AccountNumber", "lb_account_number"),
    ("BankName", "lb_bank_name"),
    ("base_startDate", "dt_debut_mt_salaire_base"),
    ("salary_base", "mt_salaire_base"),
    ("pourcentage_prime_annuelles_startDate", "dt_debut_pc_prime_annuelle"),
    ("Pourcentage_prime_annuelles", "pc_prime_annuelle"),
    ("prime_annuelles_startDate", "dt_debut_prime_annuelle"),
    ("prime_annuelles", "prime_annuelle"),
    ("WorkingType", "lb_job_title"),
]
assert len(SOURCE_COLUMNS) == 69

# Clés JSON persistées dans stg.ts_employee_raw : uniquement celles lues en aval
# (minimisation A08 D-SEC-02 : ni données de santé ni champs inutilisés).
RAW_KEYS: List[str] = []
for _k, _c in SOURCE_COLUMNS:
    for _part in _k.split("|"):
        if _part not in RAW_KEYS:
            RAW_KEYS.append(_part)
assert len(RAW_KEYS) == 70

SOURCE_COLUMN_NAMES = [c for _, c in SOURCE_COLUMNS]

# Clés produites par la Function TalentSoft mais jamais persistées (santé, inutilisées)
TS_KEYS_NOT_PERSISTED = [
    "handicap", "handicap_startDate", "workaccident", "workaccident_startDate",
    "motif_visite_medical", "motif_visite_medical_startDate", "Position", "CollectiveAgreement",
    "Hors_France", "Qualification", "Qualification_startDate", "ConferredDate", "PersonalPhone",
    "Status", "PaymentMethod", "UseSameAccountForExpense", "ExpenseAccountPaymentMethod",
    "ExpenseAccountHolder", "ExpenseAccountIBANOrABA", "ExpenseAccountBicOrSwift",
    "ExpenseAccountBankName", "Country", "ExpenseAccountCountry", "date_debut_bank", "date_fin_bank",
    "SecondPeriodStartDate", "SecondPeriodEndDate", "ProbationaryPeriod_startDate", "LegacyID_startDate",
]

# ---------------------------------------------------------------------------
# A03 : 17 tables de correspondance (ordre des colonnes = ordre du fichier d'origine)
# ---------------------------------------------------------------------------
LOOKUP_TABLES: Dict[str, List[str]] = {
    "lu_hr_emp_diplome": ["cd_emp_diplome_ts", "ds_emp_diplome_ts", "cd_emp_diplome_adp", "ds_emp_diplome_adp",
                          "or_emp_diplome", "cd_emp_diplome_api_ts", "ds_emp_diplome_api_ts"],
    "lu_hr_emp_nationalite": ["id_emp_nationalite_ts", "cd_emp_nationalite_ts", "ds_emp_nationalite_ts",
                              "cd_emp_nationalite_adp", "ds_emp_nationalite_adp", "or_emp_nationalite"],
    "lu_hr_emp_pays": ["id_emp_pays_ts", "cd_emp_pays_ts", "ds_emp_pays_ts", "id_emp_pays_adp",
                       "ds_emp_pays_adp", "or_emp_pays"],
    "lu_hr_emp_situation_matrimoniale": ["cd_emp_situation_matrimoniale_ts", "ds_emp_situation_matrimoniale_ts",
                                         "cd_emp_situation_matrimoniale_adp", "ds_emp_situation_matrimoniale_adp",
                                         "or_emp_situation_matrimoniale"],
    "lu_hr_emp_titre": ["cd_emp_titre_ts", "lb_emp_titre_ts", "ds_emp_titre_ts", "cd_emp_titre_adp",
                        "ds_emp_titre_adp", "or_emp_titre"],
    "lu_hr_emp_sexe": ["cd_emp_sexe_ts", "ds_emp_sexe_ts", "cd_emp_sexe_adp", "ds_emp_sexe_adp", "or_emp_sexe"],
    "lu_hr_emp_work_accident": ["cd_emp_work_accident_ts", "ds_emp_work_accident_ts", "cd_emp_work_accident_adp",
                                "ds_emp_work_accident_adp", "or_emp_work_accident"],
    "lu_hr_ppc_convention_collective": ["cd_ppc_convention_collective_ts", "ds_ppc_convention_collective_ts",
                                        "cd_ppc_convention_collective_adp", "ds_ppc_convention_collective_adp",
                                        "or_ppc_convention_collective"],
    "lu_hr_ppc_classification": ["cd_convention_collective_lvl2_ts", "lb_convention_collective_lvl2_ts",
                                 "cd_classification_adp", "lb_classification_adp", "cd_coefficient_adp",
                                 "cd_ppc_convention_collective_adp", "cd_ppc_convention_collective_ts"],
    "lu_hr_ppc_classes": ["cd_occupational_category_ts", "lb_fr_occupational_category_ts",
                          "lb_en_occupational_category_ts", "cd_categorie_cotisant_adp", "lb_categorie_cotisant_adp",
                          "cd_classe_remuneration_adp", "lb_classe_remuneration_adp"],
    "lu_hr_geo_entite": ["cd_geo_entite_ts", "ds_geo_entite_ts", "cd_geo_parent_entite_ts", "cd_geo_entite_adp",
                         "ds_geo_entite_adp", "or_geo_entite_adp"],
    "lu_hr_ppc_contrat": ["cd_ppc_contrat_ts", "ds_ppc_contrat_ts", "cd_ppc_contrat_adp", "ds_ppc_contrat_adp",
                          "cd_ppc_contrat_type", "ds_ppc_contrat_type", "or_ppc_contrat"],
    "lu_hr_ppc_contrat2": ["cd_contract_ts", "lb_contract_ts", "cd_nature_contract_adp", "lb_nature_contract_adp",
                           "cd_type_contract_adp", "lb_type_contract_adp", "cd_type_collaboration_adp",
                           "lb_type_collaboration_adp"],
    "lu_hr_ppc_cost_center": ["cd_ppc_cost_center_ts", "ds_ppc_cost_center_ts", "cd_ppc_cost_center_adp",
                              "ds_ppc_cost_center_adp", "or_ppc_cost_center"],
    "lu_hr_ppc_location_country": ["cd_ppc_location_country_ts", "ds_ppc_location_country_ts",
                                   "cd_ppc_location_country_adp", "ds_ppc_location_country_adp",
                                   "or_ppc_location_country"],
    "lu_hr_ppc_raison_debut_contrat": ["cd_ppc_raison_debut_contrat_ts", "ds_ppc_raison_debut_contrat_ts",
                                       "cd_ppc_raison_debut_contrat_adp", "ds_ppc_raison_debut_contrat_adp",
                                       "or_ppc_raison_debut_contrat"],
    "lu_hr_organisationalstructure": ["cd_organisationalstructure_ts", "ds_organisationalstructure_ts",
                                      "cd_organisationalstructure_adp", "ds_organisationalstructure_adp",
                                      "or_organisationalstructure"],
}
assert len(LOOKUP_TABLES) == 17
LOOKUP_TECH_COLUMNS = ["_loaded_at_utc", "_loaded_by", "_change_ref"]

# Jointures (R-SQL-010 à 024) : alias, table, colonne source, clé de la table,
# colonnes ramenées (nom dans la table -> nom en staging), dans l'ordre de PROC2.
LOOKUP_JOINS: List[dict] = [
    {"rule": "R-SQL-010", "table": "lu_hr_emp_titre", "src": "cd_emp_titre_ts", "key": "cd_emp_titre_ts",
     "out": [("cd_emp_titre_ts", "cd_emp_titre_ts_lu"), ("cd_emp_titre_adp", "cd_emp_titre_adp")]},
    {"rule": "R-SQL-011", "table": "lu_hr_emp_sexe", "src": "cd_emp_sexe_ts", "key": "cd_emp_sexe_ts",
     "out": [("cd_emp_sexe_ts", "cd_emp_sexe_ts_lu"), ("cd_emp_sexe_adp", "cd_emp_sexe_adp")]},
    {"rule": "R-SQL-012", "table": "lu_hr_emp_nationalite", "src": "cd_emp_nationalite_ts",
     "key": "cd_emp_nationalite_ts",
     "out": [("cd_emp_nationalite_ts", "cd_emp_nationalite_ts_lu"), ("cd_emp_nationalite_adp", "cd_emp_nationalite_adp")]},
    {"rule": "R-SQL-013", "table": "lu_hr_emp_situation_matrimoniale", "src": "cd_emp_situation_matrimoniale_ts",
     "key": "cd_emp_situation_matrimoniale_ts",
     "out": [("cd_emp_situation_matrimoniale_ts", "cd_emp_situation_matrimoniale_ts_lu"),
             ("cd_emp_situation_matrimoniale_adp", "cd_emp_situation_matrimoniale_adp")]},
    {"rule": "R-SQL-014", "table": "lu_hr_emp_pays", "src": "cd_naissance_pays_ts", "key": "cd_emp_pays_ts",
     "out": [("cd_emp_pays_ts", "cd_naissance_pays_ts_lu"), ("id_emp_pays_adp", "id_naissance_pays_adp")]},
    {"rule": "R-SQL-015", "table": "lu_hr_emp_pays", "src": "cd_adresse_pays_ts", "key": "cd_emp_pays_ts",
     "out": [("cd_emp_pays_ts", "cd_adresse_pays_ts_lu"), ("id_emp_pays_adp", "id_adresse_pays_adp")]},
    {"rule": "R-SQL-016", "table": "lu_hr_emp_pays", "src": "lb_fiscal_country", "key": "cd_emp_pays_ts",
     "out": [("cd_emp_pays_ts", "cd_fiscal_pays_ts_lu"), ("id_emp_pays_adp", "id_fiscal_pays_adp")]},
    {"rule": "R-SQL-017", "table": "lu_hr_ppc_raison_debut_contrat", "src": "cd_ppc_raison_debut_contrat_ts",
     "key": "cd_ppc_raison_debut_contrat_ts",
     "out": [("cd_ppc_raison_debut_contrat_ts", "cd_ppc_raison_debut_contrat_ts_lu"),
             ("cd_ppc_raison_debut_contrat_adp", "cd_ppc_raison_debut_contrat_adp")]},
    {"rule": "R-SQL-018", "table": "lu_hr_ppc_contrat2", "src": "cd_ppc_contrat_ts", "key": "cd_contract_ts",
     "out": [("cd_contract_ts", "cd_ppc_contrat_ts_lu"), ("cd_nature_contract_adp", "cd_nature_contract_adp"),
             ("cd_type_contract_adp", "cd_type_contract_adp"),
             ("cd_type_collaboration_adp", "cd_type_collaboration_adp")]},
    {"rule": "R-SQL-019", "table": "lu_hr_ppc_classification", "src": "cd_ppc_convention_collective_lvl2_ts",
     "key": "cd_convention_collective_lvl2_ts",
     "out": [("cd_convention_collective_lvl2_ts", "cd_ppc_convention_collective_lvl2_ts_lu"),
             ("cd_classification_adp", "cd_classification_adp"), ("cd_coefficient_adp", "cd_coefficient_adp"),
             ("cd_ppc_convention_collective_adp", "cd_ppc_convention_collective_adp")]},
    {"rule": "R-SQL-020", "table": "lu_hr_ppc_classes", "src": "cd_occupational_category",
     "key": "cd_occupational_category_ts",
     "out": [("cd_occupational_category_ts", "cd_occupational_category_ts_lu"),
             ("cd_categorie_cotisant_adp", "cd_categorie_cotisant_adp"),
             ("cd_classe_remuneration_adp", "cd_classe_remuneration_adp")]},
    {"rule": "R-SQL-021", "table": "lu_hr_geo_entite", "src": "cd_geo_entite_ts", "key": "cd_geo_entite_ts",
     "out": [("cd_geo_entite_ts", "cd_geo_entite_ts_lu"), ("cd_geo_entite_adp", "cd_geo_entite_adp")]},
    {"rule": "R-SQL-022", "table": "lu_hr_ppc_location_country", "src": "cd_ppc_location_country_ts",
     "key": "cd_ppc_location_country_ts",
     "out": [("cd_ppc_location_country_ts", "cd_ppc_location_country_ts_lu"),
             ("cd_ppc_location_country_adp", "cd_ppc_location_country_adp")]},
    {"rule": "R-SQL-023", "table": "lu_hr_ppc_cost_center", "src": "cd_ppc_cost_center_ts",
     "key": "cd_ppc_cost_center_ts",
     "out": [("cd_ppc_cost_center_ts", "cd_ppc_cost_center_ts_lu"), ("cd_ppc_cost_center_adp", "cd_ppc_cost_center_adp")]},
    {"rule": "R-SQL-024", "table": "lu_hr_organisationalstructure", "src": "cd_organisationalstructure_ts",
     "key": "cd_organisationalstructure_ts",
     "out": [("cd_organisationalstructure_ts", "cd_organisationalstructure_ts_lu"),
             ("cd_organisationalstructure_adp", "cd_organisationalstructure_adp"),
             ("ds_organisationalstructure_adp", "ds_organisationalstructure_adp")]},
]
USED_LOOKUP_TABLES = sorted({j["table"] for j in LOOKUP_JOINS})  # 13 tables (D-REF-02 : 4 non utilisées)
ENRICHED_LOOKUP_COLUMNS = [out for j in LOOKUP_JOINS for _, out in j["out"]]

# ---------------------------------------------------------------------------
# Rejets (A02 § 5, PROC3)
# ---------------------------------------------------------------------------
NULL_REJECT_COLUMNS = ["id_payroll", "cd_matricule_it", "id_numero_securite_sociale"]

# Ordre exact de la clause VALUES de PROC3 (Format_Rejet) ; kind : len8 / date / date10 / pct / amount
FORMAT_REJECT_COLUMNS: List[Tuple[str, str]] = [
    ("id_payroll", "len8"),
    ("dt_naissance", "date"),
    ("dt_situation_start", "date"),
    ("dt_entree_societe", "date10"),
    ("dt_entree_groupe", "date"),
    ("dt_anciennete", "date"),
    ("dt_sortie", "date"),
    ("dt_company_start", "date"),
    ("dt_ppc_debut_contrat", "date"),
    ("dt_ppc_fin_contrat", "date"),
    ("dt_ppc_debut_periode_essaie", "date"),
    ("dt_ppc_fin_periode_essaie", "date"),
    ("pc_ppc_fte", "pct"),
    ("dt_debut_mt_salaire_base", "date"),
    ("dt_debut_pc_prime_annuelle", "date"),
    ("pc_prime_annuelle", "pct"),
    ("dt_debut_prime_annuelle", "date"),
    ("prime_annuelle", "amount"),
]
# R-SQL-037/038 : somme des indicateurs SANS pc_ppc_fte (anomalie de l'existant conservée, D-PAR-06)
FORMAT_SUM_EXCLUDED = {"pc_ppc_fte"}

# Référence_Rejet : (fichier, nom de colonne journalisé, valeur testée, colonne _lu, code ADP contrôlé|None)
REFERENCE_REJECTS: List[Tuple[str, str, str, str, object]] = [
    ("titre.csv", "cd_emp_titre", "cd_emp_titre_ts", "cd_emp_titre_ts_lu", "cd_emp_titre_adp"),
    ("sexe.csv", "cd_emp_sexe", "cd_emp_sexe_ts", "cd_emp_sexe_ts_lu", "cd_emp_sexe_adp"),
    ("nationalite.csv", "cd_emp_nationalite", "cd_emp_nationalite_ts", "cd_emp_nationalite_ts_lu",
     "cd_emp_nationalite_adp"),
    ("situation_matrimoniale.csv", "cd_emp_situation_matrimoniale", "cd_emp_situation_matrimoniale_ts",
     "cd_emp_situation_matrimoniale_ts_lu", "cd_emp_situation_matrimoniale_adp"),
    ("pays.csv", "cd_naissance_pays", "cd_naissance_pays_ts", "cd_naissance_pays_ts_lu", "id_naissance_pays_adp"),
    ("pays.csv", "cd_adresse_pays", "cd_adresse_pays_ts", "cd_adresse_pays_ts_lu", "id_adresse_pays_adp"),
    ("raison_debut_contrat.csv", "cd_ppc_raison_debut_contrat", "cd_ppc_raison_debut_contrat_ts",
     "cd_ppc_raison_debut_contrat_ts_lu", "cd_ppc_raison_debut_contrat_adp"),
    ("contrat_v2.csv", "cd_nature_contrat", "cd_ppc_contrat_ts", "cd_ppc_contrat_ts_lu", "cd_nature_contract_adp"),
    ("contrat_v2.csv", "cd_type_contrat", "cd_ppc_contrat_ts", "cd_ppc_contrat_ts_lu", "cd_type_contract_adp"),
    ("contrat_v2.csv", "cd_type_collaboration_adp", "cd_ppc_contrat_ts", "cd_ppc_contrat_ts_lu",
     "cd_type_collaboration_adp"),
    ("classes.csv", "cd_occupational_category", "cd_occupational_category", "cd_occupational_category_ts_lu", None),
    ("entite.csv", "cd_geo_entite", "cd_geo_entite_ts", "cd_geo_entite_ts_lu", "cd_geo_entite_adp"),
    ("location_country.csv", "cd_ppc_location_country", "cd_ppc_location_country_ts",
     "cd_ppc_location_country_ts_lu", "cd_ppc_location_country_adp"),
    ("cost_center.csv", "cd_ppc_cost_center", "cd_ppc_cost_center_ts", "cd_ppc_cost_center_ts_lu",
     "cd_ppc_cost_center_adp"),
    ("organisationalstructure.csv", "cd_organisationalstructure", "cd_organisationalstructure_ts",
     "cd_organisationalstructure_ts_lu", "cd_organisationalstructure_adp"),
]
# Pays fiscal (R-SQL-042) : indicateur calculé mais jamais émis ; classification commentée.

REJECT_TYPE_NULL = "Null_Rejet"
REJECT_TYPE_FORMAT = "Format_Rejet"
REJECT_TYPE_REFERENCE = "Référence_Rejet"

# ---------------------------------------------------------------------------
# R-SQL-056 : 72 colonnes de sortie (ordre exact) et typage
# ---------------------------------------------------------------------------
ADP_COLUMNS: List[str] = [
    "dt_transforme", "id_unique", "id_payroll", "cd_matricule_it", "lb_nom", "lb_prenom", "lb_nom_naissance",
    "cd_emp_titre_adp", "cd_emp_sexe_adp", "cd_emp_nationalite_adp", "id_numero_securite_sociale",
    "dt_debut_situation_matrimoniale", "cd_emp_situation_matrimoniale_adp", "dt_naissance", "lb_naissance_ville",
    "lb_naissance_cp", "lb_naissance_departement", "id_naissance_pays_adp", "lb_adresse_cp", "lb_adresse_ville",
    "id_adresse_pays_adp", "lb_email_professionnel", "lb_email_personnel", "lb_telephone_professionnel",
    "lb_telephone_personnel", "lb_fiscal_ville", "id_fiscal_pays_adp", "lb_fiscal_cp", "lb_fiscal_street",
    "lb_fiscal_streetnumber", "lb_fiscal_additional_address_information", "lb_fiscal_streetnumber_complement",
    "dt_entree_societe", "dt_entree_groupe", "dt_anciennete", "dt_sortie", "dt_company_start",
    "cd_ppc_raison_debut_contrat_adp", "cd_ppc_raison_fin_contrat_ts", "dt_ppc_debut_contrat",
    "dt_ppc_fin_contrat", "cd_nature_contract_adp", "cd_type_contract_adp", "cd_type_collaboration_adp",
    "cd_recours_cdd", "cd_categorie_cotisant_adp", "cd_classe_remuneration_adp", "dt_ppc_debut_periode_essaie",
    "dt_ppc_fin_periode_essaie", "cd_classification_adp", "cd_coefficient_adp", "cd_organisationalstructure_adp",
    "dt_situation_start", "cd_ppc_convention_collective_adp", "cd_geo_entite_adp", "cd_ppc_location_country_adp",
    "pc_ppc_fte", "id_numero_ordre", "cd_ppc_cost_center_adp", "pc_imputation", "dt_debut_mt_salaire_base",
    "mt_salaire_mensuel", "dt_debut_pc_prime_annuelle", "pc_prime_annuelle", "dt_debut_prime_annuelle",
    "prime_annuelle", "lb_job_title", "lb_account_holder", "lb_iban", "lb_bic", "lb_account_number",
    "lb_bank_name",
]
assert len(ADP_COLUMNS) == 72
# R-SQL-051 : CONVERT(DATE, x) ; dt_entree_societe via SUBSTRING(x,1,10)
ADP_DATE_COLUMNS = [
    "dt_entree_societe", "dt_entree_groupe", "dt_anciennete", "dt_sortie", "dt_company_start",
    "dt_ppc_debut_contrat", "dt_ppc_fin_contrat", "dt_ppc_debut_periode_essaie", "dt_ppc_fin_periode_essaie",
    "dt_situation_start", "dt_debut_mt_salaire_base", "dt_debut_pc_prime_annuelle", "dt_debut_prime_annuelle",
]
ADP_DECIMAL_COLUMNS = {"mt_salaire_mensuel": "DECIMAL(23,13)"}
# D-CMP-02 (À CONFIRMER) : 71 colonnes métier comparées = R-SQL-056 sans dt_transforme
COMPARE_COLUMNS = [c for c in ADP_COLUMNS if c != "dt_transforme"]
BUSINESS_KEY = "id_unique"  # D-CMP-01 (À CONFIRMER)


def adp_column_type(col: str) -> str:
    if col == "dt_transforme" or col in ADP_DATE_COLUMNS:
        return "DATE"
    return ADP_DECIMAL_COLUMNS.get(col, "STRING")


# ---------------------------------------------------------------------------
# DDL : nom logique -> liste (colonne, type) ; les tables ctl.* et log.* reprennent 05 (fait foi)
# ---------------------------------------------------------------------------
def _cols(spec: str) -> List[Tuple[str, str]]:
    out = []
    for line in spec.strip().splitlines():
        line = line.strip().rstrip(",")
        if line:
            name, typ = line.split(None, 1)
            out.append((name, typ.strip()))
    return out


TABLES: Dict[str, List[Tuple[str, str]]] = {}

TABLES["cfg.process"] = _cols("""
process_code STRING
enabled BOOLEAN
acquisition_mode STRING
publication_target STRING
key_columns ARRAY<STRING>
compare_columns ARRAY<STRING>
reference_table STRING
staging_table STRING
max_absent_pct DECIMAL(9,4)
max_modified_pct DECIMAL(9,4)
max_reject_pct DECIMAL(9,4)
updated_at_utc TIMESTAMP
""")
for _t, _c in LOOKUP_TABLES.items():
    TABLES[f"cfg.{_t}"] = [(c, "STRING") for c in _c] + [
        ("_loaded_at_utc", "TIMESTAMP"), ("_loaded_by", "STRING"), ("_change_ref", "STRING")]

TABLES["stg.ts_employee_raw"] = (
    [("run_id", "STRING"), ("business_date", "DATE"), ("_extracted_at_utc", "TIMESTAMP")]
    + [(k, "STRING") for k in RAW_KEYS])
TABLES["stg.ts_employee_source"] = (
    [("run_id", "STRING"), ("business_date", "DATE")] + [(c, "STRING") for c in SOURCE_COLUMN_NAMES])
TABLES["stg.ts_employee_enriched"] = (
    [("run_id", "STRING"), ("business_date", "DATE"), ("dt_staging", "DATE")]
    + [(c, "STRING") for c in SOURCE_COLUMN_NAMES] + [(c, "STRING") for c in ENRICHED_LOOKUP_COLUMNS])
TABLES["stg.ts_employee_reject"] = _cols("""
run_id STRING
business_date DATE
dt_rejet DATE
id_unique STRING
id_payroll STRING
cd_type_rejet STRING
lb_nom_colonne_rejet STRING
lb_valeur_rejet STRING
fl_rejet INT
lb_fichier_mapping STRING
ds_rejet STRING
""")
TABLES["stg.ts_employee_adp"] = (
    [("run_id", "STRING"), ("business_date", "DATE")]
    + [(c, adp_column_type(c)) for c in ADP_COLUMNS] + [("row_hash", "STRING")])
TABLES["ref.ts_employee_adp"] = (
    [(c, adp_column_type(c)) for c in ADP_COLUMNS]
    + [("row_hash", "STRING"), ("_last_run_id", "STRING"), ("_validated_at_utc", "TIMESTAMP")])

TABLES["pub.adp_outbox"] = _cols("""
run_id STRING
employee_key STRING
change_category STRING
row_hash STRING
status STRING
attempt_count INT
last_attempt_at_utc TIMESTAMP
last_outcome_code STRING
gate_decision_id STRING
created_at_utc TIMESTAMP
updated_at_utc TIMESTAMP
""")

# --- PROPOSÉ (contrat ADP final absent) : intention durable et tentatives par opération ---
TABLES["pub.adp_operation"] = _cols("""
op_id STRING
run_id STRING
employee_key STRING
seq_no INT
op_code STRING
op_variant STRING
depends_on ARRAY<STRING>
target_hash STRING
payload_hash STRING
status STRING
attempt_count INT
last_outcome_class STRING
last_http_status INT
gate_decision_id STRING
manual_decision_id STRING
created_at_utc TIMESTAMP
updated_at_utc TIMESTAMP
""")
TABLES["pub.adp_attempt"] = _cols("""
attempt_id STRING
op_id STRING
run_id STRING
attempt_no INT
phase STRING
outcome_class STRING
http_status INT
error_class STRING
response_digest STRING
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
""")
TABLES["pub.adp_reconciliation"] = _cols("""
reconciliation_id STRING
op_id STRING
run_id STRING
observed_at_utc TIMESTAMP
method STRING
fields_checked ARRAY<STRING>
result STRING
conclusion STRING
""")
TABLES["log.adp_manual_decision"] = _cols("""
decision_id STRING
op_id STRING
run_id STRING
employee_key STRING
decision STRING
decided_by STRING
decided_at_utc TIMESTAMP
basis STRING
reference STRING
""")

TABLES["ctl.process_lock"] = _cols("""
process_code STRING
lock_status STRING
holder_run_id STRING
acquired_at_utc TIMESTAMP
lease_expires_at_utc TIMESTAMP
released_at_utc TIMESTAMP
""")
TABLES["ctl.run"] = _cols("""
run_id STRING
process_code STRING
run_mode STRING
business_date DATE
run_seq_in_day INT
trigger_time_utc TIMESTAMP
pipeline_name STRING
workspace_id STRING
workspace_name STRING
status STRING
compared_against_run_id STRING
promoted_at_utc TIMESTAMP
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
error_code STRING
error_summary STRING
target_run_id STRING
""")
TABLES["ctl.run_step"] = _cols("""
run_id STRING
step_code STRING
attempt_no INT
status STRING
rows_in BIGINT
rows_out BIGINT
content_hash STRING
message STRING
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
""")
TABLES["log.comparison_summary"] = _cols("""
run_id STRING
process_code STRING
business_date DATE
compared_against_run_id STRING
nb_staging BIGINT
nb_reference BIGINT
nb_new BIGINT
nb_modified BIGINT
nb_unchanged BIGINT
nb_absent BIGINT
nb_rejected BIGINT
pct_absent DECIMAL(9,4)
pct_modified DECIMAL(9,4)
pct_rejected DECIMAL(9,4)
is_first_load BOOLEAN
created_at_utc TIMESTAMP
""")
TABLES["log.comparison_detail"] = _cols("""
run_id STRING
process_code STRING
business_key STRING
category STRING
changed_columns ARRAY<STRING>
old_row_hash STRING
new_row_hash STRING
created_at_utc TIMESTAMP
""")
TABLES["log.anomaly"] = _cols("""
run_id STRING
process_code STRING
control_code STRING
severity STRING
metric_value DECIMAL(18,4)
threshold_value DECIMAL(18,4)
outcome STRING
message STRING
created_at_utc TIMESTAMP
""")
TABLES["log.reject"] = _cols("""
run_id STRING
process_code STRING
dt_rejet DATE
id_unique STRING
id_payroll STRING
cd_type_rejet STRING
lb_nom_colonne_rejet STRING
lb_valeur_rejet STRING
lb_fichier_mapping STRING
ds_rejet STRING
created_at_utc TIMESTAMP
""")
TABLES["log.run_summary"] = _cols("""
run_id STRING
process_code STRING
business_date DATE
dt_resume DATE
cd_source STRING
cd_target STRING
mt_lignes_inserts BIGINT
mt_lignes_rejets BIGINT
mt_lignes_lues BIGINT
pc_lignes_rejets DECIMAL(12,4)
ds_type_rejets STRING
nb_outbox BIGINT
nb_sent BIGINT
nb_no_change BIGINT
nb_failed BIGINT
nb_not_pushed BIGINT
created_at_utc TIMESTAMP
""")
TABLES["log.adp_gate_decision"] = _cols("""
decision_id STRING
run_id STRING
evaluated_at_utc TIMESTAMP
workspace_name STRING
workspace_id STRING
is_for_pipeline BOOLEAN
notebook_name STRING
rule_version STRING
decision BOOLEAN
reason_code STRING
""")
TABLES["log.adp_publication_log"] = _cols("""
log_id STRING
source_system STRING
run_id STRING
log_date DATE
ts_adp TIMESTAMP
matricule_id STRING
nom STRING
mode STRING
field STRING
status_code STRING
reason_code STRING
response_code STRING
error_message STRING
log_message STRING
employee_key STRING
step_name STRING
http_method STRING
endpoint_template STRING
attempt_no INT
gate_decision_id STRING
_import_id STRING
_source_file_uri STRING
_source_row_number INT
_row_hash STRING
_ingested_at_utc TIMESTAMP
""")
TABLES["log.log_file_import"] = _cols("""
import_id STRING
source_uri STRING
file_name STRING
file_size BIGINT
file_modified_at_utc TIMESTAMP
file_format STRING
status STRING
detail STRING
rows_read INT
rows_loaded INT
rows_rejected INT
target_table STRING
executed_by STRING
workspace_name STRING
processed_at_utc TIMESTAMP
""")
TABLES["log.log_file_import_reject"] = _cols("""
import_id STRING
source_uri STRING
source_row_number INT
reject_reason STRING
rejected_at_utc TIMESTAMP
""")
TABLES["log.notification"] = _cols("""
run_id STRING
notification_type STRING
channel STRING
recipients_ref STRING
status STRING
sent_at_utc TIMESTAMP
error_summary STRING
""")
TABLES["log.cell_execution"] = _cols("""
event_id STRING
run_id STRING
operation_id STRING
notebook_name STRING
notebook_id STRING
cell_code STRING
cell_seq INT
status STRING
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
duration_ms BIGINT
workspace_name STRING
workspace_id STRING
lakehouse_name STRING
is_for_pipeline BOOLEAN
error_type STRING
error_code STRING
error_message STRING
error_location STRING
logged_at_utc TIMESTAMP
""")

# --- Flux 2 de l'existant (rémunération, Pipeline_SIRH-SFTP) en tables, sans fichier (hris.remuneration) ---
REMUNERATION_SNAPSHOT_COLUMNS = ["section", "employeenumber", "date", "type", "amount", "currency", "periodicity",
                                 "start_date", "end_date", "calculation_method", "value", "base_salary_effective_date",
                                 "src_section_code", "src_category", "src_indexed_on"]
TABLES["cfg.ts_remuneration_mapping"] = _cols("""
section_code STRING
legacy_section STRING
type_value STRING
calculation_method STRING
effective_date_rule STRING
date_format STRING
currency_case STRING
_loaded_at_utc TIMESTAMP
_loaded_by STRING
_change_ref STRING
""")
TABLES["stg.ts_remuneration_snapshot"] = (
    [("run_id", "STRING"), ("business_date", "DATE")] + [(c, "STRING") for c in REMUNERATION_SNAPSHOT_COLUMNS])
TABLES["ref.ts_remuneration_snapshot"] = (
    [(c, "STRING") for c in REMUNERATION_SNAPSHOT_COLUMNS]
    + [("_source_run_id", "STRING"), ("_validated_at_utc", "TIMESTAMP")])
TABLES["pub.ts_remuneration_change"] = _cols("""
run_id STRING
business_date DATE
import_type STRING
line_no INT
rule STRING
employeenumber STRING
date STRING
type STRING
amount STRING
currency STRING
periodicity STRING
extra_augm_n_1 STRING
start_date STRING
end_date STRING
calculation_method STRING
value STRING
base_salary_effective_date STRING
""")
TABLES["log.ts_remuneration_import"] = _cols("""
run_id STRING
import_type STRING
file_name STRING
rows INT
payload_sha256 STRING
gate_outcome STRING
gate_reason STRING
status STRING
ts_import_id BIGINT
ts_import_status STRING
reference_advanced BOOLEAN
message STRING
created_at_utc TIMESTAMP
""")

SCHEMAS =["cfg", "stg", "ref", "pub", "ctl", "log"]

# Tables à données personnelles (accès restreint, 05 § 4 ; D-SEC-03)
PERSONAL_DATA_TABLES = [
    "stg.ts_employee_raw", "stg.ts_employee_source", "stg.ts_employee_enriched", "stg.ts_employee_reject",
    "stg.ts_employee_adp", "ref.ts_employee_adp", "log.comparison_detail", "log.reject",
    "log.adp_publication_log", "pub.adp_outbox", "pub.adp_operation",
    "stg.ts_remuneration_snapshot", "ref.ts_remuneration_snapshot", "pub.ts_remuneration_change",
]

PROPOSED_TABLES = ["pub.adp_operation", "pub.adp_attempt", "pub.adp_reconciliation", "log.adp_manual_decision"]


class Tables:
    """Résolution nom logique -> nom physique.

    * Sans `lakehouse` : noms partiels (`ctl.run`) qui exigent un Lakehouse par défaut attaché (tests locaux).
    * Avec `lakehouse` : noms à 4 parties `` `workspace`.`lakehouse`.schema.table ``, qui fonctionnent SANS
      Lakehouse par défaut (vérifié dans Fabric : reports/probe_4part_names.json ; les noms partiels échouent sans
      lien par défaut : reports/probe_no_default_lakehouse.json). Aucune dépendance à une association que la
      synchronisation Git pourrait ne pas conserver.
    * `test_schema` redirige toutes les tables vers un schéma d'isolement unique : `<schéma_test>.<schéma>__<table>`.
    """

    def __init__(self, test_schema: str = "", lakehouse: str = "", workspace: str = ""):
        if bool(lakehouse) != bool(workspace):
            raise ValueError("lakehouse et workspace vont ensemble")
        for part in (test_schema, lakehouse, workspace):
            if "`" in part or any(ch in part for ch in ("\n", "\r")):
                raise ValueError("nom invalide")
        self.test_schema, self.lakehouse, self.workspace = test_schema, lakehouse, workspace

    @property
    def prefix(self) -> str:
        return f"`{self.workspace}`.`{self.lakehouse}`." if self.lakehouse else ""

    def __call__(self, logical: str) -> str:
        if logical not in TABLES:
            raise KeyError(logical)
        schema, name = logical.split(".", 1)
        if not self.test_schema:
            return f"{self.prefix}{schema}.{name}"
        return f"{self.prefix}{self.test_schema}.{schema}__{name}"

    def extra(self, name: str) -> str:
        """Table annexe du schéma de test (ex. rapport de campagne)."""
        if not self.test_schema:
            raise ValueError("table annexe réservée au schéma de test")
        return f"{self.prefix}{self.test_schema}.{name}"

    def schemas(self) -> List[str]:
        names = [self.test_schema] if self.test_schema else list(SCHEMAS)
        return [f"{self.prefix}{n}" for n in names]


def tables_for_notebook(context, test_schema: str = "", lakehouse: str = "") -> "Tables":
    """Tables du notebook courant : workspace lu dans le contexte d'exécution RÉEL ; Lakehouse fourni par paramètre
    (VL_HRIS.lakehouse_name) ou, à défaut, Lakehouse par défaut attaché. Refuse de deviner."""
    def get(key):
        try:
            return context.get(key)
        except Exception:
            return None
    ws = get("currentWorkspaceName")
    lh = lakehouse or get("defaultLakehouseName")
    if not ws or not lh:
        raise ValueError("LAKEHOUSE_NOT_RESOLVED : fournir p_lakehouse_name (VL_HRIS.lakehouse_name)")
    return Tables(test_schema, lh, ws)


def create_table_sql(physical: str, logical: str) -> str:
    cols = ",\n  ".join(f"`{n}` {t}" for n, t in TABLES[logical])
    return f"CREATE TABLE IF NOT EXISTS {physical} (\n  {cols}\n) USING DELTA"


def normalize_type(t: str) -> str:
    """Normalise un type Spark (simpleString) pour comparer DDL attendu et schéma constaté."""
    t = t.strip().lower().replace(" ", "")
    return {"int": "int", "integer": "int", "bigint": "bigint", "long": "bigint", "string": "string",
            "boolean": "boolean", "date": "date", "timestamp": "timestamp"}.get(t, t)
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.ts_acquire =============
register_hris_module('hris.ts_acquire', 'd381880b82ee3b6a7d071d8f416315677188498410fe5a9652a813409531b2ba', r'''"""hris.ts_acquire — portage de la Function TalentSoft (ts-function-app/function_app.py, PRJ @ 5541490).

Parité (A01, R-TS-007 à R-TS-030) : endpoints, règle de « valeur courante » dépendant de l'ordre,
conversion « falsy », sentinelle 1900-01-01, Hors_France, Level sans %, filtre « Motul France »,
retries 429/503/réseau avec backoff, disjoncteur global.

Écarts volontaires avec l'existant (extraction plus rapide, mêmes données) : le filtre « Motul France » n'est plus
une phase 1 séquentielle mais s'évalue salarié par salarié dans la fenêtre parallèle, et la situation et le poste
lus pour le filtre sont réutilisés pour le détail (aucun appel en double) ; plus de lots de 75 ni de pause entre
lots (fenêtre glissante) ; concurrence configurable (par défaut 8 salariés et 10 requêtes simultanées, contre 2 et
5) ; jeton renouvelé selon son âge ; les trois entités de santé ne sont plus appelées (jamais persistées).

Améliorations techniques (A08 D-TS-02, À CONFIRMER, sans effet sur les données si l'API se comporte
comme attendu) : appels synchrones bornés (pas d'asyncio dans un notebook) ; renouvellement du jeton
sur 401 en LECTURE ; prise en compte de Retry-After ; aucune redirection suivie ; comptage de tous
les échecs et exclusions pour le contrôle de complétude (aucun échec n'est plus silencieux) ;
aucune URL, aucun matricule, aucun corps de réponse journalisé.

Source : le tenant TalentSoft de PRODUCTION, quel que soit l'environnement d'exécution
(consigne du mandat). Aucune autre URL n'est acceptée et il n'existe aucun repli.
"""
from __future__ import annotations

import datetime
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from typing import Any, Callable, Dict, List, Optional, Tuple
from urllib.parse import urlparse

from hris.common import TS_PRD_BASE_URL, HrisError, SecretValue
from hris.model import RAW_KEYS
from hris.parity import json_value_repr

EMPLOYEE_CONCURRENCY = 8  # salariés traités en même temps (ts:L254 : 2 par lot dans l'existant)
PER_HOST_LIMIT = 10       # requêtes HTTP simultanées vers TalentSoft (ts:L100 : 5 dans l'existant)
TOKEN_MAX_AGE_SECONDS = 300  # jeton renouvelé au-delà (l'existant le renouvelait à chaque lot de 75)
MOTUL_FRANCE = "Motul France"
# Entités de détail lues pour les salariés retenus. « Situation » et la position n'en font pas partie :
# elles sont lues une seule fois, pour le filtre « Motul France », puis réutilisées pour le détail.
HUB_ENTITIES = {
    "keydate": "KeyDates", "legacyid": "EmployeeIDs",
    "qualification": "Qualification", "mainpostaladdress": "MainPostalAddress",
    "contactinfo": "ContactInformation", "maritalstatus": "FamilyStatus", "contract": "Contract",
    "bankdetails": "BankDetails", "probation": "ProbationaryPeriod", "salary": "Compensation",
    "civilregistration": "CivilRegistration", "identity": "Identity",
}
# Les entités de santé (Disability, WorkAccident, MedicalCheckUp) ne sont plus appelées : l'existant les lisait
# par parité mais leur résultat n'était jamais persisté (A08 D-SEC-02, seules les clés RAW_KEYS sont conservées).
# Les fonctions process_* correspondantes restent disponibles (voir PROCESSORS).


def assert_prd_source(base_url: str) -> str:
    """Refuse toute source autre que le tenant TalentSoft de production."""
    normalized = (base_url or "").strip().rstrip("/")
    if normalized != TS_PRD_BASE_URL:
        raise HrisError("TS_SOURCE_NOT_PRD", "seul le tenant TalentSoft de production est autorisé")
    parsed = urlparse(normalized)
    if parsed.scheme != "https":
        raise HrisError("TS_SOURCE_NOT_HTTPS")
    return normalized


# ---------------------------------------------------------------------------
# Disjoncteur (ts:L32-65) — parité
# ---------------------------------------------------------------------------
class CircuitBreaker:
    def __init__(self, failure_threshold: int = 5, timeout: float = 60, clock: Callable[[], float] = time.time):
        self.failure_count = 0
        self.failure_threshold = failure_threshold
        self.timeout = timeout
        self.last_failure_time: Optional[float] = None
        self.state = "CLOSED"
        self._clock = clock
        self._lock = threading.Lock()

    def record_failure(self) -> None:
        with self._lock:
            self.failure_count += 1
            self.last_failure_time = self._clock()
            if self.failure_count >= self.failure_threshold:
                self.state = "OPEN"

    def record_success(self) -> None:
        with self._lock:
            self.failure_count = 0
            self.state = "CLOSED"

    def can_attempt(self) -> bool:
        with self._lock:
            if self.state == "CLOSED":
                return True
            if self.state == "OPEN":
                if self._clock() - (self.last_failure_time or 0) > self.timeout:
                    self.state = "HALF_OPEN"
                    return True
                return False
            return True


class AcquisitionStats:
    """Compteurs sans donnée personnelle (aucun matricule, aucune URL)."""

    def __init__(self):
        self._lock = threading.Lock()
        self.counters: Counter = Counter()
        self.endpoint_status: Dict[str, Counter] = defaultdict(Counter)
        self.total_count: Optional[int] = None

    def inc(self, key: str, n: int = 1) -> None:
        with self._lock:
            self.counters[key] += n

    def status(self, endpoint: str, status: Any) -> None:
        with self._lock:
            self.endpoint_status[endpoint][str(status)] += 1

    def as_dict(self) -> dict:
        return {"total_count": self.total_count, "counters": dict(self.counters),
                "endpoint_status": {k: dict(v) for k, v in sorted(self.endpoint_status.items())}}


# ---------------------------------------------------------------------------
# Client HTTP (synchrone, borné) — équivalent de make_request_with_retry (ts:L126-183)
# ---------------------------------------------------------------------------
class NetworkError(Exception):
    """Erreur réseau au sens de l'existant (aiohttp.ClientError, TimeoutError, OSError)."""


class TsClient:
    def __init__(self, base_url: str, client_id: SecretValue, client_secret: SecretValue,
                 session_factory: Callable[[], Any], stats: AcquisitionStats,
                 sleep: Callable[[float], None] = time.sleep,
                 timeout: Tuple[float, float] = (30.0, 60.0),
                 network_errors: Tuple[type, ...] = (OSError,),
                 breaker: Optional[CircuitBreaker] = None,
                 max_retry_after: float = 60.0,
                 max_concurrency: int = PER_HOST_LIMIT,
                 clock: Callable[[], float] = time.monotonic):
        self.base_url = assert_prd_source(base_url)
        self._client_id = client_id
        self._client_secret = client_secret
        self._session = session_factory()
        self.stats = stats
        self._sleep = sleep
        self._timeout = timeout
        self._network_errors = network_errors
        self.breaker = breaker or CircuitBreaker()
        self._token: Optional[SecretValue] = None
        self._token_lock = threading.Lock()
        self._refresh_lock = threading.Lock()
        self._clock = clock
        self._token_obtained_at = 0.0
        self._max_retry_after = max_retry_after
        self.max_concurrency = max(1, int(max_concurrency))
        self._limiter = threading.BoundedSemaphore(self.max_concurrency)

    def __repr__(self) -> str:
        return "TsClient(<prd>)"

    def close(self) -> None:
        try:
            self._session.close()
        except Exception:
            pass

    # -- jeton (ts:L331-364) ----------------------------------------------------
    def get_token(self) -> None:
        data = {"client_id": self._client_id.reveal(), "client_secret": self._client_secret.reveal(),
                "grant_type": "client_credentials"}
        token_data, status = self._request("POST", "/api/token", {"Accept": "application/json"},
                                           max_retries=3, initial_delay=2, form=data, endpoint="token")
        if status == 200 and token_data and token_data.get("access_token"):
            with self._token_lock:
                self._token = SecretValue("Bearer " + token_data.get("access_token"))
                self._token_obtained_at = self._clock()
            self.stats.inc("token_obtained")
            return
        raise HrisError("TS_AUTH_FAILED", f"statut {status}")

    def _token_is_stale(self) -> bool:
        with self._token_lock:
            return self._token is not None and self._clock() - self._token_obtained_at > TOKEN_MAX_AGE_SECONDS

    def _auth_headers(self) -> dict:
        # Renouvellement périodique (remplace le renouvellement par lot de l'existant) : un seul fil renouvelle,
        # les autres attendent puis réutilisent le nouveau jeton. Appelé hors sémaphore : aucun blocage croisé.
        if self._token_is_stale():
            with self._refresh_lock:
                if self._token_is_stale():
                    self.get_token()
        with self._token_lock:
            token = self._token.reveal() if self._token else None
        if not token:
            raise HrisError("TS_NO_TOKEN")
        return {"Authorization": token, "Accept": "application/json"}

    def get(self, path: str, endpoint: str, max_retries: int = 3, initial_delay: float = 1) -> Tuple[Any, Any]:
        data, status = self._request("GET", path, None, max_retries, initial_delay, endpoint=endpoint)
        if status == 401:  # amélioration D-TS-02 : jeton renouvelé, lecture idempotente rejouée une fois
            self.stats.inc("token_refresh_on_401")
            self.get_token()
            data, status = self._request("GET", path, None, max_retries, initial_delay, endpoint=endpoint)
        return data, status

    def _request(self, method: str, path: str, headers: Optional[dict], max_retries: int,
                 initial_delay: float, form: Optional[dict] = None, endpoint: str = "") -> Tuple[Any, Any]:
        if not self.breaker.can_attempt():
            self.stats.inc("circuit_open_skip")
            self.stats.status(endpoint, "CIRCUIT_OPEN")
            return None, None
        url = self.base_url + path
        for attempt in range(max_retries):
            try:
                hdrs = headers if headers is not None else self._auth_headers()
                with self._limiter:
                    if method == "GET":
                        response = self._session.get(url, headers=hdrs, timeout=self._timeout,
                                                     allow_redirects=False)
                    else:
                        response = self._session.post(url, headers=hdrs, data=form, timeout=self._timeout,
                                                      allow_redirects=False)
                status = response.status_code
                if status == 200:
                    try:
                        result = response.json()
                    except Exception:
                        self.stats.status(endpoint, "200_INVALID_JSON")
                        return None, None  # parité : exception inattendue -> (None, None)
                    self.breaker.record_success()
                    self.stats.status(endpoint, 200)
                    return result, status
                if status in (429, 503):
                    self.stats.inc(f"http_{status}")
                    if attempt < max_retries - 1:
                        delay = initial_delay * (2 ** attempt)
                        retry_after = _retry_after_seconds(response)
                        if retry_after is not None:
                            delay = min(max(delay, retry_after), self._max_retry_after)
                        self._sleep(delay)
                        continue
                    # parité : dernier essai -> statut perdu, (None, None) après la boucle
                    self.stats.status(endpoint, f"{status}_EXHAUSTED")
                else:
                    self.stats.status(endpoint, status)
                    return None, status
            except self._network_errors as exc:
                self.breaker.record_failure()
                self.stats.inc("network_error")
                if attempt < max_retries - 1:
                    self._sleep(initial_delay * (2 ** attempt))
                else:
                    self.stats.status(endpoint, f"NETWORK_{type(exc).__name__}")
                    return None, None
            except HrisError:
                raise
            except Exception as exc:  # parité : erreur inattendue -> (None, None), sans retry
                self.stats.status(endpoint, f"ERROR_{type(exc).__name__}")
                return None, None
        return None, None


def _retry_after_seconds(response: Any) -> Optional[float]:
    try:
        value = response.headers.get("Retry-After")
    except Exception:
        return None
    if value is None:
        return None
    try:
        return max(0.0, float(value))
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Traitements des réponses — copie fidèle de l'existant (ts:L619-1629)
# ---------------------------------------------------------------------------
def _d(value: str) -> datetime.date:
    return datetime.datetime.strptime(value, "%Y-%m-%d").date()


def initialize_employee_fields(employee_number: str) -> dict:
    return {
        "employeeNumber": employee_number, "hubSocialSecurityNumber": "", "referenceAdmissionDate": "",
        "WorkingType": "", "Situation_starDate": "",
        "GeographicOrganisationStructure": "", "LegalStructure": "", "CostCenter": "", "CollectiveAgreement": "",
        "PPC_FTE": "", "OccupationalCategory": "", "Coefficient": "", "OrganisationalStructure": "",
        "salary_base": "0", "base_startDate": "", "Pourcentage_prime_annuelles": "0",
        "pourcentage_prime_annuelles_startDate": "", "prime_annuelles": "0", "prime_annuelles_startDate": "",
        "address_postalCode": "", "adresse_city": "", "adresse_countryCode": "", "Hors_France": "0",
        "GroupStartDate": "", "GroupEndDate": "", "CompanyStartDate": "", "RecalculatedSeniorityDate": "",
        "handicap": "", "handicap_startDate": "", "workaccident": "", "workaccident_startDate": "",
        "motif_visite_medical": "", "motif_visite_medical_startDate": "",
        "LegacyID": "", "LegacyID_startDate": "", "Matricule_IT": "",
        "Qualification": "", "Qualification_startDate": "", "ConferredDate": "",
        "LegalTownOrCity": "", "LegalCountry": "", "LegalPostalCode": "", "LegalStreet": "", "LegalStreetNumber": "",
        "LegalAdditionalAddressInformation": "", "LegalStreetNumberComplement": "",
        "BusinessMobile": "", "PersonalMobile": "", "BusinessPhone": "", "PersonalPhone": "", "BusinessEmail": "",
        "PersonalEmail": "", "ContactInformation_startDate": "",
        "MaritalStatus": "", "MaritalStatus_startDate": "",
        "NatureOfContract": "", "HiringReason": "", "date_debut_contract": "", "date_fin_contract": "",
        "FixedTermContractReason": "", "TerminationReason": "", "Status": "",
        "AccountHolder": "", "IBANOrABA": "", "BicOrSwift": "", "AccountNumber": "", "BankName": "",
        "PaymentMethod": "", "UseSameAccountForExpense": "", "ExpenseAccountPaymentMethod": "",
        "ExpenseAccountHolder": "", "ExpenseAccountIBANOrABA": "", "ExpenseAccountBicOrSwift": "",
        "ExpenseAccountBankName": "", "Country": "", "ExpenseAccountCountry": "", "date_debut_bank": "",
        "date_fin_bank": "",
        "FirstPeriodStartDate": "", "FirstPeriodEndDate": "", "SecondPeriodStartDate": "", "SecondPeriodEndDate": "",
        "ProbationaryPeriod_startDate": "",
        "Date_de_naissance": "", "Ville_de_naissance": "", "Pays_de_naissance": "", "code_postal_lieu_de_naissance": "",
        "region_ou_etat_lieu_de_naissance": "", "CivilRegistration_startDate": "",
        "Title": "", "Gender": "", "LastName": "", "FirstName": "", "Nationality": "", "BirthName": "",
        "Identity_startDate": "",
        "numero_ordre": "1", "pourcentage_imputation": "100",
    }


def process_situation_data(data):
    if not data:
        return {"WorkingType": "", "Situation_starDate": "", "Position": ""}
    latest_date = "0001-01-01"
    result = {"WorkingType": "", "Situation_starDate": "", "Position": ""}
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            result["Situation_starDate"] = latest_date
            if item["field"] == "WorkingType":
                result["WorkingType"] = item["value"] if item["value"] else ""
            elif item["field"] == "Position":
                result["Position"] = item["value"] if item["value"] else ""
    return result


def process_position_data(data):
    if not data:
        return {}
    result = {"GeographicOrganisationStructure": "", "LegalStructure": "", "CostCenter": "",
              "CollectiveAgreement": "", "PPC_FTE": "", "OccupationalCategory": "", "Coefficient": "",
              "pourcentage_imputation": "100", "OrganisationalStructure": ""}
    for item in data:
        field = item.get("field", "")
        value = item.get("value", "")
        if field == "GeographicOrganisationStructure":
            result["GeographicOrganisationStructure"] = value if value else ""
        elif field == "LegalStructure":
            result["LegalStructure"] = value if value else ""
        elif field == "CostCenter":
            result["CostCenter"] = value if value else ""
        elif field == "CollectiveAgreement":
            result["CollectiveAgreement"] = value if value else ""
        elif field == "FullTimeEquivalentPercent":
            result["PPC_FTE"] = value if value else ""
        elif field == "OccupationalCategory":
            result["OccupationalCategory"] = value if value else ""
        elif field == "Coefficient":
            result["Coefficient"] = value if value else ""
        elif field == "OrganisationalStructure":
            result["OrganisationalStructure"] = value if value else ""
        elif field == "Level":
            if value:
                level_value = str(value).replace('%', '').strip()
                result["pourcentage_imputation"] = level_value if level_value else "100"
    return result


def process_salary_data(data):
    if not data:
        return {"salary_base": "0", "base_startDate": "", "Pourcentage_prime_annuelles": "0",
                "pourcentage_prime_annuelles_startDate": "", "prime_annuelles": "0", "prime_annuelles_startDate": ""}
    base_date = "0001-01-01"
    bonus_date = "0001-01-01"
    bonus_percentage_date = "0001-01-01"
    result = {"salary_base": "0", "base_startDate": "", "Pourcentage_prime_annuelles": "0",
              "pourcentage_prime_annuelles_startDate": "", "prime_annuelles": "0", "prime_annuelles_startDate": ""}
    for item in data:
        start_date = item["startDate"]
        if "ANNUAL_BASE_SALARY" in item["key"] and "Amount" in item["field"]:
            if _d(start_date) > _d(base_date):
                base_date = start_date
                result["salary_base"] = item["value"]
                result["base_startDate"] = start_date
        elif "CONTRACTUAL_BONUS_PERCENTAGE" in item["key"] and "Amount" in item["field"]:
            if _d(start_date) > _d(bonus_percentage_date):
                bonus_percentage_date = start_date
                result["Pourcentage_prime_annuelles"] = item["value"]
                result["pourcentage_prime_annuelles_startDate"] = start_date
        elif "CONTRACTUAL BONUS" in item["key"] and "Amount" in item["field"]:
            if _d(start_date) > _d(bonus_date):
                bonus_date = start_date
                result["prime_annuelles"] = item["value"]
                result["prime_annuelles_startDate"] = start_date
    return result


def process_postal_address_data(data):
    if not data:
        return {"address_postalCode": "", "adresse_city": "", "adresse_countryCode": "", "Hors_France": "0"}
    result = {"address_postalCode": data.get("postalCode", ""), "adresse_city": data.get("city", ""),
              "adresse_countryCode": data.get("countryCode", ""), "Hors_France": "0"}
    if result["adresse_countryCode"] and result["adresse_countryCode"].upper() != "FR":
        result["Hors_France"] = "1"
    return result


def _latest(data, fields: Dict[str, str], date_key: Optional[str] = None, defaults: Optional[dict] = None,
            date_on_field: Optional[Dict[str, str]] = None):
    """Règle générique « valeur courante » (R-TS-018) : un seul maximum startDate (>=) commun
    à tous les champs, items parcourus dans l'ordre de l'API."""
    result = dict(defaults or {})
    for out in fields.values():
        result.setdefault(out, "")
    if date_key:
        result.setdefault(date_key, "")
    latest_date = "0001-01-01"
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            if date_key:
                result[date_key] = latest_date
            out = fields.get(item["field"])
            if out is not None:
                result[out] = item["value"] if item["value"] else ""
                if date_on_field and item["field"] in date_on_field:
                    result[date_on_field[item["field"]]] = latest_date
    return result


def process_keydate_data(data):
    return _latest(data or [], {"GroupStartDate": "GroupStartDate", "GroupEndDate": "GroupEndDate",
                                "CompanyStartDate": "CompanyStartDate",
                                "RecalculatedSeniorityDate": "RecalculatedSeniorityDate"})


def process_disability_data(data):
    return _latest(data or [], {"RecognizedAsDisabled": "handicap"}, date_key="handicap_startDate")


def process_workaccident_data(data):
    return _latest(data or [], {"DateOfOccurrence": "workaccident"}, date_key="workaccident_startDate")


def process_medicalcheckup_data(data):
    return _latest(data or [], {"Reason": "motif_visite_medical"}, date_key="motif_visite_medical_startDate")


def process_legacyid_data(data):
    return _latest(data or [], {"LegacyID": "LegacyID", "Initials": "Matricule_IT"},
                   defaults={"LegacyID_startDate": ""}, date_on_field={"LegacyID": "LegacyID_startDate"})


def process_qualification_data(data):
    return _latest(data or [], {"Qualification": "Qualification", "ConferredDate": "ConferredDate"},
                   defaults={"Qualification_startDate": ""}, date_on_field={"Qualification": "Qualification_startDate"})


def process_mainpostaladdress_data(data):
    return _latest(data or [], {"TownOrCity": "LegalTownOrCity", "Country": "LegalCountry",
                                "PostalCode": "LegalPostalCode", "Street": "LegalStreet",
                                "StreetNumber": "LegalStreetNumber",
                                "AdditionalAddressInformation": "LegalAdditionalAddressInformation",
                                "StreetNumberComplement": "LegalStreetNumberComplement"})


def process_contactinfo_data(data):
    return _latest(data or [], {"BusinessMobile": "BusinessMobile", "PersonalMobile": "PersonalMobile",
                                "BusinessPhone": "BusinessPhone", "PersonalPhone": "PersonalPhone",
                                "BusinessEmail": "BusinessEmail", "PersonalEmail": "PersonalEmail"},
                   date_key="ContactInformation_startDate")


def process_maritalstatus_data(data):
    return _latest(data or [], {"MaritalStatus": "MaritalStatus"}, defaults={"MaritalStatus_startDate": ""},
                   date_on_field={"MaritalStatus": "MaritalStatus_startDate"})


def process_contract_data(data):
    if not data:
        return {"NatureOfContract": "", "HiringReason": "", "date_debut_contract": "", "date_fin_contract": "",
                "FixedTermContractReason": "", "TerminationReason": "", "Status": ""}
    latest_date = "0001-01-01"
    result = {"NatureOfContract": "", "HiringReason": "", "date_debut_contract": "", "date_fin_contract": "",
              "FixedTermContractReason": "", "TerminationReason": "", "Status": ""}
    fields = {"NatureOfContract", "HiringReason", "FixedTermContractReason", "TerminationReason", "Status"}
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            result["date_debut_contract"] = latest_date
            result["date_fin_contract"] = (item["endDate"] if item["endDate"] and str(item["endDate"]) != "None"
                                           else "1900-01-01")
            if item["field"] in fields:
                result[item["field"]] = item["value"] if item["value"] else ""
    return result


def bank_detail_data(data):
    keys = ["AccountHolder", "IBANOrABA", "BicOrSwift", "AccountNumber", "BankName", "PaymentMethod",
            "UseSameAccountForExpense", "ExpenseAccountPaymentMethod", "ExpenseAccountHolder",
            "ExpenseAccountIBANOrABA", "ExpenseAccountBicOrSwift", "ExpenseAccountBankName", "Country",
            "ExpenseAccountCountry"]
    result = {k: "" for k in keys}
    result.update({"date_debut_bank": "", "date_fin_bank": ""})
    if not data:
        return result
    latest_date = "0001-01-01"
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            result["date_debut_bank"] = item["startDate"]
            if item.get("endDate"):
                result["date_fin_bank"] = item["endDate"]
            if item["field"] in keys:
                result[item["field"]] = item["value"] if item["value"] else ""
    return result


def process_probation_data(data):
    return _latest(data or [], {"FirstPeriodStartDate": "FirstPeriodStartDate",
                                "FirstPeriodEndDate": "FirstPeriodEndDate",
                                "SecondPeriodStartDate": "SecondPeriodStartDate",
                                "SecondPeriodEndDate": "SecondPeriodEndDate"},
                   date_key="ProbationaryPeriod_startDate")


def process_civilregistration_data(data):
    return _latest(data or [], {"DateOfBirth": "Date_de_naissance", "TownOrCity": "Ville_de_naissance",
                                "Country": "Pays_de_naissance", "PostalCode": "code_postal_lieu_de_naissance",
                                "AreaOrState": "region_ou_etat_lieu_de_naissance"},
                   date_key="CivilRegistration_startDate")


def process_identity_data(data):
    return _latest(data or [], {"Title": "Title", "Gender": "Gender", "LastName": "LastName",
                                "FirstName": "FirstName", "Nationality": "Nationality", "BirthName": "BirthName"},
                   date_key="Identity_startDate")


PROCESSORS = {
    "keydate": process_keydate_data, "disability": process_disability_data,
    "workaccident": process_workaccident_data, "medicalcheckup": process_medicalcheckup_data,
    "legacyid": process_legacyid_data, "qualification": process_qualification_data,
    "mainpostaladdress": process_mainpostaladdress_data, "contactinfo": process_contactinfo_data,
    "maritalstatus": process_maritalstatus_data, "contract": process_contract_data,
    "bankdetails": bank_detail_data, "probation": process_probation_data, "salary": process_salary_data,
    "civilregistration": process_civilregistration_data, "identity": process_identity_data,
}


# ---------------------------------------------------------------------------
# Orchestration de l'extraction (ts:L186-616)
# ---------------------------------------------------------------------------
class TalentSoftExtractor:
    def __init__(self, client: TsClient, employee_concurrency: int = EMPLOYEE_CONCURRENCY):
        self.client = client
        self.stats = client.stats
        self.employee_concurrency = max(1, int(employee_concurrency))
        self.request_concurrency = client.max_concurrency   # une seule source de vérité : la limite du client

    # -- liste (ts:L367-405) -------------------------------------------------------
    def get_all_employees(self) -> List[dict]:
        data, status = self.client.get("/api/v1/directory/employees?count=1&offset=0", "directory_count")
        if status != 200 or not data:
            self.stats.inc("directory_count_failed")
            return []
        if not isinstance(data, dict) or "totalCount" not in data:
            self.stats.inc("directory_schema_unexpected")
        total_count = data.get("totalCount", 0) if isinstance(data, dict) else 0
        self.stats.total_count = total_count
        if not total_count or total_count <= 0:
            return []
        data, status = self.client.get(f"/api/v1/directory/employees?count={total_count}&offset=0", "directory_list")
        if status != 200 or not data:
            self.stats.inc("directory_list_failed")
            return []
        if not isinstance(data, dict) or "results" not in data:
            self.stats.inc("directory_schema_unexpected")
        employees = data.get("results", []) if isinstance(data, dict) else []
        self.stats.inc("directory_received", len(employees))
        return employees

    def _classify_unavailable(self, endpoint: str, status: Any) -> None:
        """Absence légitime (200 vide), 404 (à qualifier, D-TS-01) ou échec technique."""
        if status == 200:
            self.stats.inc(f"{endpoint}_empty")
        elif status == 404:
            self.stats.inc(f"{endpoint}_404")
        else:
            self.stats.inc(f"{endpoint}_unavailable")

    def get_situation(self, employee_number) -> dict:
        try:
            data, status = self.client.get(f"/api/hub/v2/employees/{employee_number}/Situation",
                                           "phase1_situation", max_retries=2, initial_delay=1)
            if status == 200 and data:
                return process_situation_data(data)
            self._classify_unavailable("phase1_situation", status)
            return {}
        except HrisError:
            raise
        except Exception:
            self.stats.inc("phase1_situation_error")
            return {}

    def get_position(self, position_code, date, endpoint: str) -> dict:
        try:
            if "T" not in date:
                date = f"{date}T00:00:00Z"
            data, status = self.client.get(f"/api/hub/v2.0/positions/{position_code}?date={date}",
                                           endpoint, max_retries=2, initial_delay=1)
            if status == 200 and data:
                return process_position_data(data)
            self._classify_unavailable(endpoint, status)
            return {}
        except HrisError:
            raise
        except Exception:
            self.stats.inc(f"{endpoint}_error")
            return {}

    # -- filtre « Motul France » (ts:L217-246), évalué salarié par salarié --------------
    def qualify(self, employee_number) -> Optional[Tuple[dict, dict]]:
        """Situation la plus récente -> code de poste -> poste à cette date -> LegalStructure.

        Renvoie (situation, poste) pour un salarié « Motul France », sinon None (exclusion comptée).
        Mêmes règles et mêmes compteurs que la phase 1 de l'existant, mais un salarié à la fois dans la
        fenêtre parallèle, et sans relire plus tard la situation et le poste des salariés retenus.
        """
        situation = self.get_situation(employee_number)
        if not situation:
            self.stats.inc("excluded_no_situation")  # échec technique OU absence : voir contrôles
            return None
        position_code = situation.get("Position")
        situation_date = situation.get("Situation_starDate")
        if not (position_code and situation_date):
            self.stats.inc("excluded_missing_position_or_date")
            return None
        position_data = self.get_position(position_code, situation_date, "phase1_position")
        if not position_data:
            self.stats.inc("excluded_no_position_data")
            return None
        if str(position_data.get("LegalStructure")) != MOTUL_FRANCE:
            self.stats.inc("excluded_not_motul_france")
            return None
        self.stats.inc("retained_phase1")
        return situation, position_data

    # -- détail d'un salarié retenu (ts:L440-616) ---------------------------------------
    def get_employee_data(self, employee: dict, pool: ThreadPoolExecutor) -> dict:
        employee_number = employee.get("employeeNumber")
        if not employee_number:
            self.stats.inc("excluded_no_employee_number")
            return {}
        qualified = self.qualify(employee_number)
        if qualified is None:
            return {}
        situation_data, position_data = qualified
        try:
            employee_data = {employee_number: initialize_employee_fields(str(employee_number))}
            employee_data[employee_number].update(situation_data)
            employee_data[employee_number].update(position_data)
            individual = employee.get("individual", {}) or {}
            nir = (individual.get("extras", {}) or {}).get("hubSocialSecurityNumber")
            if nir:
                employee_data[employee_number]["hubSocialSecurityNumber"] = str(nir)
            if employee.get("referenceAdmissionDate"):
                employee_data[employee_number]["referenceAdmissionDate"] = str(employee["referenceAdmissionDate"])
            username = (individual.get("user", {}) or {}).get("username", "")
            endpoints = {key: (f"/api/hub/v2/employees/{employee_number}/{entity}", f"hub_{key}")
                         for key, entity in HUB_ENTITIES.items()}
            if username:
                endpoints["postaladdress"] = (f"/api/v1.0/directory/individuals/{username}/postal-addresses",
                                              "postaladdress")
            futures = {key: pool.submit(self.client.get, path, ep) for key, (path, ep) in endpoints.items()}
            responses = {}
            for key, fut in futures.items():
                try:
                    responses[key] = fut.result()
                except HrisError:
                    raise
                except Exception as exc:
                    responses[key] = exc
            for key, result in responses.items():
                if isinstance(result, Exception):
                    self.stats.inc(f"phase2_exception_{key}")
                    continue
                if not isinstance(result, tuple) or len(result) != 2:
                    self.stats.inc(f"phase2_invalid_{key}")
                    continue
                data, status = result
                if status == 200:
                    if data:
                        try:
                            if key == "postaladdress":
                                if isinstance(data, list) and len(data) > 0:
                                    employee_data[employee_number].update(process_postal_address_data(data[0]))
                            else:
                                employee_data[employee_number].update(PROCESSORS[key](data))
                        except Exception:
                            self.stats.inc(f"processing_error_{key}")  # parité : on continue
                    else:
                        self.stats.inc(f"phase2_empty_{key}")
                else:
                    self.stats.inc(f"phase2_endpoint_not_ok_{key}")
            return employee_data
        except HrisError:
            raise
        except Exception:
            self.stats.inc("employee_error_defaults")  # parité : enregistrement par défaut (ts:L613-616)
            return {employee_number: initialize_employee_fields(str(employee_number))}

    def _extract_one(self, employee: dict, request_pool: ThreadPoolExecutor) -> dict:
        try:
            return self.get_employee_data(employee, request_pool)
        except HrisError:
            raise
        except Exception:
            self.stats.inc("employee_dropped")
            return {}

    def run(self) -> List[dict]:
        """Extraction complète de l'annuaire, filtrée sur « Motul France » ; renvoie les enregistrements bruts.

        Chaque salarié de l'annuaire est d'abord qualifié (situation puis poste : deux appels), puis, s'il est
        « Motul France », détaillé. Les salariés avancent en fenêtre glissante, sans lots ni pause : la file de
        requêtes ne se vide qu'à la toute fin. Une erreur fatale (HrisError) annule les salariés non commencés.
        """
        self.client.get_token()
        employees = self.get_all_employees()
        if not employees:
            self.stats.inc("no_employees")
            return []
        all_employee_data: Dict[Any, dict] = {}
        with ThreadPoolExecutor(max_workers=self.request_concurrency) as request_pool, \
                ThreadPoolExecutor(max_workers=self.employee_concurrency) as employee_pool:
            try:
                for result in employee_pool.map(lambda emp: self._extract_one(emp, request_pool), employees):
                    if isinstance(result, dict) and result:
                        all_employee_data.update(result)
            except BaseException:
                employee_pool.shutdown(wait=True, cancel_futures=True)
                raise
        retained = self.stats.counters["retained_phase1"]
        self.stats.inc("duplicates_merged", retained - len(all_employee_data))
        self.stats.inc("records", len(all_employee_data))
        return list(all_employee_data.values())


def to_raw_row(record: dict) -> Dict[str, Optional[str]]:
    """Projection sur les clés persistées, valeurs au format JSON_VALUE (R-SQL-002)."""
    return {key: json_value_repr(record.get(key)) for key in RAW_KEYS}


# ---------------------------------------------------------------------------
# Contrôles de complétude (D-QUA-01, D-QUA-04, D-TS-02 — À CONFIRMER)
# ---------------------------------------------------------------------------
def completeness_anomalies(stats: dict) -> List[dict]:
    """Traduit les compteurs en contrôles. BLOCKING : l'extraction ne peut pas être présentée
    comme complète. Aucun matricule ni aucune URL dans les messages."""
    c = stats.get("counters", {})
    total = stats.get("total_count")
    received = c.get("directory_received", 0)
    out = []

    def add(code, severity, metric, outcome, message):
        out.append({"control_code": code, "severity": severity, "metric_value": metric,
                    "threshold_value": None, "outcome": outcome, "message": message})

    if total is None or c.get("directory_count_failed") or c.get("directory_list_failed"):
        add("EXTRACTION_INCOMPLETE", "BLOCKING", None, "FAIL", "annuaire TalentSoft illisible (totalCount ou liste)")
    elif received != total:
        add("EXTRACTION_INCOMPLETE", "BLOCKING", received, "FAIL",
            f"reçus {received} pour totalCount {total}")
    else:
        add("EXTRACTION_INCOMPLETE", "BLOCKING", received, "PASS", f"reçus = totalCount = {total}")
    if c.get("directory_schema_unexpected"):
        add("SCHEMA_CHANGE", "BLOCKING", c["directory_schema_unexpected"], "FAIL", "réponse annuaire inattendue")
    records = c.get("records", 0)
    add("EMPTY_EXTRACTION", "BLOCKING", records, "FAIL" if records == 0 else "PASS",
        "aucun salarié retenu" if records == 0 else f"{records} salariés retenus")
    tech_phase1 = c.get("phase1_situation_unavailable", 0) + c.get("phase1_situation_error", 0) + \
        c.get("phase1_position_unavailable", 0) + c.get("phase1_position_error", 0)
    add("TS_FILTER_TECHNICAL_FAILURE", "BLOCKING", tech_phase1, "FAIL" if tech_phase1 else "PASS",
        "salariés non qualifiables (Situation/Position indisponible) : le filtre Motul France est incertain"
        if tech_phase1 else "filtre Motul France évalué pour tous les salariés")
    nf_phase1 = c.get("phase1_situation_404", 0) + c.get("phase1_position_404", 0)
    add("TS_FILTER_NOT_FOUND", "WARNING", nf_phase1, "FAIL" if nf_phase1 else "PASS",
        f"{nf_phase1} réponses 404 au filtrage (salariés exclus, statut à qualifier D-TS-01)")
    failures = 0
    not_found = 0
    for ep, statuses in stats.get("endpoint_status", {}).items():
        if not (ep.startswith("hub_") or ep == "postaladdress"):
            continue
        for st, n in statuses.items():
            if st == "200":
                continue
            if st == "404":
                not_found += n
            else:
                failures += n
    processing = sum(v for k, v in c.items() if k.startswith("processing_error_") or k.startswith("phase2_exception_")
                     or k.startswith("phase2_invalid_"))
    employee_level = c.get("employee_error_defaults", 0) + c.get("employee_dropped", 0)
    add("TS_ENDPOINT_FAILURE", "BLOCKING", failures, "FAIL" if failures else "PASS",
        f"{failures} appels de détail en échec après reprises" if failures else "tous les appels de détail aboutis")
    add("TS_ENDPOINT_NOT_FOUND", "WARNING", not_found, "FAIL" if not_found else "PASS", f"{not_found} réponses 404")
    add("TS_PROCESSING_ERROR", "BLOCKING", processing + employee_level, "FAIL" if (processing + employee_level) else "PASS",
        "erreurs de traitement de réponses (valeurs par défaut appliquées par parité)"
        if (processing + employee_level) else "aucune erreur de traitement")
    if c.get("circuit_open_skip"):
        add("TS_CIRCUIT_OPEN", "BLOCKING", c["circuit_open_skip"], "FAIL", "requêtes ignorées par le disjoncteur")
    return out
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.transform_ref =============
register_hris_module('hris.transform_ref', '2fe8bb63d4b2f4b141e4dac0b7f9f6af9ea5c2273caf7fafde1927caf9c1caae', r'''"""hris.transform_ref — implémentation de RÉFÉRENCE (Python pur) des procédures SQL de production.

Chaque fonction reproduit une procédure de A09 (PROC1 à PROC5) ligne à ligne, avec les sémantiques
T-SQL émulées par hris.parity. Elle sert :
  * aux tests unitaires par règle R-SQL (exécutés en local, sans Spark) ;
  * de référence pour le test de parité de l'implémentation Spark (hris.transform), exécuté
    sur données fictives en local et dans Fabric.
"""
from __future__ import annotations

import datetime as dt
import decimal
from typing import Dict, Iterable, List, Optional

from hris import parity as P
from hris.model import (ADP_COLUMNS, ADP_DATE_COLUMNS, COMPARE_COLUMNS, FORMAT_REJECT_COLUMNS,
                        FORMAT_SUM_EXCLUDED, LOOKUP_JOINS, NULL_REJECT_COLUMNS, REFERENCE_REJECTS,
                        REJECT_TYPE_FORMAT, REJECT_TYPE_NULL, REJECT_TYPE_REFERENCE, SOURCE_COLUMNS)

HASH_SEPARATOR = "\u001f"
HASH_NULL = "\u0000"


# --- PROC1 : lecture JSON (R-SQL-001 à 004) ----------------------------------------
def parse_source_row(raw: Dict[str, Optional[str]]) -> Dict[str, Optional[str]]:
    out = {}
    for key, col in SOURCE_COLUMNS:
        if key == "BusinessMobile|BusinessPhone":
            mobile, phone = raw.get("BusinessMobile"), raw.get("BusinessPhone")
            value = mobile if mobile is not None else phone  # COALESCE : NULL seulement, pas ''
        elif key == "date_fin_contract":
            value = raw.get(key)
            value = None if value is None else value.replace("None", "")
        else:
            value = raw.get(key)
        out[col] = P.nvarchar(value, 50)
    return out


# --- PROC2 : staging, 15 LEFT JOIN (R-SQL-010 à 024) ---------------------------------
def enrich_rows(source_rows: Iterable[dict], lookups: Dict[str, List[dict]]) -> List[dict]:
    """LEFT JOIN successifs : une correspondance dupliquée DUPLIQUE la ligne (parité R-SQL-058)."""
    rows = [dict(r) for r in source_rows]
    for join in LOOKUP_JOINS:
        table = lookups.get(join["table"], [])
        next_rows = []
        for row in rows:
            matches = [lu for lu in table if P.sql_eq(row.get(join["src"]), lu.get(join["key"]))]
            if not matches:
                r = dict(row)
                for _, out in join["out"]:
                    r[out] = None
                next_rows.append(r)
            for lu in matches:
                r = dict(row)
                for lu_col, out in join["out"]:
                    r[out] = lu.get(lu_col)
                next_rows.append(r)
        rows = next_rows
    return rows


# --- PROC3 : rejets (R-SQL-030 à 042) --------------------------------------------------
def _format_flag(value: Optional[str], kind: str) -> int:
    if kind == "len8":
        return 0 if P.sql_len(value) == 8 else 1
    if kind == "date":
        return 0 if (P.sql_is_empty(value) or P.isdate(value) == 1) else 1
    if kind == "date10":
        return 0 if (P.sql_is_empty(value) or P.isdate(None if value is None else value[:10]) == 1) else 1
    if kind == "pct":
        return P.pct_flag(value, 100)
    if kind == "amount":
        return P.pct_flag(value, 1000000)
    raise ValueError(kind)


def _format_ds(col: str, value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    if col[:2] == "dt":
        return f"Format Date non respecté au YYYY-MM-DD pour {col} = {value}"
    if col[:2] == "pc":
        return f"Format Pourcentage non respecté du 0 au 100 pour {col} = {value}"
    return None


def reject_rows(enriched: Iterable[dict], dt_rejet: dt.date) -> List[dict]:
    enriched = list(enriched)
    out: List[dict] = []
    # Null_Rejet
    for r in enriched:
        flags = {c: (1 if P.sql_is_empty(r.get(c)) else 0) for c in NULL_REJECT_COLUMNS}
        if not any(flags.values()):
            continue
        for c in NULL_REJECT_COLUMNS:
            if flags[c] == 1:
                out.append(_reject(dt_rejet, r, REJECT_TYPE_NULL, c, r.get(c), "", f"{c} ne peut pas être NULL"))
    # Format_Rejet
    for r in enriched:
        flags = {c: _format_flag(r.get(c), kind) for c, kind in FORMAT_REJECT_COLUMNS}
        total = sum(v for c, v in flags.items() if c not in FORMAT_SUM_EXCLUDED)
        if total <= 0:
            continue
        for c, _ in FORMAT_REJECT_COLUMNS:
            if flags[c] == 1:
                out.append(_reject(dt_rejet, r, REJECT_TYPE_FORMAT, c, r.get(c), "", _format_ds(c, r.get(c))))
    # Référence_Rejet
    for r in enriched:
        flags = []
        for fichier, nom, val_col, lu_col, adp_col in REFERENCE_REJECTS:
            value = r.get(val_col)
            missing = r.get(lu_col) is None or (adp_col is not None and r.get(adp_col) is None)
            flags.append(1 if (P.sql_is_not_empty(value) and missing) else 0)
        if sum(flags) <= 0:
            continue
        for flag, (fichier, nom, val_col, _, _) in zip(flags, REFERENCE_REJECTS):
            if flag == 1:
                value = r.get(val_col)
                out.append(_reject(dt_rejet, r, REJECT_TYPE_REFERENCE, nom, value, fichier,
                                   f"Mapping non trouvé entre TS et ADP pour {nom} = {value}"))
    return out


def _reject(dt_rejet, r, typ, col, value, fichier, ds):
    return {"dt_rejet": dt_rejet, "id_unique": r.get("id_unique"), "id_payroll": r.get("id_payroll"),
            "cd_type_rejet": typ, "lb_nom_colonne_rejet": col, "lb_valeur_rejet": value, "fl_rejet": 1,
            "lb_fichier_mapping": fichier, "ds_rejet": ds}


# --- PROC4 : transformation ADP (R-SQL-050 à 058) ------------------------------------
def adp_rows(enriched: Iterable[dict], rejects: Iterable[dict], dt_transforme: dt.date) -> List[dict]:
    rejected = {r["id_unique"] for r in rejects if r.get("id_unique") is not None}
    out = []
    for r in enriched:
        if r.get("id_unique") is not None and r["id_unique"] in rejected:
            continue  # NOT EXISTS sur id_unique (un id_unique NULL n'est jamais exclu)
        row: Dict[str, object] = {"dt_transforme": dt_transforme}
        for col in ADP_COLUMNS[1:]:
            if col in ADP_DATE_COLUMNS:
                value = r.get(col)
                if col == "dt_entree_societe" and value is not None:
                    value = value[:10]
                row[col] = P.convert_date(value)  # ValueError = échec du run (parité)
            elif col == "mt_salaire_mensuel":
                row[col] = P.monthly_salary(r.get("mt_salaire_base"))  # ValueError = échec (D-PAR-07)
            else:
                row[col] = r.get(col)
        out.append(row)
    return out


def hash_value(value) -> str:
    if value is None:
        return HASH_NULL
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, decimal.Decimal):
        return format(value, "f")
    return str(value)


def row_hash(row: dict, columns: List[str] = COMPARE_COLUMNS) -> str:
    """sha2(concat_ws('\\u001f', colonnes comparées, NULL -> '\\u0000'), 256) (A07 § 4)."""
    from hris.common import sha256_hex
    return sha256_hex(HASH_SEPARATOR.join(hash_value(row.get(c)) for c in columns))


# --- PROC5 : résumé (R-SQL-060 à 066) ------------------------------------------------
def summary(inserted_keys: Iterable[str], rejects: Iterable[dict], dt_resume: dt.date) -> dict:
    rejects = list(rejects)
    n_rej = len({r["id_unique"] for r in rejects if r.get("id_unique") is not None})
    n_ins = len({k for k in inserted_keys if k is not None})
    lues = n_ins + n_rej
    pc = decimal.Decimal(0) if lues == 0 else (decimal.Decimal(n_rej) / decimal.Decimal(lues)).quantize(
        decimal.Decimal("0.0001"), rounding=decimal.ROUND_HALF_UP)
    pairs = sorted({(r["cd_type_rejet"], r.get("ds_rejet")) for r in rejects},
                   key=lambda p: (p[0], p[1] or ""))  # STRING_AGG non ordonné : tri explicite (A02 § 8)
    ds_type = ",".join(p[0] for p in pairs) if pairs else None
    return {"dt_resume": dt_resume, "cd_source": "TalentSoft", "cd_target": "ADP", "mt_lignes_inserts": n_ins,
            "mt_lignes_rejets": n_rej, "mt_lignes_lues": lues, "pc_lignes_rejets": pc,
            "ds_type_rejets": P.utf8_truncate(ds_type, 50)}
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.spark_io =============
register_hris_module('hris.spark_io', 'd6aececfa4dfaec8f69b974eba9fe21320ec7eecd66848473edafe628fdfae0f', r'''"""hris.spark_io — écritures Delta idempotentes, journal des étapes, anomalies, socle.

Toutes les écritures passent par Spark (un seul moteur d'écriture par table, A06 F2).
Idempotence : `replaceWhere run_id = …` pour les tables à grain « run », MERGE sur clé naturelle
pour les autres. Aucune table préexistante n'est supprimée, vidée ni remplacée : la création
est `CREATE TABLE IF NOT EXISTS` et un schéma incompatible est signalé, jamais corrigé d'office.
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Any, Dict, Iterable, List, Optional

from hris.common import HrisError, PROCESS_CODE, safe_error, sanitize, to_utc_naive, utc_now
from hris.model import TABLES, Tables, create_table_sql, normalize_type

_RUN_ID = re.compile(r"^[A-Za-z0-9_\-]{1,100}$")


def check_run_id(run_id: str) -> str:
    """Les identifiants de run entrent dans des prédicats replaceWhere : format strict."""
    if not run_id or not _RUN_ID.match(run_id):
        raise HrisError("PARAM_INVALID", "run_id invalide")
    return run_id


def sql_str(value: str) -> str:
    return "'" + str(value).replace("\\", "\\\\").replace("'", "\\'") + "'"


def spark_schema(logical: str):
    from pyspark.sql import types as T
    mapping = {"STRING": T.StringType(), "DATE": T.DateType(), "TIMESTAMP": T.TimestampType(),
               "INT": T.IntegerType(), "BIGINT": T.LongType(), "BOOLEAN": T.BooleanType(),
               "ARRAY<STRING>": T.ArrayType(T.StringType())}
    fields = []
    for name, typ in TABLES[logical]:
        t = typ.upper()
        if t.startswith("DECIMAL"):
            p, s = re.findall(r"\d+", t)
            dtype = T.DecimalType(int(p), int(s))
        else:
            dtype = mapping[t]
        fields.append(T.StructField(name, dtype, True))
    return T.StructType(fields)


def conform(df, logical: str):
    """Sélectionne et type les colonnes dans l'ordre du DDL (colonnes absentes -> NULL)."""
    from pyspark.sql import functions as F
    schema = spark_schema(logical)
    cols = []
    present = set(df.columns)
    for f in schema.fields:
        if f.name in present:
            cols.append(F.col(f"`{f.name}`").cast(f.dataType).alias(f.name))
        else:
            cols.append(F.lit(None).cast(f.dataType).alias(f.name))
    return df.select(*cols)


def rows_df(spark, rows: List[dict], logical: str):
    schema = spark_schema(logical)
    data = [tuple(_py(r.get(f.name), f.dataType) for f in schema.fields) for r in rows]
    return spark.createDataFrame(data, schema=schema)


def _py(value, dtype):
    from pyspark.sql import types as T
    if value is None:
        return None
    if isinstance(dtype, T.TimestampType) and isinstance(value, dt.datetime):
        return to_utc_naive(value)
    if isinstance(dtype, T.DecimalType):
        import decimal
        return decimal.Decimal(str(value))
    if isinstance(dtype, T.IntegerType) or isinstance(dtype, T.LongType):
        return int(value)
    return value


def replace_run(spark, df, physical: str, logical: str, run_id: str, extra_predicate: str = "") -> None:
    check_run_id(run_id)
    predicate = f"run_id = {sql_str(run_id)}" + (f" AND ({extra_predicate})" if extra_predicate else "")
    (conform(df, logical).write.format("delta").mode("overwrite")
     .option("replaceWhere", predicate).saveAsTable(physical))


def append(df, physical: str, logical: str) -> None:
    conform(df, logical).write.format("delta").mode("append").saveAsTable(physical)


def merge_rows(spark, rows: List[dict], physical: str, logical: str, keys: List[str],
               update: bool = True) -> None:
    """MERGE idempotent sur clé naturelle (insert si absent, mise à jour sinon)."""
    if not rows:
        return
    from delta.tables import DeltaTable
    src = rows_df(spark, rows, logical)
    cond = " AND ".join(f"t.`{k}` <=> s.`{k}`" for k in keys)
    m = DeltaTable.forName(spark, physical).alias("t").merge(src.alias("s"), cond)
    if update:
        m = m.whenMatchedUpdateAll()
    m.whenNotMatchedInsertAll().execute()


# ---------------------------------------------------------------------------
# Socle : schémas et tables (NB_HRIS_SETUP), contrôle de conformité
# ---------------------------------------------------------------------------
def ensure_schemas_and_tables(spark, T: Tables, logicals: Optional[Iterable[str]] = None) -> List[dict]:
    """Crée ce qui manque ; vérifie le schéma de ce qui existe. Ne modifie jamais une table existante."""
    report = []
    for schema in T.schemas():
        spark.sql(f"CREATE SCHEMA IF NOT EXISTS {schema}")  # schéma déjà qualifié par Tables.schemas()
    for logical in (logicals or TABLES.keys()):
        physical = T(logical)
        existed = spark.catalog.tableExists(physical)
        if not existed:
            spark.sql(create_table_sql(physical, logical))
            report.append({"table": physical, "action": "CREATED", "detail": ""})
            continue
        report.append({"table": physical, "action": "EXISTS", "detail": check_table_schema(spark, physical, logical)})
    return report


def check_table_schema(spark, physical: str, logical: str) -> str:
    actual = {f.name.lower(): normalize_type(f.dataType.simpleString()) for f in spark.table(physical).schema.fields}
    problems = []
    for name, typ in TABLES[logical]:
        exp = normalize_type(typ)
        got = actual.get(name.lower())
        if got is None:
            problems.append(f"colonne manquante {name}")
        elif got != exp:
            problems.append(f"type {name} {got}!={exp}")
    return "CONFORME" if not problems else "INCOMPATIBLE: " + "; ".join(problems)


# ---------------------------------------------------------------------------
# Journal des étapes (ctl.run_step) et anomalies (log.anomaly)
# ---------------------------------------------------------------------------
class StepLogger:
    """`STARTED` au début, `SUCCEEDED` / `FAILED` / `SKIPPED` à la fin (03 § 4). attempt_no = nombre
    de tentatives antérieures de la même étape pour ce run + 1."""

    def __init__(self, spark, T: Tables, run_id: str, step_code: str):
        self.spark, self.T, self.run_id, self.step_code = spark, T, check_run_id(run_id), step_code
        self.started = utc_now()
        prior = (spark.table(T("ctl.run_step"))
                 .where(f"run_id = {sql_str(run_id)} AND step_code = {sql_str(step_code)} AND status = 'STARTED'")
                 .count())
        self.attempt_no = int(prior) + 1
        self._write("STARTED", None, None, None, "")

    def _write(self, status, rows_in, rows_out, content_hash, message, counters_only=False):
        row = {"run_id": self.run_id, "step_code": self.step_code, "attempt_no": self.attempt_no,
               "status": status, "rows_in": rows_in, "rows_out": rows_out, "content_hash": content_hash,
               "message": sanitize(message, 8000, mask_numbers=not counters_only), "started_at_utc": self.started,
               "ended_at_utc": None if status == "STARTED" else utc_now()}
        append(rows_df(self.spark, [row], "ctl.run_step"), self.T("ctl.run_step"), "ctl.run_step")

    def succeeded(self, rows_in=None, rows_out=None, content_hash=None, message="", counters_only=False):
        """counters_only=True : message construit par le code à partir de compteurs (nombres conservés)."""
        self._write("SUCCEEDED", rows_in, rows_out, content_hash, message, counters_only)

    def skipped(self, message=""):
        self._write("SKIPPED", None, None, None, message)

    def failed(self, exc: BaseException):
        self._write("FAILED", None, None, None, safe_error(exc))


def step_started(spark, T: Tables, run_id: str, step_code: str) -> bool:
    return spark.table(T("ctl.run_step")).where(
        f"run_id = {sql_str(run_id)} AND step_code = {sql_str(step_code)}").limit(1).count() > 0


def step_succeeded(spark, T: Tables, run_id: str, step_code: str) -> bool:
    return spark.table(T("ctl.run_step")).where(
        f"run_id = {sql_str(run_id)} AND step_code = {sql_str(step_code)} AND status = 'SUCCEEDED'"
    ).limit(1).count() > 0


def write_anomalies(spark, T: Tables, run_id: str, anomalies: List[dict], owned_codes: Iterable[str]) -> None:
    """Remplace, pour ce run, les contrôles dont l'étape est propriétaire (rejouable sans doublon)."""
    owned = sorted(set(owned_codes) | {a["control_code"] for a in anomalies})
    now = utc_now()
    rows = [{"run_id": run_id, "process_code": PROCESS_CODE, "created_at_utc": now,
             "message": sanitize(a.get("message", ""), 500), **{k: a.get(k) for k in
             ("control_code", "severity", "metric_value", "threshold_value", "outcome")}} for a in anomalies]
    predicate = "control_code IN (" + ",".join(sql_str(c) for c in owned) + ")" if owned else "1=0"
    replace_run(spark, rows_df(spark, rows, "log.anomaly"), T("log.anomaly"), "log.anomaly", run_id, predicate)


def blocking_failures(spark, T: Tables, run_id: str) -> List[str]:
    rows = (spark.table(T("log.anomaly"))
            .where(f"run_id = {sql_str(run_id)} AND severity = 'BLOCKING' AND outcome = 'FAIL'")
            .select("control_code").distinct().collect())
    return sorted(r.control_code for r in rows)


# ---------------------------------------------------------------------------
# Registre des runs (ctl.run)
# ---------------------------------------------------------------------------
def get_run(spark, T: Tables, run_id: str) -> Optional[dict]:
    rows = spark.table(T("ctl.run")).where(f"run_id = {sql_str(check_run_id(run_id))}").collect()
    if len(rows) > 1:
        raise HrisError("RUN_DUPLICATE", "plusieurs lignes ctl.run pour le même run")
    return rows[0].asDict() if rows else None


def update_run(spark, T: Tables, run_id: str, **fields: Any) -> None:
    """MERGE de mise à jour de ctl.run (colonnes nommées uniquement)."""
    from delta.tables import DeltaTable
    from pyspark.sql import functions as F
    check_run_id(run_id)
    if not fields:
        return
    schema = {n: t for n, t in TABLES["ctl.run"]}
    sets = {}
    for k, v in fields.items():
        if k not in schema:
            raise KeyError(k)
        if isinstance(v, dt.datetime):
            v = to_utc_naive(v)
        if k == "error_summary" and v is not None:
            v = sanitize(v, 1000)
        sets[k] = F.lit(v).cast(spark_schema("ctl.run")[k].dataType)
    (DeltaTable.forName(spark, T("ctl.run")).alias("t")
     .update(condition=F.col("t.run_id") == F.lit(run_id), set=sets))


def set_once_promoted(spark, T: Tables, run_id: str, when: dt.datetime) -> None:
    """promoted_at_utc écrit une seule fois (garde d'idempotence de la promotion)."""
    from delta.tables import DeltaTable
    from pyspark.sql import functions as F
    (DeltaTable.forName(spark, T("ctl.run")).alias("t")
     .update(condition=(F.col("t.run_id") == F.lit(run_id)) & F.col("t.promoted_at_utc").isNull(),
             set={"promoted_at_utc": F.lit(to_utc_naive(when)).cast("timestamp")}))
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.transform =============
register_hris_module('hris.transform', '4ec98aabe890c0f7120d801b5468250db4824daef6a17a4de1b25823d2193357', r'''"""hris.transform — implémentation Spark des procédures SQL de production (A02, A09).

Doit produire exactement les résultats de hris.transform_ref (test de parité sur données fictives).
Guide T-SQL -> Spark appliqué (A02 § 8) : espaces de fin (rtrim), NULL vs '', troncature à 50,
CONVERT(DATE,'') = 1900-01-01, émulations ISDATE / ISNUMERIC, décimales.

Écart relevé par rapport à A02 § 6 : en Spark 3.5, `decimal(12,2) / 12` donne DECIMAL(16,6)
(littéral typé au plus juste) ; le diviseur est donc converti explicitement en DECIMAL(10,0)
pour obtenir DECIMAL(23,13) comme T-SQL (R-SQL-054).
"""
from __future__ import annotations

from typing import Dict, List, Tuple

from hris.common import HrisError, PROCESS_CODE, sha256_hex, utc_now
from hris.model import (ADP_COLUMNS, ADP_DATE_COLUMNS, COMPARE_COLUMNS, FORMAT_REJECT_COLUMNS,
                        FORMAT_SUM_EXCLUDED, LOOKUP_JOINS, LOOKUP_TABLES, NULL_REJECT_COLUMNS, REFERENCE_REJECTS,
                        REJECT_TYPE_FORMAT, REJECT_TYPE_NULL, REJECT_TYPE_REFERENCE, SOURCE_COLUMNS,
                        USED_LOOKUP_TABLES, Tables)

SENSITIVE_REJECT_COLUMNS = ["dt_naissance", "id_numero_securite_sociale", "lb_iban", "lb_bic",
                            "lb_account_number", "lb_account_holder", "lb_bank_name"]  # D-SEC-03

_WS_TRIM = r"^[ \t\r\n]+|[ \t\r\n]+$"
# Gardes identiques à hris.parity.DATE_PATTERNS ; l'analyse elle-même est stricte (calendrier réel)
_SPARK_DATE_PARSE = {
    "ymd_dash": ("text", "yyyy-M-d"),
    "ymd_compact": ("text", "yyyyMMdd"),
    "ymd_iso_time": ("first10", "yyyy-MM-dd"),
    "ymd_space_time": ("before_space", "yyyy-M-d"),
    "mdy_slash": ("text", "M/d/yyyy"),
}
_ISNUMERIC = r"^[+-]?[$€£¥]?[ ]*[+-]?(\d[\d,]*\.?\d*|\.\d*|,[\d,]*\.?\d*)?([eEdD][+-]?\d+)?$"
_STRICT_DECIMAL = r"^[ \t]*([+-]?)(\d+\.?\d*|\.\d+)[ \t]*$"


def _F():
    from pyspark.sql import functions as F
    return F


def ensure_session_settings(spark) -> None:
    """Réglages requis par les émulations : fuseau UTC et analyseur de dates « CORRECTED »
    (les chaînes non conformes donnent NULL au lieu d'une exception de compatibilité)."""
    spark.conf.set("spark.sql.session.timeZone", "UTC")
    spark.conf.set("spark.sql.legacy.timeParserPolicy", "CORRECTED")


# ---------------------------------------------------------------------------
# Sémantiques T-SQL en expressions Spark (miroir de hris.parity)
# ---------------------------------------------------------------------------
def is_empty(c):
    F = _F()
    return F.rtrim(c) == F.lit("")


def is_not_empty(c):
    F = _F()
    return F.rtrim(c) != F.lit("")


def parse_date(c):
    """Partie date reconnue (même motif que hris.parity.parse_tsql_date) ou NULL."""
    from hris.parity import DATE_PATTERNS
    F = _F()
    text = F.regexp_replace(c, _WS_TRIM, "")
    out = F.lit(None).cast("date")
    for name, guard, _extract, _order in reversed(DATE_PATTERNS):
        source, fmt = _SPARK_DATE_PARSE[name]
        arg = {"text": text, "first10": F.substring(text, 1, 10),
               "before_space": F.regexp_extract(text, r"^(\S+) ", 1)}[source]
        parsed = F.try_to_timestamp(arg, F.lit(fmt)).cast("date")
        out = F.when(text.rlike(guard), parsed).otherwise(out)
    return out


def isdate(c):
    F = _F()
    d = parse_date(c)
    return F.when(d.isNotNull() & (F.year(d) >= 1753), F.lit(1)).otherwise(F.lit(0))


def convert_date(c):
    """CONVERT(DATE, x) : NULL -> NULL ; blanc -> 1900-01-01 ; non convertible -> erreur."""
    F = _F()
    d = parse_date(c)
    blank = F.regexp_replace(c, _WS_TRIM, "") == F.lit("")
    return (F.when(c.isNull(), F.lit(None).cast("date"))
            .when(blank, F.lit("1900-01-01").cast("date"))
            .when(d.isNotNull(), d)
            .otherwise(F.raise_error(F.lit("CONVERSION_DATE")).cast("date")))


def isnumeric(c):
    F = _F()
    text = F.regexp_replace(c, _WS_TRIM, "")
    body = F.regexp_extract(text, _ISNUMERIC, 1)
    exp = F.regexp_extract(text, _ISNUMERIC, 2)
    valid = text.rlike(_ISNUMERIC) & ~((exp != "") & ~body.rlike(r"\d"))
    return (F.when(c.isNull(), F.lit(0)).when(text == "", F.lit(0))
            .when(valid, F.lit(1)).otherwise(F.lit(0)))


def cast_numeric_12_2(c):
    """CAST(x AS NUMERIC(12,2)) strict ; échec -> erreur (parité : le run échoue)."""
    F = _F()
    sign = F.regexp_extract(c, _STRICT_DECIMAL, 1)
    num = F.regexp_extract(c, _STRICT_DECIMAL, 2)
    wide = F.concat(sign, num).cast("decimal(38,10)")
    rounded = F.round(wide, 2)
    ok = c.rlike(_STRICT_DECIMAL) & wide.isNotNull() & (F.abs(rounded) < F.lit(10000000000))
    return (F.when(c.isNull(), F.lit(None).cast("decimal(12,2)"))
            .when(ok, rounded.cast("decimal(12,2)"))
            .otherwise(F.raise_error(F.lit("CONVERSION_NUMERIC")).cast("decimal(12,2)")))


def pct_flag(c, upper: int):
    F = _F()
    value = cast_numeric_12_2(c)
    return (F.when(isnumeric(c) == 0, F.lit(0))
            .when((value >= F.lit(0)) & (value <= F.lit(upper)), F.lit(0)).otherwise(F.lit(1)))


# ---------------------------------------------------------------------------
# PROC1 : stg.ts_employee_raw -> stg.ts_employee_source
# ---------------------------------------------------------------------------
def source_df(raw):
    F = _F()
    cols = [F.col("run_id"), F.col("business_date")]
    for key, col in SOURCE_COLUMNS:
        if key == "BusinessMobile|BusinessPhone":
            value = F.coalesce(F.col("BusinessMobile"), F.col("BusinessPhone"))
        elif key == "date_fin_contract":
            value = F.expr("replace(`date_fin_contract`, 'None', '')")
        else:
            value = F.col(f"`{key}`")
        cols.append(F.substring(value, 1, 50).alias(col))
    return raw.select(*cols)


# ---------------------------------------------------------------------------
# Contrôles des correspondances (D-QUA-03, À CONFIRMER : bloquant)
# ---------------------------------------------------------------------------
def mapping_controls(spark, T: Tables) -> Tuple[List[dict], str]:
    F = _F()
    anomalies = []
    digests = []
    keys = {j["table"]: j["key"] for j in LOOKUP_JOINS}
    for table in USED_LOOKUP_TABLES:
        physical = T(f"cfg.{table}")
        if not spark.catalog.tableExists(physical):
            anomalies.append(_anom("MAPPING_TABLE_MISSING", 1, f"{table} absente"))
            continue
        df = spark.table(physical)
        missing = [c for c in LOOKUP_TABLES[table] if c not in df.columns]
        if missing:
            anomalies.append(_anom("MAPPING_SCHEMA", len(missing), f"{table} : colonnes attendues absentes"))
            continue
        n = df.count()
        if n == 0:
            anomalies.append(_anom("MAPPING_TABLE_EMPTY", 0, f"{table} vide (référentiel TalentSoft non synchronisé)"))
            continue
        # Partie TalentSoft synchronisée par l'API, partie ADP saisie par le métier : sans aucun code ADP, la table
        # est vide pour la paie (sinon tous les salariés seraient rejetés sans blocage explicite).
        adp = sorted({lu for j in LOOKUP_JOINS if j["table"] == table for lu, _ in j["out"] if lu != j["key"]})
        if not df.where(" OR ".join(f"(`{c}` IS NOT NULL AND trim(`{c}`) <> '')" for c in adp)).limit(1).count():
            anomalies.append(_anom("MAPPING_TABLE_EMPTY", 0, f"{table} : aucun code ADP renseigné (NB_HRIS_REFERENCE_ADMIN)"))
            continue
        dup = (df.where(F.col(keys[table]).isNotNull()).groupBy(F.rtrim(F.col(keys[table])).alias("k"))
               .count().where("count > 1").count())
        if dup:
            anomalies.append(_anom("MAPPING_KEY_DUPLICATE", dup, f"{table} : {dup} clés en double"))
        row_digest = F.sha2(F.concat_ws("\u001f", *[F.coalesce(F.col(c), F.lit("\u0000"))
                                                    for c in LOOKUP_TABLES[table]]), 256)
        hashes = sorted(r[0] for r in df.select(row_digest).collect())
        digests.append(f"{table}:{sha256_hex(''.join(hashes))}")
    controls = ["MAPPING_TABLE_MISSING", "MAPPING_SCHEMA", "MAPPING_TABLE_EMPTY", "MAPPING_KEY_DUPLICATE"]
    found = {a["control_code"] for a in anomalies}
    for code in controls:
        if code not in found:
            anomalies.append({"control_code": code, "severity": "BLOCKING", "metric_value": 0,
                              "threshold_value": None, "outcome": "PASS", "message": "13 tables contrôlées"})
    return anomalies, sha256_hex("|".join(digests))


def _anom(code, metric, message):
    return {"control_code": code, "severity": "BLOCKING", "metric_value": metric, "threshold_value": None,
            "outcome": "FAIL", "message": message}


# ---------------------------------------------------------------------------
# PROC2 : staging (15 LEFT JOIN)
# ---------------------------------------------------------------------------
def enriched_df(spark, T: Tables, source, dt_staging, loader=None):
    """`loader(table)` -> DataFrame de correspondance (défaut : table Delta cfg.<table>)."""
    F = _F()
    loader = loader or (lambda table: spark.table(T(f"cfg.{table}")))
    df = source.withColumn("dt_staging", F.lit(dt_staging).cast("date"))
    for i, join in enumerate(LOOKUP_JOINS):
        lu = loader(join["table"])
        alias = f"lu{i}"
        sel = [F.rtrim(F.col(join["key"])).alias(f"__k{i}")] + \
              [F.col(lu_col).alias(out) for lu_col, out in join["out"]]
        lu = lu.select(*sel).alias(alias)
        df = df.join(lu, F.rtrim(df[join["src"]]) == F.col(f"{alias}.__k{i}"), "left").drop(f"__k{i}")
    return df


# ---------------------------------------------------------------------------
# PROC3 : rejets
# ---------------------------------------------------------------------------
def reject_df(enriched, dt_rejet):
    F = _F()
    base_cols = [F.lit(dt_rejet).cast("date").alias("dt_rejet"), F.col("id_unique"), F.col("id_payroll")]

    def explode_family(entries, typ, where_any):
        arr = F.array(*[F.struct(F.lit(i).alias("ord"), F.lit(fichier).alias("lb_fichier_mapping"),
                                 F.lit(nom).alias("lb_nom_colonne_rejet"), val.alias("lb_valeur_rejet"),
                                 flag.alias("fl_rejet"), ds.alias("ds_rejet"))
                        for i, (fichier, nom, val, flag, ds) in enumerate(entries)])
        return (enriched.where(where_any)
                .select(*base_cols, F.posexplode(arr).alias("pos", "r"))
                .where(F.col("r.fl_rejet") == 1)
                .select("dt_rejet", "id_unique", "id_payroll", F.lit(typ).alias("cd_type_rejet"),
                        F.col("r.lb_nom_colonne_rejet").alias("lb_nom_colonne_rejet"),
                        F.col("r.lb_valeur_rejet").alias("lb_valeur_rejet"), F.col("r.fl_rejet").alias("fl_rejet"),
                        F.col("r.lb_fichier_mapping").alias("lb_fichier_mapping"),
                        F.col("r.ds_rejet").alias("ds_rejet")))

    # Null_Rejet
    null_flags = {c: F.when(is_empty(F.col(c)), 1).otherwise(0) for c in NULL_REJECT_COLUMNS}
    null_entries = [("", c, F.col(c), null_flags[c], F.concat(F.lit(c), F.lit(" ne peut pas être NULL")))
                    for c in NULL_REJECT_COLUMNS]
    null_any = None
    for c in NULL_REJECT_COLUMNS:
        cond = is_empty(F.col(c))
        null_any = cond if null_any is None else (null_any | cond)
    nulls = explode_family(null_entries, REJECT_TYPE_NULL, F.coalesce(null_any, F.lit(False)))

    # Format_Rejet
    def fflag(col, kind):
        c = F.col(col)
        if kind == "len8":
            return F.when(F.length(F.rtrim(c)) == 8, 0).otherwise(1)
        if kind == "date":
            return F.when(is_empty(c) | (isdate(c) == 1), 0).otherwise(1)
        if kind == "date10":
            return F.when(is_empty(c) | (isdate(F.substring(c, 1, 10)) == 1), 0).otherwise(1)
        if kind == "pct":
            return pct_flag(c, 100)
        if kind == "amount":
            return pct_flag(c, 1000000)
        raise ValueError(kind)

    flags = {col: F.coalesce(fflag(col, kind), F.lit(1)) for col, kind in FORMAT_REJECT_COLUMNS}

    def fds(col):
        val = F.col(col)
        if col[:2] == "dt":
            return F.concat(F.lit(f"Format Date non respecté au YYYY-MM-DD pour {col} = "), val)
        if col[:2] == "pc":
            return F.concat(F.lit(f"Format Pourcentage non respecté du 0 au 100 pour {col} = "), val)
        return F.lit(None).cast("string")

    total = None
    for col, _ in FORMAT_REJECT_COLUMNS:
        if col in FORMAT_SUM_EXCLUDED:
            continue
        total = flags[col] if total is None else total + flags[col]
    fmt_entries = [("", col, F.col(col), flags[col], fds(col)) for col, _ in FORMAT_REJECT_COLUMNS]
    formats = explode_family(fmt_entries, REJECT_TYPE_FORMAT, total > 0)

    # Référence_Rejet
    ref_entries = []
    ref_total = None
    for fichier, nom, val_col, lu_col, adp_col in REFERENCE_REJECTS:
        missing = F.col(lu_col).isNull()
        if adp_col is not None:
            missing = missing | F.col(adp_col).isNull()
        flag = F.when(is_not_empty(F.col(val_col)) & missing, 1).otherwise(0)
        ref_total = flag if ref_total is None else ref_total + flag
        ref_entries.append((fichier, nom, F.col(val_col), flag,
                            F.concat(F.lit(f"Mapping non trouvé entre TS et ADP pour {nom} = "), F.col(val_col))))
    refs = explode_family(ref_entries, REJECT_TYPE_REFERENCE, ref_total > 0)
    return nulls.unionByName(formats).unionByName(refs)


def masked_log_reject(rejects, run_id: str):
    """log.reject : valeurs sensibles masquées (D-SEC-03)."""
    F = _F()
    sensitive = F.col("lb_nom_colonne_rejet").isin(*SENSITIVE_REJECT_COLUMNS)
    return rejects.select(
        F.lit(run_id).alias("run_id"), F.lit(PROCESS_CODE).alias("process_code"), "dt_rejet", "id_unique",
        "id_payroll", "cd_type_rejet", "lb_nom_colonne_rejet",
        F.when(sensitive & F.col("lb_valeur_rejet").isNotNull(), F.lit("***")).otherwise(F.col("lb_valeur_rejet"))
        .alias("lb_valeur_rejet"),
        "lb_fichier_mapping",
        F.when(sensitive & F.col("ds_rejet").contains(" = "),
               F.concat(F.regexp_replace(F.col("ds_rejet"), r" = .*$", ""), F.lit(" = ***")))
        .otherwise(F.col("ds_rejet")).alias("ds_rejet"),
        F.lit(utc_now().replace(tzinfo=None)).cast("timestamp").alias("created_at_utc"))


# ---------------------------------------------------------------------------
# PROC4 : transformation ADP + empreinte
# ---------------------------------------------------------------------------
def adp_df(enriched, rejects, dt_transforme):
    F = _F()
    rejected_keys = rejects.where(F.col("id_unique").isNotNull()).select("id_unique").distinct()
    kept = enriched.join(rejected_keys, on="id_unique", how="left_anti")
    cols = [F.lit(dt_transforme).cast("date").alias("dt_transforme")]
    for col in ADP_COLUMNS[1:]:
        if col in ADP_DATE_COLUMNS:
            src = F.substring(F.col(col), 1, 10) if col == "dt_entree_societe" else F.col(col)
            cols.append(convert_date(src).alias(col))
        elif col == "mt_salaire_mensuel":
            divisor = F.lit(12).cast("decimal(10,0)")
            cols.append((cast_numeric_12_2(F.col("mt_salaire_base")) / divisor).alias(col))
        else:
            cols.append(F.col(col))
    out = kept.select(*cols)
    return out.withColumn("row_hash", row_hash_expr(COMPARE_COLUMNS))


def row_hash_expr(columns):
    F = _F()
    return F.sha2(F.concat_ws("\u001f", *[F.coalesce(F.col(c).cast("string"), F.lit("\u0000")) for c in columns]), 256)
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.compare =============
register_hris_module('hris.compare', '44fbebdd86edfc03c67e3a93e7085e98da1800e48b52aca7ec7702ef0106e03c', r'''"""hris.compare — comparaison staging finale <-> référence du dernier run validé, puis promotion.

Règles (01 § 4.5-4.8, 03 § 7, A08 D-CMP-01 à 07 — À CONFIRMER, valeurs par défaut appliquées) :
  * catégories NEW / MODIFIED / UNCHANGED / ABSENT / REJECTED sur id_unique et row_hash ;
  * le comparatif ne journalise que des NOMS de colonnes et des empreintes ;
  * premier chargement : NORMAL + référence vide -> BLOQUÉ (REFERENCE_ABSENTE) ;
    INIT_REFERENCE accepté seulement si la référence est vide (jamais de réinitialisation) ;
  * seuils : NULL -> NOT_CONFIGURED, non bloquant ; aucun seuil n'est inventé ;
  * recalcul interdit dès qu'une écriture de promotion a eu lieu (référence ou outbox du run) ;
  * promotion : revalidation des contrôles, un seul MERGE, outbox écrite dans la même étape,
    promoted_at_utc écrit une fois.
"""
from __future__ import annotations

import decimal
from typing import Dict, List, Optional

from hris.common import HrisError, PROCESS_CODE, RUN_BLOCKED, RUN_COMPARED, RUN_PROMOTED, RUN_VALIDATED, utc_now
from hris.model import ADP_COLUMNS, BUSINESS_KEY, COMPARE_COLUMNS, Tables
from hris.spark_io import (blocking_failures, conform, get_run, replace_run, rows_df, set_once_promoted,
                           sql_str, update_run, write_anomalies)

COMPARE_CONTROLS = ["DUPLICATE_KEY", "NULL_KEY", "REFERENCE_ABSENTE", "REFERENCE_NOT_EMPTY",
                    "ABSENT_RATIO", "MODIFIED_RATIO", "REJECT_RATIO", "STAGING_MISSING"]


def _F():
    from pyspark.sql import functions as F
    return F


def last_promoted_run(spark, T: Tables) -> Optional[str]:
    rows = (spark.table(T("ctl.run")).where(f"process_code = {sql_str(PROCESS_CODE)} AND promoted_at_utc IS NOT NULL")
            .orderBy(_F().col("promoted_at_utc").desc()).select("run_id").limit(1).collect())
    return rows[0].run_id if rows else None


def promotion_writes_exist(spark, T: Tables, run_id: str) -> bool:
    run = get_run(spark, T, run_id) or {}
    if run.get("promoted_at_utc"):
        return True
    if spark.table(T("ctl.run_step")).where(
            f"run_id = {sql_str(run_id)} AND step_code = 'PROMOTE_WRITE'").limit(1).count():
        return True  # marqueur écrit juste avant la première écriture de promotion
    if spark.table(T("ref.ts_employee_adp")).where(f"_last_run_id = {sql_str(run_id)}").limit(1).count():
        return True
    return spark.table(T("pub.adp_outbox")).where(f"run_id = {sql_str(run_id)}").limit(1).count() > 0


def load_thresholds(spark, T: Tables) -> Dict[str, Optional[decimal.Decimal]]:
    rows = spark.table(T("cfg.process")).where(f"process_code = {sql_str(PROCESS_CODE)}").collect()
    if not rows:
        return {"max_absent_pct": None, "max_modified_pct": None, "max_reject_pct": None}
    r = rows[0]
    return {"max_absent_pct": r.max_absent_pct, "max_modified_pct": r.max_modified_pct,
            "max_reject_pct": r.max_reject_pct}


def _pct(num: int, den: int) -> Optional[decimal.Decimal]:
    if den == 0:
        return None
    return (decimal.Decimal(num) * 100 / decimal.Decimal(den)).quantize(decimal.Decimal("0.0001"))


def threshold_controls(summary: dict, thresholds: dict) -> List[dict]:
    """D-CMP-07 : contrôle calculé et journalisé ; NOT_CONFIGURED (non bloquant) tant que le seuil est NULL."""
    out = []
    for code, metric_key, thr_key in [("ABSENT_RATIO", "pct_absent", "max_absent_pct"),
                                       ("MODIFIED_RATIO", "pct_modified", "max_modified_pct"),
                                       ("REJECT_RATIO", "pct_rejected", "max_reject_pct")]:
        metric = summary.get(metric_key)
        thr = thresholds.get(thr_key)
        if thr is None:
            out.append({"control_code": code, "severity": "WARNING", "metric_value": metric, "threshold_value": None,
                        "outcome": "NOT_CONFIGURED", "message": "seuil non fourni (D-CMP-07)"})
        else:
            fail = metric is not None and decimal.Decimal(metric) > decimal.Decimal(thr)
            out.append({"control_code": code, "severity": "BLOCKING", "metric_value": metric, "threshold_value": thr,
                        "outcome": "FAIL" if fail else "PASS", "message": "seuil configuré"})
    return out


def compare(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    F = _F()
    run = get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if promotion_writes_exist(spark, T, run_id):
        # Recalcul interdit : la référence a (au moins en partie) été promue pour ce run.
        prev = spark.table(T("log.comparison_summary")).where(f"run_id = {sql_str(run_id)}").collect()
        return {"decision": run["status"], "recomputed": False,
                **({k: prev[0][k] for k in ("nb_new", "nb_modified", "nb_absent", "nb_rejected")} if prev else {})}

    staging = spark.table(T("stg.ts_employee_adp")).where(f"run_id = {sql_str(run_id)}").drop("run_id", "business_date")
    rejected = (spark.table(T("stg.ts_employee_reject")).where(f"run_id = {sql_str(run_id)} AND id_unique IS NOT NULL")
                .select(F.col("id_unique").alias("k")).distinct())
    ref = spark.table(T("ref.ts_employee_adp"))
    anomalies: List[dict] = []
    nb_staging_rows = staging.count()
    nb_null = staging.where(F.col(BUSINESS_KEY).isNull()).count()
    nb_dup = staging.where(F.col(BUSINESS_KEY).isNotNull()).groupBy(BUSINESS_KEY).count().where("count > 1").count()
    nb_reference = ref.count()
    nb_rejected = rejected.count()
    is_first_load = nb_reference == 0
    staging_ran = spark.table(T("ctl.run_step")).where(
        f"run_id = {sql_str(run_id)} AND step_code = 'TRANSFORM' AND status = 'SUCCEEDED'").limit(1).count() > 0

    anomalies.append({"control_code": "STAGING_MISSING", "severity": "BLOCKING",
                      "metric_value": nb_staging_rows, "threshold_value": None,
                      "outcome": "PASS" if staging_ran else "FAIL",
                      "message": "transformation aboutie" if staging_ran else "aucune transformation aboutie pour ce run"})
    anomalies.append({"control_code": "NULL_KEY", "severity": "BLOCKING", "metric_value": nb_null,
                      "threshold_value": None, "outcome": "FAIL" if nb_null else "PASS",
                      "message": f"{nb_null} lignes sans id_unique"})
    anomalies.append({"control_code": "DUPLICATE_KEY", "severity": "BLOCKING", "metric_value": nb_dup,
                      "threshold_value": None, "outcome": "FAIL" if nb_dup else "PASS",
                      "message": f"{nb_dup} clés en double (correspondances dupliquées ?)"})
    if run_mode == "NORMAL":
        anomalies.append({"control_code": "REFERENCE_ABSENTE", "severity": "BLOCKING", "metric_value": nb_reference,
                          "threshold_value": None, "outcome": "FAIL" if is_first_load else "PASS",
                          "message": "référence vide : premier chargement par INIT_REFERENCE (D-CMP-06)"
                          if is_first_load else "référence présente"})
    elif run_mode == "INIT_REFERENCE":
        anomalies.append({"control_code": "REFERENCE_NOT_EMPTY", "severity": "BLOCKING", "metric_value": nb_reference,
                          "threshold_value": None, "outcome": "PASS" if is_first_load else "FAIL",
                          "message": "référence vide : initialisation autorisée" if is_first_load
                          else "référence non vide : INIT_REFERENCE refusé (pas de réinitialisation)"})

    st = staging.where(F.col(BUSINESS_KEY).isNotNull()).alias("s")
    rf = ref.alias("r")
    joined = st.join(rf, F.col(f"s.{BUSINESS_KEY}") == F.col(f"r.{BUSINESS_KEY}"), "full_outer")
    key = F.coalesce(F.col(f"s.{BUSINESS_KEY}"), F.col(f"r.{BUSINESS_KEY}"))
    changed = F.array_compact(F.array(*[
        F.when(~F.col(f"s.{c}").eqNullSafe(F.col(f"r.{c}")), F.lit(c)) for c in COMPARE_COLUMNS]))
    detail = joined.select(
        key.alias("business_key"),
        F.when(F.col(f"r.{BUSINESS_KEY}").isNull(), "NEW")
        .when(F.col(f"s.{BUSINESS_KEY}").isNull(), "ABSENT_OR_REJECTED")
        .when(F.col("s.row_hash") != F.col("r.row_hash"), "MODIFIED").otherwise("UNCHANGED").alias("category"),
        F.when(F.col(f"s.{BUSINESS_KEY}").isNotNull() & F.col(f"r.{BUSINESS_KEY}").isNotNull()
               & (F.col("s.row_hash") != F.col("r.row_hash")), changed).alias("changed_columns"),
        F.col("r.row_hash").alias("old_row_hash"), F.col("s.row_hash").alias("new_row_hash"))
    detail = (detail.join(rejected, detail.business_key == rejected.k, "left")
              .withColumn("category", F.when(F.col("category") == "ABSENT_OR_REJECTED",
                                             F.when(F.col("k").isNotNull(), "REJECTED").otherwise("ABSENT"))
                          .otherwise(F.col("category"))).drop("k"))
    # Clés rejetées absentes de la référence : catégorie REJECTED également
    rej_only = (rejected.join(ref.select(F.col(BUSINESS_KEY).alias("rk")), rejected.k == F.col("rk"), "left_anti")
                .select(F.col("k").alias("business_key"), F.lit("REJECTED").alias("category"),
                        F.lit(None).cast("array<string>").alias("changed_columns"),
                        F.lit(None).cast("string").alias("old_row_hash"), F.lit(None).cast("string").alias("new_row_hash")))
    detail = detail.unionByName(rej_only).cache()
    counts = {r.category: r["count"] for r in detail.groupBy("category").count().collect()}
    summary = {"nb_staging": nb_staging_rows, "nb_reference": nb_reference,
               "nb_new": counts.get("NEW", 0), "nb_modified": counts.get("MODIFIED", 0),
               "nb_unchanged": counts.get("UNCHANGED", 0), "nb_absent": counts.get("ABSENT", 0),
               "nb_rejected": counts.get("REJECTED", 0)}
    summary["pct_absent"] = _pct(summary["nb_absent"], nb_reference)
    summary["pct_modified"] = _pct(summary["nb_modified"], nb_reference)
    summary["pct_rejected"] = _pct(summary["nb_rejected"], nb_staging_rows + nb_rejected)
    anomalies += threshold_controls(summary, load_thresholds(spark, T))

    now = utc_now()
    against = last_promoted_run(spark, T)
    replace_run(spark, rows_df(spark, [{"run_id": run_id, "process_code": PROCESS_CODE,
                                        "business_date": run["business_date"], "compared_against_run_id": against,
                                        "is_first_load": is_first_load, "created_at_utc": now, **summary}],
                               "log.comparison_summary"),
                T("log.comparison_summary"), "log.comparison_summary", run_id)
    detail_out = (detail.where("category <> 'UNCHANGED'")
                  .withColumn("run_id", F.lit(run_id)).withColumn("process_code", F.lit(PROCESS_CODE))
                  .withColumn("created_at_utc", F.lit(now.replace(tzinfo=None)).cast("timestamp")))
    replace_run(spark, detail_out, T("log.comparison_detail"), "log.comparison_detail", run_id)
    detail.unpersist()
    write_anomalies(spark, T, run_id, anomalies, COMPARE_CONTROLS)

    blocking = blocking_failures(spark, T, run_id)
    decision = RUN_BLOCKED if blocking else RUN_VALIDATED
    update_run(spark, T, run_id, status=decision, compared_against_run_id=against,
               error_code=("BLOCKING_ANOMALY" if blocking else None),
               error_summary=(",".join(blocking) if blocking else None))
    return {"decision": decision, "recomputed": True, "blocking": blocking,
            **{k: summary[k] for k in ("nb_new", "nb_modified", "nb_absent", "nb_rejected", "nb_unchanged")}}


def revalidate(spark, T: Tables, run_id: str) -> List[str]:
    """Avant de lever un blocage ou de promouvoir : les contrôles bloquants du run sont relus et les
    seuils réévalués avec la configuration COURANTE. Renvoie la liste des contrôles en échec."""
    failures = set(blocking_failures(spark, T, run_id))
    rows = spark.table(T("log.comparison_summary")).where(f"run_id = {sql_str(run_id)}").collect()
    if not rows:
        failures.add("COMPARISON_MISSING")
        return sorted(failures)
    summary = rows[0].asDict()
    for a in threshold_controls(summary, load_thresholds(spark, T)):
        if a["severity"] == "BLOCKING" and a["outcome"] == "FAIL":
            failures.add(a["control_code"])
    return sorted(failures)


def promote(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    F = _F()
    from delta.tables import DeltaTable
    run = get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if run.get("promoted_at_utc"):
        # Déjà promu : aucun nouveau MERGE (rejouer la promotion d'un ancien run ramènerait la
        # référence en arrière si un run plus récent a été promu depuis).
        nb = spark.table(T("pub.adp_outbox")).where(f"run_id = {sql_str(run_id)}").count()
        return {"promoted": True, "already_promoted": True, "nb_outbox": nb, "mode": run_mode}
    if run["status"] != RUN_VALIDATED:
        return {"promoted": False, "reason": f"STATUS_{run['status']}", "nb_outbox": 0}
    failures = revalidate(spark, T, run_id)
    if failures:
        update_run(spark, T, run_id, status=RUN_BLOCKED, error_code="REVALIDATION_FAILED",
                   error_summary=",".join(failures))
        return {"promoted": False, "reason": "REVALIDATION_FAILED", "failures": failures, "nb_outbox": 0}
    current_base = last_promoted_run(spark, T)
    if current_base != run.get("compared_against_run_id"):
        update_run(spark, T, run_id, status=RUN_BLOCKED, error_code="REFERENCE_MOVED",
                   error_summary="la référence a été promue par un autre run depuis la comparaison")
        return {"promoted": False, "reason": "REFERENCE_MOVED", "nb_outbox": 0}
    summary = spark.table(T("log.comparison_summary")).where(f"run_id = {sql_str(run_id)}").collect()[0]
    if summary.nb_staging == 0:
        raise HrisError("PROMOTE_EMPTY_STAGING", "staging vide : promotion refusée (protection de la référence)")
    now = utc_now()
    from hris.spark_io import append
    append(rows_df(spark, [{"run_id": run_id, "step_code": "PROMOTE_WRITE", "attempt_no": 1, "status": "STARTED",
                            "message": "écritures de promotion engagées : recalcul du comparatif interdit",
                            "started_at_utc": now}], "ctl.run_step"), T("ctl.run_step"), "ctl.run_step")

    # 1. Outbox (même étape, avant la publication) — jamais en INIT_REFERENCE
    nb_outbox = 0
    if run_mode == "NORMAL":
        changes = (spark.table(T("log.comparison_detail"))
                   .where(f"run_id = {sql_str(run_id)} AND category IN ('NEW','MODIFIED')")
                   .select(F.lit(run_id).alias("run_id"), F.col("business_key").alias("employee_key"),
                           F.col("category").alias("change_category"), F.col("new_row_hash").alias("row_hash"),
                           F.lit("PENDING").alias("status"), F.lit(0).alias("attempt_count"),
                           F.lit(now.replace(tzinfo=None)).cast("timestamp").alias("created_at_utc"),
                           F.lit(now.replace(tzinfo=None)).cast("timestamp").alias("updated_at_utc")))
        nb_outbox = changes.count()
        (DeltaTable.forName(spark, T("pub.adp_outbox")).alias("t")
         .merge(conform(changes, "pub.adp_outbox").alias("s"),
                "t.run_id = s.run_id AND t.employee_key = s.employee_key")
         .whenNotMatchedInsertAll().execute())

    # 2. Référence : un seul MERGE (transaction Delta atomique sur la table)
    staging = (spark.table(T("stg.ts_employee_adp")).where(f"run_id = {sql_str(run_id)}")
               .select(*ADP_COLUMNS, "row_hash").withColumn("_action", F.lit("UPSERT")))
    keep = (spark.table(T("stg.ts_employee_reject")).where(f"run_id = {sql_str(run_id)} AND id_unique IS NOT NULL")
            .select("id_unique").distinct())
    keep = keep.join(staging.select("id_unique"), "id_unique", "left_anti")
    keep_rows = keep.select(*[F.col(c) if c == "id_unique" else F.lit(None).cast(staging.schema[c].dataType).alias(c)
                              for c in ADP_COLUMNS], F.lit(None).cast("string").alias("row_hash"),
                            F.lit("KEEP").alias("_action"))
    source = staging.unionByName(keep_rows)
    values = {c: f"s.`{c}`" for c in ADP_COLUMNS}
    values.update({"row_hash": "s.row_hash", "_last_run_id": sql_str(run_id),
                   "_validated_at_utc": f"timestamp'{now.replace(tzinfo=None).isoformat(sep=' ')}'"})
    (DeltaTable.forName(spark, T("ref.ts_employee_adp")).alias("t")
     .merge(source.alias("s"), f"t.{BUSINESS_KEY} = s.{BUSINESS_KEY}")
     .whenMatchedUpdate(condition="s._action = 'UPSERT' AND t.row_hash <> s.row_hash", set=values)
     .whenNotMatchedInsert(condition="s._action = 'UPSERT'", values=values)
     .whenNotMatchedBySourceDelete()
     .execute())

    # 3. Garde d'idempotence
    set_once_promoted(spark, T, run_id, now)
    update_run(spark, T, run_id, status=RUN_PROMOTED)
    return {"promoted": True, "nb_outbox": nb_outbox, "mode": run_mode}
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.run_control =============
register_hris_module('hris.run_control', '3dc91ff068c000a4fe73d5a5825dbc08944bed3b9e1a8641aad9d46e4e262e5f', r'''"""hris.run_control — modes START / END / FAIL de NB_HRIS_RUN_CONTROL.

Verrou applicatif ctl.process_lock (A07 § 2) : MERGE conditionnel puis relecture ; un conflit de
concurrence Delta ou une relecture différente = verrou non obtenu -> SKIPPED_CONCURRENT (pas d'échec,
aucune écriture métier). Bail : p_lock_lease_minutes (A08 D-RUN-04, À CONFIRMER).
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Optional

from hris.common import (PROCESS_CODE, RUN_BLOCKED, RUN_FAILED, RUN_MODES, RUN_NOT_PUBLISHED_NON_PRD,
                         RUN_PROMOTED, RUN_PUBLISHED, RUN_PUBLISHED_PARTIAL, RUN_SKIPPED_CONCURRENT, RUN_STARTED,
                         RUN_SUCCEEDED, HrisError, business_date_of, parse_trigger_time, sanitize, to_utc_naive,
                         utc_now)
from hris.model import Tables
from hris.parity import utf8_truncate
from hris.spark_io import check_run_id, get_run, merge_rows, replace_run, rows_df, sql_str, update_run

_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _F():
    from pyspark.sql import functions as F
    return F


def acquire_lock(spark, T: Tables, run_id: str, lease_minutes: int, now: Optional[dt.datetime] = None) -> bool:
    from delta.tables import DeltaTable
    F = _F()
    now = now or utc_now()
    lease = now + dt.timedelta(minutes=lease_minutes)
    src = rows_df(spark, [{"process_code": PROCESS_CODE, "lock_status": "HELD", "holder_run_id": run_id,
                           "acquired_at_utc": now, "lease_expires_at_utc": lease, "released_at_utc": None}],
                  "ctl.process_lock")
    now_lit = F.lit(to_utc_naive(now)).cast("timestamp")
    try:
        (DeltaTable.forName(spark, T("ctl.process_lock")).alias("t")
         .merge(src.alias("s"), "t.process_code = s.process_code")
         .whenMatchedUpdate(condition=(F.col("t.lock_status") == "RELEASED") | (F.col("t.lease_expires_at_utc") < now_lit)
                            | (F.col("t.holder_run_id") == F.col("s.holder_run_id")),
                            set={"lock_status": "s.lock_status", "holder_run_id": "s.holder_run_id",
                                 "acquired_at_utc": "s.acquired_at_utc",
                                 "lease_expires_at_utc": "s.lease_expires_at_utc", "released_at_utc": "s.released_at_utc"})
         .whenNotMatchedInsertAll().execute())
    except Exception as exc:  # conflit Delta (Concurrent*Exception) = verrou non obtenu
        if "Concurrent" in type(exc).__name__ or "Concurrent" in str(exc)[:500]:
            return False
        raise
    rows = spark.table(T("ctl.process_lock")).where(f"process_code = {sql_str(PROCESS_CODE)}").collect()
    return len(rows) == 1 and rows[0].lock_status == "HELD" and rows[0].holder_run_id == run_id


def release_lock(spark, T: Tables, run_id: str) -> bool:
    from delta.tables import DeltaTable
    F = _F()
    holder = spark.table(T("ctl.process_lock")).where(
        f"process_code = {sql_str(PROCESS_CODE)} AND holder_run_id = {sql_str(run_id)} AND lock_status = 'HELD'").count()
    if not holder:
        return False
    (DeltaTable.forName(spark, T("ctl.process_lock")).alias("t")
     .update(condition=(F.col("t.process_code") == PROCESS_CODE) & (F.col("t.holder_run_id") == run_id)
             & (F.col("t.lock_status") == "HELD"),
             set={"lock_status": F.lit("RELEASED"), "released_at_utc": F.lit(to_utc_naive(utc_now())).cast("timestamp")}))
    return True


def start(spark, T: Tables, *, run_id: str, process_code: str, run_mode: str, target_run_id: str,
          trigger_time: str, business_timezone: str, lock_lease_minutes: int, business_date_override: str,
          pipeline_name: str, workspace_id: Optional[str], workspace_name: Optional[str]) -> dict:
    check_run_id(run_id)
    if process_code != PROCESS_CODE:
        raise HrisError("PROCESS_NOT_ENABLED", "seul TS_ADP_EMPLOYEE est actif (flux SFTP bloqué D-FIL-01)")
    if run_mode not in RUN_MODES:
        raise HrisError("PARAM_INVALID", "p_run_mode")
    if run_mode == "RETRY_PUBLICATION":
        if not target_run_id:
            raise HrisError("PARAM_MISSING", "p_target_run_id obligatoire en RETRY_PUBLICATION")
        target = get_run(spark, T, check_run_id(target_run_id))
        if target is None or not target.get("promoted_at_utc"):
            raise HrisError("RETRY_TARGET_INVALID", "le run cible n'existe pas ou n'a pas été promu")
    elif target_run_id:
        raise HrisError("PARAM_INVALID", "p_target_run_id réservé au mode RETRY_PUBLICATION")
    if not isinstance(lock_lease_minutes, int) or lock_lease_minutes <= 0:
        raise HrisError("PARAM_INVALID", "p_lock_lease_minutes")
    trigger_utc = parse_trigger_time(trigger_time)
    if business_date_override:
        # tests uniquement ; refusé en NORMAL dans un workspace dont le nom contient PRD (02 § 5, À CONFIRMER)
        if run_mode == "NORMAL" and isinstance(workspace_name, str) and "PRD" in workspace_name:
            raise HrisError("OVERRIDE_REFUSED", "p_business_date_override interdit en NORMAL dans ce workspace")
        if not _DATE.match(business_date_override):
            raise HrisError("PARAM_INVALID", "p_business_date_override (AAAA-MM-JJ)")
        business_date = dt.date.fromisoformat(business_date_override)
    else:
        business_date = business_date_of(trigger_utc, business_timezone)

    existing = get_run(spark, T, run_id)
    if existing and existing["status"] not in (RUN_STARTED, RUN_SKIPPED_CONCURRENT):
        # RunStart rejoué pour un run déjà avancé : on ne réinitialise rien
        return {"status": "OK", "run_id": run_id, "business_date": str(existing["business_date"]),
                "run_seq_in_day": existing["run_seq_in_day"], "resumed": True}

    if not acquire_lock(spark, T, run_id, lock_lease_minutes):
        merge_rows(spark, [{"run_id": run_id, "process_code": PROCESS_CODE, "run_mode": run_mode,
                            "business_date": business_date, "trigger_time_utc": trigger_utc,
                            "pipeline_name": pipeline_name, "workspace_id": workspace_id,
                            "workspace_name": workspace_name, "status": RUN_SKIPPED_CONCURRENT,
                            "started_at_utc": utc_now(), "ended_at_utc": utc_now(),
                            "target_run_id": target_run_id or None}],
                   T("ctl.run"), "ctl.run", ["run_id"])
        return {"status": RUN_SKIPPED_CONCURRENT, "run_id": run_id, "business_date": str(business_date)}

    seq = (spark.table(T("ctl.run"))
           .where(f"process_code = {sql_str(PROCESS_CODE)} AND business_date = DATE'{business_date.isoformat()}' "
                  f"AND status <> '{RUN_SKIPPED_CONCURRENT}' AND run_id <> {sql_str(run_id)}").count()) + 1
    merge_rows(spark, [{"run_id": run_id, "process_code": PROCESS_CODE, "run_mode": run_mode,
                        "business_date": business_date, "run_seq_in_day": seq, "trigger_time_utc": trigger_utc,
                        "pipeline_name": pipeline_name, "workspace_id": workspace_id,
                        "workspace_name": workspace_name, "status": RUN_STARTED, "started_at_utc": utc_now(),
                        "target_run_id": target_run_id or None}],
               T("ctl.run"), "ctl.run", ["run_id"])
    return {"status": "OK", "run_id": run_id, "business_date": business_date.isoformat(), "run_seq_in_day": seq,
            "run_mode": run_mode}


def final_status(run: dict) -> str:
    status = run["status"]
    if status in (RUN_BLOCKED, RUN_FAILED, RUN_SKIPPED_CONCURRENT, RUN_NOT_PUBLISHED_NON_PRD, RUN_PUBLISHED_PARTIAL):
        return status
    if status in (RUN_PUBLISHED, RUN_PROMOTED):
        return RUN_SUCCEEDED
    return RUN_FAILED  # run arrêté avant la fin de la séquence attendue


def run_summary(spark, T: Tables, run: dict) -> dict:
    """log.run_summary : métriques historiques (R-SQL-060 à 066) calculées sur le delta du run
    (NEW + MODIFIED = insérés ; REJECTED = rejetés) + compteurs de publication."""
    F = _F()
    run_id = run["run_id"]
    source_run = run.get("target_run_id") or run_id
    det = spark.table(T("log.comparison_detail")).where(f"run_id = {sql_str(source_run)}")
    inserts = det.where("category IN ('NEW','MODIFIED')").select("business_key").distinct().count()
    rejects_n = det.where("category = 'REJECTED'").select("business_key").distinct().count()
    lues = inserts + rejects_n
    import decimal
    pc = decimal.Decimal(0) if lues == 0 else (decimal.Decimal(rejects_n) / decimal.Decimal(lues)).quantize(
        decimal.Decimal("0.0001"))
    pairs = sorted({(r.cd_type_rejet, r.ds_rejet or "") for r in
                    spark.table(T("stg.ts_employee_reject")).where(f"run_id = {sql_str(source_run)}")
                    .select("cd_type_rejet", "ds_rejet").distinct().collect()})
    ds_type = utf8_truncate(",".join(p[0] for p in pairs), 50) if pairs else None
    ob = {r.status: r["count"] for r in spark.table(T("pub.adp_outbox")).where(f"run_id = {sql_str(source_run)}")
          .groupBy("status").count().collect()}
    trig = run.get("trigger_time_utc")
    return {"run_id": run_id, "process_code": PROCESS_CODE, "business_date": run.get("business_date"),
            "dt_resume": trig.date() if trig else None, "cd_source": "TalentSoft", "cd_target": "ADP",
            "mt_lignes_inserts": inserts, "mt_lignes_rejets": rejects_n, "mt_lignes_lues": lues,
            "pc_lignes_rejets": pc, "ds_type_rejets": ds_type, "nb_outbox": sum(ob.values()),
            "nb_sent": ob.get("SENT", 0), "nb_no_change": ob.get("NO_CHANGE", 0), "nb_failed": ob.get("FAILED", 0),
            "nb_not_pushed": ob.get("NOT_PUSHED_NON_PRD", 0), "created_at_utc": utc_now()}


def end(spark, T: Tables, run_id: str, env_code: str) -> dict:
    run = get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    status = final_status(run)
    summary = run_summary(spark, T, run)
    replace_run(spark, rows_df(spark, [summary], "log.run_summary"), T("log.run_summary"), "log.run_summary", run_id)
    update_run(spark, T, run_id, status=status, ended_at_utc=utc_now())
    released = release_lock(spark, T, run_id)
    notification = {  # contenu préparé, sans donnée personnelle ; envoi bloqué (D-NOT-01/02)
        "type": "SUMMARY" if status in (RUN_SUCCEEDED, RUN_NOT_PUBLISHED_NON_PRD) else "ERROR",
        "env": env_code, "run_id": run_id, "business_date": str(run.get("business_date")), "status": status,
        "inserts": summary["mt_lignes_inserts"], "rejects": summary["mt_lignes_rejets"],
        "outbox": summary["nb_outbox"], "error_code": run.get("error_code"), "sent": False,
        "send_blocked_reason": "D-NOT-01/D-NOT-02 non décidés"}
    return {"status": status, "run_id": run_id, "lock_released": released, "notification": notification}


def fail(spark, T: Tables, run_id: str, error_message: str) -> dict:
    run = get_run(spark, T, run_id)
    if run is None:
        return {"status": RUN_FAILED, "run_id": run_id, "lock_released": False, "note": "run inconnu"}
    if run["status"] == RUN_SKIPPED_CONCURRENT:
        return {"status": RUN_SKIPPED_CONCURRENT, "run_id": run_id, "lock_released": False}
    update_run(spark, T, run_id, status=RUN_FAILED, ended_at_utc=utc_now(), error_code="PIPELINE_FAILURE",
               error_summary=sanitize(error_message, 1000))
    released = release_lock(spark, T, run_id)
    return {"status": RUN_FAILED, "run_id": run_id, "lock_released": released,
            "notification": {"type": "ERROR", "run_id": run_id, "sent": False,
                             "send_blocked_reason": "D-NOT-01/D-NOT-02 non décidés"}}
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.adp_gate =============
register_hris_module('hris.adp_gate', 'e78ceb69b7fd6fdceccc8018e16154788bc7e3bc55436de34f03145142fbd6f8', r'''"""hris.adp_gate — condition d'environnement de la publication ADP (D-ENV-01).

Règle (inchangée, littérale) : la condition d'environnement est satisfaite si et seulement si le
nom RÉEL du workspace d'exécution du notebook de publication contient la sous-chaîne "PRD"
(sensible à la casse, recherche par sous-chaîne, aucune normalisation).

Trois issues distinctes (révision demandée par le mandat) :
  * ENV_SATISFIED     : nom connu contenant "PRD" -> condition d'environnement satisfaite. Cela
                        N'AUTORISE PAS l'envoi : les autres validations (qualité, métier, sécurité,
                        authentification, contrat, activation) restent nécessaires (hris.adp_publish).
  * DENIED_NON_PRD    : nom connu ne contenant pas "PRD" -> absence normale de publication
                        (statut d'outbox NOT_PUSHED_NON_PRD).
  * BLOCKED_UNRESOLVED: nom non résolu, vide ou d'un type inattendu -> publication bloquée, les
                        éléments restent en attente (PENDING) et seront réévalués après résolution.

Le contexte provient exclusivement de notebookutils.runtime.context lu DANS le notebook de
publication. Aucun paramètre, aucune variable, aucun simulateur ne peut forcer la décision :
les contextes simulés ne sont injectés que par les tests isolés.

Ambiguïtés signalées, règle NON modifiée (D-ENV-02, à trancher avant toute activation) :
"PREPRD" et "NOTPRD" satisfont la condition ; "prd" en minuscules ne la satisfait pas.
"""
from __future__ import annotations

import dataclasses
import datetime as dt
import uuid
from typing import Any, Callable, Mapping, Optional

PRD_MARKER = "PRD"
RULE_VERSION = "contains:PRD:case-sensitive:v2"

REASON_CONTAINS_PRD = "WORKSPACE_NAME_CONTAINS_PRD"
REASON_WITHOUT_PRD = "WORKSPACE_NAME_WITHOUT_PRD"
REASON_UNRESOLVED = "WORKSPACE_NAME_UNRESOLVED"
REASON_EMPTY = "WORKSPACE_NAME_EMPTY"
REASON_INVALID_TYPE = "WORKSPACE_NAME_INVALID_TYPE"

ENV_SATISFIED = "ENV_SATISFIED"
DENIED_NON_PRD = "DENIED_NON_PRD"
BLOCKED_UNRESOLVED = "BLOCKED_UNRESOLVED"


def _utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


@dataclasses.dataclass(frozen=True)
class WorkspaceContext:
    name: Any
    workspace_id: Optional[str]
    is_for_pipeline: Optional[bool]
    notebook_name: Optional[str]
    resolution_error: Optional[str]


@dataclasses.dataclass(frozen=True)
class GateDecision:
    decision_id: str
    outcome: str
    env_condition_met: bool
    reason_code: str
    workspace_name: Optional[str]
    workspace_id: Optional[str]
    is_for_pipeline: Optional[bool]
    notebook_name: Optional[str]
    rule_version: str
    evaluated_at_utc: dt.datetime

    @property
    def allowed(self) -> bool:  # compatibilité : True = condition d'environnement satisfaite
        return self.env_condition_met


def resolve_current_workspace(context_provider: Callable[[], Mapping[str, Any]]) -> WorkspaceContext:
    """En Fabric : ``context_provider = lambda: notebookutils.runtime.context`` (A06 F3)."""
    try:
        ctx = context_provider()
    except Exception as exc:  # contexte indisponible : échec fermé
        return WorkspaceContext(None, None, None, None, type(exc).__name__)

    def _get(key: str) -> Any:
        try:
            return ctx[key]
        except Exception:
            return None

    return WorkspaceContext(_get("currentWorkspaceName"), _get("currentWorkspaceId"), _get("isForPipeline"),
                            _get("currentNotebookName"), None)


def decide_adp_push(workspace_name: Any) -> tuple:
    """Règle pure. Renvoie (condition satisfaite, motif)."""
    if workspace_name is None:
        return False, REASON_UNRESOLVED
    if not isinstance(workspace_name, str):
        return False, REASON_INVALID_TYPE
    if workspace_name.strip() == "":
        return False, REASON_EMPTY
    if PRD_MARKER in workspace_name:
        return True, REASON_CONTAINS_PRD
    return False, REASON_WITHOUT_PRD


def outcome_of(reason: str) -> str:
    if reason == REASON_CONTAINS_PRD:
        return ENV_SATISFIED
    if reason == REASON_WITHOUT_PRD:
        return DENIED_NON_PRD
    return BLOCKED_UNRESOLVED  # non résolu / vide / type invalide : réévaluable


def evaluate_gate(context_provider: Callable[[], Mapping[str, Any]],
                  clock: Callable[[], dt.datetime] = _utc_now) -> GateDecision:
    ws = resolve_current_workspace(context_provider)
    met, reason = decide_adp_push(ws.name)
    return GateDecision(
        decision_id=str(uuid.uuid4()), outcome=outcome_of(reason), env_condition_met=met, reason_code=reason,
        workspace_name=ws.name if isinstance(ws.name, str) else None,
        workspace_id=ws.workspace_id if isinstance(ws.workspace_id, str) else None,
        is_for_pipeline=ws.is_for_pipeline if isinstance(ws.is_for_pipeline, bool) else None,
        notebook_name=ws.notebook_name if isinstance(ws.notebook_name, str) else None,
        rule_version=RULE_VERSION, evaluated_at_utc=clock())
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.adp_transport =============
register_hris_module('hris.adp_transport', '453008e434ed5d3259b35d281f9cf8b2e30ee31c53eeb00b6677176a15b2fa73', r'''"""hris.adp_transport — transport HTTP de publication ADP, classification des résultats, instrumentation.

Principes (mandat § 8) :
  * aucun retry automatique (y compris bibliothèque), aucune redirection suivie, aucun renvoi
    implicite après 401, timeouts explicites, aucun nouvel envoi après un résultat incertain ;
  * la non-transmission n'est PROUVÉE que si l'échec survient avant l'émission du premier octet de
    la requête (phase CONNECT : DNS, TCP, TLS/mTLS) — jamais d'après le seul nom d'une exception ;
  * une réponse 2xx ne suffit pas : l'acceptation suit le contrat de l'endpoint.

Le contrat final de publication est ABSENT : les listes de statuts ci-dessous sont PROVISOIRES,
utilisables pour les simulations uniquement (docs/CONTRAT_ADP_POINTS_D_USAGE.md). Le transport réel
refuse de se construire tant que ADP_ACTIVATION_APPROVED est faux (constante de code, aucun paramètre).
"""
from __future__ import annotations

import dataclasses
import hashlib
import http.client
import json
import os
import shutil
import socket
import ssl
import tempfile
import threading
import time
from typing import Any, Callable, Dict, List, Optional, Tuple
from urllib.parse import quote, urlencode, urlparse

from hris.common import HrisError, SecretValue

# ---------------------------------------------------------------------------
# Verrous d'activation (constantes de CODE : ni paramètre, ni variable, ni simulateur)
# ---------------------------------------------------------------------------
ADP_BASE_URL = "https://accounts.eu.adp.com/"  # cible unique ADP PRD (A05 § 4.1)
ADP_ACTIVATION_APPROVED = False      # CONTRAT-ADP:C00 contrat final + D-ADP-01/02, D-ENV-02, D-SEC-01, D-CMP-07, D-NET-01
ADP_READ_POLICY_ESTABLISHED = False  # CONTRAT-ADP:C14 politique des lectures ADP réelles non établie
ADP_HOST_PATTERNS = ("adp.com", "geo.api.gouv.fr")  # hôtes surveillés par la sentinelle

NOT_TRANSMITTED = "NOT_TRANSMITTED"  # non-transmission prouvée
ACCEPTED = "ACCEPTED"                # acceptation selon le contrat de l'endpoint
REJECTED = "REJECTED"                # rejet confirmé
UNCERTAIN = "UNCERTAIN"              # résultat incertain : aucun nouvel envoi automatique

# CONTRAT-ADP:C02 statuts 2xx acceptés et forme de corps attendue (PROVISOIRE : 200/201/202 + JSON,
# d'après la seule gestion explicite de l'existant, change_pay_distribution L4190)
PROVISIONAL_ACCEPTED_2XX = frozenset({200, 201, 202})
# CONTRAT-ADP:C03 statuts 4xx valant rejet confirmé (PROVISOIRE)
PROVISIONAL_REJECT_4XX = frozenset({400, 403, 404, 405, 409, 412, 415, 422})


@dataclasses.dataclass
class TransportResult:
    outcome: str
    reason: str
    phase: str
    http_status: Optional[int] = None
    http_reason: Optional[str] = None
    error_class: Optional[str] = None
    response_text: Optional[str] = dataclasses.field(default=None, repr=False)  # donnée personnelle possible
    body_digest: Optional[str] = None

    def __repr__(self) -> str:
        return (f"TransportResult({self.outcome}, {self.reason}, phase={self.phase}, "
                f"http={self.http_status}, error={self.error_class})")


def classify_response(status: int, body: str, content_type: str = "") -> Tuple[str, str]:
    """Classification d'une réponse HTTP REÇUE à une requête de publication (PROVISOIRE)."""
    if 200 <= status < 300:
        if status not in PROVISIONAL_ACCEPTED_2XX:
            return UNCERTAIN, "2XX_UNEXPECTED_STATUS"          # CONTRAT-ADP:C02
        try:
            parsed = json.loads(body) if body and body.strip() else None
        except ValueError:
            return UNCERTAIN, "2XX_BODY_NOT_JSON"               # CONTRAT-ADP:C02
        if parsed is None and status != 202:
            return UNCERTAIN, "2XX_EMPTY_BODY"                  # CONTRAT-ADP:C02
        return ACCEPTED, f"HTTP_{status}_ACCEPTED_PROVISIONAL"
    if 300 <= status < 400:
        return UNCERTAIN, "REDIRECT_NOT_FOLLOWED"               # CONTRAT-ADP:C04
    if status == 401:
        return REJECTED, "AUTH_REJECTED_NO_RESEND"              # CONTRAT-ADP:C05
    if status == 429:
        return REJECTED, "RATE_LIMITED_NO_RETRY"                # CONTRAT-ADP:C06
    if status in PROVISIONAL_REJECT_4XX:
        return REJECTED, f"HTTP_{status}_REJECTED_PROVISIONAL"  # CONTRAT-ADP:C03
    if 400 <= status < 500:
        return UNCERTAIN, f"HTTP_{status}_UNCLASSIFIED"
    return UNCERTAIN, f"HTTP_{status}_SERVER_ERROR"            # CONTRAT-ADP:C04 (5xx : traitement partiel possible)


def digest(text: Optional[str]) -> Optional[str]:
    return None if text is None else hashlib.sha256(text.encode("utf-8", "replace")).hexdigest()


# Caractères conservés tels quels par requests (requests.utils.requote_uri) : réservés RFC 3986 et « % ».
_REQUESTS_SAFE = "!#$%&'()*+,/:;=?@[]~"


def request_target(base_path: str, path: str) -> str:
    """Cible HTTP envoyée comme le faisait l'existant (requests) : les caractères interdits, dont les espaces des
    filtres OData (« idValue eq '…' »), sont encodés ; les réservés sont conservés. http.client refuse un chemin
    contenant un espace (InvalidURL) : sans cet encodage, toute lecture filtrée échouait."""
    return quote(base_path + path, safe=_REQUESTS_SAFE)


def mtls_context(cert_pem: str, key_pem: str) -> ssl.SSLContext:
    """Contexte TLS client chargé depuis les secrets du coffre. ssl.load_cert_chain n'accepte que des chemins : le
    certificat et la clé sont écrits dans un dossier privé (0700, fichiers 0600) le temps du chargement, puis effacés
    aussitôt — l'existant les laissait dans le dossier temporaire partagé."""
    context = ssl.create_default_context()
    folder = tempfile.mkdtemp(prefix="hris-adp-")
    try:
        cert_path, key_path = os.path.join(folder, "client.cer"), os.path.join(folder, "client.key")
        for path, content in ((cert_path, cert_pem), (key_path, key_pem)):
            handle = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            with os.fdopen(handle, "w") as stream:
                stream.write(content)
        context.load_cert_chain(certfile=cert_path, keyfile=key_path)
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    return context


class OAuthToken:
    """Jeton ADP client_credentials, comme login() de l'existant (L407-452) : POST auth/oauth/v2/token en mTLS, corps
    x-www-form-urlencoded. Renouvelé une minute avant expiration ; jamais affiché ni journalisé (l'existant l'imprimait)."""

    def __init__(self, ssl_context, client_id: SecretValue, client_secret: SecretValue, timeout_seconds: float,
                 base_url: str = ADP_BASE_URL, clock: Callable[[], float] = time.monotonic,
                 connection_factory: Optional[Callable[[], Any]] = None):
        parsed = urlparse(base_url)
        self.host, self.base_path = parsed.hostname, parsed.path or "/"
        self.timeout = float(timeout_seconds)
        self._ctx, self._client_id, self._client_secret, self._clock = ssl_context, client_id, client_secret, clock
        self._connection_factory = connection_factory or (
            lambda: http.client.HTTPSConnection(self.host, 443, context=self._ctx, timeout=self.timeout))
        self._token: Optional[SecretValue] = None
        self._expires_at = 0.0

    def __repr__(self) -> str:
        return "OAuthToken(***)"

    def __call__(self) -> str:
        if self._token is None or self._clock() >= self._expires_at:
            self._fetch()
        return self._token.reveal()

    def _fetch(self) -> None:
        body = urlencode({"client_id": self._client_id.reveal(), "client_secret": self._client_secret.reveal(),
                          "grant_type": "client_credentials"})
        conn = self._connection_factory()
        try:
            conn.request("POST", request_target(self.base_path, "auth/oauth/v2/token"), body=body,
                         headers={"Content-Type": "application/x-www-form-urlencoded", "Accept": "application/json"})
            resp = conn.getresponse()
            text = resp.read(AdpHttpTransport.MAX_BODY).decode("utf-8", "replace")
        finally:
            conn.close()
        if not 200 <= resp.status < 300:
            raise HrisError("ADP_AUTH_FAILED", f"statut {resp.status}")
        try:
            data = json.loads(text)
        except ValueError:
            raise HrisError("ADP_AUTH_FAILED", "réponse non JSON") from None
        if not data.get("access_token"):
            raise HrisError("ADP_AUTH_FAILED", "jeton absent")
        self._token = SecretValue(data["access_token"])
        self._expires_at = self._clock() + max(60.0, float(data.get("expires_in") or 3600) - 60.0)


# ---------------------------------------------------------------------------
# Transport réel (construit seulement si toutes les autorisations sont réunies)
# ---------------------------------------------------------------------------
class AdpHttpTransport:
    """POST/GET vers ADP en http.client : pas de pool à reprise, pas de redirection, phases explicites.
    mTLS par le contexte fourni (mtls_context) ; jeton porteur fourni par un callable (jamais journalisé)."""

    MAX_BODY = 1_000_000

    def __init__(self, ssl_context, bearer: Callable[[], str], timeout_seconds: float,
                 base_url: str = ADP_BASE_URL, _allow_construction: bool = False):
        if not (ADP_ACTIVATION_APPROVED and _allow_construction):
            raise HrisError("ADP_ACTIVATION_NOT_APPROVED", "transport ADP réel non autorisé dans cette version")
        parsed = urlparse(base_url)
        self.host = parsed.hostname
        self.base_path = parsed.path or "/"
        self.timeout = float(timeout_seconds)
        self._bearer = bearer
        self._ctx = ssl_context

    def __repr__(self) -> str:
        return "AdpHttpTransport(<prd>)"

    def _connection(self):
        return http.client.HTTPSConnection(self.host, 443, context=self._ctx, timeout=self.timeout)

    def post(self, path: str, body: bytes) -> TransportResult:
        conn = self._connection()
        try:
            try:
                conn.connect()  # DNS + TCP + TLS : aucun octet de la requête HTTP n'est émis
            except Exception as exc:
                return TransportResult(NOT_TRANSMITTED, "CONNECT_FAILED", "CONNECT", error_class=type(exc).__name__)
            headers = {"Content-Type": "application/json", "Accept": "application/json",
                       "Authorization": "Bearer " + self._bearer(), "Content-Length": str(len(body))}
            try:
                conn.request("POST", request_target(self.base_path, path), body=body, headers=headers)
            except Exception as exc:
                return TransportResult(UNCERTAIN, "SEND_INTERRUPTED", "SEND", error_class=type(exc).__name__)
            try:
                resp = conn.getresponse()
                raw = resp.read(self.MAX_BODY + 1)
            except socket.timeout as exc:
                return TransportResult(UNCERTAIN, "TIMEOUT_AFTER_SEND", "RECEIVE", error_class=type(exc).__name__)
            except Exception as exc:
                return TransportResult(UNCERTAIN, "RECEIVE_FAILED", "RECEIVE", error_class=type(exc).__name__)
            text = raw[: self.MAX_BODY].decode("utf-8", "replace")
            outcome, reason = classify_response(resp.status, text, resp.getheader("Content-Type", ""))
            return TransportResult(outcome, reason, "DONE", http_status=resp.status, http_reason=resp.reason,
                                   response_text=text, body_digest=digest(text))
        finally:
            try:
                conn.close()
            except Exception:
                pass

    def get(self, path: str) -> Tuple[Optional[int], Optional[str]]:
        """Lecture ADP (politique distincte, CONTRAT-ADP:C14). Pas de redirection suivie."""
        if not ADP_READ_POLICY_ESTABLISHED:
            raise HrisError("ADP_READ_POLICY_NOT_ESTABLISHED")
        conn = self._connection()
        try:
            conn.request("GET", request_target(self.base_path, path), headers={
                "Accept": "application/json", "Authorization": "Bearer " + self._bearer()})
            resp = conn.getresponse()
            return resp.status, resp.read(self.MAX_BODY).decode("utf-8", "replace")
        finally:
            conn.close()


# ---------------------------------------------------------------------------
# Instrumentation : comptage applicatif et sentinelle réseau
# ---------------------------------------------------------------------------
class RecordingTransport:
    """Enveloppe d'un transport (réel ou fictif) qui enregistre chaque appel de publication (POST)
    séparément des lectures (GET). Aucune donnée de requête n'est conservée hors mémoire du test."""

    def __init__(self, inner):
        self.inner = inner
        self.post_calls: List[str] = []
        self.get_calls: List[str] = []

    def post(self, path: str, body: bytes) -> TransportResult:
        self.post_calls.append(path)
        return self.inner.post(path, body)

    def get(self, path: str):
        self.get_calls.append(path)
        return self.inner.get(path)


class NetworkSentinel:
    """Intercepte, pour la durée d'un bloc, toute tentative de résolution DNS ou de connexion vers un
    hôte ADP (ou geo.api.gouv.fr) : la tentative est ENREGISTRÉE puis BLOQUÉE avant tout trafic réseau.
    Couvre socket.getaddrinfo, socket.create_connection, socket.socket.connect et
    http.client.HTTPConnection.connect (donc requests/urllib3 et http.client).
    Limite : ne voit pas un trafic émis hors du processus Python du driver (exécuteurs Spark, autres
    processus) ; la publication ne s'exécute que sur le driver."""

    def __init__(self, patterns=ADP_HOST_PATTERNS):
        self.patterns = tuple(p.lower() for p in patterns)
        self.attempts: List[Dict[str, str]] = []
        self._lock = threading.Lock()
        self._saved: Dict[str, Any] = {}

    def _watched(self, host: Any) -> bool:
        if not isinstance(host, (str, bytes)):
            return False
        h = host.decode() if isinstance(host, bytes) else host
        h = h.lower().rstrip(".")
        return any(h == p or h.endswith("." + p) or h.endswith(p) for p in self.patterns)

    def _record(self, host: str, via: str) -> None:
        with self._lock:
            self.attempts.append({"host": str(host), "via": via})
        raise ConnectionRefusedError(f"HRIS network sentinel: connexion vers un hôte ADP bloquée ({via})")

    def __enter__(self):
        sentinel = self
        self._saved = {"getaddrinfo": socket.getaddrinfo, "create_connection": socket.create_connection,
                       "socket_connect": socket.socket.connect, "http_connect": http.client.HTTPConnection.connect}
        orig_gai, orig_cc = self._saved["getaddrinfo"], self._saved["create_connection"]
        orig_sc, orig_hc = self._saved["socket_connect"], self._saved["http_connect"]

        def getaddrinfo(host, *a, **k):
            if sentinel._watched(host):
                sentinel._record(host, "getaddrinfo")
            return orig_gai(host, *a, **k)

        def create_connection(address, *a, **k):
            if sentinel._watched(address[0]):
                sentinel._record(address[0], "create_connection")
            return orig_cc(address, *a, **k)

        def socket_connect(sock, address, *a, **k):
            if isinstance(address, tuple) and address and sentinel._watched(address[0]):
                sentinel._record(address[0], "socket.connect")
            return orig_sc(sock, address, *a, **k)

        def http_connect(conn, *a, **k):
            if sentinel._watched(getattr(conn, "host", None)):
                sentinel._record(conn.host, "http.client.connect")
            return orig_hc(conn, *a, **k)

        socket.getaddrinfo = getaddrinfo
        socket.create_connection = create_connection
        socket.socket.connect = socket_connect
        http.client.HTTPConnection.connect = http_connect
        return self

    def __exit__(self, *exc):
        socket.getaddrinfo = self._saved["getaddrinfo"]
        socket.create_connection = self._saved["create_connection"]
        socket.socket.connect = self._saved["socket_connect"]
        http.client.HTTPConnection.connect = self._saved["http_connect"]
        return False
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.adp_ops =============
register_hris_module('hris.adp_ops', '58585e7444141ad671315c630cc21636577579a49f7c1203f9808ed76779a78e', r'''"""hris.adp_ops — opérations de publication ADP : identité, intention durable, exécution, réconciliation.

PROPOSITION soumise à validation : le contrat final de publication et de reprise est absent.
Chaque point qui en dépend porte un marqueur `CONTRAT-ADP:Cxx` recensé dans
docs/CONTRAT_ADP_POINTS_D_USAGE.md (cohérence vérifiée par tests/unit/test_contract_markers.py).

Modèle :
  * une OPÉRATION = un événement métier ADP (POST) pour un salarié, une variante (P1..P9, P8 détaillé)
    et une valeur cible ; ses TENTATIVES sont distinctes (pub.adp_attempt) ;
  * identité stable : op_id = H(version, run source, salarié, variante, empreinte de la cible) —
    le même événement rejoué garde le même op_id ; une nouvelle modification légitime (autre run,
    y compris un retour à une valeur déjà rencontrée) a un op_id différent ; aucune déduplication
    à vie sur (salarié, code d'opération) ;
  * l'intention (INTENT_RECORDED) est persistée AVANT tout transport ; si cette persistance échoue,
    rien n'est envoyé ; une intention sans issue persistée est traitée comme INCERTAINE ;
  * aucun nouvel envoi automatique d'une opération INCERTAINE : réconciliation ou décision manuelle.
"""
from __future__ import annotations

import dataclasses
import datetime as dt
import uuid
from typing import Any, Callable, Dict, Iterable, List, Optional, Protocol, Sequence

from hris.adp_transport import ACCEPTED, NOT_TRANSMITTED, REJECTED, UNCERTAIN, TransportResult
from hris.common import HrisError, canonical_json, sha256_hex

OP_IDENTITY_VERSION = "op-v1"  # CONTRAT-ADP:C01

# Statuts d'opération (PROPOSÉS) — CONTRAT-ADP:C11
OP_INTENT_RECORDED = "INTENT_RECORDED"
OP_ACCEPTED = "ACCEPTED"              # = « SENT » : acceptation selon le contrat, pas application
OP_REJECTED = "REJECTED"
OP_NOT_SENT = "NOT_SENT"              # non-transmission prouvée (ou renvoi autorisé par décision)
OP_UNCERTAIN = "UNCERTAIN"
OP_BLOCKED_DEPENDENCY = "BLOCKED_DEPENDENCY"
OP_SUPERSEDED = "SUPERSEDED"
OP_APPLIED_OBSERVED = "APPLIED_OBSERVED"
OP_ABANDONED = "ABANDONED"
OP_DONE = {OP_ACCEPTED, OP_APPLIED_OBSERVED}
OP_OPEN_UNCERTAIN = {OP_INTENT_RECORDED, OP_UNCERTAIN}
OP_RETRYABLE = {OP_REJECTED, OP_NOT_SENT, OP_BLOCKED_DEPENDENCY}

# Statuts d'outbox (A07 § 6, inchangés)
OB_PENDING, OB_IN_PROGRESS, OB_SENT = "PENDING", "IN_PROGRESS", "SENT"
OB_NO_CHANGE, OB_FAILED, OB_NOT_PUSHED, OB_SUPERSEDED = "NO_CHANGE", "FAILED", "NOT_PUSHED_NON_PRD", "SUPERSEDED"

_TRANSPORT_TO_OP = {ACCEPTED: OP_ACCEPTED, REJECTED: OP_REJECTED, NOT_TRANSMITTED: OP_NOT_SENT,
                    UNCERTAIN: OP_UNCERTAIN}

# Décisions manuelles (PROPOSÉES) — CONTRAT-ADP:C10
DEC_AUTHORIZE_RESEND = "AUTHORIZE_RESEND"   # incertaine -> NOT_SENT (nouvelle tentative possible en reprise)
DEC_MARK_APPLIED = "MARK_APPLIED"           # incertaine -> APPLIED_OBSERVED (constat humain)
DEC_ABANDON = "ABANDON"                     # -> ABANDONED (aucun envoi)
DECISION_EFFECT = {DEC_AUTHORIZE_RESEND: OP_NOT_SENT, DEC_MARK_APPLIED: OP_APPLIED_OBSERVED,
                   DEC_ABANDON: OP_ABANDONED}


def _utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


@dataclasses.dataclass
class PlannedOp:
    """Opération produite par la logique ADP portée (hris.adp_port). payload, target et log
    contiennent des données personnelles : uniquement en mémoire, jamais persistés en clair."""
    op_code: str                      # P1 .. P9
    op_variant: str                   # ex. P8_SALARY, P8_CONTRACT_A (CONTRAT-ADP:C12)
    endpoint: str                     # chemin relatif, sans identifiant
    payload: dict = dataclasses.field(repr=False)
    target: dict = dataclasses.field(repr=False)   # valeur métier cible (sans partie volatile)
    log: dict = dataclasses.field(repr=False)      # colonnes du log de l'existant (mode, field, nom...)
    depends_on: Sequence[str] = ()    # variantes du même salarié qui doivent être acceptées avant
    read_back: Optional[Callable[[Any], Optional[dict]]] = dataclasses.field(default=None, repr=False)


def target_hash(target: dict) -> str:
    return sha256_hex(canonical_json(target))


def op_identity(source_run_id: str, employee_key: str, op_variant: str, tgt_hash: str) -> str:
    """CONTRAT-ADP:C01 — identité proposée de l'événement (à valider avant activation réelle)."""
    return sha256_hex("|".join([OP_IDENTITY_VERSION, source_run_id, employee_key, op_variant, tgt_hash]))


# ---------------------------------------------------------------------------
# Dépôts (implémentés en Spark dans le notebook, en mémoire dans les tests)
# ---------------------------------------------------------------------------
class OperationRepository(Protocol):
    def get(self, op_id: str) -> Optional[dict]: ...
    def open_uncertain_for_employee(self, employee_key: str) -> List[dict]: ...
    def for_item(self, run_id: str, employee_key: str) -> List[dict]: ...
    def upsert(self, op: dict) -> None: ...            # durable avant retour, sinon exception


class AttemptRepository(Protocol):
    def append(self, attempt: dict) -> None: ...       # durable avant retour, sinon exception


class PublicationLog(Protocol):
    def append(self, rows: List[dict]) -> None: ...


@dataclasses.dataclass
class OpResult:
    op_id: str
    variant: str
    status: str
    detail: str
    transport: Optional[TransportResult] = None


class OperationExecutor:
    """Exécute les opérations d'UN salarié, dans l'ordre, avec intention durable avant transport."""

    def __init__(self, *, ops: OperationRepository, attempts: AttemptRepository, transport,
                 source_run_id: str, publish_run_id: str, employee_key: str, gate_decision_id: str,
                 allow_new_attempt: bool, max_attempts: int, clock: Callable[[], dt.datetime] = _utc_now):
        self.ops, self.attempts, self.transport = ops, attempts, transport
        self.source_run_id, self.publish_run_id = source_run_id, publish_run_id
        self.employee_key, self.gate_decision_id = employee_key, gate_decision_id
        self.allow_new_attempt, self.max_attempts = allow_new_attempt, max_attempts
        self.clock = clock
        self.results: List[OpResult] = []
        self.status_by_variant: Dict[str, str] = {}
        self.seq = 0
        self.halt_reason: Optional[str] = None
        self.log_rows: List[dict] = []   # lignes du log fonctionnel (format de l'existant), en mémoire

    def log(self, row: dict) -> None:
        """Ligne de log.adp_publication_log produite par la logique portée (une par POST, rejet NIR
        ou information), conservée même si le traitement du salarié s'interrompt ensuite."""
        self.log_rows.append(row)

    def submit(self, op: PlannedOp) -> OpResult:
        self.seq += 1
        thash = target_hash(op.target)
        op_id = op_identity(self.source_run_id, self.employee_key, op.op_variant, thash)
        existing = self.ops.get(op_id)
        if self.halt_reason:
            return self._record(op, op_id, existing, OP_NOT_SENT if not existing else existing["status"],
                                f"HALTED_{self.halt_reason}", persist=False)
        if existing:
            st = existing["status"]
            if st in OP_DONE or st in (OP_SUPERSEDED, OP_ABANDONED):
                return self._record(op, op_id, existing, st, "ALREADY_TERMINAL", persist=False)
            if st in OP_OPEN_UNCERTAIN:
                return self._record(op, op_id, existing, OP_UNCERTAIN, "UNCERTAIN_AWAITING_RECONCILIATION",
                                    persist=False)  # aucun nouvel envoi automatique
            if st in OP_RETRYABLE and not (self.allow_new_attempt or existing.get("manual_decision_id")):
                return self._record(op, op_id, existing, st, "RETRY_REQUIRES_EXPLICIT_REPRISE", persist=False)
            if int(existing.get("attempt_count") or 0) >= self.max_attempts:
                return self._record(op, op_id, existing, st, "MAX_ATTEMPTS_REACHED", persist=False)
        # Dépendances (CONTRAT-ADP:C08)
        for dep in op.depends_on:
            dep_status = self.status_by_variant.get(dep)
            if dep_status is not None and dep_status not in OP_DONE:
                row = self._op_row(op, op_id, thash, existing, OP_BLOCKED_DEPENDENCY, None)
                self.ops.upsert(row)
                return self._record(op, op_id, row, OP_BLOCKED_DEPENDENCY, f"DEPENDS_ON_{dep}", persist=False)
        # 1. Intention durable AVANT transport : une exception ici interrompt tout (rien n'est envoyé)
        attempt_no = int((existing or {}).get("attempt_count") or 0) + 1
        intent = self._op_row(op, op_id, thash, existing, OP_INTENT_RECORDED, None, attempt_no=attempt_no)
        self.ops.upsert(intent)
        attempt_id = str(uuid.uuid4())
        started = self.clock()
        self.attempts.append({"attempt_id": attempt_id, "op_id": op_id, "run_id": self.publish_run_id,
                              "attempt_no": attempt_no, "phase": "INTENT", "outcome_class": None,
                              "http_status": None, "error_class": None, "response_digest": None,
                              "started_at_utc": started, "ended_at_utc": None})
        # 2. Transport (aucun retry, aucune redirection, aucun renvoi)
        try:
            result = self.transport.post(op.endpoint, canonical_payload(op.payload))
        except Exception as exc:  # exception non classée par le transport : on ne conclut pas
            result = TransportResult(UNCERTAIN, "TRANSPORT_EXCEPTION", "UNKNOWN", error_class=type(exc).__name__)
        status = _TRANSPORT_TO_OP[result.outcome]
        # 3. Issue persistée ; si cette écriture échoue, l'intention reste et vaudra « incertain »
        self.attempts.append({"attempt_id": attempt_id, "op_id": op_id, "run_id": self.publish_run_id,
                              "attempt_no": attempt_no, "phase": result.phase, "outcome_class": result.outcome,
                              "http_status": result.http_status, "error_class": result.error_class,
                              "response_digest": result.body_digest, "started_at_utc": started,
                              "ended_at_utc": self.clock()})
        final = self._op_row(op, op_id, thash, intent, status, result, attempt_no=attempt_no)
        self.ops.upsert(final)
        if result.reason == "AUTH_REJECTED_NO_RESEND":
            self.halt_reason = "AUTH_REJECTED"   # CONTRAT-ADP:C05 : jeton refusé, la suite n'est pas envoyée
        return self._record(op, op_id, final, status, result.reason, persist=False, transport=result)

    def _op_row(self, op, op_id, thash, existing, status, result, attempt_no=None):
        now = self.clock()
        base = dict(existing or {})
        base.update({"op_id": op_id, "run_id": self.source_run_id, "employee_key": self.employee_key,
                     "seq_no": base.get("seq_no") or self.seq, "op_code": op.op_code, "op_variant": op.op_variant,
                     "depends_on": list(op.depends_on), "target_hash": thash,
                     "payload_hash": sha256_hex(canonical_payload(op.payload).decode("utf-8")),
                     "status": status, "gate_decision_id": self.gate_decision_id,
                     "created_at_utc": base.get("created_at_utc") or now, "updated_at_utc": now})
        if attempt_no is not None:
            base["attempt_count"] = attempt_no
        if result is not None:
            base["last_outcome_class"] = result.outcome
            base["last_http_status"] = result.http_status
        return base

    def _record(self, op, op_id, row, status, detail, persist, transport=None) -> OpResult:
        self.status_by_variant[op.op_variant] = status
        res = OpResult(op_id, op.op_variant, status, detail, transport)
        self.results.append(res)
        return res


def canonical_payload(payload: dict) -> bytes:
    """Corps JSON envoyé (même sérialisation que l'existant : json.dumps, ASCII échappé).
    CONTRAT-ADP:C15 aucune clé d'idempotence n'est transmise à ADP (support inconnu) : pas de garantie
    « exactement une fois » de bout en bout ; seule la non-réémission automatique est garantie ici."""
    import json
    return json.dumps(payload).encode("utf-8")


def item_status(results: Sequence[OpResult]) -> str:
    """Statut d'outbox (A07) déduit des opérations d'un salarié (D-ADP-07, À CONFIRMER).
    CONTRAT-ADP:C16 SENT = au moins une opération ACCEPTÉE selon le contrat (pas une application observée)."""
    statuses = {r.status for r in results}
    if statuses & OP_OPEN_UNCERTAIN:
        return OB_IN_PROGRESS            # incertain : jamais resélectionné automatiquement
    if statuses & OP_RETRYABLE:
        return OB_FAILED
    if statuses & OP_DONE:
        return OB_SENT
    return OB_NO_CHANGE


# ---------------------------------------------------------------------------
# Obsolescence (SUPERSEDED) — CONTRAT-ADP:C09
# ---------------------------------------------------------------------------
def supersession_candidates(item: dict, newer_items: Iterable[dict], item_ops: Iterable[dict]) -> bool:
    """Un élément plus ancien peut devenir SUPERSEDED si et seulement si :
      * un élément PLUS RÉCENT pour le même salarié porte l'état complet le plus récent et n'est pas
        lui-même incertain (PENDING, FAILED, SENT ou NO_CHANGE) ;
      * aucune opération de l'élément ancien n'est INCERTAINE ou en cours (jamais d'obsolescence
        automatique d'un résultat incertain) ;
    Les opérations déjà ACCEPTED de l'élément ancien restent ACCEPTED (déjà faites)."""
    if item.get("status") not in (OB_PENDING, OB_FAILED):
        return False
    if any(o.get("status") in OP_OPEN_UNCERTAIN for o in item_ops):
        return False
    return any(n.get("status") in (OB_PENDING, OB_FAILED, OB_SENT, OB_NO_CHANGE) for n in newer_items)


# ---------------------------------------------------------------------------
# Réconciliation — CONTRAT-ADP:C07
# ---------------------------------------------------------------------------
RECON_MATCH = "MATCH_ALL_TARGET_FIELDS"
RECON_PARTIAL = "PARTIAL_READ"
RECON_DIFFERENT = "DIFFERENT_FROM_TARGET"
RECON_READ_FAILED = "READ_FAILED"


def reconcile(op_row: dict, planned: PlannedOp, reader) -> dict:
    """Compare l'état lu chez ADP à la cible COMPLÈTE de l'opération.
    * tous les champs cibles lus et égaux -> APPLIED_OBSERVED (application observée, pas garantie) ;
    * lecture partielle -> aucune conclusion (une lecture partielle ne prouve pas tous les champs) ;
    * état différent -> aucune conclusion de non-traitement (un GET différent ne prouve rien) :
      l'opération reste INCERTAINE et requiert une décision manuelle tracée."""
    if planned.read_back is None:
        return {"result": RECON_PARTIAL, "conclusion": "NO_READBACK_DEFINED", "fields": []}
    try:
        observed = planned.read_back(reader)
    except Exception as exc:
        return {"result": RECON_READ_FAILED, "conclusion": "UNCHANGED", "fields": [], "error": type(exc).__name__}
    if observed is None:
        return {"result": RECON_READ_FAILED, "conclusion": "UNCHANGED", "fields": []}
    fields = sorted(planned.target.keys())
    missing = [f for f in fields if f not in observed]
    if missing:
        return {"result": RECON_PARTIAL, "conclusion": "UNCHANGED", "fields": fields}
    if all(observed[f] == planned.target[f] for f in fields):
        return {"result": RECON_MATCH, "conclusion": OP_APPLIED_OBSERVED, "fields": fields}
    return {"result": RECON_DIFFERENT, "conclusion": "MANUAL_DECISION_REQUIRED", "fields": fields}


def apply_manual_decision(op_row: dict, decision: str, decided_by: str, basis: str, reference: str = "",
                          clock: Callable[[], dt.datetime] = _utc_now) -> tuple:
    """Décision humaine tracée (auteur, fondement). Renvoie (op mise à jour, ligne de décision)."""
    if decision not in DECISION_EFFECT:
        raise HrisError("DECISION_INVALID", decision)
    if not decided_by or not basis or not basis.strip():
        raise HrisError("DECISION_INCOMPLETE", "auteur et fondement obligatoires")
    if op_row["status"] in OP_DONE and decision != DEC_ABANDON:
        raise HrisError("DECISION_NOT_APPLICABLE", "opération déjà acceptée ou observée")
    decision_id = str(uuid.uuid4())
    now = clock()
    updated = dict(op_row, status=DECISION_EFFECT[decision], manual_decision_id=decision_id, updated_at_utc=now)
    record = {"decision_id": decision_id, "op_id": op_row["op_id"], "run_id": op_row["run_id"],
              "employee_key": op_row["employee_key"], "decision": decision, "decided_by": decided_by,
              "decided_at_utc": now, "basis": basis, "reference": reference}
    return updated, record
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.adp_port =============
register_hris_module('hris.adp_port', 'db8a126b982c68ac552b5a00bcc288fadf42508e44a9e0f1845ca081fb6a696d', r'''"""hris.adp_port — portage de la logique métier de la Function ADP (adp-function-app/function_app.py,
PRJ @ 5541490, main() L4234-5031 et fonctions appelées), pour UN salarié.

Parité par défaut (A05 § 3) : mêmes lectures (17 G3 + G1/G2/G4/G5/G6), mêmes comparaisons brutes,
mêmes payloads, mêmes valeurs de log (mode, field, nom, log_message), anomalies B1 à B27 conservées
sauf améliorations techniques actées :
  * jeton jamais journalisé ; aucun print de données personnelles (B19, B22) ;
  * session INSEE séparée, sans en-tête ADP (B22) ;
  * chaque POST passe par OperationExecutor : intention durable, aucun retry, aucun renvoi ;
  * une exception sur un salarié n'arrête pas le lot (D-ADP-01 : isolement par défaut, BLOQUANTE
    pour l'activation) — l'existant arrêtait tout (B15).
Points À ARBITRER conservés en parité (A08) : B1/B2 variantes, B3 garde salaire, B7 existence sur
erreur HTTP (BLOQUANTE D-ADP-02), B8 imputation 100, B9 FTE, B10, B11, B12, B13, B21, B23, R-ADP-012.
"""
from __future__ import annotations

import datetime
import decimal
import difflib
import json
import re
import unicodedata
from typing import Any, Callable, Dict, Optional, Tuple

from hris.adp_ops import OP_DONE, OperationExecutor, PlannedOp

EV_HIRE = "events/hr/v1/worker.hire"
EV_PERSONAL_EMAIL = "events/hr/v1/worker.personal-communication.email.change"
EV_BUSINESS_EMAIL = "events/hr/v1/worker.business-communication.email.change"
EV_BUSINESS_MOBILE = "events/hr/v1/worker.business-communication.mobile.change"
EV_PERSONAL_MOBILE = "events/hr/v1/worker.personal-communication.mobile.change"
EV_MARITAL = "events/hr/v1/worker.marital-status.change"
EV_LEGAL_ADDRESS = "events/hr/v1/worker.legal-address.change"
EV_WA_MODIFY = "events/hr/v1/worker.work-assignment.modify"
EV_PAY_DISTRIBUTION = "events/payroll/v1/worker.pay-distribution.change"

HIRE_IMPUTATION_PERCENTAGE = 100  # B8 : constante de test lue par fermeture (L367), entier


class ReadResponse:
    """Réponse de lecture (statut, corps JSON) au format utilisé par la logique portée."""

    def __init__(self, status_code: int, text: str = "", reason: str = ""):
        self.status_code = status_code
        self.text = text if isinstance(text, str) else json.dumps(text)
        self.reason = reason

    def json(self):
        return json.loads(self.text) if self.text else None


class TransportReader:
    """Lectures ADP réelles (G1..G5) via le transport, et INSEE (G6) via une session HTTP SÉPARÉE,
    sans en-tête ADP (B22 corrigé). Politique des lectures réelles non établie (CONTRAT-ADP:C14) :
    le transport refuse toute lecture tant qu'elle ne l'est pas."""

    def __init__(self, transport, insee_timeout: float = 5.0):
        self.transport = transport
        self.insee_timeout = insee_timeout

    def get(self, path):
        status, text = self.transport.get(path)
        return ReadResponse(status or 0, text or "")

    def insee(self, postal_code):
        import requests
        with requests.Session() as s:
            r = s.get("https://geo.api.gouv.fr/communes", params={"codePostal": postal_code, "fields": "nom,code"},
                      timeout=self.insee_timeout, allow_redirects=False)
            return ReadResponse(r.status_code, r.text)


# ---------------------------------------------------------------------------
# Fonctions de transformation de l'existant (L204-329, L454-458)
# ---------------------------------------------------------------------------
def normalize_string(input_string):
    normalized = unicodedata.normalize('NFD', input_string)
    without_accents = ''.join(char for char in normalized if unicodedata.category(char) != 'Mn')
    replaced_chars = without_accents.replace('-', ' ').replace("'", ' ')
    cleaned_string = re.sub(r'[^a-zA-Z0-9\s]', '', replaced_chars)
    return cleaned_string.upper().strip()


def convert_date_format(date_str):
    if date_str is None:
        return '1900-01-01'
    date_str = str(date_str).strip()
    if date_str == '' or date_str.lower() in ('nan', 'none', 'nat'):
        return '1900-01-01'
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(date_str, fmt).strftime("%Y-%m-%d")
        except Exception:
            continue
    try:
        from dateutil import parser
        return parser.parse(date_str).strftime("%Y-%m-%d")
    except Exception:
        return '1900-01-01'


def extract_error(data_str):
    """Non protégé dans l'existant (B14) : l'appelant capture l'exception (isolement)."""
    data_string = json.loads(json.dumps(json.loads(data_str), indent=4))
    return data_string['confirmMessage']['processMessages'][0]['developerMessage']['messageTxt']


def get_recordingBasisCode(remunerationType):
    return {"B": "12M", "F": "12F", "S": ""}.get(remunerationType)


def normalize_street_compl(value):
    if not value or not value.strip():
        return ""
    first = value.strip()[0].upper()
    ascii_char = "".join(c for c in unicodedata.normalize("NFD", first) if unicodedata.category(c) != "Mn")
    if ascii_char.isalpha() and ascii_char.isascii():
        return ascii_char
    return ""


def set_workSchedule(fte):
    return '001' if fte == 1 else '999'


ARRONDISSEMENT_COMMUNES = {"75056": ("750", "751"), "69123": ("6900", "6938"), "13055": ("130", "132")}


def resolve_arrondissement_insee(commune_code, postal_code):
    if commune_code not in ARRONDISSEMENT_COMMUNES:
        return commune_code
    prefix_postal, prefix_insee = ARRONDISSEMENT_COMMUNES[commune_code]
    if postal_code.startswith(prefix_postal):
        return prefix_insee + postal_code[len(prefix_postal):]
    return commune_code


def append_arrondissement_to_city(city_name, postal_code):
    city_normalized = city_name.upper().strip()
    if re.match(r'^(PARIS|LYON|MARSEILLE)\s+\d+$', city_normalized):
        return city_name
    cities = {"PARIS": "75", "LYON": "69", "MARSEILLE": "13"}
    if city_normalized not in cities:
        return city_name
    if not postal_code.startswith(cities[city_normalized]):
        return city_name
    if len(postal_code) == 5:
        return f"{city_normalized} {str(int(postal_code[-2:]))}"
    return city_name


def get_insee_code(reader, postal_code, city_name, similarity_threshold=0.8):
    """G6 via une session HTTP SÉPARÉE (B22 corrigé), timeout 5 s, mêmes règles de choix."""
    try:
        response = reader.insee(postal_code)
        if response.status_code != 200:
            return None, None, 0
        communes = response.json()
        if not communes:
            return None, None, 0
        if len(communes) == 1:
            commune = communes[0]
            return resolve_arrondissement_insee(commune['code'], postal_code), commune['nom'], 1.0
        best_match, best_ratio = None, 0
        for commune in communes:
            ratio = difflib.SequenceMatcher(None, city_name, normalize_string(commune['nom'])).ratio()
            if ratio > best_ratio:
                best_ratio, best_match = ratio, commune
        if best_match and best_ratio >= similarity_threshold:
            return resolve_arrondissement_insee(best_match['code'], postal_code), best_match['nom'], best_ratio
        return None, None, best_ratio
    except Exception:
        return None, None, 0


# ---------------------------------------------------------------------------
# Préparation de la ligne (R-ADP-003) : astype puis fillna('') comme pd.read_parquet de la CETAS
# ---------------------------------------------------------------------------
def legacy_row(row: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(row)
    for col in ("id_payroll", "lb_adresse_cp", "cd_ppc_convention_collective_adp", "cd_nature_contract_adp",
                "cd_type_contract_adp"):
        out[col] = str(out.get(col))            # astype(str) AVANT fillna : NULL -> 'None' (parité)
    nir = out.get("id_numero_securite_sociale")
    out["id_numero_securite_sociale"] = int(nir)  # B16 : échec = salarié en erreur (isolé)
    for k, v in list(out.items()):
        if v is None:
            out[k] = ''
    return out


def _ts(date_value: str, now: datetime.datetime) -> str:
    """TS(x) = date + heure courante 'THH:MM:SS.mmmZ' (L4334 etc. ; B27)."""
    return date_value + now.strftime('T%H:%M:%S.') + f"{now.microsecond // 1000:03d}Z"


# ---------------------------------------------------------------------------
# Lectures ADP (G1..G5) — même sémantique que l'existant, une requête par appel (D-ADP-11 : parité)
# ---------------------------------------------------------------------------
def _g3(reader, worker_id):
    return reader.get(f"hr/v2/workers/{worker_id}")


def _g1(reader, worker_id):
    return reader.get("hr/v2/workers?$filter=workers/workerID/idValue eq '" + worker_id + "'")


def worker_exists(reader, worker_id):
    return _g1(reader, worker_id).status_code == 200          # B7 : toute erreur = absent (D-ADP-02)


def ssn_exists_v2(reader, ssn):
    r = reader.get("hr/v2/workers?$filter=workers%2Fperson%2FidentityDocuments%5B0%5D%2FdocumentID%20eq%20%27"
                   + ssn + "%27")
    return r.status_code == 200


def get_worker_ssn(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None
    for worker in r.json()['workers']:
        for document in worker['person']['identityDocuments']:
            if document['typeCode']['codeValue'] == "SSN":
                return document['documentID']
            return False                                     # B6 : premier document seulement
    return None


def first_hire_date(reader, worker_id):
    return _g3(reader, worker_id).json()["workers"][0]["workerDates"]["firstHireDate"]


def convert_worker_id(reader, worker_id):
    return _g3(reader, worker_id).json()["workers"][0]["associateOID"]


def get_name(reader, worker_id):
    return _g3(reader, worker_id).json()["workers"][0]["person"]["legalName"]["formattedName"]


def get_worker_dates(reader, worker_id):
    data = _g3(reader, worker_id).json()
    try:
        worker = data["workers"][0]
        wa = worker.get("workAssignments", [])[0]
        start_wt, start_pb, end_pb = None, '', ''
        for field in wa.get("customFieldGroup", {}).get("dateFields", []):
            fid = field.get("itemID")
            if fid == "startWorkerTypeDate":
                start_wt = field.get("dateValue")
            elif fid == "startProbationDate":
                start_pb = field.get("dateValue")
            elif fid == "endProbationDate":
                end_pb = field.get("dateValue")
        return {"firstHireDate": worker.get("workerDates", {}).get("firstHireDate"),
                "originalHireDate": worker.get("workerDates", {}).get("originalHireDate"),
                "hireDate": wa.get("hireDate"), "seniorityDate": wa.get("seniorityDate"),
                "expectedStartDate": wa.get("expectedStartDate"), "contractStartDate": start_wt,
                "startProbationDate": start_pb, "endProbationDate": end_pb}
    except Exception:
        return None


def _first_worker(r):
    workers = r.json().get("workers", [])
    return workers[0] if workers else None


def get_worker_personalemail(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    try:
        return r.json()["workers"][0]["person"]["communication"]["emails"][0]["emailUri"]
    except (KeyError, IndexError):
        return ''


def get_worker_businessemail(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    emails = worker.get("businessCommunication", {}).get("emails", [])
    return emails[0].get("emailUri") if emails else ''


def get_business_mobile_number(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    mobiles = worker.get("businessCommunication", {}).get("mobiles", [])
    return mobiles[0].get("formattedNumber") if mobiles else ''


def get_personal_mobile_number(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    mobiles = worker.get("person", {}).get("communication", {}).get("mobiles", [])
    return mobiles[0].get("formattedNumber") if mobiles else ''


def get_marital_status_code(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    return worker.get("person", {}).get("maritalStatusCode", {}).get("codeValue") or ''


def get_marital_status_effective_date(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None
    worker = _first_worker(r)
    if not worker:
        return None
    return worker.get("person", {}).get("maritalStatusCode", {}).get("effectiveDate") or None


def get_legal_address(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code == 200:
        worker = _first_worker(r)
        if not worker:
            return "", "", "", "", "", "", "", ""
        la = worker.get("person", {}).get("legalAddress", {})
        if la:
            return (la.get("buildingNumberExtension") or "", la.get("buildingNumber") or "",
                    la.get("streetName") or "", la.get("countryCode") or "",
                    la.get("countrySubdivisionLevel2", {}).get("longName") or la.get("cityName") or "",
                    la.get("lineFive") or "", la.get("postalCode") or "",
                    la.get("countrySubdivisionLevel1", {}).get("codeValue") or "")
        return "", "", "", "", "", "", "", ""
    return "", "", "", "", "", "", ""                         # B4 : 7 valeurs -> erreur au dépaquetage


def get_worker_salary(reader, worker_id):
    r = _g1(reader, worker_id)
    if r.status_code == 200:
        try:
            wa = r.json()["workers"][0]["workAssignments"][0]
            return wa.get('baseRemuneration', {}).get('monthlyRateAmount', {}).get('amountValue')
        except (KeyError, IndexError):
            return None
    return False


def get_worker_bonus(reader, worker_id):
    r = _g1(reader, worker_id)
    if r.status_code == 200:
        wa = r.json().get('workers', [])[0].get('workAssignments', [])[0]
        value = 0
        for field in wa.get('customFieldGroup', {}).get('numberFields', []):
            if field.get('itemID') == 'REM_RE1MTS18':
                value = field.get('numberValue', 0)
                break
        return value
    return False


def get_worker_type_and_termination(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None, None
    try:
        workers = r.json().get("workers", [])
        if not workers:
            return None, None
        was = workers[0].get("workAssignments", [])
        if not was:
            return None, None
        a = was[0]
        code = a.get("workerTypeCode", {}).get("codeValue")
        if code == "00":
            end = a.get("expectedTerminationDate")
        else:
            end = None
            for field in a.get("customFieldGroup", {}).get("dateFields", []):
                if field.get("itemID") == "fixedTermContractInitialTerminationDate":
                    end = field.get("dateValue")
                    break
        return code, end
    except Exception:
        return None, None


def get_work_schedule(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code == 200:
        wa = r.json().get('workers', [])[0].get('workAssignments', [])[0]
        code = None
        for field in wa.get('customFieldGroup', {}).get('codeFields', []):
            if field.get('itemID') == 'workSchedule':
                code = field.get('codeValue', None)
                break
        return code, wa.get('fullTimeEquivalenceRatio', None)
    return False, None


def get_organisational_structure(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None
    worker = _first_worker(r)
    if not worker:
        return None
    was = worker.get("workAssignments", [])
    if not was:
        return None
    for unit in was[0].get("assignedOrganizationalUnits", []):
        if unit.get("itemID") == "departmentId":
            return unit.get("nameCode", {}).get("codeValue") or ""
    return ""


def get_worker_imputation(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None, None
    worker = _first_worker(r)
    if not worker:
        return None, None
    was = worker.get("workAssignments", [])
    if not was:
        return None, None
    ccs = was[0].get("assignmentCostCenters", [])
    if not ccs:
        return None, None
    pct = ccs[0].get("costCenterPercentage")
    return (int(pct) if pct is not None else None), ccs[0].get("costCenterID")


def get_worker_classification(reader, worker_id):
    r = _g3(reader, worker_id)
    if not (200 <= r.status_code < 400):                      # raise_for_status
        return False                                          # B5 : dépaquetage impossible -> erreur
    wa = ((r.json().get("workers") or [{}])[0].get("workAssignments") or [{}])[0]
    group = classification = coefficient = remuneration = None
    for wg in wa.get("workerGroups", []):
        if wg.get("itemID") == "collectiveAgreement":
            group = wg.get("groupCode", {}).get("codeValue")
            break
    for oc in wa.get("occupationalClassifications", []):
        if oc.get("itemID") == "classification":
            classification = oc.get("classificationCode", {}).get("codeValue")
            break
    for nf in wa.get("customFieldGroup", {}).get("numberFields", []):
        if nf.get("itemID") == "coefficient":
            try:
                coefficient = str(int(float(nf.get("numberValue", 0))))
            except (ValueError, TypeError):
                coefficient = None
            break
    for cf in wa.get("customFieldGroup", {}).get("codeFields", []):
        if cf.get("itemID") == "remunerationType":
            remuneration = cf.get("codeValue")
            break
    return (group, classification, coefficient, wa.get("jobTitle"), wa.get("jobFunctionCode", {}).get("codeValue"),
            remuneration, wa.get("baseRemuneration", {}).get("recordingBasisCode", {}).get("codeValue"))


def get_pay_distribution(reader, associate_oid):
    try:
        r = reader.get(f"payroll/v2/workers/{associate_oid}/pay-distributions")
        if r.status_code != 200:
            return "", "", "", ""
        try:
            di = r.json()["payDistributions"][0]["distributionInstructions"][0]
            holder = di["depositAccount"]["financialAccount"].get("accountName")
            iban = di["depositAccount"].get("IBAN")
            party = di["depositAccount"].get("financialParty", {})
            swift = di["depositAccount"]["financialParty"].get("SWIFTCode")
            bank = party.get("nameCode", {}).get("shortName", "")
            if not all([holder, iban, swift, bank]):
                return "", "", "", ""
            return holder, iban, swift, bank
        except (KeyError, IndexError, TypeError):
            return "", "", "", ""
    except Exception:
        return "", "", "", ""


# ---------------------------------------------------------------------------
# Payloads (A05 § 4.5) — même structure et même ordre de clés que l'existant
# ---------------------------------------------------------------------------
def _person_event(aoid, fragment):
    return {"events": [{"data": {"eventContext": {"worker": {"associateOID": aoid}},
                                 "transform": {"worker": fragment}}}]}


def _wa_event(aoid, waid, eff, wa_fragment, extra_transform=None):
    transform = {"effectiveDateTime": eff}
    if extra_transform:
        transform.update(extra_transform)
    transform["workAssignment"] = wa_fragment
    return {"events": [{"data": {"eventContext": {"associateOID": aoid, "workAssignmentID": waid},
                                 "transform": transform}}]}


def _marital_payload(aoid, code, eff_date):
    return {"events": [{"serviceCategoryCode": {"codeValue": "hr"},
                        "eventNameCode": {"codeValue": "worker.maritalStatus.change"},
                        "data": {"eventContext": {"worker": {"associateOID": aoid}},
                                 "transform": {"eventStatusCode": {"codeValue": "submit"},
                                               "worker": {"person": {"maritalStatusCode": {
                                                   "codeValue": code, "effectiveDate": eff_date}}}}}}]}


def _org_units(code):
    letters = {'A': 'administrativeAssignment1', 'B': 'administrativeAssignment2',
               'C': 'administrativeAssignment3', 'D': 'administrativeAssignment4'}
    return [{"itemID": letters[l], "nameCode": {"codeValue": l + d}} for l, d in re.findall(r'([A-D])(\d+)', code)]


# ---------------------------------------------------------------------------
# Traitement d'un salarié (main(), corps de boucle L4259-4986)
# ---------------------------------------------------------------------------
class EmployeePlanner:
    """planner(row, executor) : exécute la logique de l'existant pour un salarié. Chaque POST est
    soumis à l'exécuteur (intention durable, transport, issue) ; chaque ligne de log de l'existant
    est transmise à executor.log()."""

    def __init__(self, reader, now: Callable[[], datetime.datetime], row_index: Callable[[str], int] = lambda k: 0):
        self.reader = reader
        self.now = now
        self.row_index = row_index

    def __call__(self, raw_row: Dict[str, Any], ex: OperationExecutor) -> None:
        row = legacy_row(raw_row)
        r = self.reader
        now = self.now()
        index = self.row_index(raw_row.get("id_unique"))
        worker_id = row['id_payroll']
        matricule_it = row['cd_matricule_it']
        preferredSalutations = row['cd_emp_titre_adp']
        givenName = normalize_string(row['lb_prenom'])
        familyName1 = normalize_string(row['lb_nom'])
        marital_status = row['cd_emp_situation_matrimoniale_adp']
        marital_effectiveDate = convert_date_format(row['dt_debut_situation_matrimoniale'])
        birthName = row['lb_nom_naissance']
        birth_country_code = row['id_naissance_pays_adp']
        birth_dept = row['lb_naissance_departement']
        birth_postcode = row['lb_naissance_cp']
        birth_city_name = normalize_string(row['lb_naissance_ville'])
        email = row['lb_email_professionnel']
        emailPerso = row['lb_email_personnel']
        landline = row['lb_telephone_professionnel']
        mobile = row['lb_telephone_personnel']
        birthDate = convert_date_format(row['dt_naissance'])
        gender = row['cd_emp_sexe_adp']
        identityDocumentsId = str(int(row['id_numero_securite_sociale']))
        citizenship_country_code = row['cd_emp_nationalite_adp']
        legalStreetCompl = str(row['lb_fiscal_streetnumber_complement'] or '').strip()
        legalAddressLineOne = str(row['lb_fiscal_streetnumber'] or '').strip()
        legalAddressLineTwo = normalize_string(row['lb_fiscal_street'])
        legalAddressPostcode = str(row['lb_fiscal_cp'] or '').strip()
        legalAddressCityName = normalize_string(row['lb_fiscal_ville'])
        legalAddressCountrycode = str(row['id_fiscal_pays_adp'] or '').strip()
        legalAdditionalAddressInformation = normalize_string(row['lb_fiscal_additional_address_information'])
        organisational_structure_code = str(row['cd_organisationalstructure_adp'] or '').strip()
        account_holder = normalize_string(row['lb_account_holder'])
        iban = row['lb_iban']
        swift_code = row['lb_bic']
        bank_name = normalize_string(row['lb_bank_name'])
        assignedWorkLocations = row['cd_ppc_location_country_adp']          # R-ADP-047 (D-ADP-06)
        hireDate = convert_date_format(str(row['dt_entree_societe']))
        firstHireDate = convert_date_format(str(row['dt_entree_groupe']))
        seniorityDate = convert_date_format(str(row['dt_anciennete']))
        if not seniorityDate or seniorityDate == "" or seniorityDate == "1900-01-01":
            seniorityDate = firstHireDate
        companyStartDate = convert_date_format(str(row['dt_company_start']))
        startProbationDate = convert_date_format(str(row['dt_ppc_debut_periode_essaie']))
        endProbationDate = convert_date_format(str(row['dt_ppc_fin_periode_essaie']))
        situationStartDateTS = _ts(convert_date_format(str(row['dt_situation_start'])), now)
        collaborationType = row['cd_type_collaboration_adp']
        workerTypeCode = row['cd_nature_contract_adp']
        reasonCDD = row['cd_recours_cdd']
        contractType = row['cd_type_contract_adp'].strip(' ')
        if workerTypeCode == '20':
            reason = ''
        elif workerTypeCode == '01':
            reason = reasonCDD
        else:
            reason = '99'
        raw_reason = str(row.get('cd_ppc_raison_debut_contrat_adp', '')).strip()
        reasonHire = str(int(raw_reason)) if raw_reason.isdigit() else ''
        contractStartDate = convert_date_format(str(row['dt_ppc_debut_contrat']))
        contractStartDateTS = _ts(contractStartDate, now)
        expectedTerminationDate = convert_date_format(str(row['dt_ppc_fin_contrat']))
        if expectedTerminationDate == '1900-01-01':
            expectedTerminationDate = None
        costCenter = row['cd_ppc_cost_center_adp']
        imputationPercentage = '100'                                          # B8 (mise à jour : chaîne)
        monthlySalary = round(float(row['mt_salaire_mensuel']), 2)            # B16 : '' -> erreur
        salaryeffectiveDateTime = _ts(convert_date_format(str(row['dt_debut_mt_salaire_base'])), now)
        bonus = float(row['pc_prime_annuelle'])
        bonuseffectiveDate = convert_date_format(str(row['dt_debut_pc_prime_annuelle']))
        bonuseffectiveDateTime = _ts(bonuseffectiveDate, now)
        raw_fte = str(row.get('pc_imputation', '')).strip().replace(',', '.')
        try:
            fte = float(raw_fte) / 100                                        # B9 : FTE pris dans pc_imputation
        except ValueError:
            fte = 0.0
        terminationDate = convert_date_format(str(row['dt_sortie']))
        int(terminationDate.split("-")[0])                                    # parité : termination_year
        jobFunctionCode = str(row['cd_categorie_cotisant_adp'])
        remunerationType = row['cd_classe_remuneration_adp'].strip(' ')
        recordingBasisCode = get_recordingBasisCode(remunerationType)
        jobTitle = normalize_string(row['lb_job_title'])
        workSchedule = set_workSchedule(fte)

        # --- Décision embauche / mise à jour (R-ADP-008/009, B7 à arbitrer D-ADP-02) -------------
        if worker_exists(r, worker_id) == False and ssn_exists_v2(r, identityDocumentsId) == False:  # noqa: E712
            detail = f"Row {index} is a new hire : id_payroll = {worker_id}"
            payload = _hire_payload_impl(
                self,
                worker_id, citizenship_country_code, identityDocumentsId, givenName, familyName1, birthName,
                preferredSalutations, landline, mobile, email, marital_status, birthDate, birth_city_name,
                birth_postcode, birth_dept, birth_country_code, legalAddressLineOne, legalAddressLineTwo,
                legalAddressCityName, legalAddressCountrycode, legalAddressPostcode, hireDate, seniorityDate,
                jobFunctionCode, workerTypeCode, monthlySalary, costCenter, collaborationType, contractType,
                endProbationDate, startProbationDate, bonus, reasonHire, reason, expectedTerminationDate,
                assignedWorkLocations, emailPerso, jobTitle, gender, remunerationType,
                legalAdditionalAddressInformation, workSchedule, recordingBasisCode, matricule_it,
                contractStartDate, firstHireDate, companyStartDate, marital_effectiveDate,
                organisational_structure_code, legalStreetCompl)
            target = {"workerID": worker_id, "hire_payload_hash": _stable_hash(payload)}
            self._post(ex, "P1", "P1_HIRE_V1", EV_HIRE, payload, target, familyName1, worker_id, "INSERT", "hire",
                       detail, read_back=lambda rd: {"workerID": worker_id} if worker_exists(rd, worker_id) else None)
            return

        ssn_worker_id = get_worker_ssn(r, worker_id)
        if identityDocumentsId != ssn_worker_id:
            detail = f"SSN Mismatch for {worker_id} between TS {identityDocumentsId} and ADP {ssn_worker_id}"
            ex.log(self._log_row(worker_id, familyName1, 'UPDATE', 'SSN', 409, 'DATA_VALIDATION', None, detail,
                                 detail, step="SSN_CHECK"))
            return

        workAssignmentID = worker_id + "|" + first_hire_date(r, worker_id)
        associateOID = convert_worker_id(r, worker_id)
        dates = get_worker_dates(r, worker_id)
        current_seniority_date = dates["seniorityDate"]                      # dates None -> erreur (parité)
        formatted_name = get_name(r, worker_id)
        nom = formatted_name

        current = get_worker_personalemail(r, worker_id)
        if emailPerso != current:
            self._post(ex, "P2", "P2_PERSONAL_EMAIL", EV_PERSONAL_EMAIL,
                       _person_event(associateOID, {"person": {"communication": {"email": {"emailUri": emailPerso}}}}),
                       {"emailUri": emailPerso}, nom, worker_id, 'UPDATE', 'personalemail',
                       f"changing personal email from {current} to {emailPerso}",
                       read_back=lambda rd: {"emailUri": get_worker_personalemail(rd, worker_id)})

        current = get_worker_businessemail(r, worker_id)
        if email != current:
            self._post(ex, "P3", "P3_BUSINESS_EMAIL", EV_BUSINESS_EMAIL,
                       _person_event(associateOID, {"businessCommunication": {"email": {"emailUri": email}}}),
                       {"emailUri": email}, nom, worker_id, 'UPDATE', 'businessemail',
                       f"changing business email from {current} to {email}",
                       read_back=lambda rd: {"emailUri": get_worker_businessemail(rd, worker_id)})

        current = get_business_mobile_number(r, worker_id)
        if landline != current:
            self._post(ex, "P4", "P4_BUSINESS_MOBILE", EV_BUSINESS_MOBILE,
                       _person_event(associateOID, {"businessCommunication": {"mobile": {"formattedNumber": landline}}}),
                       {"formattedNumber": landline}, nom, worker_id, 'UPDATE', 'businessmobile',
                       f"changing business mobile  from {current} to {landline}",
                       read_back=lambda rd: {"formattedNumber": get_business_mobile_number(rd, worker_id)})

        current = get_personal_mobile_number(r, worker_id)
        if mobile != current:
            self._post(ex, "P5", "P5_PERSONAL_MOBILE", EV_PERSONAL_MOBILE,
                       _person_event(associateOID, {"person": {"communication": {"mobile": {"formattedNumber": mobile}}}}),
                       {"formattedNumber": mobile}, nom, worker_id, 'UPDATE', 'personalmobile',
                       f"changing personal mobile  from {current} to {mobile}",
                       read_back=lambda rd: {"formattedNumber": get_personal_mobile_number(rd, worker_id)})

        cur_ms = get_marital_status_code(r, worker_id)
        cur_ms_date = get_marital_status_effective_date(r, worker_id)
        if marital_status != cur_ms or marital_effectiveDate != cur_ms_date:   # B10
            self._post(ex, "P6", "P6_MARITAL_STATUS", EV_MARITAL,
                       _marital_payload(associateOID, marital_status, marital_effectiveDate),
                       {"codeValue": marital_status, "effectiveDate": marital_effectiveDate}, nom, worker_id,
                       'UPDATE', 'maritalStatusCode',
                       f"changing marital status from {cur_ms} to {marital_status} on {marital_effectiveDate}",
                       read_back=lambda rd: {"codeValue": get_marital_status_code(rd, worker_id),
                                             "effectiveDate": get_marital_status_effective_date(rd, worker_id)})

        (c_compl, c_one, c_two, c_country, c_city, c_add, c_cp, _c_bureau) = get_legal_address(r, worker_id)  # B4
        legalStreetCompl = normalize_street_compl(legalStreetCompl)                                           # B21
        city_norm = append_arrondissement_to_city(legalAddressCityName, legalAddressPostcode)
        if (legalStreetCompl != c_compl or legalAddressLineOne != c_one or legalAddressLineTwo != c_two
                or legalAddressCountrycode != c_country or city_norm != c_city
                or legalAdditionalAddressInformation != c_add or legalAddressPostcode != c_cp):
            insee, _, _ = get_insee_code(r, legalAddressPostcode, legalAddressCityName)
            longname = append_arrondissement_to_city(legalAddressCityName, legalAddressPostcode)
            payload = _person_event(associateOID, {"person": {"legalAddress": {
                "buildingNumberExtension": legalStreetCompl, "buildingNumber": legalAddressLineOne,
                "streetName": legalAddressLineTwo, "countryCode": legalAddressCountrycode,
                "countrySubdivisionLevel1": {"codeValue": legalAddressCityName},
                "countrySubdivisionLevel2": {"longName": longname, "codeValue": insee if insee else ""},
                "lineFive": legalAdditionalAddressInformation, "postalCode": legalAddressPostcode}}})
            detail = (f"changing legal address from {c_one} {c_compl} {c_two} {c_country} {c_city} {c_add} {c_cp} "
                      f"to {legalAddressLineOne} {legalStreetCompl} {legalAddressLineTwo} {legalAddressCountrycode} "
                      f"{legalAddressCityName} {legalAdditionalAddressInformation} {legalAddressPostcode}")
            target = {"buildingNumberExtension": legalStreetCompl, "buildingNumber": legalAddressLineOne,
                      "streetName": legalAddressLineTwo, "countryCode": legalAddressCountrycode, "city": longname,
                      "lineFive": legalAdditionalAddressInformation, "postalCode": legalAddressPostcode}

            def rb(rd):
                v = get_legal_address(rd, worker_id)
                if len(v) != 8:
                    return None
                return {"buildingNumberExtension": v[0], "buildingNumber": v[1], "streetName": v[2],
                        "countryCode": v[3], "city": v[4], "lineFive": v[5], "postalCode": v[6]}
            self._post(ex, "P7", "P7_LEGAL_ADDRESS", EV_LEGAL_ADDRESS, payload, target, nom, worker_id, 'UPDATE',
                       'legal address', detail, read_back=rb)  # remise à vide PUIS extraction en 4xx (L4662-4665)

        current_salary = get_worker_salary(r, worker_id)
        if current_salary is not None and monthlySalary != current_salary and salaryeffectiveDateTime != '1900-01-01':  # B3
            self._post(ex, "P8", "P8_SALARY", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, salaryeffectiveDateTime,
                                 {"baseRemuneration": {"monthlyRateAmount": {"currencyCode": "EUR",
                                                                             "amountValue": monthlySalary}}}),
                       {"amountValue": monthlySalary, "effectiveDate": salaryeffectiveDateTime[:10]}, nom, worker_id,
                       'UPDATE', 'salary', f"changing salary from {current_salary} to {monthlySalary} on {salaryeffectiveDateTime}",
                       read_back=lambda rd: {"amountValue": get_worker_salary(rd, worker_id)})

        current_bonus = get_worker_bonus(r, worker_id)
        if bonus != get_worker_bonus(r, worker_id) and bonuseffectiveDate != '1900-01-01':
            self._post(ex, "P8", "P8_BONUS", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, bonuseffectiveDateTime,
                                 {"customFieldGroup": {"numberFields": [{"itemID": "REM_RE1MTS18", "numberValue": bonus}]}}),
                       {"numberValue": bonus, "effectiveDate": bonuseffectiveDate}, nom, worker_id, 'UPDATE', 'bonus',
                       f"changing bonus from {current_bonus} to {bonus} on {bonuseffectiveDateTime}",
                       read_back=lambda rd: {"numberValue": get_worker_bonus(rd, worker_id)})

        cur_type, cur_end = get_worker_type_and_termination(r, worker_id)
        contract_variant = None
        if workerTypeCode != cur_type:
            if expectedTerminationDate != '1900-01-01' and workerTypeCode != '00':     # B2 : A pour tout non-CDI
                contract_variant = "P8_CONTRACT_A"
                fragment = {"expectedTerminationDate": expectedTerminationDate,
                            "assignmentStatus": {"reasonCode": {"codeValue": reasonHire}},
                            "workerTypeCode": {"codeValue": workerTypeCode},
                            "customFieldGroup": {
                                "codeFields": [{"itemID": "recoursReason", "codeValue": reason},
                                               {"itemID": "contractType", "codeValue": contractType}],
                                "dateFields": [{"dateValue": endProbationDate, "itemID": "endProbationDate"},
                                               {"dateValue": startProbationDate, "itemID": "startProbationDate"},
                                               {"itemID": "fixedTermContractInitialTerminationDate",
                                                "dateValue": expectedTerminationDate},
                                               {"itemID": "startWorkerTypeDate", "dateValue": contractStartDate}]}}
                payload = _wa_event(associateOID, workAssignmentID, contractStartDateTS, fragment)
            else:
                contract_variant = "P8_CONTRACT_B"
                fragment = {"expectedStartDate": contractStartDate,
                            "assignmentStatus": {"reasonCode": {"codeValue": reasonHire}},
                            "workerTypeCode": {"codeValue": workerTypeCode},
                            "assignedWorkLocations": [{"itemID": "default", "nameCode": {"codeValue": assignedWorkLocations}}],
                            "customFieldGroup": {
                                "codeFields": [{"itemID": "recoursReason", "codeValue": reason},
                                               {"itemID": "contractType", "codeValue": contractType}],
                                "dateFields": [{"itemID": "startWorkerTypeDate", "dateValue": contractStartDate}]}}
                payload = _wa_event(associateOID, workAssignmentID, contractStartDateTS, fragment,
                                    extra_transform={"expectedStartDate": contractStartDate})
            target = {"workerTypeCode": workerTypeCode, "contract_payload_hash": _stable_hash(payload, volatile=True)}
            self._post(ex, "P8", contract_variant, EV_WA_MODIFY, payload, target, nom, worker_id, 'UPDATE',
                       'contrat et affectation',
                       f"changing contract from {cur_type} to {workerTypeCode} with probation dates from "
                       f"{startProbationDate} to {endProbationDate} ",
                       read_back=lambda rd: {"workerTypeCode": get_worker_type_and_termination(rd, worker_id)[0]})

        if expectedTerminationDate != cur_end:                                          # B12
            self._post(ex, "P8", "P8_END_DATE", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, contractStartDateTS,
                                 {"expectedTerminationDate": expectedTerminationDate}),
                       {"expectedTerminationDate": expectedTerminationDate}, nom, worker_id, 'UPDATE',
                       'date de fin de contrat',
                       f"changing termination date from {cur_end} to {expectedTerminationDate}",
                       depends_on=("P8_CONTRACT_A", "P8_CONTRACT_B"))                    # CONTRAT-ADP:C08

        cur_ws, cur_fte = get_work_schedule(r, worker_id)
        if cur_fte is None:
            pass
        elif cur_fte is not None and round(float(cur_fte), 4) == round(float(fte), 4):
            pass
        elif workSchedule != cur_ws:
            self._post(ex, "P8", "P8_SCHEDULE", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS,
                                 {"customFieldGroup": {"codeFields": [{"itemID": "workSchedule", "codeValue": workSchedule}]}}),
                       {"workSchedule": workSchedule}, nom, worker_id, 'UPDATE', 'work schedule',
                       f"changing work schedule from code {cur_ws} FTE {cur_fte} to code {workSchedule} FTE {fte} "
                       f"on {situationStartDateTS}",
                       read_back=lambda rd: {"workSchedule": get_work_schedule(rd, worker_id)[0]})

        cur_org = get_organisational_structure(r, worker_id)
        if organisational_structure_code and cur_org is not None and organisational_structure_code != cur_org:  # B11
            detail = (f"changing organisational structure from {cur_org} to {organisational_structure_code} "
                      f"on {situationStartDateTS}")
            units = _org_units(organisational_structure_code)
            if not units:  # pas de POST, une ligne de log à statut vide (parité)
                ex.log(self._log_row(worker_id, nom, 'UPDATE', 'organisational structure', None, None, None, '',
                                     detail, step="P8_ORG_NO_SEGMENT"))
            else:
                self._post(ex, "P8", "P8_ORG", EV_WA_MODIFY,
                           _wa_event(associateOID, workAssignmentID, situationStartDateTS,
                                     {"assignedOrganizationalUnits": units}),
                           {"units": [u["nameCode"]["codeValue"] for u in units]}, nom, worker_id, 'UPDATE',
                           'organisational structure', detail)

        cur_imp, cur_cc = get_worker_imputation(r, worker_id)
        if imputationPercentage != str(cur_imp) or costCenter != cur_cc:
            self._post(ex, "P8", "P8_IMPUTATION", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS,
                                 {"assignmentCostCenters": [{"costCenterPercentage": imputationPercentage,
                                                             "costCenterID": costCenter}]}),
                       {"costCenterPercentage": imputationPercentage, "costCenterID": costCenter}, nom, worker_id,
                       'UPDATE', 'imputation',
                       f"changing imputation from {cur_cc} to {costCenter} and from {cur_imp} % to "
                       f"{imputationPercentage} % on {situationStartDateTS}",
                       read_back=lambda rd: dict(zip(("costCenterPercentage", "costCenterID"),
                                                     (lambda p, c: (str(p), c))(*get_worker_imputation(rd, worker_id)))))

        if seniorityDate != str(current_seniority_date):
            self._post(ex, "P8", "P8_SENIORITY", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS, {"seniorityDate": seniorityDate}),
                       {"seniorityDate": seniorityDate}, nom, worker_id, 'UPDATE', 'senioirty date',
                       f"changing seniority date from {current_seniority_date} to {seniorityDate}",
                       read_back=lambda rd: {"seniorityDate": (get_worker_dates(rd, worker_id) or {}).get("seniorityDate")})

        (_g, _c, _k, cur_title, cur_fn, cur_rem, _rb) = get_worker_classification(r, worker_id)   # B5
        if jobTitle != cur_title or jobFunctionCode != cur_fn or remunerationType != cur_rem:
            fragment = {"jobFunctionCode": {"codeValue": jobFunctionCode}, "jobTitle": jobTitle,
                        "baseRemuneration": {"recordingBasisCode": {"codeValue": recordingBasisCode}},
                        "customFieldGroup": {"indicatorFields": [{"itemID": "forcingCoefficient", "indicatorValue": True}],
                                             "codeFields": [{"itemID": "remunerationType", "codeValue": remunerationType},
                                                            {"itemID": "modeRemSATH", "codeValue": remunerationType}]}}
            self._post(ex, "P8", "P8_POSITION", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS, fragment),
                       {"jobTitle": jobTitle, "jobFunctionCode": jobFunctionCode, "remunerationType": remunerationType},
                       nom, worker_id, 'UPDATE', 'position professionnelle',
                       f"changing position from {cur_title} to {jobTitle} and from {cur_fn} to {jobFunctionCode} "
                       f"and from {cur_rem} to {remunerationType} on {situationStartDateTS}",
                       read_back=lambda rd: dict(zip(("jobTitle", "jobFunctionCode", "remunerationType"),
                                                     get_worker_classification(rd, worker_id)[3:6])))

        if terminationDate != '' and terminationDate != '1900-01-01':                    # 5.19 : information seule
            ex.log(self._log_row(worker_id, nom, 'INFO', 'GroupEndDate', '', '', '', '',
                                 f"Group End Date / Termination Date {terminationDate} i.e. date de sortie TS",
                                 step="INFO_GROUP_END_DATE"))

        c_holder, c_iban, c_swift, c_bank = get_pay_distribution(r, associateOID)
        if account_holder != c_holder or iban != c_iban or swift_code != c_swift or bank_name != c_bank:  # B13
            payload = {"events": [{"data": {"eventContext": {"worker": {"associateOID": associateOID}},
                                            "transform": {"payDistribution": {"distributionInstructions": [{
                                                "paymentMethodCode": {"codeValue": "V"}, "itemID": "1",
                                                "depositAccount": {"IBAN": iban, "financialAccount": {"accountName": account_holder},
                                                                   "financialParty": {"SWIFTCode": swift_code,
                                                                                      "nameCode": {"shortName": bank_name}}}}]}}}}]}
            self._post(ex, "P9", "P9_PAY_DISTRIBUTION", EV_PAY_DISTRIBUTION, payload,
                       {"accountName": account_holder, "IBAN": iban, "SWIFTCode": swift_code, "shortName": bank_name},
                       nom, worker_id, 'UPDATE', 'paie distribution',
                       f"changing payment distribution from {c_holder} to {account_holder} and from {c_iban}  to {iban} "
                       f"and from {c_swift} to {swift_code} and from {c_bank} to {bank_name}",
                       read_back=lambda rd: dict(zip(("accountName", "IBAN", "SWIFTCode", "shortName"),
                                                     get_pay_distribution(rd, associateOID))))

    # -- utilitaires -----------------------------------------------------------------
    def _post(self, ex: OperationExecutor, code, variant, endpoint, payload, target, nom, worker_id, mode, field,
              detail, read_back=None, depends_on=(), clear_error=False):
        res = ex.submit(PlannedOp(code, variant, endpoint, payload, target, {"mode": mode, "field": field},
                                  depends_on=depends_on, read_back=read_back))
        tr = res.transport
        if tr is None:
            return res  # aucun transport (déjà fait, incertain, dépendance, arrêt) : pas de ligne de log
        status = tr.http_status
        error_message = ''
        if str(status).startswith("4") and not clear_error and tr.response_text is not None:
            try:
                error_message = extract_error(tr.response_text)
            except Exception:
                error_message = ''  # B14 : l'existant s'arrêtait ; ici le salarié continue (D-ADP-01)
        response_code = tr.response_text if tr.response_text is not None else (tr.error_class or '')
        ex.log(self._log_row(worker_id, nom, mode, field, status, tr.http_reason or tr.reason, response_code,
                             error_message, detail, step=variant, endpoint=endpoint))
        return res

    def _log_row(self, worker_id, nom, mode, field, status, reason, response, error, detail, step, endpoint=None):
        return {"ts_adp": datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None), "matricule_id": worker_id,
                "nom": nom, "mode": mode, "field": field, "status_code": '' if status is None else str(status),
                "reason_code": '' if reason is None else str(reason), "response_code": '' if response is None else str(response),
                "error_message": error or '', "log_message": detail, "step_name": step,
                "http_method": "POST" if endpoint else None, "endpoint_template": endpoint}


def _stable_hash(payload, volatile=False):
    from hris.common import canonical_json, sha256_hex
    text = canonical_json(payload)
    if volatile:  # partie horaire de effectiveDateTime exclue de l'identité
        text = re.sub(r'(\d{4}-\d{2}-\d{2})T\d{2}:\d{2}:\d{2}\.\d{3}Z', r'\1', text)
    return sha256_hex(text)


def _hire_payload_impl(self, worker_id, citizenship, nir, given, family, birth_name, salutation, landline, mobile,
                       email, marital, birth_date, birth_city, birth_cp, birth_dept, birth_country, line_one, line_two,
                       city, country, cp, hire_date, seniority, job_fn, worker_type, salary, cost_center, collab,
                       contract_type, end_pb, start_pb, bonus, reason_hire, reason, end_date, location, email_perso,
                       job_title, gender, rem_type, line_five, schedule, recording, matricule_it, contract_start,
                       first_hire, company_start, marital_date, org_code, compl):
    """Variante 1 de hire_pending_worker (L3415-3711), toujours envoyée (B1)."""
    bureau = city
    insee, _, _ = get_insee_code(self.reader, cp, city)
    long_city = append_arrondissement_to_city(city, cp)
    units = _org_units(org_code) if org_code else [{"itemID": "departmentId", "nameCode": {"codeValue": ""}}]
    return {"events": [{"data": {"transform": {"worker": {
        "workerID": {"idValue": worker_id},
        "workerDates": {"firstHireDate": first_hire, "originalHireDate": company_start},
        "businessCommunication": {"mobiles": [{"nameCode": {"codeValue": "WORK"}, "formattedNumber": landline}],
                                  "emails": [{"emailUri": email}]},
        "person": {
            "citizenshipCountryCodes": [{"codeValue": citizenship}],
            "identityDocuments": [{"documentID": nir, "typeCode": {"codeValue": "SSN", "shortName": ""}}],
            "legalName": {"givenName": given, "middleName": '', "familyName1": family, "familyName2": birth_name,
                          "formattedName": given + " " + family,
                          "preferredSalutations": [{"sequenceNumber": 1, "salutationCode": {
                              "shortName": "", "codeValue": salutation, "longName": ""}}]},
            "communication": {"landlines": [{"itemID": "Personnel", "formattedNumber": ''}],
                              "mobiles": [{"itemID": "Personnel", "formattedNumber": mobile}],
                              "emails": [{"itemID": "Personnel", "emailUri": email_perso,
                                          "nameCode": {"codeValue": "Personnel", "shortName": "Personnel",
                                                       "longName": "Personnel"}}]},
            "maritalStatusCode": {"effectiveDate": marital_date, "codeValue": marital},
            "genderCode": {"codeValue": gender},
            "birthDate": birth_date,
            "birthPlace": {"cityName": birth_city, "postalCode": birth_cp,
                           "countrySubdivisionLevel1": {"codeValue": birth_dept}, "countryCode": birth_country},
            "legalAddress": {"buildingNumberExtension": compl, "buildingNumber": line_one, "streetName": line_two,
                             "countryCode": country, "countrySubdivisionLevel1": {"codeValue": bureau},
                             "countrySubdivisionLevel2": {"longName": long_city, "codeValue": insee if insee else ""},
                             "lineFive": line_five, "postalCode": cp}},
        "workAssignment": {
            "hireDate": hire_date, "seniorityDate": seniority, "expectedTerminationDate": end_date,
            "assignmentStatus": {"reasonCode": {"codeValue": reason_hire}},
            "jobFunctionCode": {"codeValue": job_fn},
            "assignedOrganizationalUnits": units,
            "assignedWorkLocations": [{"itemID": "default", "nameCode": {"codeValue": location}}],
            "jobTitle": job_title,
            "workerGroups": [{"groupCode": {"shortName": "", "codeValue": "94", "longName": ""}}],   # R-ADP-012
            "workerTypeCode": {"codeValue": worker_type},
            "workArrangementCode": {"codeValue": "900"},
            "remunerationBasisCode": {"codeValue": "1"},
            "baseRemuneration": {"recordingBasisCode": {"codeValue": recording},
                                 "monthlyRateAmount": {"currencyCode": "EUR", "amountValue": salary}},
            "assignmentCostCenters": [{"costCenterPercentage": HIRE_IMPUTATION_PERCENTAGE, "costCenterID": cost_center,
                                       "costCenterName": ""}],
            "customFieldGroup": {
                "amountFields": [{"itemID": "internshipCompensation", "amountValue": salary}],
                "indicatorFields": [{"itemID": "forcingCoefficient", "indicatorValue": True}],
                "codeFields": [{"itemID": "collaborationType", "codeValue": collab},
                               {"itemID": "modeRemSATH", "codeValue": rem_type},
                               {"itemID": "recoursReason", "codeValue": reason},
                               {"itemID": "contractType", "codeValue": contract_type},
                               {"itemID": "TLM", "codeValue": "W"},
                               {"itemID": "remunerationType", "codeValue": rem_type},
                               {"itemID": "workSchedule", "codeValue": schedule}],
                "dateFields": [{"dateValue": end_pb, "itemID": "endProbationDate"},
                               {"dateValue": start_pb, "itemID": "startProbationDate"},
                               {"itemID": "fixedTermContractInitialTerminationDate", "dateValue": end_date},
                               {"itemID": "startWorkerTypeDate", "dateValue": contract_start}],
                "numberFields": [{"itemID": "REM_RE1MTS18", "numberValue": bonus}],
                "stringFields": [{"itemID": "SAD_LIB34", "stringValue": matricule_it}]}},
        "eventReasonCode": {"codeValue": "ACQ"}}}}}]}

''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.adp_publish =============
register_hris_module('hris.adp_publish', '60c09780945f4bb6bd6b748b05f0283444aa39bcf153f54ca241f06a21ce651b', r'''"""hris.adp_publish — orchestration de la publication ADP d'un run (outbox -> opérations).

Ordre imposé :
  1. décision du garde-fou (nom réel du workspace) écrite AVANT toute autre action ;
  2. DENIED_NON_PRD : éléments PENDING -> NOT_PUSHED_NON_PRD ; aucun secret, aucun jeton, aucun appel ;
     BLOCKED_UNRESOLVED : rien n'est modifié (réévaluable) ; aucun appel ;
  3. ENV_SATISFIED : la condition d'environnement ne suffit pas — les blocages d'activation
     (contrat, décisions A08, contrôles du run) sont revalidés ; s'il en reste, aucun appel ;
  4. sinon seulement : fabrique du transport (secrets + jeton), puis traitement salarié par salarié.

Sélection sans recalcul : NORMAL -> éléments PENDING du run ; RETRY_PUBLICATION -> éléments PENDING
ou FAILED du run cible (jamais IN_PROGRESS = incertain, qui exige réconciliation ou décision).
"""
from __future__ import annotations

import dataclasses
import datetime as dt
from typing import Any, Callable, Dict, List, Mapping, Optional, Protocol, Sequence

from hris.adp_gate import BLOCKED_UNRESOLVED, DENIED_NON_PRD, ENV_SATISFIED, GateDecision, evaluate_gate
from hris.adp_ops import (OB_FAILED, OB_IN_PROGRESS, OB_NO_CHANGE, OB_NOT_PUSHED, OB_PENDING, OB_SENT,
                          OB_SUPERSEDED, OP_OPEN_UNCERTAIN, OP_SUPERSEDED, OperationExecutor, item_status,
                          supersession_candidates)
from hris.adp_transport import ADP_ACTIVATION_APPROVED

STATUS_NOT_PUBLISHED_NON_PRD = "NOT_PUBLISHED_NON_PRD"
STATUS_BLOCKED_UNRESOLVED = "BLOCKED_WORKSPACE_UNRESOLVED"
STATUS_BLOCKED_ACTIVATION = "BLOCKED_ACTIVATION"
STATUS_NOTHING = "NOTHING_TO_PUBLISH"
STATUS_PUBLISHED = "PUBLISHED"
STATUS_PARTIAL = "PUBLISHED_PARTIAL"
STATUS_HALTED = "HALTED_AUTH_REJECTED"


def _utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


class OutboxRepository(Protocol):
    def load(self, run_id: str, statuses: Sequence[str], max_attempts: int) -> List[dict]: ...
    def newer_items(self, run_id: str, employee_key: str) -> List[dict]: ...
    def update(self, run_id: str, updates: List[dict]) -> None: ...


class GateLog(Protocol):
    def record(self, run_id: str, decision: GateDecision) -> None: ...


def code_activation_blockers() -> List[str]:
    """Blocages d'activation dans CETTE version du code (aucun paramètre ne peut les lever)."""
    blockers = []
    if not ADP_ACTIVATION_APPROVED:
        blockers += ["CONTRAT_ADP_FINAL_ABSENT", "D-ADP-01", "D-ADP-02", "D-ENV-02", "D-SEC-01", "D-NET-01",
                     "MANDAT_DEV_SANS_PUBLICATION_REELLE"]
    return blockers


@dataclasses.dataclass
class PublicationContext:
    publish_run_id: str           # run en cours (pipeline)
    source_run_id: str            # run dont l'outbox est publiée (= publish_run_id en NORMAL)
    mode: str                     # NORMAL | RETRY_PUBLICATION
    max_attempts: int
    pause_seconds: float = 1.0    # parité R-ADP-006 (D-ADP-10)


def run_publication(ctx: PublicationContext, *, context_provider: Callable[[], Mapping[str, Any]],
                    outbox: OutboxRepository, gate_log: GateLog, ops, attempts, publication_log,
                    activation_blockers: Callable[[], List[str]], transport_factory: Callable[[], Any],
                    planner: Callable[..., Any], rows_by_key: Callable[[], Dict[str, dict]],
                    sleep: Callable[[float], None], clock: Callable[[], dt.datetime] = _utc_now) -> dict:
    if not ctx.publish_run_id or not ctx.source_run_id:
        raise ValueError("run_id obligatoire")
    if ctx.max_attempts < 1:
        raise ValueError("max_attempts doit être >= 1")
    if ctx.mode not in ("NORMAL", "RETRY_PUBLICATION"):
        raise ValueError("mode de publication invalide")

    decision = evaluate_gate(context_provider, clock)
    gate_log.record(ctx.publish_run_id, decision)            # 1. tracé AVANT toute autre action

    # CONTRAT-ADP:C13 reprise : seuls PENDING/FAILED sous max_attempts ; jamais IN_PROGRESS (incertain)
    selectable = (OB_PENDING,) if ctx.mode == "NORMAL" else (OB_PENDING, OB_FAILED)
    items = outbox.load(ctx.source_run_id, selectable, ctx.max_attempts)
    summary = {"publish_run_id": ctx.publish_run_id, "source_run_id": ctx.source_run_id, "mode": ctx.mode,
               "gate_decision_id": decision.decision_id, "gate_outcome": decision.outcome,
               "gate_reason_code": decision.reason_code, "env_condition_met": decision.env_condition_met,
               "workspace_name": decision.workspace_name, "items_selected": len(items),
               "transport_built": False, "sent": 0, "no_change": 0, "failed": 0, "uncertain": 0,
               "blocked_prior_uncertain": 0, "superseded": 0, "not_pushed": 0, "activation_blockers": []}

    if decision.outcome == DENIED_NON_PRD:                     # 2. absence normale de publication
        now = clock()
        pending = [it for it in items if it["status"] == OB_PENDING]
        outbox.update(ctx.source_run_id, [{"employee_key": it["employee_key"], "status": OB_NOT_PUSHED,
                                           "gate_decision_id": decision.decision_id, "updated_at_utc": now,
                                           "last_outcome_code": decision.reason_code} for it in pending])
        summary.update(not_pushed=len(pending), status=STATUS_NOT_PUBLISHED_NON_PRD)
        return summary
    if decision.outcome == BLOCKED_UNRESOLVED:                 # réévaluable : rien n'est classé
        summary.update(status=STATUS_BLOCKED_UNRESOLVED)
        return summary
    assert decision.outcome == ENV_SATISFIED
    blockers = list(activation_blockers())                     # 3. revalidation des autres conditions
    if blockers:
        summary.update(status=STATUS_BLOCKED_ACTIVATION, activation_blockers=blockers)
        return summary
    if not items:
        summary.update(status=STATUS_NOTHING)
        return summary

    transport = transport_factory()                            # 4. secrets + jeton ici seulement
    summary["transport_built"] = True
    rows = rows_by_key()
    halted = False
    try:
        for it in items:
            key = it["employee_key"]
            now = clock()
            # Opérations incertaines d'un run quelconque pour ce salarié : aucun envoi
            prior = [o for o in ops.open_uncertain_for_employee(key)]
            if prior:
                own = any(o["run_id"] == ctx.source_run_id for o in prior)
                outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_IN_PROGRESS if own else it["status"],
                                                   "last_outcome_code": "BLOCKED_PRIOR_UNCERTAIN",
                                                   "updated_at_utc": now, "gate_decision_id": decision.decision_id}])
                summary["blocked_prior_uncertain"] += 1
                continue
            # Obsolescence d'un élément ancien par un élément plus récent (jamais si incertain)
            newer = outbox.newer_items(ctx.source_run_id, key)
            item_ops = ops.for_item(ctx.source_run_id, key)
            if supersession_candidates(it, newer, item_ops):
                for o in item_ops:
                    if o["status"] not in ("ACCEPTED", "APPLIED_OBSERVED"):
                        ops.upsert(dict(o, status=OP_SUPERSEDED, updated_at_utc=now))
                outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_SUPERSEDED,
                                                   "last_outcome_code": "SUPERSEDED_BY_NEWER_RUN",
                                                   "updated_at_utc": now}])
                summary["superseded"] += 1
                continue
            row = rows.get(key)
            if row is None:
                outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_FAILED,
                                                   "last_outcome_code": "STAGING_ROW_MISSING", "updated_at_utc": now}])
                summary["failed"] += 1
                continue
            sleep(ctx.pause_seconds)
            attempt = int(it.get("attempt_count") or 0) + 1
            outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_IN_PROGRESS, "attempt_count": attempt,
                                               "last_attempt_at_utc": now, "gate_decision_id": decision.decision_id,
                                               "updated_at_utc": now}])
            executor = OperationExecutor(ops=ops, attempts=attempts, transport=transport,
                                         source_run_id=ctx.source_run_id, publish_run_id=ctx.publish_run_id,
                                         employee_key=key, gate_decision_id=decision.decision_id,
                                         allow_new_attempt=(ctx.mode == "RETRY_PUBLICATION"),
                                         max_attempts=ctx.max_attempts, clock=clock)
            port_error = None
            try:
                planner(row, executor)
            except Exception as exc:  # erreur de la logique portée : salarié isolé (D-ADP-01, À CONFIRMER)
                port_error = f"PORT_ERROR_{type(exc).__name__}"
            if executor.log_rows:
                publication_log.append([dict(r, run_id=ctx.publish_run_id, employee_key=key, attempt_no=attempt,
                                             gate_decision_id=decision.decision_id) for r in executor.log_rows])
            status = item_status(executor.results)
            if port_error and status in (OB_SENT, OB_NO_CHANGE):
                status = OB_FAILED  # traitement incomplet : reprise explicite requise
            outbox.update(ctx.source_run_id, [{"employee_key": key, "status": status, "updated_at_utc": clock(),
                                               "last_outcome_code": port_error or _last_code(executor)}])
            summary[{OB_SENT: "sent", OB_NO_CHANGE: "no_change", OB_FAILED: "failed",
                     OB_IN_PROGRESS: "uncertain"}[status]] += 1
            if executor.halt_reason:
                halted = True
                break
    finally:
        close = getattr(transport, "close", None)
        if close:
            close()
    if halted:
        summary["status"] = STATUS_HALTED
    elif summary["failed"] or summary["uncertain"] or summary["blocked_prior_uncertain"]:
        summary["status"] = STATUS_PARTIAL
    else:
        summary["status"] = STATUS_PUBLISHED
    return summary


def _last_code(executor) -> str:
    for r in reversed(executor.results):
        detail = getattr(r, "detail", None)
        if detail:
            return str(detail)[:100]
    return "NO_OPERATION"
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.adp_spark =============
register_hris_module('hris.adp_spark', '7259828dc901296c2b2911a281b52ce9e6c0f8e1ffef998cc73ac0447d567129', r'''"""hris.adp_spark — dépôts Delta de la publication ADP et fabriques du transport réel.

Chaque upsert d'opération est un commit Delta distinct (MERGE) : l'intention est durable AVANT
le transport. Le transport réel et le lecteur réel ne se construisent que si l'activation est
approuvée DANS LE CODE (hris.adp_transport.ADP_ACTIVATION_APPROVED) — jamais par paramètre.
"""
from __future__ import annotations

import datetime as dt
import re
import uuid
from typing import Any, Dict, List, Optional

from hris import adp_transport as TR
from hris.adp_gate import GateDecision
from hris.common import HrisError, PROCESS_CODE, SecretValue, utc_now
from hris.model import Tables
from hris.spark_io import append, merge_rows, rows_df, sql_str


def _F():
    from pyspark.sql import functions as F
    return F


class SparkOutbox:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def load(self, run_id, statuses, max_attempts):
        F = _F()
        rows = (self.spark.table(self.T("pub.adp_outbox")).where(F.col("run_id") == run_id)
                .where(F.col("status").isin(list(statuses)))
                .where(F.coalesce(F.col("attempt_count"), F.lit(0)) < max_attempts)
                .orderBy("employee_key").collect())
        return [r.asDict() for r in rows]

    def newer_items(self, run_id, employee_key):
        """Éléments du même salarié issus de runs promus APRÈS le run donné."""
        F = _F()
        runs = self.spark.table(self.T("ctl.run")).select("run_id", "promoted_at_utc")
        mine = runs.where(F.col("run_id") == run_id).collect()
        if not mine or mine[0].promoted_at_utc is None:
            return []
        later = runs.where(F.col("promoted_at_utc") > F.lit(mine[0].promoted_at_utc)).select("run_id")
        rows = (self.spark.table(self.T("pub.adp_outbox")).where(F.col("employee_key") == employee_key)
                .join(later, "run_id").collect())
        return [r.asDict() for r in rows]

    def update(self, run_id, updates):
        if not updates:
            return
        from delta.tables import DeltaTable
        cols = ["employee_key", "status", "attempt_count", "last_attempt_at_utc", "last_outcome_code",
                "gate_decision_id", "updated_at_utc"]
        rows = [dict({c: u.get(c) for c in cols}, run_id=run_id) for u in updates]
        src = rows_df(self.spark, rows, "pub.adp_outbox")
        (DeltaTable.forName(self.spark, self.T("pub.adp_outbox")).alias("t")
         .merge(src.alias("s"), "t.run_id = s.run_id AND t.employee_key = s.employee_key")
         .whenMatchedUpdate(set={
             "status": "coalesce(s.status, t.status)",
             "attempt_count": "coalesce(s.attempt_count, t.attempt_count)",
             "last_attempt_at_utc": "coalesce(s.last_attempt_at_utc, t.last_attempt_at_utc)",
             "last_outcome_code": "coalesce(s.last_outcome_code, t.last_outcome_code)",
             "gate_decision_id": "coalesce(s.gate_decision_id, t.gate_decision_id)",
             "updated_at_utc": "coalesce(s.updated_at_utc, t.updated_at_utc)"})
         .execute())


class SparkOps:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def get(self, op_id):
        rows = self.spark.table(self.T("pub.adp_operation")).where(_F().col("op_id") == op_id).collect()
        return rows[0].asDict() if rows else None

    def open_uncertain_for_employee(self, employee_key):
        F = _F()
        rows = (self.spark.table(self.T("pub.adp_operation")).where(F.col("employee_key") == employee_key)
                .where(F.col("status").isin("INTENT_RECORDED", "UNCERTAIN")).collect())
        return [r.asDict() for r in rows]

    def for_item(self, run_id, employee_key):
        F = _F()
        rows = (self.spark.table(self.T("pub.adp_operation"))
                .where((F.col("run_id") == run_id) & (F.col("employee_key") == employee_key)).collect())
        return [r.asDict() for r in rows]

    def upsert(self, op):
        merge_rows(self.spark, [op], self.T("pub.adp_operation"), "pub.adp_operation", ["op_id"])


class SparkAttempts:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def append(self, attempt):
        append(rows_df(self.spark, [attempt], "pub.adp_attempt"), self.T("pub.adp_attempt"), "pub.adp_attempt")


class SparkGateLog:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def record(self, run_id, decision: GateDecision):
        row = {"decision_id": decision.decision_id, "run_id": run_id, "evaluated_at_utc": decision.evaluated_at_utc,
               "workspace_name": decision.workspace_name, "workspace_id": decision.workspace_id,
               "is_for_pipeline": decision.is_for_pipeline, "notebook_name": decision.notebook_name,
               "rule_version": decision.rule_version, "decision": decision.env_condition_met,
               "reason_code": decision.reason_code}
        append(rows_df(self.spark, [row], "log.adp_gate_decision"), self.T("log.adp_gate_decision"),
               "log.adp_gate_decision")


class SparkPublicationLog:
    def __init__(self, spark, T: Tables, log_date: Optional[dt.date]):
        self.spark, self.T, self.log_date = spark, T, log_date

    def append(self, rows: List[dict]):
        now = utc_now()
        out = [dict(r, log_id=str(uuid.uuid4()), source_system="FABRIC", log_date=self.log_date,
                    _ingested_at_utc=now) for r in rows]
        append(rows_df(self.spark, out, "log.adp_publication_log"), self.T("log.adp_publication_log"),
               "log.adp_publication_log")


# ---------------------------------------------------------------------------
# Fabriques réelles (refusent tant que l'activation n'est pas approuvée dans le code)
# ---------------------------------------------------------------------------
SECRET_NAMES = {"client_id": "client-id-adp", "client_secret": "client-secret-adp",
                "cert_pem": "MOTUL-SA-cer", "key_pem": "motulsa-auth-key"}


def format_pem(raw: str, label: str) -> str:
    """Équivalent de format_cert / format_rsa_key (L5068-5106), PKCS#1, un seul bloc (B24, parité)."""
    match = re.search(rf"-----BEGIN {label}-----(.*)-----END {label}-----", raw.strip(), re.DOTALL)
    if not match:
        raise HrisError("PEM_INVALID", label)
    body = re.sub(r"\s+", "", match.group(1))
    return f"-----BEGIN {label}-----\n" + "\n".join(body[i:i + 64] for i in range(0, len(body), 64)) + \
        f"\n-----END {label}-----\n"


class RealAdpTransportFactory:
    """Lecture des secrets ADP dans le coffre (mêmes noms que l'existant), contexte mTLS, jeton OAuth, transport.
    Appelée seulement si le garde-fou ET tous les blocages d'activation sont levés ; refuse tant que l'activation n'est
    pas approuvée dans le code (comme l'existant, dont l'activité ADP du pipeline est désactivée)."""

    def __init__(self, kv_url: str, timeout_seconds: float, get_secret, *, context_builder=None, token_class=None,
                 transport_class=None):
        self.kv_url, self.timeout, self._get_secret = kv_url, timeout_seconds, get_secret
        self._context_builder = context_builder or TR.mtls_context
        self._token_class = token_class or TR.OAuthToken
        self._transport_class = transport_class or TR.AdpHttpTransport
        self.built = 0

    def __call__(self):
        self.built += 1
        if not TR.ADP_ACTIVATION_APPROVED:
            raise HrisError("ADP_ACTIVATION_NOT_APPROVED", "aucune lecture de secret ADP ni aucun jeton")
        if not self.kv_url:
            raise HrisError("PARAM_MISSING", "p_kv_url_adp (VL_HRIS.kv_url_adp)")
        values = {key: self._get_secret(self.kv_url, name) for key, name in SECRET_NAMES.items()}
        empty = [SECRET_NAMES[key] for key, value in values.items() if not value]
        if empty:
            raise HrisError("ADP_SECRET_EMPTY", ",".join(empty))
        context = self._context_builder(format_pem(values["cert_pem"], "CERTIFICATE"),
                                        format_pem(values["key_pem"], "RSA PRIVATE KEY"))
        token = self._token_class(context, SecretValue(values["client_id"]), SecretValue(values["client_secret"]),
                                  self.timeout)
        token()                                   # authentification avant tout envoi : un refus n'émet aucun POST
        return self._transport_class(context, token, self.timeout, _allow_construction=True)
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.steps =============
register_hris_module('hris.steps', '0643eecdf0d9af61343b1a29b01d62af4486d6e6d7ff0aceb7a9a070c6b290b3', r'''"""hris.steps — étapes du pipeline (logique des notebooks), appelées à l'identique par les notebooks et par
les tests d'intégration Fabric (schéma isolé). Aucune sortie de donnée personnelle : compteurs seulement."""
from __future__ import annotations

import datetime as dt
import json
import time
from typing import Any, Callable, Dict, List, Mapping, Optional

from hris import compare as C
from hris import spark_io as IO
from hris import transform as S
from hris import adp_publish as PUB
from hris.common import HrisError, PROCESS_CODE, safe_error, utc_now
from hris.model import COMPARE_COLUMNS, Tables
from hris.ts_acquire import completeness_anomalies, to_raw_row

TS_OWNED_CONTROLS = ["EXTRACTION_INCOMPLETE", "SCHEMA_CHANGE", "EMPTY_EXTRACTION", "TS_FILTER_TECHNICAL_FAILURE",
                     "TS_FILTER_NOT_FOUND", "TS_ENDPOINT_FAILURE", "TS_ENDPOINT_NOT_FOUND", "TS_PROCESSING_ERROR",
                     "TS_CIRCUIT_OPEN", "TS_SECRET_RESOLUTION"]
TRANSFORM_OWNED = ["MAPPING_TABLE_MISSING", "MAPPING_SCHEMA", "MAPPING_TABLE_EMPTY", "MAPPING_KEY_DUPLICATE",
                   "TRANSFORM_NOT_EXECUTED", "CONVERSION_ERROR"]


def _F():
    from pyspark.sql import functions as F
    return F


def seed(spark, T: Tables) -> dict:
    """Graines insérées seulement si absentes (cfg.process, verrou, correspondance du flux rémunération)."""
    out = {"process": False, "lock": False, "remuneration_mapping": False}
    if spark.table(T("cfg.ts_remuneration_mapping")).limit(1).count() == 0:
        from hris.remuneration import DEFAULT_MAPPING, mapping_rows
        rows = [dict(r, _loaded_at_utc=utc_now(), _loaded_by="NB_HRIS_SETUP",
                     _change_ref="graine : hypothèses H-REM à confirmer par le mode VALIDATE")
                for r in mapping_rows(DEFAULT_MAPPING)]
        IO.append(IO.rows_df(spark, rows, "cfg.ts_remuneration_mapping"), T("cfg.ts_remuneration_mapping"),
                  "cfg.ts_remuneration_mapping")
        out["remuneration_mapping"] = True
    if spark.table(T("cfg.process")).where(f"process_code = '{PROCESS_CODE}'").count() == 0:
        IO.append(IO.rows_df(spark, [{
            "process_code": PROCESS_CODE, "enabled": True, "acquisition_mode": "TS_API", "publication_target": "ADP_API",
            "key_columns": ["id_unique"], "compare_columns": COMPARE_COLUMNS, "reference_table": "ref.ts_employee_adp",
            "staging_table": "stg.ts_employee_adp", "max_absent_pct": None, "max_modified_pct": None,
            "max_reject_pct": None, "updated_at_utc": utc_now()}], "cfg.process"), T("cfg.process"), "cfg.process")
        out["process"] = True
    if spark.table(T("ctl.process_lock")).where(f"process_code = '{PROCESS_CODE}'").count() == 0:
        IO.append(IO.rows_df(spark, [{"process_code": PROCESS_CODE, "lock_status": "RELEASED",
                                      "released_at_utc": utc_now()}], "ctl.process_lock"),
                  T("ctl.process_lock"), "ctl.process_lock")
        out["lock"] = True
    return out


# ---------------------------------------------------------------------------
# Acquisition : stockage d'une extraction et contrôles (l'extraction réseau est dans le notebook)
# ---------------------------------------------------------------------------
SOURCE_REAL = "TS_PRD_API"      # extraction réelle du tenant TalentSoft de production
SOURCE_FIXTURE = "FIXTURE"       # données fictives (tests isolés uniquement)


def store_acquisition(spark, T: Tables, run: dict, records: List[dict], stats: dict, step: IO.StepLogger,
                      *, source: str) -> dict:
    """Écrit stg.ts_employee_raw et les contrôles. `source` est OBLIGATOIRE et tracé dans ctl.run_step : il permet de
    prouver, dans le Lakehouse, qu'un run est une extraction réelle (TS_PRD_API) et non un jeu fictif (FIXTURE)."""
    if source not in (SOURCE_REAL, SOURCE_FIXTURE):
        raise HrisError("PARAM_INVALID", "source")
    if source == SOURCE_FIXTURE and not getattr(T, "test_schema", ""):
        raise HrisError("FIXTURE_OUTSIDE_TEST_SCHEMA", "des données fictives ne s'écrivent que dans le schéma de test")
    run_id = run["run_id"]
    extracted_at = utc_now()
    raw_rows = [dict(to_raw_row(r), run_id=run_id, business_date=run["business_date"], _extracted_at_utc=extracted_at)
                for r in records]
    distinct_keys = len({r.get("employeeNumber") for r in records})
    IO.replace_run(spark, IO.rows_df(spark, raw_rows, "stg.ts_employee_raw"), T("stg.ts_employee_raw"),
                   "stg.ts_employee_raw", run_id)
    anomalies = completeness_anomalies(stats)
    IO.write_anomalies(spark, T, run_id, anomalies, TS_OWNED_CONTROLS)
    failed = [a["control_code"] for a in anomalies if a["severity"] == "BLOCKING" and a["outcome"] == "FAIL"]
    IO.update_run(spark, T, run_id, status="ACQUIRED")
    step.succeeded(rows_out=len(raw_rows), counters_only=True, message=json.dumps({
        "source": source, "scope": "LegalStructure=Motul France", "mode": "full_extraction_no_period_filter",
        "rows_persisted": len(raw_rows), "distinct_employee_numbers": distinct_keys,
        "stats": stats, "blocking_failures": failed}))
    return {"status": "ACQUIRED" if not failed else "ACQUIRED_INCOMPLETE", "rows": len(raw_rows),
            "distinct_keys": distinct_keys, "total_count": stats.get("total_count"), "blocking_failures": failed}


# ---------------------------------------------------------------------------
# Transformation
# ---------------------------------------------------------------------------
def run_transform(spark, T: Tables, run_id: str) -> dict:
    F = _F()
    run = IO.get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if run["run_mode"] == "RETRY_PUBLICATION":
        return {"run_id": run_id, "status": "SKIPPED_RETRY_PUBLICATION"}
    if run["status"] not in ("ACQUIRED", "TRANSFORMED"):
        return {"run_id": run_id, "status": "SKIPPED_STATUS_" + str(run["status"])}
    step = IO.StepLogger(spark, T, run_id, "TRANSFORM")
    try:
        d_run = run["trigger_time_utc"].date()  # date UTC du run (parité GETDATE()), figée au démarrage
        pred = f"run_id = {IO.sql_str(run_id)}"
        bdate = F.lit(run["business_date"]).cast("date")
        raw = spark.table(T("stg.ts_employee_raw")).where(pred)
        n_raw = raw.count()
        IO.replace_run(spark, S.source_df(raw), T("stg.ts_employee_source"), "stg.ts_employee_source", run_id)
        anomalies, mapping_hash = S.mapping_controls(spark, T)
        blocked = [a["control_code"] for a in anomalies if a["outcome"] == "FAIL"]
        if blocked:
            anomalies.append({"control_code": "TRANSFORM_NOT_EXECUTED", "severity": "BLOCKING", "metric_value": None,
                              "threshold_value": None, "outcome": "FAIL",
                              "message": "enrichissement non exécuté : correspondances invalides"})
            IO.write_anomalies(spark, T, run_id, anomalies, TRANSFORM_OWNED)
            IO.update_run(spark, T, run_id, status="TRANSFORMED")
            step.succeeded(rows_in=n_raw, rows_out=0, content_hash=mapping_hash, counters_only=True,
                           message=json.dumps({"blocked_by": blocked}))
            return {"run_id": run_id, "status": "BLOCKED_MAPPINGS", "rows_raw": n_raw, "blocked_by": blocked}
        source = spark.table(T("stg.ts_employee_source")).where(pred)
        IO.replace_run(spark, S.enriched_df(spark, T, source, d_run), T("stg.ts_employee_enriched"),
                       "stg.ts_employee_enriched", run_id)
        enriched = spark.table(T("stg.ts_employee_enriched")).where(pred)
        try:
            rejects = S.reject_df(enriched, d_run).withColumn("run_id", F.lit(run_id)).withColumn("business_date", bdate)
            IO.replace_run(spark, rejects, T("stg.ts_employee_reject"), "stg.ts_employee_reject", run_id)
            rejects = spark.table(T("stg.ts_employee_reject")).where(pred)
            IO.replace_run(spark, S.masked_log_reject(rejects, run_id), T("log.reject"), "log.reject", run_id)
            adp = S.adp_df(enriched, rejects, d_run).withColumn("run_id", F.lit(run_id)).withColumn("business_date", bdate)
            IO.replace_run(spark, adp, T("stg.ts_employee_adp"), "stg.ts_employee_adp", run_id)
        except Exception as exc:
            if "CONVERSION_" in str(exc):
                anomalies.append({"control_code": "CONVERSION_ERROR", "severity": "BLOCKING", "metric_value": None,
                                  "threshold_value": None, "outcome": "FAIL",
                                  "message": "erreur de conversion T-SQL reproduite (échec du run, D-PAR-02/07)"})
                IO.write_anomalies(spark, T, run_id, anomalies, TRANSFORM_OWNED)
            raise
        IO.write_anomalies(spark, T, run_id, anomalies, TRANSFORM_OWNED)
        rej_t = spark.table(T("stg.ts_employee_reject")).where(pred)
        counts = {"rows_raw": n_raw, "rows_enriched": enriched.count(), "reject_lines": rej_t.count(),
                  "rejected_keys": rej_t.select("id_unique").distinct().count(),
                  "rows_adp": spark.table(T("stg.ts_employee_adp")).where(pred).count()}
        IO.update_run(spark, T, run_id, status="TRANSFORMED")
        step.succeeded(rows_in=n_raw, rows_out=counts["rows_adp"], content_hash=mapping_hash,
                       message=json.dumps(counts), counters_only=True)
        return dict(counts, run_id=run_id, status="TRANSFORMED")
    except Exception as exc:
        step.failed(exc)
        IO.update_run(spark, T, run_id, error_code="TRANSFORM_FAILED", error_summary=safe_error(exc))
        raise


def _check_mode(run, run_mode):
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if run["run_mode"] != run_mode:
        raise HrisError("PARAM_INVALID", "mode différent du mode enregistré pour ce run")


def run_compare(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    run = IO.get_run(spark, T, run_id)
    _check_mode(run, run_mode)
    if run_mode == "RETRY_PUBLICATION":
        return {"run_id": run_id, "decision": "SKIPPED_RETRY_PUBLICATION"}
    if run["status"] not in ("TRANSFORMED", "COMPARED", "BLOCKED", "VALIDATED"):
        return {"run_id": run_id, "decision": "SKIPPED_STATUS_" + str(run["status"])}
    step = IO.StepLogger(spark, T, run_id, "COMPARE")
    try:
        result = dict(C.compare(spark, T, run_id, run_mode), run_id=run_id)
        step.succeeded(message=json.dumps(result, default=str), counters_only=True)
        return result
    except Exception as exc:
        step.failed(exc)
        raise


def run_promote(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    run = IO.get_run(spark, T, run_id)
    _check_mode(run, run_mode)
    if run_mode == "RETRY_PUBLICATION":
        return {"run_id": run_id, "promoted": False, "reason": "SKIPPED_RETRY_PUBLICATION"}
    step = IO.StepLogger(spark, T, run_id, "PROMOTE")
    try:
        result = dict(C.promote(spark, T, run_id, run_mode), run_id=run_id)
        if result.get("promoted"):
            step.succeeded(rows_out=result.get("nb_outbox"), message=json.dumps(result, default=str), counters_only=True)
        else:
            step.skipped(json.dumps(result, default=str))
        return result
    except Exception as exc:
        step.failed(exc)
        raise


# ---------------------------------------------------------------------------
# Publication
# ---------------------------------------------------------------------------
STATUS_MAP = {PUB.STATUS_NOT_PUBLISHED_NON_PRD: ("NOT_PUBLISHED_NON_PRD", None),
              PUB.STATUS_BLOCKED_UNRESOLVED: ("BLOCKED", "PUBLICATION_GATE_UNRESOLVED"),
              PUB.STATUS_BLOCKED_ACTIVATION: ("BLOCKED", "PUBLICATION_ACTIVATION_BLOCKED"),
              PUB.STATUS_NOTHING: ("PUBLISHED", None), PUB.STATUS_PUBLISHED: ("PUBLISHED", None),
              PUB.STATUS_PARTIAL: ("PUBLISHED_PARTIAL", None), PUB.STATUS_HALTED: ("PUBLISHED_PARTIAL", "AUTH_REJECTED")}


class _NoSentinel:
    """Sentinelle neutre (publication approuvée) : même interface, aucune interception."""
    attempts: List[Dict[str, str]] = []

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def run_publish(spark, T: Tables, run_id: str, run_mode: str, *, context_provider: Callable[[], Mapping[str, Any]],
                transport_factory, planner, max_attempts: int, pause_seconds: float,
                extra_blockers: Callable[[], List[str]] = lambda: [], sleep: Callable[[float], None] = time.sleep) -> dict:
    """Publication du run (ou du run cible en RETRY_PUBLICATION) sous sentinelle réseau bloquante."""
    from hris.adp_spark import SparkAttempts, SparkGateLog, SparkOps, SparkOutbox, SparkPublicationLog
    from hris.adp_transport import ADP_ACTIVATION_APPROVED, NetworkSentinel
    run = IO.get_run(spark, T, run_id)
    _check_mode(run, run_mode)
    result = {"run_id": run_id, "mode": run_mode}
    if run_mode == "INIT_REFERENCE":
        return dict(result, status="SKIPPED_INIT_REFERENCE")          # TG-03 : aucune étape de publication
    if run_mode == "NORMAL" and run["status"] != "PROMOTED":
        return dict(result, status="SKIPPED_NOT_PROMOTED_" + str(run["status"]))
    source_run_id = run["target_run_id"] if run_mode == "RETRY_PUBLICATION" else run_id
    source_run = IO.get_run(spark, T, source_run_id) or {}

    def run_blockers():
        blockers = list(C.revalidate(spark, T, source_run_id))
        if all(v is None for v in C.load_thresholds(spark, T).values()):
            blockers.append("D-CMP-07_SEUILS_NON_FOURNIS")
        return blockers

    def rows_by_key():
        rows = spark.table(T("stg.ts_employee_adp")).where(f"run_id = {IO.sql_str(source_run_id)}").collect()
        return {r["id_unique"]: r.asDict() for r in rows}

    calls = {"factory": 0}

    def counted_factory():
        calls["factory"] += 1
        return transport_factory()

    ctx = PUB.PublicationContext(publish_run_id=run_id, source_run_id=source_run_id, mode=run_mode,
                                 max_attempts=max_attempts, pause_seconds=pause_seconds)
    step = IO.StepLogger(spark, T, run_id, "PUBLISH_ADP")
    # Sentinelle bloquante tant que l'activation n'est pas approuvée (preuve de non-publication) ; une fois approuvée,
    # elle bloquerait les envois légitimes et l'INSEE : elle n'est plus posée, le garde-fou et les blocages décident.
    sentinel = NetworkSentinel() if not ADP_ACTIVATION_APPROVED else _NoSentinel()
    try:
        with sentinel:
            summary = PUB.run_publication(
                ctx, context_provider=context_provider, outbox=SparkOutbox(spark, T), gate_log=SparkGateLog(spark, T),
                ops=SparkOps(spark, T), attempts=SparkAttempts(spark, T),
                publication_log=SparkPublicationLog(spark, T, source_run.get("business_date")),
                activation_blockers=lambda: PUB.code_activation_blockers() + run_blockers() + list(extra_blockers()),
                transport_factory=counted_factory, planner=planner, rows_by_key=rows_by_key, sleep=sleep)
        proof = {"gate_outcome": summary["gate_outcome"], "gate_reason_code": summary["gate_reason_code"],
                 "workspace_name": summary["workspace_name"], "transport_factory_calls": calls["factory"],
                 "transport_built": summary["transport_built"], "sentinel_blocking": not ADP_ACTIVATION_APPROVED,
                 "sentinel_attempts": len(sentinel.attempts),
                 "sentinel_hosts": sorted({a["host"] for a in sentinel.attempts}),
                 "adp_activation_approved": ADP_ACTIVATION_APPROVED}
        run_status, error_code = STATUS_MAP[summary["status"]]
        fields = {"status": run_status}
        if error_code:
            fields.update(error_code=error_code, error_summary=",".join(summary.get("activation_blockers") or []))
        IO.update_run(spark, T, run_id, **fields)
        step.succeeded(rows_in=summary["items_selected"], counters_only=True,
                       message=json.dumps({"summary": summary, "proof": proof}, default=str))
        return dict(result, status=summary["status"], summary=summary, proof=proof)
    except Exception as exc:
        step.failed(exc)
        IO.update_run(spark, T, run_id, error_code="PUBLISH_FAILED", error_summary=safe_error(exc))
        raise
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.legacy_logs =============
register_hris_module('hris.legacy_logs', '23e7cd84abb562bf3097d258b4c17af8515a2074394ff3ecd146c2ab8d6c1409', r'''"""hris.legacy_logs — lecture des logs historiques de l'ancienne Function ADP (import unique, D-LOG-00).

Reprise des fonctions du prototype NB_HRIS_IMPORT_LOGS_HISTORIQUES (TEST_2/notebooks) sous forme
testable. Lecture seule : aucun fichier n'est produit, modifié ni supprimé. Les lignes rejetées sont
tracées sans valeur. Les zéros de tête de matricule_id ne sont pas recomplétés.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import os
import re
import uuid
from typing import Dict, List, Optional, Tuple

EXPECTED_COLUMNS = ["ts_adp", "matricule_id", "nom", "mode", "field", "status_code",
                    "reason_code", "response_code", "error_message", "log_message"]
CANDIDATE_DELIMITERS = [";", "\t", ","]
TS_FORMATS = ["%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f",
              "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"]
DATE_FOLDER = re.compile(r"(?:^|/)(\d{8})(?:/|$)")
SOURCE_SYSTEM = "LEGACY_ADP_FUNCTION"


def list_files(root: str) -> List[dict]:
    found = []
    for dirpath, _, filenames in os.walk(root):
        for name in filenames:
            full = os.path.join(dirpath, name)
            st = os.stat(full)
            found.append({"full": full, "rel": os.path.relpath(full, root).replace(os.sep, "/"), "name": name,
                          "size": int(st.st_size),
                          "mtime": dt.datetime.fromtimestamp(st.st_mtime, tz=dt.timezone.utc)})
    return sorted(found, key=lambda f: f["rel"])


def parse_ts(value: Optional[str]) -> Optional[dt.datetime]:
    text = (value or "").strip()
    for fmt in TS_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def read_delimited(path: str) -> Tuple[Optional[List[str]], Optional[List[List[str]]], str]:
    with open(path, "rb") as handle:
        raw = handle.read()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None, None, "ENCODING_NOT_SUPPORTED"
    first_line = text.splitlines()[0] if text else ""
    for delimiter in CANDIDATE_DELIMITERS:
        header = [h.strip() for h in next(csv.reader([first_line], delimiter=delimiter))]
        if set(EXPECTED_COLUMNS).issubset(header):
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter, quotechar='"'))
            return header, rows[1:], f"text/{encoding}/{delimiter!r}"
    return None, None, "HEADER_MISMATCH"


def read_excel(path: str, sheet_name: str = "") -> Tuple[Optional[List[str]], Optional[List[List[str]]], str]:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return None, None, "OPENPYXL_MISSING"
    import pandas as pd
    df = pd.read_excel(path, sheet_name=sheet_name or 0, dtype=str, engine="openpyxl")
    header = [str(c).strip() for c in df.columns]
    if not set(EXPECTED_COLUMNS).issubset(header):
        return None, None, "HEADER_MISMATCH"
    return header, df.fillna("").astype(str).values.tolist(), "xlsx"


def folder_date(rel_path: str) -> Optional[dt.date]:
    match = DATE_FOLDER.search(rel_path)
    if not match:
        return None
    try:
        return dt.datetime.strptime(match.group(1), "%Y%m%d").date()
    except ValueError:
        return None


def process_folder(local_root: str, uri_prefix: str, extensions: set, already_imported: set, import_id: str,
                   executed_by: Optional[str], workspace_name: Optional[str], sheet_name: str = "",
                   now: Optional[dt.datetime] = None, target_table: str = "log.adp_publication_log"):
    """Analyse un dossier ; renvoie (enregistrements, rejets sans valeur, registre des fichiers)."""
    now = now or dt.datetime.now(dt.timezone.utc)
    records, rejects, ledger = [], [], []
    for f in list_files(local_root):
        source_uri = f"{uri_prefix}/{f['rel']}"
        entry = {"import_id": import_id, "source_uri": source_uri, "file_name": f["name"], "file_size": f["size"],
                 "file_modified_at_utc": f["mtime"], "file_format": None, "status": None, "detail": None,
                 "rows_read": 0, "rows_loaded": 0, "rows_rejected": 0, "target_table": target_table,
                 "executed_by": executed_by, "workspace_name": workspace_name, "processed_at_utc": now}
        extension = os.path.splitext(f["name"])[1].lower()
        if source_uri in already_imported:
            entry.update(status="SKIPPED_ALREADY_IMPORTED")
        elif extension not in extensions:
            entry.update(status="SKIPPED_UNSUPPORTED", detail=f"extension {extension or '(aucune)'}")
        else:
            try:
                header, rows, fmt = read_excel(f["full"], sheet_name) if extension == ".xlsx" else read_delimited(f["full"])
            except Exception as exc:
                header, rows, fmt = None, None, f"READ_ERROR:{type(exc).__name__}"
            entry["file_format"] = fmt
            if header is None:
                entry.update(status="REJECTED_SCHEMA", detail=fmt)
            else:
                position = {name: header.index(name) for name in EXPECTED_COLUMNS}
                log_date_from_folder = folder_date(f["rel"])
                for row_number, row in enumerate(rows, start=1):
                    if not any(str(cell).strip() for cell in row):
                        continue
                    entry["rows_read"] += 1
                    values = {c: str(row[i]) if i < len(row) else "" for c, i in position.items()}
                    ts = parse_ts(values["ts_adp"])
                    reason = ("COLUMN_COUNT_MISMATCH" if len(row) != len(header)
                              else "TS_ADP_INVALID" if ts is None else None)
                    if reason:
                        entry["rows_rejected"] += 1
                        rejects.append({"import_id": import_id, "source_uri": source_uri,
                                        "source_row_number": row_number, "reject_reason": reason,
                                        "rejected_at_utc": now})
                        continue
                    payload = "\u001f".join(values[c] for c in EXPECTED_COLUMNS)
                    records.append({
                        "log_id": str(uuid.uuid4()), "source_system": SOURCE_SYSTEM, "run_id": None,
                        "log_date": log_date_from_folder or ts.date(), "ts_adp": ts,
                        **{c: values[c] for c in EXPECTED_COLUMNS if c != "ts_adp"},
                        "employee_key": None, "step_name": None, "http_method": None, "endpoint_template": None,
                        "attempt_no": None, "gate_decision_id": None, "_import_id": import_id,
                        "_source_file_uri": source_uri, "_source_row_number": row_number,
                        "_row_hash": hashlib.sha256(payload.encode("utf-8")).hexdigest(), "_ingested_at_utc": now})
                    entry["rows_loaded"] += 1
                entry["status"] = "IMPORTED"
        ledger.append(entry)
    return records, rejects, ledger
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.cell_log =============
register_hris_module('hris.cell_log', '753f215b9da90c692337f729a62af15d6baadf7882559746639b6f9346aeddc9', r'''"""hris.cell_log — journal des cellules de code exécutées (table log.cell_execution).

Une ligne par cellule de code terminée, avec le contexte non sensible du notebook :
  * SUCCEEDED : le code Python de la cellule s'est terminé sans exception. Les écritures Delta lancées
    par la cellule sont alors terminées ; un DataFrame seulement déclaré (calcul Spark différé) n'est PAS
    validé par ce statut. Traces optionnelles (paramètre p_log_cell_success des notebooks).
  * FAILED : toujours journalisé, AVANT que la cellule ne propage l'erreur.
Une cellule qui n'a pas été exécutée ne produit aucune ligne.

Coût : aucune action Spark (ni count, ni collect) n'est lancée pour produire un événement. Les événements
restent en mémoire et sont écrits en UNE seule écriture Delta (append) : à la clôture du notebook pour les
réussites, immédiatement en cas d'échec (réussites en attente + échec). Tant que cette écriture n'a pas eu
lieu, une trace n'est PAS persistée.

Robustesse : une écriture de journal en échec ne lève jamais d'exception, l'erreur métier n'est donc jamais
masquée ; un diagnostic expurgé est affiché et `write_failed` passe à True (limite de traçabilité).
Contenu : identifiants techniques, horodatages, statut, type d'erreur, message assaini (littéraux entre
guillemets masqués, puis hris.common.sanitize) et emplacement fichier:fonction:ligne. Jamais de secret, de
jeton, de ligne RH ni de valeur de paramètre.
"""
from __future__ import annotations

import os
import re
import time
import traceback
import uuid
from typing import Any, Callable, List, Optional

from hris.common import HrisError, safe_error, sanitize, utc_now

LOGICAL_TABLE = "log.cell_execution"
STATUS_SUCCEEDED = "SUCCEEDED"   # mêmes statuts que ctl.run_step
STATUS_FAILED = "FAILED"

_QUOTED = [re.compile(r"'[^'\n]*'"), re.compile(r'"[^"\n]*"')]


def context_value(context: Any, key: str) -> Optional[str]:
    """Valeur du contexte d'exécution Fabric (notebookutils.runtime.context), None si absente ou illisible."""
    try:
        value = context.get(key)
    except Exception:
        return None
    return None if value in (None, "") else str(value)


def _context_flag(context: Any, key: str) -> Optional[bool]:
    try:
        value = context.get(key)
    except Exception:
        return None
    if value is None or value == "":
        return None
    return value if isinstance(value, bool) else str(value).strip().lower() == "true"


def redact_error(error: BaseException) -> str:
    """Raison publiable d'une erreur. HrisError : code + message construit par le code (safe_error).
    Autres erreurs : classe + message dont les littéraux entre guillemets (valeurs possibles) sont masqués,
    puis assaini (URL, e-mails, IBAN, jetons, nombres de 4 chiffres ou plus)."""
    if isinstance(error, HrisError):
        return safe_error(error)
    text = " ".join(str(error).split())
    for pattern in _QUOTED:
        text = pattern.sub("'***'", text)
    if isinstance(error, ImportError) and error.name:   # nom de module : identifiant de code, jamais une donnée
        text = f"{text} (module {error.name})"
    return f"{type(error).__name__}: {sanitize(text, 300)}"


def error_location(error: BaseException) -> str:
    """Dernière frame de la bibliothèque hris (sinon la plus profonde) : fichier:fonction:ligne, sans valeur."""
    frames = traceback.extract_tb(error.__traceback__) if error.__traceback__ is not None else []
    if not frames:
        return ""
    chosen = next((f for f in reversed(frames) if f.filename.replace("\\", "/").startswith("hris/")), frames[-1])
    name = chosen.filename.replace("\\", "/")
    name = name if name.startswith("hris/") else os.path.basename(name)
    return sanitize(f"{name}:{chosen.name}:{chosen.lineno}", 200, mask_numbers=False)


def report_unlogged_failure(cell_code: str, error: BaseException) -> None:
    """Échec survenu avant l'ouverture du journal (destination inconnue) : diagnostic expurgé affiché."""
    print(f"[cell_log] ❌ cellule {cell_code} en échec avant l'ouverture du journal : {redact_error(error)} "
          f"— trace NON persistée (destination de journal inconnue)")


class CellLog:
    """Journal des cellules d'UNE exécution de notebook : start() / success() / failure() puis flush()."""

    def __init__(self, context: Any, *, lakehouse_name: str, run_id: str = "", log_success: bool = True,
                 writer: Optional[Callable[[List[dict]], None]] = None,
                 clock: Callable[[], float] = time.perf_counter):
        self.base = {
            "run_id": (run_id or "")[:100] or None,
            "operation_id": context_value(context, "activityId"),
            "notebook_name": context_value(context, "currentNotebookName"),
            "notebook_id": context_value(context, "currentNotebookId"),
            "workspace_name": context_value(context, "currentWorkspaceName"),
            "workspace_id": context_value(context, "currentWorkspaceId"),
            "lakehouse_name": lakehouse_name or None,
            "is_for_pipeline": _context_flag(context, "isForPipeline"),
        }
        self.log_success = bool(log_success)
        self.write_failed = False
        self.events_written = 0
        self.writes = 0
        self._writer = writer
        self._clock = clock
        self._pending: List[dict] = []
        self._current = None
        self._seq = 0

    def start(self, cell_code: str) -> None:
        """Ouvre la cellule `cell_code` (identifiant stable) : horodatage de début, aucune écriture."""
        self._current = (cell_code, utc_now(), self._clock())

    def success(self) -> None:
        """Clôt la cellule ouverte en SUCCEEDED ; gardée en mémoire (si l'option est active) jusqu'au flush."""
        event = self._close(STATUS_SUCCEEDED, None)
        if event is not None and self.log_success:
            self._pending.append(event)

    def failure(self, error: BaseException) -> None:
        """Clôt la cellule ouverte en FAILED et écrit tout de suite (réussites en attente + échec).
        Ne lève jamais d'exception : la cellule propage ensuite l'erreur d'origine par `raise`."""
        try:
            event = self._close(STATUS_FAILED, error)
            if event is None:  # échec hors cellule ouverte : tracé quand même, jamais ignoré
                self.start("UNKNOWN_CELL")
                event = self._close(STATUS_FAILED, error)
            self._pending.append(event)
        except Exception as exc:  # pragma: no cover - garde ultime : ne jamais masquer l'erreur métier
            print(f"[cell_log] ⚠ événement d'échec non construit : {type(exc).__name__}")
            return
        self.flush()

    def flush(self) -> None:
        """Écrit les événements en attente en UNE écriture. Ne lève jamais d'exception."""
        if not self._pending:
            return
        events, self._pending = self._pending, []
        now = utc_now()
        for event in events:
            event["logged_at_utc"] = now
        if self._writer is None:
            self._report_unwritten(events, "aucune destination de journal configurée")
            return
        try:
            self._writer(events)
            self.writes += 1
            self.events_written += len(events)
        except Exception as exc:
            self.write_failed = True
            self._report_unwritten(events, redact_error(exc))

    def _close(self, status: str, error: Optional[BaseException]) -> Optional[dict]:
        if self._current is None:
            return None
        cell_code, started_at, started_clock = self._current
        self._current = None
        self._seq += 1
        event = dict(self.base, event_id=str(uuid.uuid4()), cell_code=cell_code, cell_seq=self._seq,
                     status=status, started_at_utc=started_at, ended_at_utc=utc_now(),
                     duration_ms=max(0, int(round((self._clock() - started_clock) * 1000))),
                     error_type=None, error_code=None, error_message=None, error_location=None)
        if error is not None:
            event.update(error_type=type(error).__name__,
                         error_code=error.code if isinstance(error, HrisError) else None,
                         error_message=redact_error(error), error_location=error_location(error))
        return event

    def _report_unwritten(self, events: List[dict], reason: str) -> None:
        cells = ", ".join(f"{e['cell_code']}={e['status']}" + (f"({e['error_type']})" if e["error_type"] else "")
                          for e in events)
        print(f"[cell_log] ⚠ journal NON persisté ({len(events)} événement(s)) : {reason} ; cellules : {cells}")


def spark_writer(spark, tables) -> Callable[[List[dict]], None]:
    """Écriture Delta en ajout dans log.cell_execution (nom à 4 parties résolu par `tables`).
    La table est créée par NB_HRIS_SETUP ; à défaut, saveAsTable en mode append la crée au même schéma."""
    from hris import spark_io as IO

    def write(events: List[dict]) -> None:
        IO.append(IO.rows_df(spark, events, LOGICAL_TABLE), tables(LOGICAL_TABLE), LOGICAL_TABLE)
    return write
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.reference_api =============
register_hris_module('hris.reference_api', '0cc94f15aced3436e95bc761518740b1729a41c2f0bb26fce66d5eb15c38970d', r'''"""hris.reference_api — tables de correspondance cfg.lu_hr_* sans aucun fichier : partie TalentSoft lue par l'API.

Constat (sonde en lecture seule du 13/10/2026 sur le tenant de production, métadonnées seulement) :
  * les 49 référentiels du Core HR (GET /api/hub/v2.0/libraries/{type}) portent clientCode, ShortName (libellé traduit
    par langue), FullName, IsActive (Y/N), ParentClientCode et quelques champs propres au type ; AUCUN code paie ou ADP ;
  * les fiches salariés portent le clientCode du référentiel (« Motul France » est le clientCode de LegalStructure) :
    c'est la clé de jointure de la transformation ;
  * le sexe, la nationalité et les pays ne sont pas des référentiels : leurs codes sont relevés dans l'extraction.

Conséquence, sans aucun CSV (A08 D-FIL-00, D-REF-01) :
  * la partie TalentSoft des 17 tables (code, libellés, parent) est synchronisée à chaque run depuis l'API et
    l'extraction : codes nouveaux ajoutés, libellés mis à jour ;
  * la partie ADP (codes, libellés et ordre ADP) n'existe nulle part ailleurs que dans ces tables : elle est saisie par
    le métier (NB_HRIS_REFERENCE_ADMIN, `apply_assignments`). La synchronisation n'écrit JAMAIS une colonne ADP et ne
    supprime jamais une ligne ; un code TalentSoft sans code ADP fait rejeter les salariés concernés (Référence_Rejet),
    comme une correspondance absente dans l'existant.
Aucune valeur n'est journalisée : noms de tables, comptes et statuts seulement.
"""
from __future__ import annotations

import datetime as dt
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple
from urllib.parse import quote

from hris import spark_io as IO
from hris.common import HrisError, SEVERITY_INFO, SEVERITY_WARNING, utc_now
from hris.model import LOOKUP_JOINS, LOOKUP_TABLES, LOOKUP_TECH_COLUMNS, SOURCE_COLUMNS, USED_LOOKUP_TABLES, create_table_sql

LIBRARY_PATH = "/api/hub/v2.0/libraries/{library}"   # droit requis : ADM_API_EXTRACTION_API_ACCESS (comme l'extraction)
LABEL_CULTURES = ("fr-fr", "en-gb")                   # libellé : français, sinon anglais, sinon la première langue
MAX_JOIN_KEY = 50                                     # la transformation tronque les valeurs source à 50 caractères


@dataclass(frozen=True)
class CatalogueSpec:
    """Origine de la partie TalentSoft d'une table : référentiel (ou None) et clés de l'extraction relevées."""
    table: str
    library: Optional[str]
    observed_keys: Tuple[str, ...]
    code: str
    columns: Tuple[Tuple[str, str], ...] = ()         # (colonne, attribut : label | label_fr | label_en | full | parent)

    @property
    def ts_columns(self) -> List[str]:
        return [self.code] + [column for column, _ in self.columns]


CATALOGUE: Dict[str, CatalogueSpec] = {spec.table: spec for spec in (
    CatalogueSpec("lu_hr_emp_titre", "Title", ("Title",), "cd_emp_titre_ts",
                  (("lb_emp_titre_ts", "label"), ("ds_emp_titre_ts", "full"))),
    CatalogueSpec("lu_hr_emp_sexe", None, ("Gender",), "cd_emp_sexe_ts"),
    CatalogueSpec("lu_hr_emp_nationalite", None, ("Nationality",), "cd_emp_nationalite_ts"),
    CatalogueSpec("lu_hr_emp_situation_matrimoniale", "MaritalStatus", ("MaritalStatus",),
                  "cd_emp_situation_matrimoniale_ts", (("ds_emp_situation_matrimoniale_ts", "label"),)),
    CatalogueSpec("lu_hr_emp_pays", None, ("Pays_de_naissance", "adresse_countryCode", "LegalCountry"), "cd_emp_pays_ts"),
    CatalogueSpec("lu_hr_ppc_raison_debut_contrat", "HiringReason", ("HiringReason",), "cd_ppc_raison_debut_contrat_ts",
                  (("ds_ppc_raison_debut_contrat_ts", "label"),)),
    CatalogueSpec("lu_hr_ppc_contrat2", "NatureOfContract", ("NatureOfContract",), "cd_contract_ts",
                  (("lb_contract_ts", "label"),)),
    CatalogueSpec("lu_hr_ppc_classification", "Coefficient", ("Coefficient",), "cd_convention_collective_lvl2_ts",
                  (("lb_convention_collective_lvl2_ts", "label"),)),
    CatalogueSpec("lu_hr_ppc_classes", "OccupationalCategory", ("OccupationalCategory",), "cd_occupational_category_ts",
                  (("lb_fr_occupational_category_ts", "label_fr"), ("lb_en_occupational_category_ts", "label_en"))),
    CatalogueSpec("lu_hr_geo_entite", "LegalStructure", ("LegalStructure",), "cd_geo_entite_ts",
                  (("ds_geo_entite_ts", "label"), ("cd_geo_parent_entite_ts", "parent"))),
    CatalogueSpec("lu_hr_ppc_location_country", "GeographicOrganisationStructure", ("GeographicOrganisationStructure",),
                  "cd_ppc_location_country_ts", (("ds_ppc_location_country_ts", "label"),)),
    CatalogueSpec("lu_hr_ppc_cost_center", "CostCenter", ("CostCenter",), "cd_ppc_cost_center_ts",
                  (("ds_ppc_cost_center_ts", "label"),)),
    CatalogueSpec("lu_hr_organisationalstructure", "OrganisationalStructure", ("OrganisationalStructure",),
                  "cd_organisationalstructure_ts", (("ds_organisationalstructure_ts", "label"),)),
    # Tables chargées par l'existant mais jamais jointes (D-REF-02)
    CatalogueSpec("lu_hr_emp_diplome", "Diploma", (), "cd_emp_diplome_ts", (("ds_emp_diplome_ts", "label"),)),
    CatalogueSpec("lu_hr_emp_work_accident", "WorkAccident", (), "cd_emp_work_accident_ts",
                  (("ds_emp_work_accident_ts", "label"),)),
    CatalogueSpec("lu_hr_ppc_convention_collective", "CollectiveAgreement", (), "cd_ppc_convention_collective_ts",
                  (("ds_ppc_convention_collective_ts", "label"),)),
    CatalogueSpec("lu_hr_ppc_contrat", "NatureOfContract", (), "cd_ppc_contrat_ts", (("ds_ppc_contrat_ts", "label"),)),
)}
assert sorted(CATALOGUE) == sorted(LOOKUP_TABLES)


def adp_columns(table: str) -> List[str]:
    """Colonnes ADP qui comptent pour la paie : celles que la transformation ramène (tables utilisées), sinon toutes les
    colonnes « *_adp » hors ordre d'affichage."""
    joined = sorted({lu for join in LOOKUP_JOINS if join["table"] == table for lu, _ in join["out"] if lu != join["key"]})
    return joined or [c for c in LOOKUP_TABLES[table] if c.endswith("_adp") and not c.startswith("or_")]


def business_columns(table: str) -> List[str]:
    """Colonnes saisies par le métier : tout sauf la partie TalentSoft synchronisée (code, libellés, parent)."""
    managed = set(CATALOGUE[table].ts_columns)
    return [c for c in LOOKUP_TABLES[table] if c not in managed]


_KEY_TO_COLUMN = dict(SOURCE_COLUMNS)
for _spec in CATALOGUE.values():
    assert set(_spec.ts_columns) <= set(LOOKUP_TABLES[_spec.table]), _spec.table
    assert not set(_spec.ts_columns) & set(adp_columns(_spec.table)), _spec.table
    if _spec.table in USED_LOOKUP_TABLES:   # clés relevées = exactement les colonnes source jointes par la transformation
        assert sorted(_KEY_TO_COLUMN[k] for k in _spec.observed_keys) == sorted(
            join["src"] for join in LOOKUP_JOINS if join["table"] == _spec.table), _spec.table


def _norm(code: Optional[str]) -> Optional[str]:
    """Clé comparée comme dans la jointure : espaces de fin ignorés ; vide = pas de clé."""
    if code is None:
        return None
    text = str(code).rstrip(" ")
    return text or None


def _filled(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


# ---------------------------------------------------------------------------
# Lecture d'un référentiel (lignes clientCode / field / culture / value)
# ---------------------------------------------------------------------------
class LibraryUnavailable(HrisError):
    def __init__(self, library: str, status: Any):
        super().__init__("TS_LIBRARY_UNAVAILABLE", f"{library} : statut {status}")
        self.library = library


def fetch_library(client, library: str) -> List[dict]:
    """GET du référentiel par le client d'acquisition (jeton, reprises et disjoncteur compris)."""
    rows, status = client.get(LIBRARY_PATH.format(library=quote(library, safe="")), "library_nodes")
    if status != 200 or not isinstance(rows, list):
        raise LibraryUnavailable(library, status)
    return rows


def _pick(values: Mapping[str, Any], order: Sequence[str]) -> Optional[str]:
    for culture in order:
        if _filled(values.get(culture)):
            return str(values[culture]).strip()
    for culture in sorted(values):
        if _filled(values[culture]):
            return str(values[culture]).strip()
    return None


def library_nodes(rows: Iterable[Mapping[str, Any]], today: Optional[dt.date] = None) -> Dict[str, Dict[str, Optional[str]]]:
    """Nœuds d'un référentiel : dernière version (date d'effet <= aujourd'hui) de chaque champ et de chaque langue."""
    limit = (today or utc_now().date()).isoformat()
    latest: Dict[Tuple[str, str, str], Tuple[str, Any]] = {}
    for row in rows:
        code = row.get("clientCode")
        if not _filled(code):
            continue
        date = str(row.get("startEffectDate") or "")[:10]
        if date and date > limit:
            continue                                   # version future : pas encore en vigueur
        key = (str(code), str(row.get("field") or ""), str(row.get("culture") or "").lower())
        if key not in latest or date >= latest[key][0]:
            latest[key] = (date, row.get("value"))
    fields: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for (code, name, culture), (_, value) in latest.items():
        fields.setdefault(code, {}).setdefault(name, {})[culture] = value
    nodes = {}
    for code, values in fields.items():
        short = values.get("ShortName", {})
        label = _pick(short, LABEL_CULTURES)
        nodes[code] = {
            "label": label,
            "label_fr": _pick({"fr-fr": short.get("fr-fr")}, ("fr-fr",)),
            "label_en": _pick({"en-gb": short.get("en-gb")}, ("en-gb",)),
            "full": _pick(values.get("FullName", {}), LABEL_CULTURES) or label,
            "parent": _pick(values.get("ParentClientCode", {}), ("",)),
            "active": _pick(values.get("IsActive", {}), ("",)),
        }
    return nodes


# ---------------------------------------------------------------------------
# Fusion de la partie TalentSoft dans une table (jamais de colonne ADP, jamais de suppression)
# ---------------------------------------------------------------------------
@dataclass
class TableSync:
    table: str
    library: Optional[str]
    status: str = "OK"            # OK | SKIPPED_LIBRARY_UNAVAILABLE | SKIPPED_SCHEMA | SKIPPED_DUPLICATE_KEYS
    reason: str = ""
    library_codes: int = 0
    observed_codes: int = 0
    observed_not_in_library: int = 0
    inserted: int = 0
    updated: int = 0
    kept_absent: int = 0          # lignes existantes dont le code n'est ni dans le référentiel ni dans l'extraction
    codes_over_50: int = 0        # codes que la jointure ne peut pas retrouver (valeur source tronquée à 50)
    rows_after: int = 0
    without_adp: int = 0          # lignes sans aucun code ADP utile : à compléter dans NB_HRIS_REFERENCE_ADMIN
    written: bool = False


def merge_catalogue(spec: CatalogueSpec, rows: Sequence[Mapping[str, Any]], nodes: Mapping[str, Mapping[str, Any]],
                    observed: Iterable[str], technical: Mapping[str, Any]) -> Tuple[List[dict], TableSync]:
    """Ajoute les codes nouveaux et met à jour les libellés TalentSoft ; ne touche à aucune autre colonne."""
    sync = TableSync(spec.table, spec.library)
    columns = LOOKUP_TABLES[spec.table]
    out = [dict(row) for row in rows]
    index: Dict[str, int] = {}
    duplicates = 0
    for position, row in enumerate(out):
        key = _norm(row.get(spec.code))
        if key is None:
            continue
        if key in index:
            duplicates += 1
        index[key] = position
    if duplicates:
        sync.status, sync.reason = "SKIPPED_DUPLICATE_KEYS", f"{duplicates} clé(s) en double à corriger d'abord"
        sync.rows_after = len(out)
        return list(rows), sync
    observed_set = {k for k in (_norm(v) for v in observed) if k}
    library_set = {k for k in (_norm(c) for c in nodes) if k}
    sync.library_codes, sync.observed_codes = len(library_set), len(observed_set)
    sync.observed_not_in_library = len(observed_set - library_set) if spec.library else 0
    by_key = {_norm(code): attrs for code, attrs in nodes.items() if _norm(code)}
    for key in sorted(library_set | observed_set):
        attrs = by_key.get(key, {})
        wanted = {column: attrs.get(attribute) for column, attribute in spec.columns}
        if key in index:
            row = out[index[key]]
            changes = {c: v for c, v in wanted.items() if v is not None and row.get(c) != v}
            if changes:
                row.update(changes, **technical)
                sync.updated += 1
        else:
            row = {c: None for c in columns}
            row.update({spec.code: key}, **{c: v for c, v in wanted.items()}, **technical)
            out.append(row)
            index[key] = len(out) - 1
            sync.inserted += 1
    sync.kept_absent = len(set(index) - library_set - observed_set)
    sync.codes_over_50 = sum(1 for key in index if len(key) > MAX_JOIN_KEY)
    useful = adp_columns(spec.table)
    sync.without_adp = sum(1 for row in out if _norm(row.get(spec.code)) and not any(_filled(row.get(c)) for c in useful))
    sync.rows_after = len(out)
    return out, sync


def observed_from_records(records: Iterable[Mapping[str, Any]], keys: Sequence[str]) -> List[str]:
    """Codes distincts relevés dans des enregistrements bruts (version Python de SparkReferenceStore.observed)."""
    return sorted({_norm(r.get(k)) for r in records for k in keys if _norm(r.get(k))})


# ---------------------------------------------------------------------------
# Saisie des codes ADP par le métier (NB_HRIS_REFERENCE_ADMIN)
# ---------------------------------------------------------------------------
@dataclass
class AssignmentResult:
    table: str
    changes: List[Tuple[str, str, Optional[str], Optional[str]]] = field(default_factory=list)  # code, colonne, avant, après
    errors: List[str] = field(default_factory=list)


def apply_assignments(table: str, rows: Sequence[Mapping[str, Any]], assignments: Mapping[str, Mapping[str, Any]],
                      technical: Mapping[str, Any]) -> Tuple[List[dict], AssignmentResult]:
    """Renseigne des colonnes métier (codes, libellés, ordre ADP) pour des codes TalentSoft existants.
    Tout ou rien : à la moindre erreur, aucune ligne n'est modifiée."""
    result = AssignmentResult(table)
    if table not in CATALOGUE:
        result.errors.append(f"table inconnue : {table}")
        return list(rows), result
    spec = CATALOGUE[table]
    editable = set(business_columns(table))
    out = [dict(row) for row in rows]
    index: Dict[str, int] = {}
    for position, row in enumerate(out):
        key = _norm(row.get(spec.code))
        if key is None:
            continue
        if key in index:
            result.errors.append(f"clé en double dans la table : {key}")
        index[key] = position
    if not isinstance(assignments, Mapping):
        result.errors.append("les affectations doivent être un objet {code TalentSoft: {colonne: valeur}}")
        return list(rows), result
    for code, values in assignments.items():
        key = _norm(code)
        if key not in index:
            result.errors.append(f"code TalentSoft absent de la table : {code}")
            continue
        if not isinstance(values, Mapping) or not values:
            result.errors.append(f"aucune colonne pour {code}")
            continue
        for column, value in values.items():
            if column not in editable:
                result.errors.append(f"colonne non modifiable pour {table} : {column}")
            elif value is not None and (isinstance(value, bool) or not isinstance(value, (str, int))):
                result.errors.append(f"valeur non textuelle pour {code}.{column}")
    if result.errors:
        return list(rows), result
    for code, values in assignments.items():
        row = out[index[_norm(code)]]
        changed = False
        for column, value in values.items():
            new = None if value is None or str(value) == "" else str(value)
            if row.get(column) != new:
                result.changes.append((_norm(code), column, row.get(column), new))
                row[column] = new
                changed = True
        if changed:
            row.update(technical)
    return out, result


def missing_adp(table: str, rows: Sequence[Mapping[str, Any]], usage: Mapping[str, int]) -> List[dict]:
    """Codes TalentSoft sans aucun code ADP utile, les plus employés d'abord (usage = salariés du dernier run)."""
    spec = CATALOGUE[table]
    useful = adp_columns(table)
    label_columns = [c for c, attribute in spec.columns if attribute in ("label", "label_fr", "full")]
    out = []
    for row in rows:
        key = _norm(row.get(spec.code))
        if key and not any(_filled(row.get(c)) for c in useful):
            label = next((row.get(c) for c in label_columns if _filled(row.get(c))), None)
            out.append({"code": key, "label": label, "employees": int(usage.get(key, 0))})
    return sorted(out, key=lambda item: (-item["employees"], item["code"]))


# ---------------------------------------------------------------------------
# Saisie par grille collée : copier-coller depuis Excel ou un résultat de requête SQL (aucun fichier)
# ---------------------------------------------------------------------------
GRID_CODE_ALIAS = "code"                                 # en-tête accepté à la place de la colonne de code de la table
GRID_EMPTY_TEXTS = ("", "NULL")                          # cellule vide, ou NULL affiché par un éditeur SQL : rien à écrire
GRID_INFO_COLUMNS = ("libelle_ts", "occurrences")        # colonnes d'aide du modèle de REPORT, ignorées à la lecture
GRID_SECTION = re.compile(r"#\s*(lu_hr_[A-Za-z0-9_]+)")  # « # lu_hr_emp_titre » ouvre la grille d'une table


@dataclass
class GridTable:
    table: str
    assignments: Dict[str, Dict[str, str]] = field(default_factory=dict)   # code de la grille -> {colonne: valeur}
    columns: List[str] = field(default_factory=list)                       # colonnes modifiables lues
    ignored_columns: List[str] = field(default_factory=list)               # colonnes TalentSoft ou inconnues
    rows: int = 0
    rows_without_value: int = 0                                            # code sans aucune valeur : rien à écrire


def _grid_cell(cell: str) -> str:
    text = cell.strip()
    if len(text) >= 2 and text[0] == text[-1] == '"':   # cellule mise entre guillemets par Excel
        text = text[1:-1].replace('""', '"').strip()
    return text


def _grid_cells(line: str) -> List[str]:
    """Cellules d'une ligne : tabulations (copier-coller), sinon points-virgules (en-têtes tapés à la main)."""
    separator = "\t" if "\t" in line or ";" not in line else ";"
    cells = [_grid_cell(cell) for cell in line.split(separator)]
    while cells and cells[-1] == "":
        cells.pop()                                      # Excel copie un rectangle : cellules vides en fin de ligne
    return cells


def _grid_header(table: str, names: Sequence[str]) -> Tuple[int, List[Tuple[int, str]], List[str], List[str]]:
    """Position du code TalentSoft, colonnes modifiables (position, nom), colonnes ignorées, erreurs."""
    spec = CATALOGUE[table]
    editable = set(business_columns(table))
    errors = [f"en-tête en double : {name}" for name, count in Counter(names).items() if name and count > 1]
    codes = [i for i, name in enumerate(names) if name in (spec.code, GRID_CODE_ALIAS)]
    if len(codes) != 1:
        errors.append(f"une seule colonne de code TalentSoft attendue ({spec.code} ou {GRID_CODE_ALIAS})")
    used, ignored = [], []
    for i, name in enumerate(names):
        if i in codes:
            continue
        if name in editable:
            used.append((i, name))
        elif name.endswith("_adp") or name.startswith("or_"):
            errors.append(f"colonne inconnue pour {table} : {name}")     # faute de frappe probable sur une colonne ADP
        elif name not in GRID_INFO_COLUMNS:
            ignored.append(name or f"colonne {i + 1} sans nom")
    if not used:
        errors.append(f"aucune colonne modifiable pour {table} ({', '.join(business_columns(table))})")
    return (codes[0] if len(codes) == 1 else -1), used, ignored, errors


def parse_grid(text: str, default_table: str = "") -> Tuple[Dict[str, GridTable], List[str]]:
    """Lit une grille collée dans le notebook : une ligne d'en-têtes (noms des colonnes de la table, comme dans
    l'existant), puis une ligne par code TalentSoft ; séparateur tabulation, celui d'un copier-coller depuis Excel ou un
    éditeur SQL (point-virgule accepté dans une ligne sans tabulation). Plusieurs tables à la suite si chacune est
    précédée d'une ligne « # lu_hr_... » ; sinon la table est `default_table`. Les autres lignes qui commencent par #
    sont des commentaires. Une cellule vide ou NULL ne change rien ; les colonnes TalentSoft sont ignorées."""
    grids: Dict[str, GridTable] = {}
    errors: List[str] = []
    current: Optional[GridTable] = None
    header: Optional[Tuple[int, List[Tuple[int, str]], int]] = None   # code, colonnes, largeur
    skipping = False
    if default_table:
        if default_table not in CATALOGUE:
            return grids, [f"table inconnue : {default_table}"]
        current = grids.setdefault(default_table, GridTable(default_table))
    for number, line in enumerate(str(text or "").lstrip("﻿").splitlines(), start=1):
        if not line.strip("\t ;"):
            continue
        stripped = line.strip()
        if stripped.startswith("#"):
            match = GRID_SECTION.match(stripped)
            if match:
                table, header, skipping, current = match.group(1).lower(), None, False, None   # existant : organisationalStructure
                if table not in CATALOGUE:
                    errors.append(f"ligne {number} : table inconnue : {table}")
                    skipping = True
                elif default_table and table != default_table:
                    errors.append(f"ligne {number} : la grille {table} ne correspond pas à p_table ({default_table})")
                    skipping = True
                else:
                    current = grids.setdefault(table, GridTable(table))
            continue
        if skipping:
            continue
        if current is None:
            errors.append(f"ligne {number} : table non désignée (p_table, ou ligne « # lu_hr_... » avant la grille)")
            skipping = True
            continue
        if header is None:
            names = [cell.strip("[]` ").lower() for cell in _grid_cells(line)]
            code_index, used, ignored, problems = _grid_header(current.table, names)
            if problems:
                errors.extend(f"ligne {number} : {message}" for message in problems)
                skipping = True
                continue
            header = (code_index, used, len(names))
            current.columns += [name for _, name in used if name not in current.columns]
            current.ignored_columns += [name for name in ignored if name not in current.ignored_columns]
            continue
        code_index, used, width = header
        cells = _grid_cells(line)
        if len(cells) > width:
            errors.append(f"ligne {number} : {len(cells)} cellules pour {width} en-têtes")
            continue
        cells += [""] * (width - len(cells))
        code = _norm(cells[code_index])
        values = {name: cells[i] for i, name in used if cells[i] not in GRID_EMPTY_TEXTS}
        current.rows += 1
        if code is None:
            if values:
                errors.append(f"ligne {number} : valeurs sans code TalentSoft")
            continue
        if not values:
            current.rows_without_value += 1
            continue
        previous = current.assignments.setdefault(code, {})
        conflict = sorted(c for c in values if c in previous and previous[c] != values[c])
        if conflict:
            errors.append(f"ligne {number} : code {code} déjà saisi avec d'autres valeurs ({', '.join(conflict)})")
            continue
        previous.update(values)
    if not errors and not any(grid.rows for grid in grids.values()):
        errors.append("grille vide : aucune ligne de code TalentSoft")
    return grids, errors


@dataclass
class GridResolution:
    assignments: Dict[str, Dict[str, str]] = field(default_factory=dict)   # code de la table -> {colonne: valeur}
    unknown: List[str] = field(default_factory=list)                       # codes absents de la table : ignorés
    case_matched: List[Tuple[str, str]] = field(default_factory=list)      # (code de la grille, code de la table)
    errors: List[str] = field(default_factory=list)


def resolve_grid(table: str, rows: Sequence[Mapping[str, Any]], grid: Mapping[str, Mapping[str, str]]) -> GridResolution:
    """Rapproche les codes de la grille de ceux de la table : à l'identique, sinon à la casse près quand un seul code
    correspond. Un code absent de la table est ignoré et listé, car une grille reprise de l'existant peut contenir des
    codes que TalentSoft n'emploie plus ; les autres contrôles restent ceux de `apply_assignments`."""
    spec = CATALOGUE[table]
    exact = {key for key in (_norm(r.get(spec.code)) for r in rows) if key}
    folded: Dict[str, List[str]] = {}
    for key in sorted(exact):
        folded.setdefault(key.casefold(), []).append(key)
    result = GridResolution()
    origin: Dict[str, str] = {}
    for code in sorted(grid):
        if code in exact:
            target = code
        else:
            candidates = folded.get(code.casefold(), [])
            if len(candidates) > 1:
                result.errors.append(f"code {code} ambigu à la casse près : {', '.join(candidates)}")
                continue
            if not candidates:
                result.unknown.append(code)
                continue
            target = candidates[0]
            result.case_matched.append((code, target))
        values = result.assignments.setdefault(target, {})
        conflict = sorted(c for c in grid[code] if c in values and values[c] != grid[code][c])
        if conflict:
            result.errors.append(f"codes {origin[target]} et {code} : valeurs différentes pour {target} "
                                 f"({', '.join(conflict)})")
            continue
        values.update(grid[code])
        origin.setdefault(target, code)
    return result


def grid_template(table: str, missing: Sequence[Mapping[str, Any]]) -> List[str]:
    """Modèle de grille (REPORT) : code TalentSoft, aides (libellé, occurrences), colonnes ADP à remplir."""
    columns = adp_columns(table)
    lines = ["\t".join([CATALOGUE[table].code, *GRID_INFO_COLUMNS, *columns])]
    for item in missing:
        label = re.sub(r"\s+", " ", str(item.get("label") or "")).strip()
        lines.append("\t".join([item["code"], label, str(item.get("employees", 0))] + [""] * len(columns)))
    return lines


GRID_BANNER = ("# Grille à remplir (colonnes *_adp), à coller dans p_grid de T4_NB_HRIS_REFERENCE_ADMIN : CHECK, puis "
               "APPLY. Copier d'ici à la fin.")


def adp_grid_report(store, tables: Sequence[str], usage: Callable[[str], Mapping[str, int]],
                    max_lines: int = 100) -> Tuple[Dict[str, Any], List[str]]:
    """Synthèse par table et grille à remplir, au format de p_grid, les codes les plus employés d'abord
    (REPORT de NB_HRIS_REFERENCE_ADMIN et STATUS de NB_HRIS_BOOTSTRAP). `usage(table)` : occurrences par code."""
    summary: Dict[str, Any] = {}
    lines: List[str] = [GRID_BANNER]
    for table in tables:
        state = store.read(table)
        if not state.schema_ok:
            summary[table] = state.schema_detail[:200]
            lines.append(f"# {table} : table non conforme, {state.schema_detail[:200]}")
            continue
        missing = missing_adp(table, state.rows, usage(table))
        concerned = sum(item["employees"] for item in missing)
        summary[table] = {"codes": len(state.rows), "without_adp": len(missing), "occurrences_without_adp": concerned}
        lines.append(f"# {table} : {len(state.rows)} codes TalentSoft, {len(missing)} sans code ADP "
                     f"({concerned} occurrence(s) dans l'extraction)")
        if missing:
            lines += grid_template(table, missing[:max_lines])
            if len(missing) > max_lines:
                lines.append(f"# {len(missing) - max_lines} autre(s) code(s) non affiché(s) : augmenter p_max_lines")
    return summary, lines


# ---------------------------------------------------------------------------
# Stockage Delta (Spark)
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class TableState:
    exists: bool
    schema_ok: bool = True
    schema_detail: str = ""
    rows: Tuple[dict, ...] = ()


class SparkReferenceStore:
    """Tables cfg.lu_hr_* du Lakehouse (noms à 4 parties résolus par `tables`)."""

    def __init__(self, spark, tables):
        self.spark, self.tables = spark, tables

    def read(self, table: str) -> TableState:
        logical = f"cfg.{table}"
        physical = self.tables(logical)
        if not self.spark.catalog.tableExists(physical):
            return TableState(exists=False, schema_detail="ABSENTE")
        detail = IO.check_table_schema(self.spark, physical, logical)
        if detail != "CONFORME":
            return TableState(exists=True, schema_ok=False, schema_detail=detail)
        columns = LOOKUP_TABLES[table] + LOOKUP_TECH_COLUMNS
        rows = tuple(r.asDict() for r in self.spark.table(physical).select(*[f"`{c}`" for c in columns]).collect())
        return TableState(exists=True, schema_detail=detail, rows=rows)

    def write(self, table: str, rows: Sequence[Mapping[str, Any]]) -> None:
        """Remplacement complet (nouvelle version Delta ; la précédente reste lisible par voyage dans le temps)."""
        logical = f"cfg.{table}"
        physical = self.tables(logical)
        if not self.spark.catalog.tableExists(physical):
            self.spark.sql(create_table_sql(physical, logical))
        frame = IO.conform(IO.rows_df(self.spark, [dict(r) for r in rows], logical), logical)
        frame.write.format("delta").mode("overwrite").saveAsTable(physical)

    def observed(self, run_id: str, keys: Sequence[str]) -> List[str]:
        """Codes distincts de l'extraction du run pour ces clés JSON (stg.ts_employee_raw)."""
        if not keys:
            return []
        from pyspark.sql import functions as F
        raw = self.spark.table(self.tables("stg.ts_employee_raw")).where(F.col("run_id") == F.lit(IO.check_run_id(run_id)))
        values = None
        for key in keys:
            part = raw.select(F.rtrim(F.col(f"`{key}`")).alias("code"))   # comme la jointure : espaces de fin
            values = part if values is None else values.unionByName(part)
        return sorted(r.code for r in values.where("code IS NOT NULL AND code <> ''").distinct().collect())

    def latest_run(self) -> Optional[str]:
        """Run de la dernière extraction présente dans stg.ts_employee_raw."""
        rows = (self.spark.table(self.tables("stg.ts_employee_raw")).groupBy("run_id")
                .agg({"_extracted_at_utc": "max"}).orderBy("max(_extracted_at_utc)", ascending=False).limit(1).collect())
        return rows[0].run_id if rows else None

    def usage(self, run_id: str, table: str) -> Dict[str, int]:
        """Nombre d'occurrences par code TalentSoft dans l'extraction du run (clés JSON relevées pour cette table ;
        pour la table des pays, un salarié compte une fois par pays de naissance, d'adresse et fiscal)."""
        from pyspark.sql import functions as F
        keys = CATALOGUE[table].observed_keys
        if not keys:
            return {}
        raw = self.spark.table(self.tables("stg.ts_employee_raw")).where(F.col("run_id") == F.lit(IO.check_run_id(run_id)))
        counts: Counter = Counter()
        for key in keys:
            for r in raw.groupBy(F.rtrim(F.col(f"`{key}`")).alias("code")).count().collect():
                if r.code:
                    counts[r.code] += int(r["count"])
        return dict(counts)


# ---------------------------------------------------------------------------
# Synchronisation de la partie TalentSoft (à chaque run, dans l'acquisition)
# ---------------------------------------------------------------------------
def synchronize_catalogue(fetch: Callable[[str], List[dict]], store, observed: Callable[[Sequence[str]], List[str]], *,
                          loaded_by: str, now: Optional[dt.datetime] = None, write: bool = True,
                          today: Optional[dt.date] = None) -> dict:
    """Lit les référentiels et les codes relevés, fusionne la partie TalentSoft table par table, écrit ce qui change.
    Un référentiel illisible ou une table non conforme est sauté sans bloquer les autres (contenu conservé)."""
    now = now or utc_now()
    technical = {"_loaded_at_utc": now, "_loaded_by": loaded_by or None}
    cache: Dict[str, Any] = {}
    results: List[TableSync] = []
    for table, spec in CATALOGUE.items():
        state = store.read(table)
        if state.exists and not state.schema_ok:
            results.append(TableSync(table, spec.library, "SKIPPED_SCHEMA", state.schema_detail[:300]))
            continue
        nodes: Mapping[str, Any] = {}
        if spec.library:
            if spec.library not in cache:
                try:
                    cache[spec.library] = library_nodes(fetch(spec.library), today)
                except Exception as exc:          # le référentiel manque : la table garde son contenu
                    cache[spec.library] = exc
            if isinstance(cache[spec.library], Exception):
                results.append(TableSync(table, spec.library, "SKIPPED_LIBRARY_UNAVAILABLE",
                                         type(cache[spec.library]).__name__ + ": "
                                         + str(getattr(cache[spec.library], "safe_message", ""))[:200]))
                continue
            nodes = cache[spec.library]
        change_ref = "talentsoft-api:" + (spec.library or "extraction:" + "+".join(spec.observed_keys))
        rows, sync = merge_catalogue(spec, list(state.rows), nodes, observed(spec.observed_keys),
                                     dict(technical, _change_ref=change_ref))
        if write and sync.status == "OK" and (sync.inserted or sync.updated or not state.exists):
            store.write(table, rows)
            sync.written = True
        results.append(sync)
    used = [r for r in results if r.table in USED_LOOKUP_TABLES]
    return {"status": "OK" if all(r.status == "OK" for r in results) else "PARTIAL",
            "writes_enabled": write,
            "tables_written": [r.table for r in results if r.written],
            "tables_skipped": {r.table: r.status for r in results if r.status != "OK"},
            "inserted": sum(r.inserted for r in results), "updated": sum(r.updated for r in results),
            "used_codes_without_adp": sum(r.without_adp for r in used),
            "used_tables_without_any_adp": [r.table for r in used if r.status == "OK" and r.rows_after
                                            and r.without_adp == r.rows_after],
            "tables": [asdict(r) for r in results]}


CATALOGUE_CONTROLS = ["REFERENCE_CATALOGUE", "REFERENCE_ADP_MISSING"]


def catalogue_anomalies(report: Optional[dict], error: str = "") -> List[dict]:
    """Contrôles pour log.anomaly, jamais BLOCKING : le blocage reste décidé par les contrôles de TRANSFORM sur l'état
    réel des tables (table vide ou sans aucun code ADP)."""
    def row(code, severity, metric, failed, message):
        return {"control_code": code, "severity": severity, "metric_value": metric, "threshold_value": None,
                "outcome": "FAIL" if failed else "PASS", "message": message}
    if report is None:
        return [row("REFERENCE_CATALOGUE", SEVERITY_WARNING, 1, True, "référentiels TalentSoft non synchronisés : " + error)]
    skipped = report["tables_skipped"]
    missing = report["used_codes_without_adp"]
    return [row("REFERENCE_CATALOGUE", SEVERITY_WARNING, len(skipped), bool(skipped),
                "tables non synchronisées : " + (", ".join(f"{t} ({s})" for t, s in skipped.items()) or "aucune")),
            row("REFERENCE_ADP_MISSING", SEVERITY_INFO, missing, missing > 0,
                f"{missing} code(s) TalentSoft sans code ADP dans les 13 tables utilisées (NB_HRIS_REFERENCE_ADMIN)")]


def step_message(report: dict, limit: int = 7800) -> str:
    """Message de ctl.run_step : comptes par table, sans valeur."""
    import json
    compact = {k: report[k] for k in ("status", "writes_enabled", "tables_written", "tables_skipped", "inserted",
                                       "updated", "used_codes_without_adp", "used_tables_without_any_adp")}
    compact["tables"] = {t["table"]: [t["status"], t["inserted"], t["updated"], t["rows_after"], t["without_adp"]]
                         for t in report["tables"]}
    return json.dumps(compact, ensure_ascii=False)[:limit]
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.remuneration =============
register_hris_module('hris.remuneration', '163d1b3e4bffe13b138b6f8f506c7a74dd3aefc9a34e1de822d59655f1622785', r'''"""hris.remuneration — flux 2 de l'existant (pipeline Synapse Pipeline_SIRH-SFTP), en tables, sans aucun fichier.

Existant : trois exports CSV du Core HR lus sur le SFTP TalentSoft (/ExportCarriere/Data/COMPENSATION_HUB_ADMIN_
ANNUALBASESALARY, _BONUSPERCENT, _BONUSTARGETVALUE), archivés en Parquet ; vues SFTP.vw_BaseSalary_Changes et
SFTP.vw_Bonus_Changes entre l'archive du jour (J) et la précédente (J-1) ; deux fichiers CSV déposés sur le SFTP
(/ImportCarriere/Data/RemunSalaryHistoIE_InsertAndUpdate_fr-FR_1_<ddMMyyyy>.csv et RemunTargetBonusHistoIE_…_2_…)
que TalentSoft importe (35 et 36 imports réussis relevés le 13/10/2026 par l'API d'imports).

Cible (A08 D-FIL-00 : aucun fichier) :
  * lecture par l'API d'exports du Core HR, jeu « Compensation » (section, dates, montant, catégorie, devise,
    périodicité, indexation), contenu lu en mémoire ;
  * instantané au format des exports de l'existant en table (stg.ts_remuneration_snapshot) ; J-1 = référence du
    dernier envoi accepté (ref.ts_remuneration_snapshot) et non « l'avant-dernière archive », ce qui supprime deux
    défauts de l'existant (réémission le même jour, changements perdus après un échec) ;
  * changements calculés par les règles des deux vues reproduites à l'identique, particularités comprises
    (A04 R-SFTP-010 à 027) : comparaison décimale pour le salaire, comparaison de CHAÎNES pour les bonus, clé
    employeenumber seule (lignes multipliées en cas de doublon), devise reprise du salaire de base du jour, date de
    clôture « veille du nouveau début » au format dd/MM/yyyy 00:00:00, date de clôture NULL sur disparition ;
  * lignes à importer en table (pub.ts_remuneration_change), envoi par l'API d'imports (types RemunSalaryHistoIE et
    RemunTargetBonusHistoIE, mêmes colonnes et même mise en forme que les fichiers de l'existant). Le contenu exigé par
    l'API d'import est construit en mémoire : aucun fichier n'est écrit, ni dans le Lakehouse ni sur un disque.

Correspondance export -> colonnes de l'existant : table cfg.ts_remuneration_mapping (graines = hypothèses à valider
par le mode VALIDATE du notebook, qui compare la STRUCTURE du contenu produit à celle du dernier import de l'existant).
Envoi : seulement depuis un workspace dont le nom contient « PRD » (même règle que la publication ADP) ET si la
variable ts_remuneration_import_enabled vaut true. Aucune valeur (matricule, montant) n'est journalisée.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import time
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, localcontext
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from hris import spark_io as IO
from hris.adp_gate import DENIED_NON_PRD, ENV_SATISFIED, evaluate_gate
from hris.common import HrisError, safe_error, sha256_hex, utc_now
from hris.model import REMUNERATION_SNAPSHOT_COLUMNS

BASE_SALARY, BONUS_PERCENT, BONUS_TARGET_VALUE = "BASE_SALARY", "BONUS_PERCENT", "BONUS_TARGET_VALUE"
SECTIONS = (BASE_SALARY, BONUS_PERCENT, BONUS_TARGET_VALUE)
IMPORT_SALARY, IMPORT_BONUS = "RemunSalaryHistoIE", "RemunTargetBonusHistoIE"     # codes des types d'import TalentSoft
IMPORT_TYPES = (IMPORT_SALARY, IMPORT_BONUS)
IMPORT_SECTIONS = {IMPORT_SALARY: (BASE_SALARY,), IMPORT_BONUS: (BONUS_PERCENT, BONUS_TARGET_VALUE)}
IMPORT_FILE_ORDER = {IMPORT_SALARY: 1, IMPORT_BONUS: 2}                           # « _1_ » / « _2_ » de l'existant
IMPORT_CULTURE = "fr-FR"
SALARY_COLUMNS = ["employeenumber", "date", "type", "amount", "currency", "periodicity", "extra_augm_n-1"]
BONUS_COLUMNS = ["employeenumber", "type", "start_date", "end_date", "calculation_method", "value", "currency",
                 "base_salary_effective_date", "periodicity"]
IMPORT_COLUMNS = {IMPORT_SALARY: SALARY_COLUMNS, IMPORT_BONUS: BONUS_COLUMNS}
SNAPSHOT_COLUMNS = {BASE_SALARY: ["employeenumber", "date", "type", "amount", "currency", "periodicity"],
                    BONUS_PERCENT: BONUS_COLUMNS, BONUS_TARGET_VALUE: BONUS_COLUMNS}
CODE_COLUMNS = ("type", "currency", "periodicity", "calculation_method")         # comparées en valeurs par VALIDATE
EXPORT_COLUMNS = ["Section", "StartDate", "EndDate", "Amount", "Category", "ContractNumber", "Currency", "IndexedOn",
                  "IsEligibleFor", "Periodicity"]
EXPORT_DONE, EXPORT_FAILED = "Successful", ("Failed", "Canceled")
EXPORT_STATUS_NAMES = {0: "Pending", 1: "Running", 2: "Canceled", 3: "Successful", 4: "Failed"}
IMPORT_TERMINAL_OK, IMPORT_TERMINAL_KO = ("Success",), ("Failed", "Error", "Canceled", "Cancelled")

# Statuts d'envoi par type d'import (log.ts_remuneration_import)
SENT, EMPTY, NOT_PUSHED_NON_PRD, BLOCKED_UNRESOLVED, DISABLED, FAILED, INIT_REFERENCE = (
    "SENT", "EMPTY", "NOT_PUSHED_NON_PRD", "BLOCKED_UNRESOLVED", "DISABLED", "FAILED", "INIT_REFERENCE")
ADVANCING = (SENT, EMPTY, NOT_PUSHED_NON_PRD, INIT_REFERENCE)   # la référence avance : changements remis ou sans objet


# ---------------------------------------------------------------------------
# Correspondance export -> format des exports de l'existant (cfg.ts_remuneration_mapping)
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class MappingRule:
    section_code: str               # code de section du jeu Compensation (ex. ANNUAL_BASE_SALARY)
    legacy_section: str             # BASE_SALARY | BONUS_PERCENT | BONUS_TARGET_VALUE
    type_value: str = ""            # colonne « type » ; vide = code de section
    calculation_method: str = ""    # colonne « calculation_method » (bonus)
    effective_date_rule: str = ""   # base_salary_effective_date : BASE_SALARY_DATE | START_DATE | INDEXED_ON
    date_format: str = "%d/%m/%Y"   # format des dates (strftime) ; l'existant lisait les dates en style 103
    currency_case: str = "upper"    # upper | raw


DEFAULT_MAPPING = [   # hypothèses à confirmer par le mode VALIDATE avant d'activer l'envoi (H-REM-01 à 04)
    MappingRule("ANNUAL_BASE_SALARY", BASE_SALARY),
    MappingRule("CONTRACTUAL_BONUS_PERCENTAGE", BONUS_PERCENT, calculation_method="PERCENTAGE",
                effective_date_rule="BASE_SALARY_DATE"),
    MappingRule("CONTRACTUAL BONUS", BONUS_TARGET_VALUE, calculation_method="AMOUNT",
                effective_date_rule="BASE_SALARY_DATE"),
]


def section_key(value: Any) -> str:
    """« {{ANNUAL_BASE_SALARY}} », « Annual_Base_Salary » -> « ANNUAL_BASE_SALARY » (accolades et casse ignorées)."""
    return re.sub(r"\s+", " ", str(value or "").replace("{", "").replace("}", "")).strip().upper()


def mapping_rows(rules: Iterable[MappingRule]) -> List[dict]:
    return [dict(rule.__dict__) for rule in rules]


def rules_from_rows(rows: Iterable[Mapping[str, Any]]) -> List[MappingRule]:
    out = []
    for row in rows:
        values = {k: (row.get(k) if row.get(k) is not None else MappingRule.__dataclass_fields__[k].default)
                  for k in MappingRule.__dataclass_fields__}
        if values["legacy_section"] not in SECTIONS:
            raise HrisError("REMUNERATION_MAPPING_INVALID", "legacy_section inconnue : " + str(values["legacy_section"]))
        out.append(MappingRule(**values))
    return out


# ---------------------------------------------------------------------------
# Lecture du contenu d'export (en mémoire : JSON, CSV, XLSX ou ZIP)
# ---------------------------------------------------------------------------
def _norm_header(name: Any) -> str:
    text = unicodedata.normalize("NFD", str(name or ""))
    return re.sub(r"[^a-z0-9]", "", "".join(c for c in text if unicodedata.category(c) != "Mn").lower())


HEADER_ALIASES = {
    "employeenumber": ("employeenumber", "employeeid", "matricule", "employee", "employeenumberhub"),
    "section": ("section",), "startdate": ("startdate", "datededebut"), "enddate": ("enddate", "datedefin"),
    "amount": ("amount", "montant"), "category": ("category", "categorie"),
    "contractnumber": ("contractnumber", "linkedcontract"), "currency": ("currency", "devise"),
    "indexedon": ("indexedon", "indexedwith"), "iseligiblefor": ("iseligiblefor",),
    "periodicity": ("periodicity", "frequency", "frequence"),
}


def parse_export(content: bytes, content_type: str = "") -> List[dict]:
    """Lignes de l'export, clés normalisées (HEADER_ALIASES). Lève TS_EXPORT_FORMAT si le format est inconnu ou si
    la colonne du matricule est introuvable (message : noms de colonnes seulement)."""
    if content[:2] == b"PK":
        archive = zipfile.ZipFile(io.BytesIO(content))
        names = archive.namelist()
        if any(n.startswith("xl/") for n in names):
            return _normalize(_xlsx_rows(content))
        inner = next((n for n in names if n.lower().endswith((".json", ".csv", ".txt"))), None)
        if inner is None:
            raise HrisError("TS_EXPORT_FORMAT", "archive sans contenu exploitable")
        return parse_export(archive.read(inner), "")
    text = content.decode("utf-8-sig", errors="strict") if _is_utf8(content) else content.decode("cp1252")
    stripped = text.lstrip()
    if stripped.startswith("[") or stripped.startswith("{"):
        data = json.loads(stripped)
        if isinstance(data, dict):
            data = next((v for v in data.values() if isinstance(v, list)), [])
        return _normalize([dict(r) for r in data if isinstance(r, dict)])
    first = stripped.splitlines()[0] if stripped else ""
    delimiter = max((";", ",", "\t"), key=first.count)
    reader = csv.DictReader(io.StringIO(stripped, newline=""), delimiter=delimiter)
    return _normalize([dict(r) for r in reader])


def _is_utf8(content: bytes) -> bool:
    try:
        content.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def _xlsx_rows(content: bytes) -> List[dict]:
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:]]


def _normalize(rows: List[dict]) -> List[dict]:
    if not rows:
        return []
    lookup = {_norm_header(k): k for k in rows[0]}
    mapping = {}
    for target, aliases in HEADER_ALIASES.items():
        source = next((lookup[a] for a in aliases if a in lookup), None)
        if source is not None:
            mapping[target] = source
    if "employeenumber" not in mapping or "section" not in mapping:
        raise HrisError("TS_EXPORT_FORMAT", "colonnes matricule/section introuvables ; colonnes reçues : "
                        + ", ".join(sorted(lookup)))
    return [{target: row.get(source) for target, source in mapping.items()} for row in rows]


# ---------------------------------------------------------------------------
# Instantané au format des exports de l'existant
# ---------------------------------------------------------------------------
def _clean(value: Any) -> Optional[str]:
    """Valeur texte ; vide -> NULL (lecture CSV -> Parquet de l'existant)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def parse_any_date(value: Any) -> Optional[dt.date]:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if match:
        return dt.date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})(\s.*)?$", text)
    if match:
        return dt.date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
    raise HrisError("CONVERSION_DATE", "date d'export illisible")


def _amount_text(value: Any) -> Optional[str]:
    text = _clean(value)
    if text is not None and re.fullmatch(r"[+-]?\d+,\d+", text):
        text = text.replace(",", ".")            # format des exports de l'existant : point décimal
    return text


def build_snapshot(export_rows: Iterable[Mapping[str, Any]], rules: Sequence[MappingRule]) -> Tuple[Dict[str, List[dict]], dict]:
    """Lignes de l'export -> trois listes au format des exports de l'existant. Sections non mappées ignorées (comptées)."""
    by_section = {section_key(rule.section_code): rule for rule in rules}
    snapshot: Dict[str, List[dict]] = {section: [] for section in SECTIONS}
    stats: Counter = Counter()
    pending_bonus: List[Tuple[MappingRule, Mapping[str, Any]]] = []
    for row in export_rows:
        rule = by_section.get(section_key(row.get("section")))
        if rule is None:
            stats["export_rows_ignored"] += 1
            continue
        stats["export_rows_mapped"] += 1
        if rule.legacy_section == BASE_SALARY:
            snapshot[BASE_SALARY].append(_base_row(rule, row))
        else:
            pending_bonus.append((rule, row))
    base_dates = defaultdict(list)
    for base in snapshot[BASE_SALARY]:
        if base["employeenumber"] and base["date"]:
            base_dates[base["employeenumber"]].append(base["_date"])
    for rule, row in pending_bonus:
        snapshot[rule.legacy_section].append(_bonus_row(rule, row, base_dates))
    for section in SECTIONS:
        for line in snapshot[section]:
            line.pop("_date", None)
        stats[section.lower() + "_rows"] = len(snapshot[section])
    return snapshot, dict(stats)


def _fmt_date(value: Optional[dt.date], rule: MappingRule) -> Optional[str]:
    return value.strftime(rule.date_format) if value else None


def _currency(value: Any, rule: MappingRule) -> Optional[str]:
    text = _clean(value)
    return text.upper() if text and rule.currency_case == "upper" else text


def _type(rule: MappingRule, row: Mapping[str, Any]) -> str:
    return rule.type_value or section_key(row.get("section"))


def _base_row(rule: MappingRule, row: Mapping[str, Any]) -> dict:
    start = parse_any_date(row.get("startdate"))
    return {"employeenumber": _clean(row.get("employeenumber")), "date": _fmt_date(start, rule),
            "type": _type(rule, row), "amount": _amount_text(row.get("amount")),
            "currency": _currency(row.get("currency"), rule), "periodicity": _clean(row.get("periodicity")),
            "src_section_code": section_key(row.get("section")), "src_category": _clean(row.get("category")),
            "src_indexed_on": _clean(row.get("indexedon")), "_date": start}


def _bonus_row(rule: MappingRule, row: Mapping[str, Any], base_dates: Mapping[str, List[dt.date]]) -> dict:
    start = parse_any_date(row.get("startdate"))
    number = _clean(row.get("employeenumber"))
    effective = None
    if rule.effective_date_rule == "BASE_SALARY_DATE":
        dates = sorted(d for d in base_dates.get(number, []) if d and (start is None or d <= start))
        effective = _fmt_date(dates[-1], rule) if dates else _fmt_date(start, rule)
    elif rule.effective_date_rule == "INDEXED_ON":
        indexed = row.get("indexedon")
        try:
            effective = _fmt_date(parse_any_date(indexed), rule)
        except HrisError:
            effective = _clean(indexed)
    elif rule.effective_date_rule == "START_DATE":
        effective = _fmt_date(start, rule)
    return {"employeenumber": number, "type": _type(rule, row), "start_date": _fmt_date(start, rule),
            "end_date": _fmt_date(parse_any_date(row.get("enddate")), rule),
            "calculation_method": rule.calculation_method or None, "value": _amount_text(row.get("amount")),
            "currency": _currency(row.get("currency"), rule), "base_salary_effective_date": effective,
            "periodicity": _clean(row.get("periodicity")), "src_section_code": section_key(row.get("section")),
            "src_category": _clean(row.get("category")), "src_indexed_on": _clean(row.get("indexedon"))}


# ---------------------------------------------------------------------------
# Règles des vues SFTP.vw_BaseSalary_Changes et SFTP.vw_Bonus_Changes (sémantique T-SQL reproduite)
# ---------------------------------------------------------------------------
_DECIMAL_TEXT = re.compile(r"^\s*[+-]?(\d+\.?\d*|\.\d+)\s*$")
_CENT = Decimal("0.01")


def to_decimal(value: Optional[str]) -> Optional[Decimal]:
    """CAST(x AS DECIMAL(20,2)) : NULL -> NULL ; nombre -> arrondi au centime (à l'écart de zéro) ; autre texte ->
    échec, comme la vue SQL qui faisait échouer le flux (le message ne contient aucune valeur)."""
    if value is None:
        return None
    if not _DECIMAL_TEXT.match(str(value)):
        raise HrisError("CONVERSION_NUMERIC", "montant non numérique dans l'instantané de rémunération")
    result = Decimal(str(value).strip()).quantize(_CENT, rounding=ROUND_HALF_UP)
    return result if result != 0 else Decimal("0.00")


def fmt_decimal(value: Optional[Decimal]) -> Optional[str]:
    """REPLACE(CAST(d AS VARCHAR(50)), '.', ',')."""
    return None if value is None else format(value, "f").replace(".", ",")


def _key(value: Optional[str]) -> Optional[str]:
    """Égalité T-SQL : NULL ne correspond à rien ; espaces de fin ignorés."""
    return None if value is None else value.rstrip(" ")


def _neq(a: Optional[str], b: Optional[str]) -> bool:
    """a != b en T-SQL : inconnu (donc faux) si l'un est NULL."""
    return a is not None and b is not None and a.rstrip(" ") != b.rstrip(" ")


def increase_pct(current: Optional[Decimal], previous: Optional[Decimal]) -> Optional[str]:
    """[extra_augm_n-1] : 0 si égaux ; sinon ROUND(((J - J1) / NULLIF(J1, 0)) * 100, 2) ; NULL si J ou J1 absent."""
    if current is None or previous is None:
        return None
    if current == previous:
        return "0,00"
    if previous == 0:
        return None
    with localcontext() as ctx:
        ctx.prec = 60
        quotient = ((current - previous) / previous).quantize(Decimal("1e-17"), rounding=ROUND_HALF_UP)
        result = (quotient * 100).quantize(_CENT, rounding=ROUND_HALF_UP)
    return fmt_decimal(result if result != 0 else Decimal("0.00"))


def _index(rows: Iterable[Mapping[str, Any]]) -> Dict[str, List[Mapping[str, Any]]]:
    out: Dict[str, List[Mapping[str, Any]]] = defaultdict(list)
    for row in rows:
        key = _key(row.get("employeenumber"))
        if key is not None:
            out[key].append(row)
    return out


def base_salary_changes(current: Sequence[Mapping[str, Any]], previous: Sequence[Mapping[str, Any]]) -> List[dict]:
    """R-SFTP-010 à 014 : J LEFT JOIN J1 sur employeenumber ; montant changé (non NULL) ou date changée."""
    previous_by = _index(previous)
    out = []
    for j in current:
        key = _key(j.get("employeenumber"))
        amount = to_decimal(j.get("amount"))
        for j1 in (previous_by.get(key) if key is not None else None) or [None]:
            amount1 = to_decimal(j1.get("amount")) if j1 is not None else None
            amount_changed = (amount if amount is not None else 0) != (amount1 if amount1 is not None else 0) \
                and j.get("amount") is not None
            date_changed = j1 is not None and _neq(j.get("date"), j1.get("date"))
            if amount_changed or date_changed:
                out.append({"employeenumber": j.get("employeenumber"), "date": j.get("date"), "type": j.get("type"),
                            "amount": fmt_decimal(amount), "currency": j.get("currency"),
                            "periodicity": j.get("periodicity"), "extra_augm_n-1": increase_pct(amount, amount1),
                            "rule": "R-SFTP-010"})
    return out


def closing_end_date(start_date: Optional[str]) -> Optional[str]:
    """CONVERT(varchar(19), DATEADD(day, -1, CONVERT(date, start, 103)), 103) + ' 00:00:00' ; NULL si start NULL."""
    if start_date is None:
        return None
    return (parse_any_date(start_date) - dt.timedelta(days=1)).strftime("%d/%m/%Y") + " 00:00:00"


def _bonus_out(row: Mapping[str, Any], end_date: Optional[str], currency: Optional[str], rule: str) -> dict:
    return {"employeenumber": row.get("employeenumber"), "type": row.get("type"), "start_date": row.get("start_date"),
            "end_date": end_date, "calculation_method": row.get("calculation_method"),
            "value": fmt_decimal(to_decimal(row.get("value"))), "currency": currency,
            "base_salary_effective_date": row.get("base_salary_effective_date"), "periodicity": row.get("periodicity"),
            "rule": rule}


def bonus_changes(percent_j: Sequence[Mapping[str, Any]], percent_j1: Sequence[Mapping[str, Any]],
                  target_j: Sequence[Mapping[str, Any]], target_j1: Sequence[Mapping[str, Any]],
                  base_j: Sequence[Mapping[str, Any]]) -> List[dict]:
    """R-SFTP-020 à 027 : UNION (dédoublonnée) des 3 blocs, pour le pourcentage puis pour le montant cible."""
    base_by = _index(base_j)
    produced: List[dict] = []
    for current, previous in ((percent_j, percent_j1), (target_j, target_j1)):
        previous_by = _index(previous)
        for j in current:
            key = _key(j.get("employeenumber"))
            currencies = [b.get("currency") for b in (base_by.get(key) if key is not None else None) or []] or [None]
            for j1 in (previous_by.get(key) if key is not None else None) or [None]:
                value_j = j.get("value") if j.get("value") is not None else "0"          # ISNULL(value, 0) : chaînes
                value_j1 = j1.get("value") if j1 is not None and j1.get("value") is not None else "0"
                start_changed = j1 is not None and _neq(j.get("start_date"), j1.get("start_date"))
                has_value = j.get("value") is not None
                for currency in currencies:
                    if (_neq(value_j, value_j1) and has_value) or (start_changed and has_value):
                        produced.append(_bonus_out(j, j.get("end_date"), currency, "R-SFTP-020"))
                    if start_changed and has_value:
                        produced.append(_bonus_out(j1, closing_end_date(j.get("start_date")), currency, "R-SFTP-021"))
                    if (j1 is not None and j1.get("start_date") is not None and j1.get("value") is not None
                            and j.get("start_date") is None and j.get("value") is None):
                        produced.append(_bonus_out(j1, None, currency, "R-SFTP-022"))   # end_date NULL (anomalie conservée)
    seen, out = set(), []
    for row in produced:                                                                # UNION : doublons exacts retirés
        signature = tuple(row.get(c) for c in BONUS_COLUMNS)
        if signature not in seen:
            seen.add(signature)
            out.append(row)
    return out


# ---------------------------------------------------------------------------
# Contenu d'import (mémoire) : même mise en forme que la copie vue -> CSV SFTP de l'existant
# ---------------------------------------------------------------------------
def _quoted(value: Any) -> str:
    if value is None:
        return ""
    return '"' + str(value).replace("\\", "\\\\").replace('"', '\\"') + '"'


def import_content(import_type: str, rows: Sequence[Mapping[str, Any]]) -> bytes:
    """Séparateur « ; », en-tête, toutes les valeurs entre guillemets, échappement par barre oblique inverse, UTF-8."""
    columns = IMPORT_COLUMNS[import_type]
    lines = [";".join(_quoted(c) for c in columns)] + [";".join(_quoted(r.get(c)) for c in columns) for r in rows]
    return ("\r\n".join(lines) + "\r\n").encode("utf-8")


def import_file_name(import_type: str, day: dt.date) -> str:
    return f"{import_type}_InsertAndUpdate_{IMPORT_CULTURE}_{IMPORT_FILE_ORDER[import_type]}_{day.strftime('%d%m%Y')}.csv"


# ---------------------------------------------------------------------------
# API TalentSoft (exports et imports), avec le jeton et la session du client d'acquisition
# ---------------------------------------------------------------------------
class TalentSoftApi:
    """Appels des API d'exports et d'imports : même hôte, même jeton (renouvelé selon l'âge) et même session que le
    client d'acquisition (hris.ts_acquire.TsClient). Aucune redirection suivie ; délais explicites."""

    def __init__(self, client, timeout: Tuple[float, float] = (30.0, 300.0)):
        self.client, self.timeout = client, timeout

    def _headers(self, accept: str = "application/json") -> dict:
        headers = dict(self.client._auth_headers())
        headers["Accept"] = accept
        return headers

    def _call(self, method: str, path: str, accept: str = "application/json", **kwargs):
        resp = self.client._session.request(method, self.client.base_url + path, headers=self._headers(accept),
                                            timeout=self.timeout, allow_redirects=False, **kwargs)
        if resp.status_code == 401:                       # jeton expiré : renouvelé, appel idempotent rejoué une fois
            self.client.get_token()
            resp = self.client._session.request(method, self.client.base_url + path, headers=self._headers(accept),
                                                timeout=self.timeout, allow_redirects=False, **kwargs)
        return resp

    def create_export(self, payload: Mapping[str, Any]) -> int:
        resp = self._call("POST", "/api/hub/v3.0/exports/datasets", json=payload)
        if resp.status_code != 200:
            raise HrisError("TS_EXPORT_REFUSED", f"statut {resp.status_code}")
        data = resp.json()
        return int(data if not isinstance(data, dict) else next(iter(data.values())))

    def export_status(self, export_id: int) -> str:
        resp = self._call("GET", f"/api/hub/v3.0/exports/{export_id}/status")
        if resp.status_code != 200:
            raise HrisError("TS_EXPORT_STATUS", f"statut {resp.status_code}")
        value = resp.json()
        return EXPORT_STATUS_NAMES.get(value, str(value)) if isinstance(value, int) else str(value)

    def export_content(self, export_id: int) -> Tuple[bytes, str]:
        resp = self._call("GET", f"/api/hub/v3.0/exports/{export_id}/file", accept="*/*")
        if resp.status_code != 200:
            raise HrisError("TS_EXPORT_DOWNLOAD", f"statut {resp.status_code}")
        return resp.content, resp.headers.get("Content-Type", "")

    def submit_import(self, import_type: str, file_name: str, content: bytes) -> int:
        """Import par l'API (le contenu est transmis en mémoire, jamais écrit). Type, culture et méthode sont aussi
        portés par le nom, comme les fichiers que l'existant déposait sur le SFTP."""
        resp = self._call("POST", "/api/administration/v1.0/imports",
                          files={"file": (file_name, content, "text/csv")},
                          data={"importTypeCode": import_type, "culture": IMPORT_CULTURE, "importMethod": "InsertUpdate"})
        if resp.status_code != 200:
            raise HrisError("TS_IMPORT_REFUSED", f"statut {resp.status_code}")
        data = resp.json()
        value = data.get("ImportId", data.get("importId")) if isinstance(data, dict) else data
        return int(value)

    def import_status(self, import_id: int) -> str:
        resp = self._call("GET", f"/api/administration/v1.0/imports/{import_id}")
        if resp.status_code != 200:
            raise HrisError("TS_IMPORT_STATUS", f"statut {resp.status_code}")
        data = resp.json()
        return str(data.get("StatusCode", data.get("statusCode")))

    def latest_imports(self) -> Dict[str, dict]:
        """Dernier import réussi de chaque type de l'existant (identifiant, date, nom), pour le mode VALIDATE."""
        resp = self._call("GET", "/api/administration/v1.0/imports")
        if resp.status_code != 200:
            raise HrisError("TS_IMPORT_LIST", f"statut {resp.status_code}")
        latest: Dict[str, dict] = {}
        for row in resp.json() or []:
            low = {str(k).lower(): v for k, v in row.items()}
            code, name = low.get("importtypecode"), str(low.get("importfilename") or "")
            if code in IMPORT_TYPES and name.startswith(code) and low.get("statuscode") in IMPORT_TERMINAL_OK:
                when = str(low.get("importcreationdatetime") or "")
                if code not in latest or when > latest[code]["when"]:
                    latest[code] = {"id": low.get("importid"), "when": when}
        return latest

    def import_file(self, import_id: int) -> bytes:
        resp = self._call("GET", f"/api/administration/v1.0/imports/{import_id}/ImportFile", accept="*/*")
        if resp.status_code != 200:
            raise HrisError("TS_IMPORT_FILE", f"statut {resp.status_code}")
        return resp.content


def export_request(only_active_employees: bool) -> dict:
    return {"ExportType": "Compensation", "IncludeEmployeesName": False, "ExportedColumns": list(EXPORT_COLUMNS),
            "ColumnFilters": {}, "RecordFilters": {"ExportOnlyCurrentRecords": True, "ExportOnlyModifiedLastDays": 0},
            "EmployeeFilters": {"ExportOnlyActiveEmployees": bool(only_active_employees), "SearchOperandId": 0}}


def extract_compensation(api, *, only_active_employees: bool = False, poll_seconds: float = 10.0,
                         timeout_seconds: float = 1800.0, sleep: Callable[[float], None] = time.sleep,
                         clock: Callable[[], float] = time.monotonic) -> List[dict]:
    export_id = api.create_export(export_request(only_active_employees))
    deadline = clock() + timeout_seconds
    while True:
        status = api.export_status(export_id)
        if status == EXPORT_DONE:
            break
        if status in EXPORT_FAILED:
            raise HrisError("TS_EXPORT_FAILED", status)
        if clock() > deadline:
            raise HrisError("TS_EXPORT_TIMEOUT", f"export {export_id} toujours {status}")
        sleep(poll_seconds)
    content, content_type = api.export_content(export_id)
    return parse_export(content, content_type)


def submit_and_follow(api, import_type: str, file_name: str, content: bytes, *, poll_seconds: float = 15.0,
                      timeout_seconds: float = 1800.0, sleep: Callable[[float], None] = time.sleep,
                      clock: Callable[[], float] = time.monotonic) -> Tuple[int, str]:
    """Envoie puis suit l'import jusqu'à un statut final ou l'échéance ; renvoie (identifiant, statut TalentSoft)."""
    import_id = api.submit_import(import_type, file_name, content)
    deadline = clock() + timeout_seconds
    while True:
        status = api.import_status(import_id)
        if status in IMPORT_TERMINAL_OK or status in IMPORT_TERMINAL_KO or clock() > deadline:
            return import_id, status
        sleep(poll_seconds)


# ---------------------------------------------------------------------------
# Comparaison de STRUCTURE avec le dernier import de l'existant (mode VALIDATE, aucune valeur personnelle)
# ---------------------------------------------------------------------------
def _shape(value: str) -> str:
    return re.sub(r"[A-Za-zÀ-ÿ]", "A", re.sub(r"\d", "9", value))


def content_structure(content: bytes) -> dict:
    text = content.decode("utf-8-sig") if _is_utf8(content) else content.decode("cp1252")
    first = text.splitlines()[0] if text else ""
    delimiter = max((";", ","), key=first.count)
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"', escapechar="\\"))
    header, body = (rows[0], [r for r in rows[1:] if any(v.strip() for v in r)]) if rows else ([], [])
    columns = {}
    for index, name in enumerate(header):
        values = [r[index] for r in body if index < len(r)]
        if name.strip().lower() in CODE_COLUMNS:
            columns[name] = {"codes": sorted(set(values))[:20]}
        else:
            columns[name] = {"shapes": sorted(set(_shape(v) for v in values))[:10]}
    return {"delimiter": delimiter, "quoted_header": first.startswith('"'), "header": header, "rows": len(body),
            "columns": columns}


def compare_structures(produced: Mapping[str, Any], legacy: Mapping[str, Any]) -> List[str]:
    """Écarts de structure entre le contenu produit et celui de l'existant (codes et formes, jamais de valeur)."""
    gaps = []
    for key in ("delimiter", "quoted_header", "header"):
        if produced.get(key) != legacy.get(key):
            gaps.append(f"{key} : produit {produced.get(key)!r} / existant {legacy.get(key)!r}")
    for name, legacy_column in (legacy.get("columns") or {}).items():
        mine = (produced.get("columns") or {}).get(name)
        if mine is None:
            continue
        if "codes" in legacy_column:
            missing = sorted(set(legacy_column["codes"]) - set(mine.get("codes", [])))
            extra = sorted(set(mine.get("codes", [])) - set(legacy_column["codes"]))
            if extra:
                gaps.append(f"{name} : codes produits absents de l'existant {extra} (existant : {legacy_column['codes']})")
            if missing and not mine.get("codes"):
                gaps.append(f"{name} : aucun code produit (existant : {legacy_column['codes']})")
        elif set(mine.get("shapes", [])) - set(legacy_column.get("shapes", [])) and legacy_column.get("shapes"):
            gaps.append(f"{name} : formes produites {mine.get('shapes')} / existant {legacy_column.get('shapes')}")
    return gaps


# ---------------------------------------------------------------------------
# Orchestration d'un run (étape REMUNERATION)
# ---------------------------------------------------------------------------
def run_remuneration(api, store, context_provider: Callable[[], Mapping[str, Any]], *, run_id: str,
                     business_date: Optional[dt.date], import_enabled: bool, only_active_employees: bool = False,
                     now: Optional[dt.datetime] = None, follow: Callable[..., Tuple[int, str]] = submit_and_follow,
                     extract: Callable[..., List[dict]] = extract_compensation) -> dict:
    """Lecture, instantané, changements, envoi (si autorisé), avance de la référence par type d'import.
    Lève TS_IMPORT_FAILED après avoir tout journalisé si un envoi a échoué (comme l'échec du flux de l'existant)."""
    now = now or utc_now()
    rules = store.mapping()
    snapshot, stats = build_snapshot(extract(api, only_active_employees=only_active_employees), rules)
    store.save_snapshot(run_id, business_date, snapshot)
    reference = store.reference()
    first_run = not any(reference.get(section) for section in SECTIONS)
    changes = {IMPORT_SALARY: [], IMPORT_BONUS: []}
    if not first_run:
        changes[IMPORT_SALARY] = base_salary_changes(snapshot[BASE_SALARY], reference.get(BASE_SALARY, []))
        changes[IMPORT_BONUS] = bonus_changes(snapshot[BONUS_PERCENT], reference.get(BONUS_PERCENT, []),
                                              snapshot[BONUS_TARGET_VALUE], reference.get(BONUS_TARGET_VALUE, []),
                                              snapshot[BASE_SALARY])
    store.save_changes(run_id, business_date, changes)
    gate = evaluate_gate(context_provider)
    logs, failures = [], []
    for import_type in IMPORT_TYPES:
        rows = changes[import_type]
        content = import_content(import_type, rows)
        entry = {"run_id": run_id, "import_type": import_type, "file_name": import_file_name(import_type, now.date()),
                 "rows": len(rows), "payload_sha256": sha256_hex(content.decode("utf-8")) if rows else None,
                 "gate_outcome": gate.outcome, "gate_reason": gate.reason_code, "ts_import_id": None,
                 "ts_import_status": None, "message": "", "created_at_utc": now}
        if first_run:
            entry.update(status=INIT_REFERENCE, message="premier run : référence initialisée, rien n'est envoyé")
        elif not rows:
            entry.update(status=EMPTY, message="aucun changement")
        elif gate.outcome == DENIED_NON_PRD:
            entry.update(status=NOT_PUSHED_NON_PRD, message="workspace non PRD : calculé, non envoyé")
        elif gate.outcome != ENV_SATISFIED:
            entry.update(status=BLOCKED_UNRESOLVED, message="nom du workspace non résolu : non envoyé, réévalué au run suivant")
        elif not import_enabled:
            entry.update(status=DISABLED, message="envoi désactivé (ts_remuneration_import_enabled) : changements conservés")
        else:
            try:
                import_id, ts_status = follow(api, import_type, entry["file_name"], content)
                entry.update(ts_import_id=import_id, ts_import_status=ts_status,
                             status=FAILED if ts_status in IMPORT_TERMINAL_KO else SENT,
                             message="import suivi jusqu'à " + str(ts_status))
            except Exception as exc:
                entry.update(status=FAILED, message=safe_error(exc))
        entry["reference_advanced"] = entry["status"] in ADVANCING
        if entry["reference_advanced"]:
            store.advance_reference(run_id, IMPORT_SECTIONS[import_type], snapshot)
        if entry["status"] == FAILED:
            failures.append(import_type)
        logs.append(entry)
    store.log_imports(run_id, logs)
    summary = {"baseSalary_nb_source": len(snapshot[BASE_SALARY]), "baseSalary_nb_inserts": len(changes[IMPORT_SALARY]),
               "bonusPercentage_nb_source": len(snapshot[BONUS_PERCENT]),
               "bonusTargetValue_nb_source": len(snapshot[BONUS_TARGET_VALUE]),
               "bonus_nb_inserts": len(changes[IMPORT_BONUS])}
    result = {"run_id": run_id, "first_run": first_run, "summary": summary, "snapshot_stats": stats,
              "gate": gate.outcome, "imports": {e["import_type"]: e["status"] for e in logs},
              "status": "FAILED" if failures else "OK"}
    if failures:
        raise RemunerationFailed(result)
    return result


class RemunerationFailed(HrisError):
    def __init__(self, result: dict):
        super().__init__("TS_IMPORT_FAILED", ",".join(t for t, s in result["imports"].items() if s == FAILED))
        self.result = result


# ---------------------------------------------------------------------------
# Stockage Delta (Spark)
# ---------------------------------------------------------------------------
SNAPSHOT_TABLE_COLUMNS = REMUNERATION_SNAPSHOT_COLUMNS


class SparkRemunerationStore:
    def __init__(self, spark, tables):
        self.spark, self.tables = spark, tables

    def mapping(self) -> List[MappingRule]:
        rows = [r.asDict() for r in self.spark.table(self.tables("cfg.ts_remuneration_mapping")).collect()]
        if not rows:
            raise HrisError("REMUNERATION_MAPPING_EMPTY", "cfg.ts_remuneration_mapping vide : lancer NB_HRIS_SETUP")
        return rules_from_rows(rows)

    def reference(self) -> Dict[str, List[dict]]:
        out: Dict[str, List[dict]] = {section: [] for section in SECTIONS}
        for row in self.spark.table(self.tables("ref.ts_remuneration_snapshot")).collect():
            data = row.asDict()
            out.setdefault(data["section"], []).append(data)
        return out

    def _rows(self, snapshot: Mapping[str, List[dict]], sections: Iterable[str]) -> List[dict]:
        return [dict({c: line.get(c) for c in SNAPSHOT_TABLE_COLUMNS if c != "section"}, section=section)
                for section in sections for line in snapshot.get(section, [])]

    def save_snapshot(self, run_id: str, business_date, snapshot: Mapping[str, List[dict]]) -> None:
        rows = [dict(r, run_id=run_id, business_date=business_date) for r in self._rows(snapshot, SECTIONS)]
        logical = "stg.ts_remuneration_snapshot"
        IO.replace_run(self.spark, IO.rows_df(self.spark, rows, logical), self.tables(logical), logical, run_id)

    def save_changes(self, run_id: str, business_date, changes: Mapping[str, List[dict]]) -> None:
        rows = []
        for import_type, lines in changes.items():
            for number, line in enumerate(lines, start=1):
                rows.append(dict({k.replace("-", "_"): v for k, v in line.items()}, run_id=run_id,
                                 business_date=business_date, import_type=import_type, line_no=number))
        logical = "pub.ts_remuneration_change"
        IO.replace_run(self.spark, IO.rows_df(self.spark, rows, logical), self.tables(logical), logical, run_id)

    def advance_reference(self, run_id: str, sections: Sequence[str], snapshot: Mapping[str, List[dict]]) -> None:
        """La référence de ces sections devient l'instantané du run (remplacement ciblé, une version Delta)."""
        logical = "ref.ts_remuneration_snapshot"
        rows = [dict(r, _source_run_id=run_id, _validated_at_utc=utc_now()) for r in self._rows(snapshot, sections)]
        predicate = "section IN (" + ",".join(IO.sql_str(s) for s in sections) + ")"
        (IO.conform(IO.rows_df(self.spark, rows, logical), logical).write.format("delta").mode("overwrite")
         .option("replaceWhere", predicate).saveAsTable(self.tables(logical)))

    def log_imports(self, run_id: str, entries: List[dict]) -> None:
        logical = "log.ts_remuneration_import"
        IO.replace_run(self.spark, IO.rows_df(self.spark, entries, logical), self.tables(logical), logical, run_id)


def step_message(result: Mapping[str, Any]) -> str:
    return json.dumps({k: result.get(k) for k in ("first_run", "summary", "snapshot_stats", "gate", "imports", "status")},
                      ensure_ascii=False, default=str)[:7800]
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= MODULE hris.bootstrap =============
register_hris_module('hris.bootstrap', 'a5cec272f5a0a22bcc7fa26a767c44803fdae95853980f54738dcb84a10a431e', r'''"""hris.bootstrap — premier chargement guidé : pipeline PL_HRIS_BOOTSTRAP, lancé à la main et jamais planifié.

Le pipeline enchaîne l'installation (NB_HRIS_SETUP), les tests sur données fictives (NB_HRIS_TESTS), PLAN, un run complet
de l'orchestrateur quotidien (PL_HRIS_ORCHESTRATOR) dans le mode choisi par PLAN, puis STATUS (NB_HRIS_BOOTSTRAP).
  * PLAN juge la référence comme la comparaison (ref.ts_employee_adp, D-CMP-06) : vide -> INIT_REFERENCE (premier
    chargement, ou reprise après un premier run bloqué, car la référence reste vide tant qu'aucun run n'est promu) ;
    présente -> NORMAL. Le pipeline est donc rejouable sans jamais réinitialiser la référence. Un verrou tenu par un run
    en cours (bail non expiré) arrête le pipeline avant tout run.
  * STATUS lit l'issue du run lancé (ctl.run, ctl.run_step, log.anomaly) et en déduit UNE action suivante. Il n'écrit
    rien. Aucune donnée personnelle : statuts, codes de contrôle et comptes seulement.
"""
from __future__ import annotations

import datetime as dt
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence

from hris import spark_io as IO
from hris.common import (PROCESS_CODE, RUN_BLOCKED, RUN_NOT_PUBLISHED_NON_PRD, RUN_SKIPPED_CONCURRENT, RUN_SUCCEEDED,
                         to_utc_naive, utc_now)

MODE_INIT, MODE_NORMAL = "INIT_REFERENCE", "NORMAL"
PLAN_OK, PLAN_LOCK_BUSY = "OK", "LOCK_BUSY"

ENABLE_SCHEDULE = "ENABLE_SCHEDULE"     # run abouti : activer la planification du pipeline quotidien
ENTER_ADP_CODES = "ENTER_ADP_CODES"     # correspondances sans code ADP : grille à saisir, puis relancer ce pipeline
FIX_REMUNERATION = "FIX_REMUNERATION"   # run salariés abouti, flux rémunération en échec
RETRY_LATER = "RETRY_LATER"             # un autre run tenait le verrou
INVESTIGATE = "INVESTIGATE"             # échec ou blocage à analyser (ctl.run_step, log.anomaly)
NEXT_ACTIONS = (ENABLE_SCHEDULE, ENTER_ADP_CODES, FIX_REMUNERATION, RETRY_LATER, INVESTIGATE)
DONE_STATUSES = (RUN_SUCCEEDED, RUN_NOT_PUBLISHED_NON_PRD)   # NOT_PUBLISHED_NON_PRD : run complet hors PRD (garde-fou)
ADP_CODES_CONTROL = "MAPPING_TABLE_EMPTY"                    # table vide ou sans aucun code ADP (B-02)
DAILY_PIPELINE = "T4_PL_HRIS_ORCHESTRATOR"
ADMIN_NOTEBOOK = "T4_NB_HRIS_REFERENCE_ADMIN"


def plan(reference_rows: int, lock: Optional[Mapping[str, Any]], now: Optional[dt.datetime] = None) -> dict:
    """Mode du run à lancer, ou refus si un run tient le verrou."""
    now = to_utc_naive(now or utc_now())
    expires = lock.get("lease_expires_at_utc") if lock else None
    busy = bool(lock) and lock.get("lock_status") == "HELD" and (expires is None or to_utc_naive(expires) > now)
    mode = MODE_INIT if int(reference_rows) == 0 else MODE_NORMAL
    reason = ("référence vide : premier chargement (INIT_REFERENCE)" if mode == MODE_INIT else
              f"référence présente ({int(reference_rows)} lignes) : premier chargement déjà fait, run courant (NORMAL)")
    result = {"status": PLAN_LOCK_BUSY if busy else PLAN_OK, "run_mode": mode, "reference_rows": int(reference_rows),
              "reason": reason}
    if busy:
        result["message"] = ("un run tient le verrou jusqu'à " + str(expires) + " (run " + str(lock.get("holder_run_id"))
                             + ") : relancer ce pipeline après sa fin")
    return result


def latest_steps(steps: Iterable[Mapping[str, Any]]) -> Dict[str, str]:
    """Dernier statut de chaque étape : le journal est en ajout (STARTED puis statut final, tentatives numérotées)."""
    best: Dict[str, tuple] = {}
    for step in steps:
        key = (int(step.get("attempt_no") or 0), step.get("status") != "STARTED",
               str(step.get("ended_at_utc") or step.get("started_at_utc") or ""))
        code = step.get("step_code")
        if code and (code not in best or key >= best[code][0]):
            best[code] = (key, step.get("status"))
    return {code: value[1] for code, value in sorted(best.items())}


def outcome(run: Optional[Mapping[str, Any]], steps: Iterable[Mapping[str, Any]],
            anomalies: Iterable[Mapping[str, Any]]) -> dict:
    """Action suivante après le run de l'orchestrateur ; `ok` seulement si le premier chargement est terminé."""
    if run is None:
        return {"ok": False, "next_action": INVESTIGATE, "run_id": None, "run_mode": None, "run_status": None,
                "steps": {}, "blocking_controls": [], "remuneration": None, "error_code": None,
                "message": "aucun run enregistré par l'orchestrateur : voir l'activité RunStart de " + DAILY_PIPELINE}
    status, mode = run.get("status"), run.get("run_mode")
    step_status = latest_steps(steps)
    blocking = sorted({a["control_code"] for a in anomalies
                       if a.get("severity") == "BLOCKING" and a.get("outcome") == "FAIL" and a.get("control_code")})
    remuneration = step_status.get("REMUNERATION")
    failed_steps = sorted(code for code, value in step_status.items() if value == "FAILED")
    if status == RUN_SKIPPED_CONCURRENT:
        action = RETRY_LATER
        message = "un autre run tenait le verrou : relancer ce pipeline plus tard"
    elif status == RUN_BLOCKED and ADP_CODES_CONTROL in blocking:
        action = ENTER_ADP_CODES
        message = ("correspondances sans code ADP (" + ADP_CODES_CONTROL + ") : coller la grille affichée par "
                   "l'activité Status dans p_grid de " + ADMIN_NOTEBOOK + " (CHECK, puis APPLY), puis relancer ce "
                   "pipeline" + _others(blocking, {ADP_CODES_CONTROL}))
    elif status in DONE_STATUSES and remuneration == "FAILED":
        action = FIX_REMUNERATION
        message = (f"run {mode} abouti ({status}) mais flux rémunération en échec (ctl.run_step, étape REMUNERATION) : "
                   "rétablir l'accès au coffre TalentSoft ou passer remuneration_enabled à false, puis relancer")
    elif status in DONE_STATUSES:
        action = ENABLE_SCHEDULE
        message = f"run {mode} abouti ({status}) : activer la planification quotidienne de {DAILY_PIPELINE}"
    else:
        action = INVESTIGATE
        message = (f"run {mode} terminé en {status}"
                   + (" ; contrôles bloquants : " + ", ".join(blocking) if blocking else "")
                   + (" ; étapes en échec : " + ", ".join(failed_steps) if failed_steps else "")
                   + (" ; erreur : " + str(run.get("error_code")) if run.get("error_code") else "")
                   + " : voir ctl.run_step et log.anomaly")
    return {"ok": action == ENABLE_SCHEDULE, "next_action": action, "run_id": run.get("run_id"), "run_mode": mode,
            "run_status": status, "steps": step_status, "blocking_controls": blocking, "remuneration": remuneration,
            "error_code": run.get("error_code"), "message": message}


def _others(blocking: Sequence[str], handled: set) -> str:
    rest = [code for code in blocking if code not in handled]
    return " ; autres contrôles bloquants : " + ", ".join(rest) if rest else ""


# ---------------------------------------------------------------------------
# Lectures Spark (aucune écriture)
# ---------------------------------------------------------------------------
def read_plan_inputs(spark, T) -> dict:
    lock = (spark.table(T("ctl.process_lock")).where(f"process_code = {IO.sql_str(PROCESS_CODE)}").collect())
    return {"reference_rows": spark.table(T("ref.ts_employee_adp")).count(),
            "lock": lock[0].asDict() if lock else None}


def find_run(spark, T, since: dt.datetime) -> Optional[dict]:
    """Run le plus récent du processus ouvert depuis `since` (déclenchement du pipeline d'amorçage)."""
    bound = to_utc_naive(since).isoformat(sep=" ", timespec="microseconds")
    rows = (spark.table(T("ctl.run"))
            .where(f"process_code = {IO.sql_str(PROCESS_CODE)} AND started_at_utc >= TIMESTAMP'{bound}'")
            .orderBy("started_at_utc", ascending=False).limit(1).collect())
    return rows[0].asDict() if rows else None


def read_status_inputs(spark, T, run_id: str) -> dict:
    predicate = f"run_id = {IO.sql_str(IO.check_run_id(run_id))}"
    steps = [r.asDict() for r in spark.table(T("ctl.run_step")).where(predicate)
             .select("step_code", "attempt_no", "status", "started_at_utc", "ended_at_utc").collect()]
    anomalies = [r.asDict() for r in spark.table(T("log.anomaly")).where(predicate)
                 .select("control_code", "severity", "outcome").collect()]
    return {"steps": steps, "anomalies": anomalies}
''')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# ============= SYNTHÈSE DU CHARGEMENT =============
try:
    hris_library_digests = {name: getattr(sys.modules[name], "__hris_digest__", None)
                            for name in sorted(sys.modules) if name.startswith("hris.")}
    print("[T4_NB_HRIS_LIB] ✓ bibliothèque HRIS chargée :", len(hris_library_digests), "modules")
except Exception:
    traceback.print_exc()
    raise

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }
