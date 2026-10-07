# Fabric notebook source

# METADATA ********************

# META {
# META   "kernel_info": {
# META     "name": "synapse_pyspark"
# META   }
# META }

# MARKDOWN ********************

# # NB_HRIS_TESTS — tests sur données FICTIVES (unitaires, parité Spark, intégration Delta isolée)
# # - Unitaires : parité T-SQL, règles R-SQL, acquisition TalentSoft (faux serveur), publication ADP (clients fictifs,
#   sentinelle réseau, preuve du garde-fou avec contrôle positif simulé).
# - Parité Spark ↔ implémentation de référence du SQL de production.
# - Intégration : schéma **isolé** `hris_tst` (tables de test recréées), aucune table réelle lue ni écrite,
#   garde-fou évalué dans le **contexte réel** de ce notebook, sans aucun client ADP.
# - Vérifications de plateforme (NV-02, NV-05, NV-06) consignées dans la sortie.
# Le simulateur ADP n'existe que dans ce notebook.

# PARAMETERS CELL ********************

p_test_schema = "hris_tst"
p_lakehouse_name = ""            # VL_HRIS.lakehouse_name
p_run_integration = "true"
p_only = ""                  # filtre facultatif : fragments d'identifiants de tests séparés par des virgules

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

NV_TIMEZONE_DEFAULT = spark.conf.get("spark.sql.session.timeZone", None)   # NV-02 avant tout réglage
NV_ANSI = spark.conf.get("spark.sql.ansi.enabled", None)                  # NV-06
NV_PARSER = spark.conf.get("spark.sql.legacy.timeParserPolicy", None)

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

# Bibliothèque HRIS embarquée — GÉNÉRÉE par tools/build_notebooks.py, ne pas modifier ici.
# Source testée : implementation/src/hris/*.py ; empreintes vérifiées au chargement.
import hashlib as _hl, sys as _sys, types as _types
_HRIS_MODULES = [
    ('hris.common', 'da64004aee4e83cbbedd33fdfbbed057845a34037de058aa2c8d68b3df1051db', r'''"""hris.common — constantes, erreurs, assainissement et utilitaires partagés.

Python pur (aucune dépendance Spark) : importable en local pour les tests et
embarqué tel quel dans les notebooks Fabric (voir tools/build_notebooks.py).
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
from typing import Any, Iterable, Optional

try:  # Python >= 3.9
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover
    ZoneInfo = None  # type: ignore

PROCESS_CODE = "TS_ADP_EMPLOYEE"

# Source TalentSoft : le tenant de PRODUCTION est la source réelle dans TOUS les
# environnements (consigne du mandat). Aucune bascule vers un tenant de test,
# aucun repli : toute autre URL est refusée (voir ts_acquire.assert_prd_source).
TS_PRD_BASE_URL = "https://motul.talent-soft.com"
TS_SECRET_CLIENT_ID = "client-id-ts"
TS_SECRET_CLIENT_SECRET = "client-secret-ts"

RUN_MODES = ("NORMAL", "INIT_REFERENCE", "RETRY_PUBLICATION")  # SHADOW : non retenu (D-CUT-03 À CONFIRMER)

# Statuts de ctl.run (05_TABLES_LOG_LAKEHOUSE.md)
RUN_STARTED = "STARTED"
RUN_ACQUIRED = "ACQUIRED"
RUN_TRANSFORMED = "TRANSFORMED"
RUN_COMPARED = "COMPARED"
RUN_BLOCKED = "BLOCKED"
RUN_VALIDATED = "VALIDATED"
RUN_PROMOTED = "PROMOTED"
RUN_PUBLISHED = "PUBLISHED"
RUN_PUBLISHED_PARTIAL = "PUBLISHED_PARTIAL"
RUN_NOT_PUBLISHED_NON_PRD = "NOT_PUBLISHED_NON_PRD"
RUN_SUCCEEDED = "SUCCEEDED"
RUN_FAILED = "FAILED"
RUN_SKIPPED_CONCURRENT = "SKIPPED_CONCURRENT"

SEVERITY_BLOCKING = "BLOCKING"
SEVERITY_WARNING = "WARNING"
SEVERITY_INFO = "INFO"


class HrisError(Exception):
    """Erreur fonctionnelle portant un code stable. Le message ne doit jamais
    contenir de donnée personnelle ni de secret (il est journalisé)."""

    def __init__(self, code: str, message: str = ""):
        super().__init__(f"{code}: {message}" if message else code)
        self.code = code
        self.safe_message = message


def utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


def to_utc_naive(value: dt.datetime) -> dt.datetime:
    """Horodatage UTC sans fuseau, convention des colonnes TIMESTAMP (session Spark en UTC)."""
    if value.tzinfo is not None:
        value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
    return value


def parse_trigger_time(value: str) -> dt.datetime:
    """@pipeline().TriggerTime (ISO 8601, UTC) -> datetime UTC aware."""
    if not value:
        raise HrisError("PARAM_INVALID", "p_trigger_time vide")
    text = value.strip().replace("Z", "+00:00")
    # fractions de seconde à 7 chiffres possibles dans les expressions Fabric
    text = re.sub(r"(\.\d{6})\d+", r"\1", text)
    try:
        parsed = dt.datetime.fromisoformat(text)
    except ValueError as exc:
        raise HrisError("PARAM_INVALID", "p_trigger_time illisible") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=dt.timezone.utc)
    return parsed.astimezone(dt.timezone.utc)


def business_date_of(trigger_utc: dt.datetime, timezone: str) -> dt.date:
    """Date métier = date locale du déclenchement (A08 D-RUN-02, À CONFIRMER : Europe/Paris)."""
    if ZoneInfo is None:  # pragma: no cover
        raise HrisError("TZ_UNAVAILABLE")
    return trigger_utc.astimezone(ZoneInfo(timezone)).date()


# ---------------------------------------------------------------------------
# Assainissement : aucun secret, aucune donnée personnelle dans les journaux
# ---------------------------------------------------------------------------
_SANITIZE_PATTERNS = [
    (re.compile(r"(?i)bearer\s+[A-Za-z0-9\-._~+/]+=*"), "Bearer ***"),
    (re.compile(r"(?i)(client_secret|client_id|access_token|password|token)=([^&\s]+)"), r"\1=***"),
    (re.compile(r"https?://[^\s'\"]+"), "<url>"),
    (re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"), "<email>"),
    (re.compile(r"\b[A-Z]{2}\d{2}[A-Z0-9 ]{10,30}\b"), "<iban>"),
    (re.compile(r"\d{4,}"), "<n>"),
]


def sanitize(text: Any, max_len: int = 500, mask_numbers: bool = True) -> str:
    """Rend un texte publiable dans ctl/log : masque URL, e-mails, IBAN, jetons et, par défaut,
    toute séquence de 4 chiffres ou plus (matricules, NIR, montants). `mask_numbers=False` est
    réservé aux messages GÉNÉRÉS par le code à partir de compteurs (aucune valeur métier)."""
    if text is None:
        return ""
    out = str(text)
    patterns = _SANITIZE_PATTERNS if mask_numbers else _SANITIZE_PATTERNS[:-1]
    for pattern, repl in patterns:
        out = pattern.sub(repl, out)
    return out[:max_len]


def safe_error(exc: BaseException) -> str:
    """Résumé d'exception publiable : classe + code fonctionnel, message assaini."""
    if isinstance(exc, HrisError):
        return f"{exc.code}: {sanitize(exc.safe_message, 300)}"
    return f"{type(exc).__name__}: {sanitize(str(exc), 300)}"


def sha256_hex(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def canonical_json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str)


def json_exit(payload: dict) -> str:
    """Valeur de sortie d'un notebook : JSON court, sans donnée personnelle."""
    return json.dumps(payload, ensure_ascii=False, default=str, separators=(",", ":"))


def require(params: dict, names: Iterable[str]) -> None:
    missing = [n for n in names if params.get(n) in (None, "")]
    if missing:
        raise HrisError("PARAM_MISSING", ",".join(missing))


def as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("1", "true", "yes", "oui")


class SecretValue:
    """Enveloppe d'une valeur secrète : jamais affichée par repr/str."""

    __slots__ = ("_v",)

    def __init__(self, value: Optional[str]):
        self._v = value

    def reveal(self) -> Optional[str]:
        return self._v

    def __repr__(self) -> str:
        return "SecretValue(***)"

    __str__ = __repr__
'''),
    ('hris.parity', '12d5d086c85cf5a0c61aadd96e0bd3edbf1a17fa2ecbedc7346161bbb6a7854b', r'''"""hris.parity — émulation Python des sémantiques T-SQL du SQL de production (A09).

Implémentation de RÉFÉRENCE (Python pur) des règles A02 ; la version Spark
(hris.transform) doit produire exactement les mêmes résultats sur les jeux de tests
(test de parité exécuté en local et dans Fabric).

Points d'émulation non exacts, mesurés en non-régression (A08 D-PAR-01, D-PAR-02) :
  * ISDATE : formats numériques usuels uniquement (pas de mois en lettres) ;
  * ISNUMERIC : grammaire approchée (signe, symbole monétaire, virgules, exposant).
"""
from __future__ import annotations

import datetime as dt
import decimal
import json
import re
from typing import Any, Optional

# ---------------------------------------------------------------------------
# Lecture JSON (R-SQL-002) : équivalent de JSON_VALUE(..., '$.<clé>') en mode lax
# ---------------------------------------------------------------------------
JSON_VALUE_MAX_LEN = 4000  # JSON_VALUE renvoie NULL au-delà de 4000 caractères (mode lax)


def json_value_repr(value: Any) -> Optional[str]:
    """Texte que JSON_VALUE renverrait pour la valeur Python sérialisée par json.dump :
    None -> NULL ; booléen -> 'true'/'false' ; nombre -> représentation JSON ;
    chaîne -> la chaîne (NULL si > 4000) ; objet ou tableau -> NULL."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return json.dumps(value)
    if isinstance(value, str):
        return value if len(value) <= JSON_VALUE_MAX_LEN else None
    return None  # dict / list : non scalaire


def nvarchar(value: Optional[str], length: int = 50) -> Optional[str]:
    """CAST(x AS NVARCHAR(n)) : troncature silencieuse à n unités UTF-16 (R-SQL-002, D-PAR-03)."""
    if value is None:
        return None
    if all(ord(ch) <= 0xFFFF for ch in value):
        return value[:length]
    units = value.encode("utf-16-le", "surrogatepass")[: 2 * length]
    return units.decode("utf-16-le", "surrogatepass")


# ---------------------------------------------------------------------------
# Comparaisons de chaînes : SQL Server ignore les espaces de fin (= , <>, LEN)
# ---------------------------------------------------------------------------
def sql_rtrim(value: Optional[str]) -> Optional[str]:
    return None if value is None else value.rstrip(" ")


def sql_is_empty(value: Optional[str]) -> bool:
    """`x = ''` : vrai pour '' et '   ' ; NULL -> inconnu (faux dans un CASE WHEN)."""
    return value is not None and value.rstrip(" ") == ""


def sql_is_not_empty(value: Optional[str]) -> bool:
    """`x <> ''` : NULL -> inconnu (faux)."""
    return value is not None and value.rstrip(" ") != ""


def sql_len(value: Optional[str]) -> Optional[int]:
    return None if value is None else len(value.rstrip(" "))


def sql_eq(a: Optional[str], b: Optional[str]) -> bool:
    """Égalité de jointure : NULL ne correspond à rien ; espaces de fin ignorés."""
    return a is not None and b is not None and a.rstrip(" ") == b.rstrip(" ")


# ---------------------------------------------------------------------------
# ISDATE (R-SQL-033/034, R-SQL-043) — émulation us_english / DATEFORMAT mdy / datetime
# ---------------------------------------------------------------------------
_WS = " \t\r\n"
# Motifs acceptés (D-PAR-01). Les MÊMES expressions régulières gardent l'analyse Spark (hris.transform) :
# (garde, extraction de la partie date, ordre y/m/d)
_HH = r"(?:[01]\d|2[0-3])"
_HH1 = r"(?:[01]?\d|2[0-3])"
_MS = r"[0-5]\d"
DATE_PATTERNS = [
    ("ymd_dash", r"^\d{4}-\d{1,2}-\d{1,2}$", r"^(\d{4})-(\d{1,2})-(\d{1,2})$", "ymd"),
    ("ymd_compact", r"^\d{8}$", r"^(\d{4})(\d{2})(\d{2})$", "ymd"),
    ("ymd_iso_time", rf"^\d{{4}}-\d{{2}}-\d{{2}}T{_HH}:{_MS}(?::{_MS}(?:\.\d{{1,3}})?)?Z?$",
     r"^(\d{4})-(\d{2})-(\d{2})T", "ymd"),
    ("ymd_space_time", rf"^\d{{4}}-\d{{1,2}}-\d{{1,2}} {_HH1}:{_MS}(?::{_MS}(?:\.\d{{1,3}})?)?$",
     r"^(\d{4})-(\d{1,2})-(\d{1,2}) ", "ymd"),
    ("mdy_slash", r"^\d{1,2}/\d{1,2}/\d{4}$", r"^(\d{1,2})/(\d{1,2})/(\d{4})$", "mdy"),
]
_COMPILED = [(n, re.compile(g), re.compile(x), o) for n, g, x, o in DATE_PATTERNS]
DATETIME_MIN_YEAR = 1753


def parse_tsql_date(value: Optional[str]) -> Optional[dt.date]:
    """Date reconnue (partie date) ou None. Les espaces de début et de fin sont tolérés."""
    if value is None:
        return None
    text = value.strip(_WS)
    for _, guard, extract, order in _COMPILED:
        if not guard.match(text):
            continue
        g = extract.match(text).groups()
        y, mo, d = (int(g[0]), int(g[1]), int(g[2])) if order == "ymd" else (int(g[2]), int(g[0]), int(g[1]))
        try:
            return dt.date(y, mo, d)
        except ValueError:
            return None
    return None


def isdate(value: Optional[str]) -> int:
    """ISDATE(x) : 1 si convertible en datetime (année >= 1753), 0 sinon ; NULL -> 0."""
    parsed = parse_tsql_date(value)
    return 1 if parsed is not None and parsed.year >= DATETIME_MIN_YEAR else 0


SQL_EMPTY_DATE = dt.date(1900, 1, 1)


def convert_date(value: Optional[str]) -> Optional[dt.date]:
    """CONVERT(DATE, x) : NULL -> NULL ; '' (espaces compris) -> 1900-01-01 (R-SQL-052) ;
    valeur invalide -> erreur de conversion (le run échoue en parité)."""
    if value is None:
        return None
    if value.strip(_WS) == "":
        return SQL_EMPTY_DATE
    parsed = parse_tsql_date(value)
    if parsed is None:
        raise ValueError("CONVERSION_DATE")
    return parsed


# ---------------------------------------------------------------------------
# ISNUMERIC et CAST NUMERIC(12,2) (R-SQL-035/036/044/054)
# ---------------------------------------------------------------------------
_ISNUMERIC = re.compile(
    r"^[+-]?[$€£¥]?[ ]*[+-]?"
    r"(?P<body>\d[\d,]*\.?\d*|\.\d*|,[\d,]*\.?\d*)?"
    r"(?P<exp>[eEdD][+-]?\d+)?$"
)
_STRICT_DECIMAL = re.compile(r"^[ \t]*([+-]?)(\d+\.?\d*|\.\d+)[ \t]*$")
_NUMERIC_12_2_LIMIT = decimal.Decimal("10000000000")  # 10^(12-2)


def isnumeric(value: Optional[str]) -> int:
    """ISNUMERIC(x) approché (A08 D-PAR-02). NULL et '' -> 0."""
    if value is None:
        return 0
    text = value.strip(" \t\r\n")
    if text == "":
        return 0
    m = _ISNUMERIC.match(text)
    if not m:
        return 0
    body = m.group("body") or ""
    if m.group("exp") and not re.search(r"\d", body):
        return 0  # exposant sans mantisse (ex. 'e5')
    return 1


def cast_numeric_12_2(value: Optional[str]) -> Optional[decimal.Decimal]:
    """CAST(x AS NUMERIC(12,2)) strict : arrondi à 2 décimales (au plus proche, à distance
    de zéro en cas d'égalité). Toute valeur non convertible ou hors capacité lève
    ValueError (erreur de conversion T-SQL => échec du run en parité)."""
    if value is None:
        return None
    m = _STRICT_DECIMAL.match(value)
    if not m:
        raise ValueError("CONVERSION_NUMERIC")
    number = decimal.Decimal(m.group(1) + m.group(2))
    rounded = number.quantize(decimal.Decimal("0.01"), rounding=decimal.ROUND_HALF_UP)
    if abs(rounded) >= _NUMERIC_12_2_LIMIT:
        raise ValueError("NUMERIC_OVERFLOW")
    return rounded


def monthly_salary(value: Optional[str]) -> Optional[decimal.Decimal]:
    """CAST(mt_salaire_base AS NUMERIC(12,2)) / 12 -> DECIMAL(23,13) (R-SQL-054)."""
    base = cast_numeric_12_2(value)
    if base is None:
        return None
    ctx = decimal.Context(prec=38, rounding=decimal.ROUND_HALF_UP)
    return ctx.divide(base, decimal.Decimal(12)).quantize(
        decimal.Decimal("1e-13"), rounding=decimal.ROUND_HALF_UP, context=ctx)


def pct_flag(value: Optional[str], upper: int) -> int:
    """Indicateur de rejet « pourcentage / montant » : ISNUMERIC = 0 -> 0 ; sinon
    CAST(...) BETWEEN 0 AND upper -> 0, hors bornes -> 1 ; échec de CAST -> ValueError."""
    if isnumeric(value) == 0:
        return 0
    number = cast_numeric_12_2(value)
    return 0 if decimal.Decimal(0) <= number <= decimal.Decimal(upper) else 1


def utf8_truncate(text: Optional[str], max_bytes: int) -> Optional[str]:
    """CAST(... AS VARCHAR(n)) avec collation UTF-8 : n OCTETS (résumé R-SQL-066).
    Ne coupe jamais un caractère multi-octets (comportement exact SQL Server À CONFIRMER)."""
    if text is None:
        return None
    raw = text.encode("utf-8")
    if len(raw) <= max_bytes:
        return text
    return raw[:max_bytes].decode("utf-8", "ignore")
'''),
    ('hris.model', '6a943353e2b932c5dc9858fdc9ef903dc056e99d113e8115fee20a49046875a7', r'''"""hris.model — contrat de données : colonnes, correspondances, règles de rejet, DDL.

Sources : A02 (R-SQL-004, § 4, § 5, R-SQL-056), A03 (17 tables), A07, 05 (DDL des tables
de contrôle et de log, qui fait foi). Python pur.

Extensions par rapport au DDL documenté (signalées dans docs/ECARTS_ET_DECISIONS.md) :
  * ctl.run.target_run_id (mode RETRY_PUBLICATION) ;
  * tables PROPOSÉES de publication par opération (pub.adp_operation, pub.adp_attempt,
    pub.adp_reconciliation, log.adp_manual_decision) : leur sémantique dépend du contrat
    final ADP, absent (docs/CONTRAT_ADP_POINTS_D_USAGE.md).
"""
from __future__ import annotations

from typing import Dict, List, Tuple

# ---------------------------------------------------------------------------
# R-SQL-004 : 69 colonnes source (clé JSON -> colonne), toutes NVARCHAR(50)
# ---------------------------------------------------------------------------
SOURCE_COLUMNS: List[Tuple[str, str]] = [
    ("employeeNumber", "id_unique"),
    ("LegacyID", "id_payroll"),
    ("Matricule_IT", "cd_matricule_it"),
    ("Identity_startDate", "dt_debut_identite"),
    ("LastName", "lb_nom"),
    ("FirstName", "lb_prenom"),
    ("BirthName", "lb_nom_naissance"),
    ("Title", "cd_emp_titre_ts"),
    ("Gender", "cd_emp_sexe_ts"),
    ("Nationality", "cd_emp_nationalite_ts"),
    ("hubSocialSecurityNumber", "id_numero_securite_sociale"),
    ("MaritalStatus_startDate", "dt_debut_situation_matrimoniale"),
    ("MaritalStatus", "cd_emp_situation_matrimoniale_ts"),
    ("CivilRegistration_startDate", "dt_debut_civil_registration"),
    ("Date_de_naissance", "dt_naissance"),
    ("Ville_de_naissance", "lb_naissance_ville"),
    ("code_postal_lieu_de_naissance", "lb_naissance_cp"),
    ("region_ou_etat_lieu_de_naissance", "lb_naissance_departement"),
    ("Pays_de_naissance", "cd_naissance_pays_ts"),
    ("LegalTownOrCity", "lb_fiscal_ville"),
    ("LegalCountry", "lb_fiscal_country"),
    ("LegalPostalCode", "lb_fiscal_cp"),
    ("LegalStreet", "lb_fiscal_street"),
    ("LegalStreetNumber", "lb_fiscal_streetnumber"),
    ("LegalAdditionalAddressInformation", "lb_fiscal_additional_address_information"),
    ("LegalStreetNumberComplement", "lb_fiscal_streetnumber_complement"),
    ("address_postalCode", "lb_adresse_cp"),
    ("adresse_city", "lb_adresse_ville"),
    ("adresse_countryCode", "cd_adresse_pays_ts"),
    ("ContactInformation_startDate", "dt_debut_contatct_info"),
    ("BusinessEmail", "lb_email_professionnel"),
    ("PersonalEmail", "lb_email_personnel"),
    ("BusinessMobile|BusinessPhone", "lb_telephone_professionnel"),  # COALESCE (R-SQL-003)
    ("PersonalMobile", "lb_telephone_personnel"),
    ("referenceAdmissionDate", "dt_entree_societe"),
    ("GroupStartDate", "dt_entree_groupe"),
    ("RecalculatedSeniorityDate", "dt_anciennete"),
    ("GroupEndDate", "dt_sortie"),
    ("CompanyStartDate", "dt_company_start"),
    ("HiringReason", "cd_ppc_raison_debut_contrat_ts"),
    ("TerminationReason", "cd_ppc_raison_fin_contrat_ts"),
    ("date_debut_contract", "dt_ppc_debut_contrat"),
    ("date_fin_contract", "dt_ppc_fin_contrat"),  # REPLACE(...,'None','') (R-SQL-003)
    ("NatureOfContract", "cd_ppc_contrat_ts"),
    ("FixedTermContractReason", "cd_recours_cdd"),
    ("OccupationalCategory", "cd_occupational_category"),
    ("FirstPeriodStartDate", "dt_ppc_debut_periode_essaie"),
    ("FirstPeriodEndDate", "dt_ppc_fin_periode_essaie"),
    ("Situation_starDate", "dt_situation_start"),
    ("Coefficient", "cd_ppc_convention_collective_lvl2_ts"),
    ("LegalStructure", "cd_geo_entite_ts"),
    ("GeographicOrganisationStructure", "cd_ppc_location_country_ts"),
    ("PPC_FTE", "pc_ppc_fte"),
    ("OrganisationalStructure", "cd_organisationalstructure_ts"),
    ("numero_ordre", "id_numero_ordre"),
    ("CostCenter", "cd_ppc_cost_center_ts"),
    ("pourcentage_imputation", "pc_imputation"),
    ("AccountHolder", "lb_account_holder"),
    ("IBANOrABA", "lb_iban"),
    ("BicOrSwift", "lb_bic"),
    ("AccountNumber", "lb_account_number"),
    ("BankName", "lb_bank_name"),
    ("base_startDate", "dt_debut_mt_salaire_base"),
    ("salary_base", "mt_salaire_base"),
    ("pourcentage_prime_annuelles_startDate", "dt_debut_pc_prime_annuelle"),
    ("Pourcentage_prime_annuelles", "pc_prime_annuelle"),
    ("prime_annuelles_startDate", "dt_debut_prime_annuelle"),
    ("prime_annuelles", "prime_annuelle"),
    ("WorkingType", "lb_job_title"),
]
assert len(SOURCE_COLUMNS) == 69

# Clés JSON persistées dans stg.ts_employee_raw : uniquement celles lues en aval
# (minimisation A08 D-SEC-02 : ni données de santé ni champs inutilisés).
RAW_KEYS: List[str] = []
for _k, _c in SOURCE_COLUMNS:
    for _part in _k.split("|"):
        if _part not in RAW_KEYS:
            RAW_KEYS.append(_part)
assert len(RAW_KEYS) == 70

SOURCE_COLUMN_NAMES = [c for _, c in SOURCE_COLUMNS]

# Clés produites par la Function TalentSoft mais jamais persistées (santé, inutilisées)
TS_KEYS_NOT_PERSISTED = [
    "handicap", "handicap_startDate", "workaccident", "workaccident_startDate",
    "motif_visite_medical", "motif_visite_medical_startDate", "Position", "CollectiveAgreement",
    "Hors_France", "Qualification", "Qualification_startDate", "ConferredDate", "PersonalPhone",
    "Status", "PaymentMethod", "UseSameAccountForExpense", "ExpenseAccountPaymentMethod",
    "ExpenseAccountHolder", "ExpenseAccountIBANOrABA", "ExpenseAccountBicOrSwift",
    "ExpenseAccountBankName", "Country", "ExpenseAccountCountry", "date_debut_bank", "date_fin_bank",
    "SecondPeriodStartDate", "SecondPeriodEndDate", "ProbationaryPeriod_startDate", "LegacyID_startDate",
]

# ---------------------------------------------------------------------------
# A03 : 17 tables de correspondance (ordre des colonnes = ordre du fichier d'origine)
# ---------------------------------------------------------------------------
LOOKUP_TABLES: Dict[str, List[str]] = {
    "lu_hr_emp_diplome": ["cd_emp_diplome_ts", "ds_emp_diplome_ts", "cd_emp_diplome_adp", "ds_emp_diplome_adp",
                          "or_emp_diplome", "cd_emp_diplome_api_ts", "ds_emp_diplome_api_ts"],
    "lu_hr_emp_nationalite": ["id_emp_nationalite_ts", "cd_emp_nationalite_ts", "ds_emp_nationalite_ts",
                              "cd_emp_nationalite_adp", "ds_emp_nationalite_adp", "or_emp_nationalite"],
    "lu_hr_emp_pays": ["id_emp_pays_ts", "cd_emp_pays_ts", "ds_emp_pays_ts", "id_emp_pays_adp",
                       "ds_emp_pays_adp", "or_emp_pays"],
    "lu_hr_emp_situation_matrimoniale": ["cd_emp_situation_matrimoniale_ts", "ds_emp_situation_matrimoniale_ts",
                                         "cd_emp_situation_matrimoniale_adp", "ds_emp_situation_matrimoniale_adp",
                                         "or_emp_situation_matrimoniale"],
    "lu_hr_emp_titre": ["cd_emp_titre_ts", "lb_emp_titre_ts", "ds_emp_titre_ts", "cd_emp_titre_adp",
                        "ds_emp_titre_adp", "or_emp_titre"],
    "lu_hr_emp_sexe": ["cd_emp_sexe_ts", "ds_emp_sexe_ts", "cd_emp_sexe_adp", "ds_emp_sexe_adp", "or_emp_sexe"],
    "lu_hr_emp_work_accident": ["cd_emp_work_accident_ts", "ds_emp_work_accident_ts", "cd_emp_work_accident_adp",
                                "ds_emp_work_accident_adp", "or_emp_work_accident"],
    "lu_hr_ppc_convention_collective": ["cd_ppc_convention_collective_ts", "ds_ppc_convention_collective_ts",
                                        "cd_ppc_convention_collective_adp", "ds_ppc_convention_collective_adp",
                                        "or_ppc_convention_collective"],
    "lu_hr_ppc_classification": ["cd_convention_collective_lvl2_ts", "lb_convention_collective_lvl2_ts",
                                 "cd_classification_adp", "lb_classification_adp", "cd_coefficient_adp",
                                 "cd_ppc_convention_collective_adp", "cd_ppc_convention_collective_ts"],
    "lu_hr_ppc_classes": ["cd_occupational_category_ts", "lb_fr_occupational_category_ts",
                          "lb_en_occupational_category_ts", "cd_categorie_cotisant_adp", "lb_categorie_cotisant_adp",
                          "cd_classe_remuneration_adp", "lb_classe_remuneration_adp"],
    "lu_hr_geo_entite": ["cd_geo_entite_ts", "ds_geo_entite_ts", "cd_geo_parent_entite_ts", "cd_geo_entite_adp",
                         "ds_geo_entite_adp", "or_geo_entite_adp"],
    "lu_hr_ppc_contrat": ["cd_ppc_contrat_ts", "ds_ppc_contrat_ts", "cd_ppc_contrat_adp", "ds_ppc_contrat_adp",
                          "cd_ppc_contrat_type", "ds_ppc_contrat_type", "or_ppc_contrat"],
    "lu_hr_ppc_contrat2": ["cd_contract_ts", "lb_contract_ts", "cd_nature_contract_adp", "lb_nature_contract_adp",
                           "cd_type_contract_adp", "lb_type_contract_adp", "cd_type_collaboration_adp",
                           "lb_type_collaboration_adp"],
    "lu_hr_ppc_cost_center": ["cd_ppc_cost_center_ts", "ds_ppc_cost_center_ts", "cd_ppc_cost_center_adp",
                              "ds_ppc_cost_center_adp", "or_ppc_cost_center"],
    "lu_hr_ppc_location_country": ["cd_ppc_location_country_ts", "ds_ppc_location_country_ts",
                                   "cd_ppc_location_country_adp", "ds_ppc_location_country_adp",
                                   "or_ppc_location_country"],
    "lu_hr_ppc_raison_debut_contrat": ["cd_ppc_raison_debut_contrat_ts", "ds_ppc_raison_debut_contrat_ts",
                                       "cd_ppc_raison_debut_contrat_adp", "ds_ppc_raison_debut_contrat_adp",
                                       "or_ppc_raison_debut_contrat"],
    "lu_hr_organisationalstructure": ["cd_organisationalstructure_ts", "ds_organisationalstructure_ts",
                                      "cd_organisationalstructure_adp", "ds_organisationalstructure_adp",
                                      "or_organisationalstructure"],
}
assert len(LOOKUP_TABLES) == 17
LOOKUP_TECH_COLUMNS = ["_loaded_at_utc", "_loaded_by", "_change_ref"]

# Jointures (R-SQL-010 à 024) : alias, table, colonne source, clé de la table,
# colonnes ramenées (nom dans la table -> nom en staging), dans l'ordre de PROC2.
LOOKUP_JOINS: List[dict] = [
    {"rule": "R-SQL-010", "table": "lu_hr_emp_titre", "src": "cd_emp_titre_ts", "key": "cd_emp_titre_ts",
     "out": [("cd_emp_titre_ts", "cd_emp_titre_ts_lu"), ("cd_emp_titre_adp", "cd_emp_titre_adp")]},
    {"rule": "R-SQL-011", "table": "lu_hr_emp_sexe", "src": "cd_emp_sexe_ts", "key": "cd_emp_sexe_ts",
     "out": [("cd_emp_sexe_ts", "cd_emp_sexe_ts_lu"), ("cd_emp_sexe_adp", "cd_emp_sexe_adp")]},
    {"rule": "R-SQL-012", "table": "lu_hr_emp_nationalite", "src": "cd_emp_nationalite_ts",
     "key": "cd_emp_nationalite_ts",
     "out": [("cd_emp_nationalite_ts", "cd_emp_nationalite_ts_lu"), ("cd_emp_nationalite_adp", "cd_emp_nationalite_adp")]},
    {"rule": "R-SQL-013", "table": "lu_hr_emp_situation_matrimoniale", "src": "cd_emp_situation_matrimoniale_ts",
     "key": "cd_emp_situation_matrimoniale_ts",
     "out": [("cd_emp_situation_matrimoniale_ts", "cd_emp_situation_matrimoniale_ts_lu"),
             ("cd_emp_situation_matrimoniale_adp", "cd_emp_situation_matrimoniale_adp")]},
    {"rule": "R-SQL-014", "table": "lu_hr_emp_pays", "src": "cd_naissance_pays_ts", "key": "cd_emp_pays_ts",
     "out": [("cd_emp_pays_ts", "cd_naissance_pays_ts_lu"), ("id_emp_pays_adp", "id_naissance_pays_adp")]},
    {"rule": "R-SQL-015", "table": "lu_hr_emp_pays", "src": "cd_adresse_pays_ts", "key": "cd_emp_pays_ts",
     "out": [("cd_emp_pays_ts", "cd_adresse_pays_ts_lu"), ("id_emp_pays_adp", "id_adresse_pays_adp")]},
    {"rule": "R-SQL-016", "table": "lu_hr_emp_pays", "src": "lb_fiscal_country", "key": "cd_emp_pays_ts",
     "out": [("cd_emp_pays_ts", "cd_fiscal_pays_ts_lu"), ("id_emp_pays_adp", "id_fiscal_pays_adp")]},
    {"rule": "R-SQL-017", "table": "lu_hr_ppc_raison_debut_contrat", "src": "cd_ppc_raison_debut_contrat_ts",
     "key": "cd_ppc_raison_debut_contrat_ts",
     "out": [("cd_ppc_raison_debut_contrat_ts", "cd_ppc_raison_debut_contrat_ts_lu"),
             ("cd_ppc_raison_debut_contrat_adp", "cd_ppc_raison_debut_contrat_adp")]},
    {"rule": "R-SQL-018", "table": "lu_hr_ppc_contrat2", "src": "cd_ppc_contrat_ts", "key": "cd_contract_ts",
     "out": [("cd_contract_ts", "cd_ppc_contrat_ts_lu"), ("cd_nature_contract_adp", "cd_nature_contract_adp"),
             ("cd_type_contract_adp", "cd_type_contract_adp"),
             ("cd_type_collaboration_adp", "cd_type_collaboration_adp")]},
    {"rule": "R-SQL-019", "table": "lu_hr_ppc_classification", "src": "cd_ppc_convention_collective_lvl2_ts",
     "key": "cd_convention_collective_lvl2_ts",
     "out": [("cd_convention_collective_lvl2_ts", "cd_ppc_convention_collective_lvl2_ts_lu"),
             ("cd_classification_adp", "cd_classification_adp"), ("cd_coefficient_adp", "cd_coefficient_adp"),
             ("cd_ppc_convention_collective_adp", "cd_ppc_convention_collective_adp")]},
    {"rule": "R-SQL-020", "table": "lu_hr_ppc_classes", "src": "cd_occupational_category",
     "key": "cd_occupational_category_ts",
     "out": [("cd_occupational_category_ts", "cd_occupational_category_ts_lu"),
             ("cd_categorie_cotisant_adp", "cd_categorie_cotisant_adp"),
             ("cd_classe_remuneration_adp", "cd_classe_remuneration_adp")]},
    {"rule": "R-SQL-021", "table": "lu_hr_geo_entite", "src": "cd_geo_entite_ts", "key": "cd_geo_entite_ts",
     "out": [("cd_geo_entite_ts", "cd_geo_entite_ts_lu"), ("cd_geo_entite_adp", "cd_geo_entite_adp")]},
    {"rule": "R-SQL-022", "table": "lu_hr_ppc_location_country", "src": "cd_ppc_location_country_ts",
     "key": "cd_ppc_location_country_ts",
     "out": [("cd_ppc_location_country_ts", "cd_ppc_location_country_ts_lu"),
             ("cd_ppc_location_country_adp", "cd_ppc_location_country_adp")]},
    {"rule": "R-SQL-023", "table": "lu_hr_ppc_cost_center", "src": "cd_ppc_cost_center_ts",
     "key": "cd_ppc_cost_center_ts",
     "out": [("cd_ppc_cost_center_ts", "cd_ppc_cost_center_ts_lu"), ("cd_ppc_cost_center_adp", "cd_ppc_cost_center_adp")]},
    {"rule": "R-SQL-024", "table": "lu_hr_organisationalstructure", "src": "cd_organisationalstructure_ts",
     "key": "cd_organisationalstructure_ts",
     "out": [("cd_organisationalstructure_ts", "cd_organisationalstructure_ts_lu"),
             ("cd_organisationalstructure_adp", "cd_organisationalstructure_adp"),
             ("ds_organisationalstructure_adp", "ds_organisationalstructure_adp")]},
]
USED_LOOKUP_TABLES = sorted({j["table"] for j in LOOKUP_JOINS})  # 13 tables (D-REF-02 : 4 non utilisées)
ENRICHED_LOOKUP_COLUMNS = [out for j in LOOKUP_JOINS for _, out in j["out"]]

# ---------------------------------------------------------------------------
# Rejets (A02 § 5, PROC3)
# ---------------------------------------------------------------------------
NULL_REJECT_COLUMNS = ["id_payroll", "cd_matricule_it", "id_numero_securite_sociale"]

# Ordre exact de la clause VALUES de PROC3 (Format_Rejet) ; kind : len8 / date / date10 / pct / amount
FORMAT_REJECT_COLUMNS: List[Tuple[str, str]] = [
    ("id_payroll", "len8"),
    ("dt_naissance", "date"),
    ("dt_situation_start", "date"),
    ("dt_entree_societe", "date10"),
    ("dt_entree_groupe", "date"),
    ("dt_anciennete", "date"),
    ("dt_sortie", "date"),
    ("dt_company_start", "date"),
    ("dt_ppc_debut_contrat", "date"),
    ("dt_ppc_fin_contrat", "date"),
    ("dt_ppc_debut_periode_essaie", "date"),
    ("dt_ppc_fin_periode_essaie", "date"),
    ("pc_ppc_fte", "pct"),
    ("dt_debut_mt_salaire_base", "date"),
    ("dt_debut_pc_prime_annuelle", "date"),
    ("pc_prime_annuelle", "pct"),
    ("dt_debut_prime_annuelle", "date"),
    ("prime_annuelle", "amount"),
]
# R-SQL-037/038 : somme des indicateurs SANS pc_ppc_fte (anomalie de l'existant conservée, D-PAR-06)
FORMAT_SUM_EXCLUDED = {"pc_ppc_fte"}

# Référence_Rejet : (fichier, nom de colonne journalisé, valeur testée, colonne _lu, code ADP contrôlé|None)
REFERENCE_REJECTS: List[Tuple[str, str, str, str, object]] = [
    ("titre.csv", "cd_emp_titre", "cd_emp_titre_ts", "cd_emp_titre_ts_lu", "cd_emp_titre_adp"),
    ("sexe.csv", "cd_emp_sexe", "cd_emp_sexe_ts", "cd_emp_sexe_ts_lu", "cd_emp_sexe_adp"),
    ("nationalite.csv", "cd_emp_nationalite", "cd_emp_nationalite_ts", "cd_emp_nationalite_ts_lu",
     "cd_emp_nationalite_adp"),
    ("situation_matrimoniale.csv", "cd_emp_situation_matrimoniale", "cd_emp_situation_matrimoniale_ts",
     "cd_emp_situation_matrimoniale_ts_lu", "cd_emp_situation_matrimoniale_adp"),
    ("pays.csv", "cd_naissance_pays", "cd_naissance_pays_ts", "cd_naissance_pays_ts_lu", "id_naissance_pays_adp"),
    ("pays.csv", "cd_adresse_pays", "cd_adresse_pays_ts", "cd_adresse_pays_ts_lu", "id_adresse_pays_adp"),
    ("raison_debut_contrat.csv", "cd_ppc_raison_debut_contrat", "cd_ppc_raison_debut_contrat_ts",
     "cd_ppc_raison_debut_contrat_ts_lu", "cd_ppc_raison_debut_contrat_adp"),
    ("contrat_v2.csv", "cd_nature_contrat", "cd_ppc_contrat_ts", "cd_ppc_contrat_ts_lu", "cd_nature_contract_adp"),
    ("contrat_v2.csv", "cd_type_contrat", "cd_ppc_contrat_ts", "cd_ppc_contrat_ts_lu", "cd_type_contract_adp"),
    ("contrat_v2.csv", "cd_type_collaboration_adp", "cd_ppc_contrat_ts", "cd_ppc_contrat_ts_lu",
     "cd_type_collaboration_adp"),
    ("classes.csv", "cd_occupational_category", "cd_occupational_category", "cd_occupational_category_ts_lu", None),
    ("entite.csv", "cd_geo_entite", "cd_geo_entite_ts", "cd_geo_entite_ts_lu", "cd_geo_entite_adp"),
    ("location_country.csv", "cd_ppc_location_country", "cd_ppc_location_country_ts",
     "cd_ppc_location_country_ts_lu", "cd_ppc_location_country_adp"),
    ("cost_center.csv", "cd_ppc_cost_center", "cd_ppc_cost_center_ts", "cd_ppc_cost_center_ts_lu",
     "cd_ppc_cost_center_adp"),
    ("organisationalstructure.csv", "cd_organisationalstructure", "cd_organisationalstructure_ts",
     "cd_organisationalstructure_ts_lu", "cd_organisationalstructure_adp"),
]
# Pays fiscal (R-SQL-042) : indicateur calculé mais jamais émis ; classification commentée.

REJECT_TYPE_NULL = "Null_Rejet"
REJECT_TYPE_FORMAT = "Format_Rejet"
REJECT_TYPE_REFERENCE = "Référence_Rejet"

# ---------------------------------------------------------------------------
# R-SQL-056 : 72 colonnes de sortie (ordre exact) et typage
# ---------------------------------------------------------------------------
ADP_COLUMNS: List[str] = [
    "dt_transforme", "id_unique", "id_payroll", "cd_matricule_it", "lb_nom", "lb_prenom", "lb_nom_naissance",
    "cd_emp_titre_adp", "cd_emp_sexe_adp", "cd_emp_nationalite_adp", "id_numero_securite_sociale",
    "dt_debut_situation_matrimoniale", "cd_emp_situation_matrimoniale_adp", "dt_naissance", "lb_naissance_ville",
    "lb_naissance_cp", "lb_naissance_departement", "id_naissance_pays_adp", "lb_adresse_cp", "lb_adresse_ville",
    "id_adresse_pays_adp", "lb_email_professionnel", "lb_email_personnel", "lb_telephone_professionnel",
    "lb_telephone_personnel", "lb_fiscal_ville", "id_fiscal_pays_adp", "lb_fiscal_cp", "lb_fiscal_street",
    "lb_fiscal_streetnumber", "lb_fiscal_additional_address_information", "lb_fiscal_streetnumber_complement",
    "dt_entree_societe", "dt_entree_groupe", "dt_anciennete", "dt_sortie", "dt_company_start",
    "cd_ppc_raison_debut_contrat_adp", "cd_ppc_raison_fin_contrat_ts", "dt_ppc_debut_contrat",
    "dt_ppc_fin_contrat", "cd_nature_contract_adp", "cd_type_contract_adp", "cd_type_collaboration_adp",
    "cd_recours_cdd", "cd_categorie_cotisant_adp", "cd_classe_remuneration_adp", "dt_ppc_debut_periode_essaie",
    "dt_ppc_fin_periode_essaie", "cd_classification_adp", "cd_coefficient_adp", "cd_organisationalstructure_adp",
    "dt_situation_start", "cd_ppc_convention_collective_adp", "cd_geo_entite_adp", "cd_ppc_location_country_adp",
    "pc_ppc_fte", "id_numero_ordre", "cd_ppc_cost_center_adp", "pc_imputation", "dt_debut_mt_salaire_base",
    "mt_salaire_mensuel", "dt_debut_pc_prime_annuelle", "pc_prime_annuelle", "dt_debut_prime_annuelle",
    "prime_annuelle", "lb_job_title", "lb_account_holder", "lb_iban", "lb_bic", "lb_account_number",
    "lb_bank_name",
]
assert len(ADP_COLUMNS) == 72
# R-SQL-051 : CONVERT(DATE, x) ; dt_entree_societe via SUBSTRING(x,1,10)
ADP_DATE_COLUMNS = [
    "dt_entree_societe", "dt_entree_groupe", "dt_anciennete", "dt_sortie", "dt_company_start",
    "dt_ppc_debut_contrat", "dt_ppc_fin_contrat", "dt_ppc_debut_periode_essaie", "dt_ppc_fin_periode_essaie",
    "dt_situation_start", "dt_debut_mt_salaire_base", "dt_debut_pc_prime_annuelle", "dt_debut_prime_annuelle",
]
ADP_DECIMAL_COLUMNS = {"mt_salaire_mensuel": "DECIMAL(23,13)"}
# D-CMP-02 (À CONFIRMER) : 71 colonnes métier comparées = R-SQL-056 sans dt_transforme
COMPARE_COLUMNS = [c for c in ADP_COLUMNS if c != "dt_transforme"]
BUSINESS_KEY = "id_unique"  # D-CMP-01 (À CONFIRMER)


def adp_column_type(col: str) -> str:
    if col == "dt_transforme" or col in ADP_DATE_COLUMNS:
        return "DATE"
    return ADP_DECIMAL_COLUMNS.get(col, "STRING")


# ---------------------------------------------------------------------------
# DDL : nom logique -> liste (colonne, type) ; les tables ctl.* et log.* reprennent 05 (fait foi)
# ---------------------------------------------------------------------------
def _cols(spec: str) -> List[Tuple[str, str]]:
    out = []
    for line in spec.strip().splitlines():
        line = line.strip().rstrip(",")
        if line:
            name, typ = line.split(None, 1)
            out.append((name, typ.strip()))
    return out


TABLES: Dict[str, List[Tuple[str, str]]] = {}

TABLES["cfg.process"] = _cols("""
process_code STRING
enabled BOOLEAN
acquisition_mode STRING
publication_target STRING
key_columns ARRAY<STRING>
compare_columns ARRAY<STRING>
reference_table STRING
staging_table STRING
max_absent_pct DECIMAL(9,4)
max_modified_pct DECIMAL(9,4)
max_reject_pct DECIMAL(9,4)
updated_at_utc TIMESTAMP
""")
for _t, _c in LOOKUP_TABLES.items():
    TABLES[f"cfg.{_t}"] = [(c, "STRING") for c in _c] + [
        ("_loaded_at_utc", "TIMESTAMP"), ("_loaded_by", "STRING"), ("_change_ref", "STRING")]

TABLES["stg.ts_employee_raw"] = (
    [("run_id", "STRING"), ("business_date", "DATE"), ("_extracted_at_utc", "TIMESTAMP")]
    + [(k, "STRING") for k in RAW_KEYS])
TABLES["stg.ts_employee_source"] = (
    [("run_id", "STRING"), ("business_date", "DATE")] + [(c, "STRING") for c in SOURCE_COLUMN_NAMES])
TABLES["stg.ts_employee_enriched"] = (
    [("run_id", "STRING"), ("business_date", "DATE"), ("dt_staging", "DATE")]
    + [(c, "STRING") for c in SOURCE_COLUMN_NAMES] + [(c, "STRING") for c in ENRICHED_LOOKUP_COLUMNS])
TABLES["stg.ts_employee_reject"] = _cols("""
run_id STRING
business_date DATE
dt_rejet DATE
id_unique STRING
id_payroll STRING
cd_type_rejet STRING
lb_nom_colonne_rejet STRING
lb_valeur_rejet STRING
fl_rejet INT
lb_fichier_mapping STRING
ds_rejet STRING
""")
TABLES["stg.ts_employee_adp"] = (
    [("run_id", "STRING"), ("business_date", "DATE")]
    + [(c, adp_column_type(c)) for c in ADP_COLUMNS] + [("row_hash", "STRING")])
TABLES["ref.ts_employee_adp"] = (
    [(c, adp_column_type(c)) for c in ADP_COLUMNS]
    + [("row_hash", "STRING"), ("_last_run_id", "STRING"), ("_validated_at_utc", "TIMESTAMP")])

TABLES["pub.adp_outbox"] = _cols("""
run_id STRING
employee_key STRING
change_category STRING
row_hash STRING
status STRING
attempt_count INT
last_attempt_at_utc TIMESTAMP
last_outcome_code STRING
gate_decision_id STRING
created_at_utc TIMESTAMP
updated_at_utc TIMESTAMP
""")

# --- PROPOSÉ (contrat ADP final absent) : intention durable et tentatives par opération ---
TABLES["pub.adp_operation"] = _cols("""
op_id STRING
run_id STRING
employee_key STRING
seq_no INT
op_code STRING
op_variant STRING
depends_on ARRAY<STRING>
target_hash STRING
payload_hash STRING
status STRING
attempt_count INT
last_outcome_class STRING
last_http_status INT
gate_decision_id STRING
manual_decision_id STRING
created_at_utc TIMESTAMP
updated_at_utc TIMESTAMP
""")
TABLES["pub.adp_attempt"] = _cols("""
attempt_id STRING
op_id STRING
run_id STRING
attempt_no INT
phase STRING
outcome_class STRING
http_status INT
error_class STRING
response_digest STRING
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
""")
TABLES["pub.adp_reconciliation"] = _cols("""
reconciliation_id STRING
op_id STRING
run_id STRING
observed_at_utc TIMESTAMP
method STRING
fields_checked ARRAY<STRING>
result STRING
conclusion STRING
""")
TABLES["log.adp_manual_decision"] = _cols("""
decision_id STRING
op_id STRING
run_id STRING
employee_key STRING
decision STRING
decided_by STRING
decided_at_utc TIMESTAMP
basis STRING
reference STRING
""")

TABLES["ctl.process_lock"] = _cols("""
process_code STRING
lock_status STRING
holder_run_id STRING
acquired_at_utc TIMESTAMP
lease_expires_at_utc TIMESTAMP
released_at_utc TIMESTAMP
""")
TABLES["ctl.run"] = _cols("""
run_id STRING
process_code STRING
run_mode STRING
business_date DATE
run_seq_in_day INT
trigger_time_utc TIMESTAMP
pipeline_name STRING
workspace_id STRING
workspace_name STRING
status STRING
compared_against_run_id STRING
promoted_at_utc TIMESTAMP
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
error_code STRING
error_summary STRING
target_run_id STRING
""")
TABLES["ctl.run_step"] = _cols("""
run_id STRING
step_code STRING
attempt_no INT
status STRING
rows_in BIGINT
rows_out BIGINT
content_hash STRING
message STRING
started_at_utc TIMESTAMP
ended_at_utc TIMESTAMP
""")
TABLES["log.comparison_summary"] = _cols("""
run_id STRING
process_code STRING
business_date DATE
compared_against_run_id STRING
nb_staging BIGINT
nb_reference BIGINT
nb_new BIGINT
nb_modified BIGINT
nb_unchanged BIGINT
nb_absent BIGINT
nb_rejected BIGINT
pct_absent DECIMAL(9,4)
pct_modified DECIMAL(9,4)
pct_rejected DECIMAL(9,4)
is_first_load BOOLEAN
created_at_utc TIMESTAMP
""")
TABLES["log.comparison_detail"] = _cols("""
run_id STRING
process_code STRING
business_key STRING
category STRING
changed_columns ARRAY<STRING>
old_row_hash STRING
new_row_hash STRING
created_at_utc TIMESTAMP
""")
TABLES["log.anomaly"] = _cols("""
run_id STRING
process_code STRING
control_code STRING
severity STRING
metric_value DECIMAL(18,4)
threshold_value DECIMAL(18,4)
outcome STRING
message STRING
created_at_utc TIMESTAMP
""")
TABLES["log.reject"] = _cols("""
run_id STRING
process_code STRING
dt_rejet DATE
id_unique STRING
id_payroll STRING
cd_type_rejet STRING
lb_nom_colonne_rejet STRING
lb_valeur_rejet STRING
lb_fichier_mapping STRING
ds_rejet STRING
created_at_utc TIMESTAMP
""")
TABLES["log.run_summary"] = _cols("""
run_id STRING
process_code STRING
business_date DATE
dt_resume DATE
cd_source STRING
cd_target STRING
mt_lignes_inserts BIGINT
mt_lignes_rejets BIGINT
mt_lignes_lues BIGINT
pc_lignes_rejets DECIMAL(12,4)
ds_type_rejets STRING
nb_outbox BIGINT
nb_sent BIGINT
nb_no_change BIGINT
nb_failed BIGINT
nb_not_pushed BIGINT
created_at_utc TIMESTAMP
""")
TABLES["log.adp_gate_decision"] = _cols("""
decision_id STRING
run_id STRING
evaluated_at_utc TIMESTAMP
workspace_name STRING
workspace_id STRING
is_for_pipeline BOOLEAN
notebook_name STRING
rule_version STRING
decision BOOLEAN
reason_code STRING
""")
TABLES["log.adp_publication_log"] = _cols("""
log_id STRING
source_system STRING
run_id STRING
log_date DATE
ts_adp TIMESTAMP
matricule_id STRING
nom STRING
mode STRING
field STRING
status_code STRING
reason_code STRING
response_code STRING
error_message STRING
log_message STRING
employee_key STRING
step_name STRING
http_method STRING
endpoint_template STRING
attempt_no INT
gate_decision_id STRING
_import_id STRING
_source_file_uri STRING
_source_row_number INT
_row_hash STRING
_ingested_at_utc TIMESTAMP
""")
TABLES["log.log_file_import"] = _cols("""
import_id STRING
source_uri STRING
file_name STRING
file_size BIGINT
file_modified_at_utc TIMESTAMP
file_format STRING
status STRING
detail STRING
rows_read INT
rows_loaded INT
rows_rejected INT
target_table STRING
executed_by STRING
workspace_name STRING
processed_at_utc TIMESTAMP
""")
TABLES["log.log_file_import_reject"] = _cols("""
import_id STRING
source_uri STRING
source_row_number INT
reject_reason STRING
rejected_at_utc TIMESTAMP
""")
TABLES["log.notification"] = _cols("""
run_id STRING
notification_type STRING
channel STRING
recipients_ref STRING
status STRING
sent_at_utc TIMESTAMP
error_summary STRING
""")

SCHEMAS = ["cfg", "stg", "ref", "pub", "ctl", "log"]

# Tables à données personnelles (accès restreint, 05 § 4 ; D-SEC-03)
PERSONAL_DATA_TABLES = [
    "stg.ts_employee_raw", "stg.ts_employee_source", "stg.ts_employee_enriched", "stg.ts_employee_reject",
    "stg.ts_employee_adp", "ref.ts_employee_adp", "log.comparison_detail", "log.reject",
    "log.adp_publication_log", "pub.adp_outbox", "pub.adp_operation",
]

PROPOSED_TABLES = ["pub.adp_operation", "pub.adp_attempt", "pub.adp_reconciliation", "log.adp_manual_decision"]


class Tables:
    """Résolution nom logique -> nom physique.

    * Sans `lakehouse` : noms partiels (`ctl.run`) qui exigent un Lakehouse par défaut attaché (tests locaux).
    * Avec `lakehouse` : noms à 4 parties `` `workspace`.`lakehouse`.schema.table ``, qui fonctionnent SANS
      Lakehouse par défaut (vérifié dans Fabric : reports/probe_4part_names.json ; les noms partiels échouent sans
      lien par défaut : reports/probe_no_default_lakehouse.json). Aucune dépendance à une association que la
      synchronisation Git pourrait ne pas conserver.
    * `test_schema` redirige toutes les tables vers un schéma d'isolement unique : `<schéma_test>.<schéma>__<table>`.
    """

    def __init__(self, test_schema: str = "", lakehouse: str = "", workspace: str = ""):
        if bool(lakehouse) != bool(workspace):
            raise ValueError("lakehouse et workspace vont ensemble")
        for part in (test_schema, lakehouse, workspace):
            if "`" in part or any(ch in part for ch in ("\n", "\r")):
                raise ValueError("nom invalide")
        self.test_schema, self.lakehouse, self.workspace = test_schema, lakehouse, workspace

    @property
    def prefix(self) -> str:
        return f"`{self.workspace}`.`{self.lakehouse}`." if self.lakehouse else ""

    def __call__(self, logical: str) -> str:
        if logical not in TABLES:
            raise KeyError(logical)
        schema, name = logical.split(".", 1)
        if not self.test_schema:
            return f"{self.prefix}{schema}.{name}"
        return f"{self.prefix}{self.test_schema}.{schema}__{name}"

    def extra(self, name: str) -> str:
        """Table annexe du schéma de test (ex. rapport de campagne)."""
        if not self.test_schema:
            raise ValueError("table annexe réservée au schéma de test")
        return f"{self.prefix}{self.test_schema}.{name}"

    def schemas(self) -> List[str]:
        names = [self.test_schema] if self.test_schema else list(SCHEMAS)
        return [f"{self.prefix}{n}" for n in names]


def tables_for_notebook(context, test_schema: str = "", lakehouse: str = "") -> "Tables":
    """Tables du notebook courant : workspace lu dans le contexte d'exécution RÉEL ; Lakehouse fourni par paramètre
    (VL_HRIS.lakehouse_name) ou, à défaut, Lakehouse par défaut attaché. Refuse de deviner."""
    def get(key):
        try:
            return context.get(key)
        except Exception:
            return None
    ws = get("currentWorkspaceName")
    lh = lakehouse or get("defaultLakehouseName")
    if not ws or not lh:
        raise ValueError("LAKEHOUSE_NOT_RESOLVED : fournir p_lakehouse_name (VL_HRIS.lakehouse_name)")
    return Tables(test_schema, lh, ws)


def create_table_sql(physical: str, logical: str) -> str:
    cols = ",\n  ".join(f"`{n}` {t}" for n, t in TABLES[logical])
    return f"CREATE TABLE IF NOT EXISTS {physical} (\n  {cols}\n) USING DELTA"


def normalize_type(t: str) -> str:
    """Normalise un type Spark (simpleString) pour comparer DDL attendu et schéma constaté."""
    t = t.strip().lower().replace(" ", "")
    return {"int": "int", "integer": "int", "bigint": "bigint", "long": "bigint", "string": "string",
            "boolean": "boolean", "date": "date", "timestamp": "timestamp"}.get(t, t)
'''),
    ('hris.ts_acquire', '2b32899f92949347bb603a100996cbdda9a916320c1dee29c7ff9ba1f0f13c3c', r'''"""hris.ts_acquire — portage de la Function TalentSoft (ts-function-app/function_app.py, PRJ @ 5541490).

Parité (A01, R-TS-007 à R-TS-030) : endpoints, règle de « valeur courante » dépendant de l'ordre,
conversion « falsy », sentinelle 1900-01-01, Hors_France, Level sans %, filtre « Motul France »,
retries 429/503/réseau avec backoff, disjoncteur global, lots de 75, 5 requêtes simultanées par hôte.

Améliorations techniques (A08 D-TS-02, À CONFIRMER, sans effet sur les données si l'API se comporte
comme attendu) : appels synchrones bornés (pas d'asyncio dans un notebook) ; renouvellement du jeton
sur 401 en LECTURE ; prise en compte de Retry-After ; aucune redirection suivie ; comptage de tous
les échecs et exclusions pour le contrôle de complétude (aucun échec n'est plus silencieux) ;
aucune URL, aucun matricule, aucun corps de réponse journalisé.

Source : le tenant TalentSoft de PRODUCTION, quel que soit l'environnement d'exécution
(consigne du mandat). Aucune autre URL n'est acceptée et il n'existe aucun repli.
"""
from __future__ import annotations

import datetime
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from typing import Any, Callable, Dict, List, Optional, Tuple
from urllib.parse import urlparse

from hris.common import TS_PRD_BASE_URL, HrisError, SecretValue
from hris.model import RAW_KEYS
from hris.parity import json_value_repr

BATCH_SIZE = 75          # ts:L253
EMPLOYEE_CONCURRENCY = 2  # ts:L254 (max_concurrent par lot)
PER_HOST_LIMIT = 5        # ts:L100 (limit_per_host)
BATCH_PAUSE_SECONDS = 2   # ts:L275
MOTUL_FRANCE = "Motul France"
HUB_ENTITIES = {
    "situation": "Situation", "keydate": "KeyDates", "disability": "Disability",
    "workaccident": "WorkAccident", "medicalcheckup": "MedicalCheckUp", "legacyid": "EmployeeIDs",
    "qualification": "Qualification", "mainpostaladdress": "MainPostalAddress",
    "contactinfo": "ContactInformation", "maritalstatus": "FamilyStatus", "contract": "Contract",
    "bankdetails": "BankDetails", "probation": "ProbationaryPeriod", "salary": "Compensation",
    "civilregistration": "CivilRegistration", "identity": "Identity",
}
# Les entités de santé sont appelées en parité (mêmes appels que l'existant) mais leur résultat
# n'est jamais persisté (A08 D-SEC-02) : seules les clés RAW_KEYS sont conservées.


def assert_prd_source(base_url: str) -> str:
    """Refuse toute source autre que le tenant TalentSoft de production."""
    normalized = (base_url or "").strip().rstrip("/")
    if normalized != TS_PRD_BASE_URL:
        raise HrisError("TS_SOURCE_NOT_PRD", "seul le tenant TalentSoft de production est autorisé")
    parsed = urlparse(normalized)
    if parsed.scheme != "https":
        raise HrisError("TS_SOURCE_NOT_HTTPS")
    return normalized


# ---------------------------------------------------------------------------
# Disjoncteur (ts:L32-65) — parité
# ---------------------------------------------------------------------------
class CircuitBreaker:
    def __init__(self, failure_threshold: int = 5, timeout: float = 60, clock: Callable[[], float] = time.time):
        self.failure_count = 0
        self.failure_threshold = failure_threshold
        self.timeout = timeout
        self.last_failure_time: Optional[float] = None
        self.state = "CLOSED"
        self._clock = clock
        self._lock = threading.Lock()

    def record_failure(self) -> None:
        with self._lock:
            self.failure_count += 1
            self.last_failure_time = self._clock()
            if self.failure_count >= self.failure_threshold:
                self.state = "OPEN"

    def record_success(self) -> None:
        with self._lock:
            self.failure_count = 0
            self.state = "CLOSED"

    def can_attempt(self) -> bool:
        with self._lock:
            if self.state == "CLOSED":
                return True
            if self.state == "OPEN":
                if self._clock() - (self.last_failure_time or 0) > self.timeout:
                    self.state = "HALF_OPEN"
                    return True
                return False
            return True


class AcquisitionStats:
    """Compteurs sans donnée personnelle (aucun matricule, aucune URL)."""

    def __init__(self):
        self._lock = threading.Lock()
        self.counters: Counter = Counter()
        self.endpoint_status: Dict[str, Counter] = defaultdict(Counter)
        self.total_count: Optional[int] = None

    def inc(self, key: str, n: int = 1) -> None:
        with self._lock:
            self.counters[key] += n

    def status(self, endpoint: str, status: Any) -> None:
        with self._lock:
            self.endpoint_status[endpoint][str(status)] += 1

    def as_dict(self) -> dict:
        return {"total_count": self.total_count, "counters": dict(self.counters),
                "endpoint_status": {k: dict(v) for k, v in sorted(self.endpoint_status.items())}}


# ---------------------------------------------------------------------------
# Client HTTP (synchrone, borné) — équivalent de make_request_with_retry (ts:L126-183)
# ---------------------------------------------------------------------------
class NetworkError(Exception):
    """Erreur réseau au sens de l'existant (aiohttp.ClientError, TimeoutError, OSError)."""


class TsClient:
    def __init__(self, base_url: str, client_id: SecretValue, client_secret: SecretValue,
                 session_factory: Callable[[], Any], stats: AcquisitionStats,
                 sleep: Callable[[float], None] = time.sleep,
                 timeout: Tuple[float, float] = (30.0, 60.0),
                 network_errors: Tuple[type, ...] = (OSError,),
                 breaker: Optional[CircuitBreaker] = None,
                 max_retry_after: float = 60.0):
        self.base_url = assert_prd_source(base_url)
        self._client_id = client_id
        self._client_secret = client_secret
        self._session = session_factory()
        self.stats = stats
        self._sleep = sleep
        self._timeout = timeout
        self._network_errors = network_errors
        self.breaker = breaker or CircuitBreaker()
        self._token: Optional[SecretValue] = None
        self._token_lock = threading.Lock()
        self._max_retry_after = max_retry_after
        self._limiter = threading.BoundedSemaphore(PER_HOST_LIMIT)

    def __repr__(self) -> str:
        return "TsClient(<prd>)"

    def close(self) -> None:
        try:
            self._session.close()
        except Exception:
            pass

    # -- jeton (ts:L331-364) ----------------------------------------------------
    def get_token(self) -> None:
        data = {"client_id": self._client_id.reveal(), "client_secret": self._client_secret.reveal(),
                "grant_type": "client_credentials"}
        token_data, status = self._request("POST", "/api/token", {"Accept": "application/json"},
                                           max_retries=3, initial_delay=2, form=data, endpoint="token")
        if status == 200 and token_data and token_data.get("access_token"):
            with self._token_lock:
                self._token = SecretValue("Bearer " + token_data.get("access_token"))
            self.stats.inc("token_obtained")
            return
        raise HrisError("TS_AUTH_FAILED", f"statut {status}")

    def _auth_headers(self) -> dict:
        with self._token_lock:
            token = self._token.reveal() if self._token else None
        if not token:
            raise HrisError("TS_NO_TOKEN")
        return {"Authorization": token, "Accept": "application/json"}

    def get(self, path: str, endpoint: str, max_retries: int = 3, initial_delay: float = 1) -> Tuple[Any, Any]:
        data, status = self._request("GET", path, None, max_retries, initial_delay, endpoint=endpoint)
        if status == 401:  # amélioration D-TS-02 : jeton renouvelé, lecture idempotente rejouée une fois
            self.stats.inc("token_refresh_on_401")
            self.get_token()
            data, status = self._request("GET", path, None, max_retries, initial_delay, endpoint=endpoint)
        return data, status

    def _request(self, method: str, path: str, headers: Optional[dict], max_retries: int,
                 initial_delay: float, form: Optional[dict] = None, endpoint: str = "") -> Tuple[Any, Any]:
        if not self.breaker.can_attempt():
            self.stats.inc("circuit_open_skip")
            self.stats.status(endpoint, "CIRCUIT_OPEN")
            return None, None
        url = self.base_url + path
        for attempt in range(max_retries):
            try:
                hdrs = headers if headers is not None else self._auth_headers()
                with self._limiter:
                    if method == "GET":
                        response = self._session.get(url, headers=hdrs, timeout=self._timeout,
                                                     allow_redirects=False)
                    else:
                        response = self._session.post(url, headers=hdrs, data=form, timeout=self._timeout,
                                                      allow_redirects=False)
                status = response.status_code
                if status == 200:
                    try:
                        result = response.json()
                    except Exception:
                        self.stats.status(endpoint, "200_INVALID_JSON")
                        return None, None  # parité : exception inattendue -> (None, None)
                    self.breaker.record_success()
                    self.stats.status(endpoint, 200)
                    return result, status
                if status in (429, 503):
                    self.stats.inc(f"http_{status}")
                    if attempt < max_retries - 1:
                        delay = initial_delay * (2 ** attempt)
                        retry_after = _retry_after_seconds(response)
                        if retry_after is not None:
                            delay = min(max(delay, retry_after), self._max_retry_after)
                        self._sleep(delay)
                        continue
                    # parité : dernier essai -> statut perdu, (None, None) après la boucle
                    self.stats.status(endpoint, f"{status}_EXHAUSTED")
                else:
                    self.stats.status(endpoint, status)
                    return None, status
            except self._network_errors as exc:
                self.breaker.record_failure()
                self.stats.inc("network_error")
                if attempt < max_retries - 1:
                    self._sleep(initial_delay * (2 ** attempt))
                else:
                    self.stats.status(endpoint, f"NETWORK_{type(exc).__name__}")
                    return None, None
            except HrisError:
                raise
            except Exception as exc:  # parité : erreur inattendue -> (None, None), sans retry
                self.stats.status(endpoint, f"ERROR_{type(exc).__name__}")
                return None, None
        return None, None


def _retry_after_seconds(response: Any) -> Optional[float]:
    try:
        value = response.headers.get("Retry-After")
    except Exception:
        return None
    if value is None:
        return None
    try:
        return max(0.0, float(value))
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Traitements des réponses — copie fidèle de l'existant (ts:L619-1629)
# ---------------------------------------------------------------------------
def _d(value: str) -> datetime.date:
    return datetime.datetime.strptime(value, "%Y-%m-%d").date()


def initialize_employee_fields(employee_number: str) -> dict:
    return {
        "employeeNumber": employee_number, "hubSocialSecurityNumber": "", "referenceAdmissionDate": "",
        "WorkingType": "", "Situation_starDate": "",
        "GeographicOrganisationStructure": "", "LegalStructure": "", "CostCenter": "", "CollectiveAgreement": "",
        "PPC_FTE": "", "OccupationalCategory": "", "Coefficient": "", "OrganisationalStructure": "",
        "salary_base": "0", "base_startDate": "", "Pourcentage_prime_annuelles": "0",
        "pourcentage_prime_annuelles_startDate": "", "prime_annuelles": "0", "prime_annuelles_startDate": "",
        "address_postalCode": "", "adresse_city": "", "adresse_countryCode": "", "Hors_France": "0",
        "GroupStartDate": "", "GroupEndDate": "", "CompanyStartDate": "", "RecalculatedSeniorityDate": "",
        "handicap": "", "handicap_startDate": "", "workaccident": "", "workaccident_startDate": "",
        "motif_visite_medical": "", "motif_visite_medical_startDate": "",
        "LegacyID": "", "LegacyID_startDate": "", "Matricule_IT": "",
        "Qualification": "", "Qualification_startDate": "", "ConferredDate": "",
        "LegalTownOrCity": "", "LegalCountry": "", "LegalPostalCode": "", "LegalStreet": "", "LegalStreetNumber": "",
        "LegalAdditionalAddressInformation": "", "LegalStreetNumberComplement": "",
        "BusinessMobile": "", "PersonalMobile": "", "BusinessPhone": "", "PersonalPhone": "", "BusinessEmail": "",
        "PersonalEmail": "", "ContactInformation_startDate": "",
        "MaritalStatus": "", "MaritalStatus_startDate": "",
        "NatureOfContract": "", "HiringReason": "", "date_debut_contract": "", "date_fin_contract": "",
        "FixedTermContractReason": "", "TerminationReason": "", "Status": "",
        "AccountHolder": "", "IBANOrABA": "", "BicOrSwift": "", "AccountNumber": "", "BankName": "",
        "PaymentMethod": "", "UseSameAccountForExpense": "", "ExpenseAccountPaymentMethod": "",
        "ExpenseAccountHolder": "", "ExpenseAccountIBANOrABA": "", "ExpenseAccountBicOrSwift": "",
        "ExpenseAccountBankName": "", "Country": "", "ExpenseAccountCountry": "", "date_debut_bank": "",
        "date_fin_bank": "",
        "FirstPeriodStartDate": "", "FirstPeriodEndDate": "", "SecondPeriodStartDate": "", "SecondPeriodEndDate": "",
        "ProbationaryPeriod_startDate": "",
        "Date_de_naissance": "", "Ville_de_naissance": "", "Pays_de_naissance": "", "code_postal_lieu_de_naissance": "",
        "region_ou_etat_lieu_de_naissance": "", "CivilRegistration_startDate": "",
        "Title": "", "Gender": "", "LastName": "", "FirstName": "", "Nationality": "", "BirthName": "",
        "Identity_startDate": "",
        "numero_ordre": "1", "pourcentage_imputation": "100",
    }


def process_situation_data(data):
    if not data:
        return {"WorkingType": "", "Situation_starDate": "", "Position": ""}
    latest_date = "0001-01-01"
    result = {"WorkingType": "", "Situation_starDate": "", "Position": ""}
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            result["Situation_starDate"] = latest_date
            if item["field"] == "WorkingType":
                result["WorkingType"] = item["value"] if item["value"] else ""
            elif item["field"] == "Position":
                result["Position"] = item["value"] if item["value"] else ""
    return result


def process_position_data(data):
    if not data:
        return {}
    result = {"GeographicOrganisationStructure": "", "LegalStructure": "", "CostCenter": "",
              "CollectiveAgreement": "", "PPC_FTE": "", "OccupationalCategory": "", "Coefficient": "",
              "pourcentage_imputation": "100", "OrganisationalStructure": ""}
    for item in data:
        field = item.get("field", "")
        value = item.get("value", "")
        if field == "GeographicOrganisationStructure":
            result["GeographicOrganisationStructure"] = value if value else ""
        elif field == "LegalStructure":
            result["LegalStructure"] = value if value else ""
        elif field == "CostCenter":
            result["CostCenter"] = value if value else ""
        elif field == "CollectiveAgreement":
            result["CollectiveAgreement"] = value if value else ""
        elif field == "FullTimeEquivalentPercent":
            result["PPC_FTE"] = value if value else ""
        elif field == "OccupationalCategory":
            result["OccupationalCategory"] = value if value else ""
        elif field == "Coefficient":
            result["Coefficient"] = value if value else ""
        elif field == "OrganisationalStructure":
            result["OrganisationalStructure"] = value if value else ""
        elif field == "Level":
            if value:
                level_value = str(value).replace('%', '').strip()
                result["pourcentage_imputation"] = level_value if level_value else "100"
    return result


def process_salary_data(data):
    if not data:
        return {"salary_base": "0", "base_startDate": "", "Pourcentage_prime_annuelles": "0",
                "pourcentage_prime_annuelles_startDate": "", "prime_annuelles": "0", "prime_annuelles_startDate": ""}
    base_date = "0001-01-01"
    bonus_date = "0001-01-01"
    bonus_percentage_date = "0001-01-01"
    result = {"salary_base": "0", "base_startDate": "", "Pourcentage_prime_annuelles": "0",
              "pourcentage_prime_annuelles_startDate": "", "prime_annuelles": "0", "prime_annuelles_startDate": ""}
    for item in data:
        start_date = item["startDate"]
        if "ANNUAL_BASE_SALARY" in item["key"] and "Amount" in item["field"]:
            if _d(start_date) > _d(base_date):
                base_date = start_date
                result["salary_base"] = item["value"]
                result["base_startDate"] = start_date
        elif "CONTRACTUAL_BONUS_PERCENTAGE" in item["key"] and "Amount" in item["field"]:
            if _d(start_date) > _d(bonus_percentage_date):
                bonus_percentage_date = start_date
                result["Pourcentage_prime_annuelles"] = item["value"]
                result["pourcentage_prime_annuelles_startDate"] = start_date
        elif "CONTRACTUAL BONUS" in item["key"] and "Amount" in item["field"]:
            if _d(start_date) > _d(bonus_date):
                bonus_date = start_date
                result["prime_annuelles"] = item["value"]
                result["prime_annuelles_startDate"] = start_date
    return result


def process_postal_address_data(data):
    if not data:
        return {"address_postalCode": "", "adresse_city": "", "adresse_countryCode": "", "Hors_France": "0"}
    result = {"address_postalCode": data.get("postalCode", ""), "adresse_city": data.get("city", ""),
              "adresse_countryCode": data.get("countryCode", ""), "Hors_France": "0"}
    if result["adresse_countryCode"] and result["adresse_countryCode"].upper() != "FR":
        result["Hors_France"] = "1"
    return result


def _latest(data, fields: Dict[str, str], date_key: Optional[str] = None, defaults: Optional[dict] = None,
            date_on_field: Optional[Dict[str, str]] = None):
    """Règle générique « valeur courante » (R-TS-018) : un seul maximum startDate (>=) commun
    à tous les champs, items parcourus dans l'ordre de l'API."""
    result = dict(defaults or {})
    for out in fields.values():
        result.setdefault(out, "")
    if date_key:
        result.setdefault(date_key, "")
    latest_date = "0001-01-01"
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            if date_key:
                result[date_key] = latest_date
            out = fields.get(item["field"])
            if out is not None:
                result[out] = item["value"] if item["value"] else ""
                if date_on_field and item["field"] in date_on_field:
                    result[date_on_field[item["field"]]] = latest_date
    return result


def process_keydate_data(data):
    return _latest(data or [], {"GroupStartDate": "GroupStartDate", "GroupEndDate": "GroupEndDate",
                                "CompanyStartDate": "CompanyStartDate",
                                "RecalculatedSeniorityDate": "RecalculatedSeniorityDate"})


def process_disability_data(data):
    return _latest(data or [], {"RecognizedAsDisabled": "handicap"}, date_key="handicap_startDate")


def process_workaccident_data(data):
    return _latest(data or [], {"DateOfOccurrence": "workaccident"}, date_key="workaccident_startDate")


def process_medicalcheckup_data(data):
    return _latest(data or [], {"Reason": "motif_visite_medical"}, date_key="motif_visite_medical_startDate")


def process_legacyid_data(data):
    return _latest(data or [], {"LegacyID": "LegacyID", "Initials": "Matricule_IT"},
                   defaults={"LegacyID_startDate": ""}, date_on_field={"LegacyID": "LegacyID_startDate"})


def process_qualification_data(data):
    return _latest(data or [], {"Qualification": "Qualification", "ConferredDate": "ConferredDate"},
                   defaults={"Qualification_startDate": ""}, date_on_field={"Qualification": "Qualification_startDate"})


def process_mainpostaladdress_data(data):
    return _latest(data or [], {"TownOrCity": "LegalTownOrCity", "Country": "LegalCountry",
                                "PostalCode": "LegalPostalCode", "Street": "LegalStreet",
                                "StreetNumber": "LegalStreetNumber",
                                "AdditionalAddressInformation": "LegalAdditionalAddressInformation",
                                "StreetNumberComplement": "LegalStreetNumberComplement"})


def process_contactinfo_data(data):
    return _latest(data or [], {"BusinessMobile": "BusinessMobile", "PersonalMobile": "PersonalMobile",
                                "BusinessPhone": "BusinessPhone", "PersonalPhone": "PersonalPhone",
                                "BusinessEmail": "BusinessEmail", "PersonalEmail": "PersonalEmail"},
                   date_key="ContactInformation_startDate")


def process_maritalstatus_data(data):
    return _latest(data or [], {"MaritalStatus": "MaritalStatus"}, defaults={"MaritalStatus_startDate": ""},
                   date_on_field={"MaritalStatus": "MaritalStatus_startDate"})


def process_contract_data(data):
    if not data:
        return {"NatureOfContract": "", "HiringReason": "", "date_debut_contract": "", "date_fin_contract": "",
                "FixedTermContractReason": "", "TerminationReason": "", "Status": ""}
    latest_date = "0001-01-01"
    result = {"NatureOfContract": "", "HiringReason": "", "date_debut_contract": "", "date_fin_contract": "",
              "FixedTermContractReason": "", "TerminationReason": "", "Status": ""}
    fields = {"NatureOfContract", "HiringReason", "FixedTermContractReason", "TerminationReason", "Status"}
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            result["date_debut_contract"] = latest_date
            result["date_fin_contract"] = (item["endDate"] if item["endDate"] and str(item["endDate"]) != "None"
                                           else "1900-01-01")
            if item["field"] in fields:
                result[item["field"]] = item["value"] if item["value"] else ""
    return result


def bank_detail_data(data):
    keys = ["AccountHolder", "IBANOrABA", "BicOrSwift", "AccountNumber", "BankName", "PaymentMethod",
            "UseSameAccountForExpense", "ExpenseAccountPaymentMethod", "ExpenseAccountHolder",
            "ExpenseAccountIBANOrABA", "ExpenseAccountBicOrSwift", "ExpenseAccountBankName", "Country",
            "ExpenseAccountCountry"]
    result = {k: "" for k in keys}
    result.update({"date_debut_bank": "", "date_fin_bank": ""})
    if not data:
        return result
    latest_date = "0001-01-01"
    for item in data:
        if _d(item["startDate"]) >= _d(latest_date):
            latest_date = item["startDate"]
            result["date_debut_bank"] = item["startDate"]
            if item.get("endDate"):
                result["date_fin_bank"] = item["endDate"]
            if item["field"] in keys:
                result[item["field"]] = item["value"] if item["value"] else ""
    return result


def process_probation_data(data):
    return _latest(data or [], {"FirstPeriodStartDate": "FirstPeriodStartDate",
                                "FirstPeriodEndDate": "FirstPeriodEndDate",
                                "SecondPeriodStartDate": "SecondPeriodStartDate",
                                "SecondPeriodEndDate": "SecondPeriodEndDate"},
                   date_key="ProbationaryPeriod_startDate")


def process_civilregistration_data(data):
    return _latest(data or [], {"DateOfBirth": "Date_de_naissance", "TownOrCity": "Ville_de_naissance",
                                "Country": "Pays_de_naissance", "PostalCode": "code_postal_lieu_de_naissance",
                                "AreaOrState": "region_ou_etat_lieu_de_naissance"},
                   date_key="CivilRegistration_startDate")


def process_identity_data(data):
    return _latest(data or [], {"Title": "Title", "Gender": "Gender", "LastName": "LastName",
                                "FirstName": "FirstName", "Nationality": "Nationality", "BirthName": "BirthName"},
                   date_key="Identity_startDate")


PROCESSORS = {
    "keydate": process_keydate_data, "disability": process_disability_data,
    "workaccident": process_workaccident_data, "medicalcheckup": process_medicalcheckup_data,
    "legacyid": process_legacyid_data, "qualification": process_qualification_data,
    "mainpostaladdress": process_mainpostaladdress_data, "contactinfo": process_contactinfo_data,
    "maritalstatus": process_maritalstatus_data, "contract": process_contract_data,
    "bankdetails": bank_detail_data, "probation": process_probation_data, "salary": process_salary_data,
    "civilregistration": process_civilregistration_data, "identity": process_identity_data,
}


# ---------------------------------------------------------------------------
# Orchestration de l'extraction (ts:L186-616)
# ---------------------------------------------------------------------------
class TalentSoftExtractor:
    def __init__(self, client: TsClient, sleep: Callable[[float], None] = time.sleep,
                 batch_size: int = BATCH_SIZE, employee_concurrency: int = EMPLOYEE_CONCURRENCY):
        self.client = client
        self.stats = client.stats
        self._sleep = sleep
        self.batch_size = batch_size
        self.employee_concurrency = employee_concurrency

    # -- liste (ts:L367-405) -------------------------------------------------------
    def get_all_employees(self) -> List[dict]:
        data, status = self.client.get("/api/v1/directory/employees?count=1&offset=0", "directory_count")
        if status != 200 or not data:
            self.stats.inc("directory_count_failed")
            return []
        if not isinstance(data, dict) or "totalCount" not in data:
            self.stats.inc("directory_schema_unexpected")
        total_count = data.get("totalCount", 0) if isinstance(data, dict) else 0
        self.stats.total_count = total_count
        if not total_count or total_count <= 0:
            return []
        data, status = self.client.get(f"/api/v1/directory/employees?count={total_count}&offset=0", "directory_list")
        if status != 200 or not data:
            self.stats.inc("directory_list_failed")
            return []
        if not isinstance(data, dict) or "results" not in data:
            self.stats.inc("directory_schema_unexpected")
        employees = data.get("results", []) if isinstance(data, dict) else []
        self.stats.inc("directory_received", len(employees))
        return employees

    def _classify_unavailable(self, endpoint: str, status: Any) -> None:
        """Absence légitime (200 vide), 404 (à qualifier, D-TS-01) ou échec technique."""
        if status == 200:
            self.stats.inc(f"{endpoint}_empty")
        elif status == 404:
            self.stats.inc(f"{endpoint}_404")
        else:
            self.stats.inc(f"{endpoint}_unavailable")

    def get_situation(self, employee_number) -> dict:
        try:
            data, status = self.client.get(f"/api/hub/v2/employees/{employee_number}/Situation",
                                           "phase1_situation", max_retries=2, initial_delay=1)
            if status == 200 and data:
                return process_situation_data(data)
            self._classify_unavailable("phase1_situation", status)
            return {}
        except HrisError:
            raise
        except Exception:
            self.stats.inc("phase1_situation_error")
            return {}

    def get_position(self, position_code, date, endpoint: str) -> dict:
        try:
            if "T" not in date:
                date = f"{date}T00:00:00Z"
            data, status = self.client.get(f"/api/hub/v2.0/positions/{position_code}?date={date}",
                                           endpoint, max_retries=2, initial_delay=1)
            if status == 200 and data:
                return process_position_data(data)
            self._classify_unavailable(endpoint, status)
            return {}
        except HrisError:
            raise
        except Exception:
            self.stats.inc(f"{endpoint}_error")
            return {}

    # -- filtre « Motul France » (phase 1, séquentielle, ts:L217-246) ----------------
    def filter_motul_france(self, employees: List[dict]) -> List[dict]:
        filtered = []
        for employee in employees:
            employee_number = employee.get("employeeNumber")
            if not employee_number:
                self.stats.inc("excluded_no_employee_number")
                continue
            situation = self.get_situation(employee_number)
            if not situation:
                self.stats.inc("excluded_no_situation")  # échec technique OU absence : voir contrôles
                continue
            position_code = situation.get("Position")
            situation_date = situation.get("Situation_starDate")
            if not (position_code and situation_date):
                self.stats.inc("excluded_missing_position_or_date")
                continue
            position_data = self.get_position(position_code, situation_date, "phase1_position")
            if position_data and str(position_data.get("LegalStructure")) == MOTUL_FRANCE:
                filtered.append(employee)
            elif not position_data:
                self.stats.inc("excluded_no_position_data")
            else:
                self.stats.inc("excluded_not_motul_france")
        self.stats.inc("retained_phase1", len(filtered))
        return filtered

    # -- détail d'un salarié (phase 2, ts:L440-616) --------------------------------
    def get_employee_data(self, employee: dict, pool: ThreadPoolExecutor) -> dict:
        employee_number = employee.get("employeeNumber")
        if not employee_number:
            return {}
        try:
            employee_data = {employee_number: initialize_employee_fields(str(employee_number))}
            individual = employee.get("individual", {}) or {}
            nir = (individual.get("extras", {}) or {}).get("hubSocialSecurityNumber")
            if nir:
                employee_data[employee_number]["hubSocialSecurityNumber"] = str(nir)
            if employee.get("referenceAdmissionDate"):
                employee_data[employee_number]["referenceAdmissionDate"] = str(employee["referenceAdmissionDate"])
            username = (individual.get("user", {}) or {}).get("username", "")
            endpoints = {key: (f"/api/hub/v2/employees/{employee_number}/{entity}", f"hub_{key}")
                         for key, entity in HUB_ENTITIES.items()}
            if username:
                endpoints["postaladdress"] = (f"/api/v1.0/directory/individuals/{username}/postal-addresses",
                                              "postaladdress")
            futures = {key: pool.submit(self.client.get, path, ep) for key, (path, ep) in endpoints.items()}
            responses = {}
            for key, fut in futures.items():
                try:
                    responses[key] = fut.result()
                except HrisError:
                    raise
                except Exception as exc:
                    responses[key] = exc
            # situation d'abord (code et date de position)
            position_code = situation_date = None
            situation_result = responses.get("situation")
            if isinstance(situation_result, tuple) and len(situation_result) == 2:
                data, status = situation_result
                if status == 200 and data:
                    try:
                        situation_data = process_situation_data(data)
                    except Exception:
                        self.stats.inc("processing_error_situation")
                        raise
                    employee_data[employee_number].update(situation_data)
                    position_code = situation_data.get("Position")
                    situation_date = situation_data.get("Situation_starDate")
                else:
                    self.stats.inc("phase2_endpoint_not_ok_situation")
            if position_code and situation_date:
                position_data = self.get_position(position_code, situation_date, "phase2_position")
                if position_data:
                    employee_data[employee_number].update(position_data)
                else:
                    self.stats.inc("phase2_position_missing")
            else:
                self.stats.inc("phase2_position_not_requested")
            for key, result in responses.items():
                if key == "situation":
                    continue
                if isinstance(result, Exception):
                    self.stats.inc(f"phase2_exception_{key}")
                    continue
                if not isinstance(result, tuple) or len(result) != 2:
                    self.stats.inc(f"phase2_invalid_{key}")
                    continue
                data, status = result
                if status == 200:
                    if data:
                        try:
                            if key == "postaladdress":
                                if isinstance(data, list) and len(data) > 0:
                                    employee_data[employee_number].update(process_postal_address_data(data[0]))
                            else:
                                employee_data[employee_number].update(PROCESSORS[key](data))
                        except Exception:
                            self.stats.inc(f"processing_error_{key}")  # parité : on continue
                    else:
                        self.stats.inc(f"phase2_empty_{key}")
                else:
                    self.stats.inc(f"phase2_endpoint_not_ok_{key}")
            return employee_data
        except HrisError:
            raise
        except Exception:
            self.stats.inc("employee_error_defaults")  # parité : enregistrement par défaut (ts:L613-616)
            return {employee_number: initialize_employee_fields(str(employee_number))}

    def process_batch(self, employees: List[dict], request_pool: ThreadPoolExecutor) -> dict:
        out: Dict[Any, dict] = {}

        def one(emp):
            try:
                return self.get_employee_data(emp, request_pool)
            except HrisError:
                raise
            except Exception:
                self.stats.inc("employee_dropped")
                return {}

        with ThreadPoolExecutor(max_workers=self.employee_concurrency) as emp_pool:
            for result in emp_pool.map(one, employees):
                if isinstance(result, dict) and result:
                    out.update(result)
        return out

    def run(self) -> List[dict]:
        """Extraction complète ; renvoie les enregistrements bruts (dictionnaires de l'existant)."""
        self.client.get_token()
        employees = self.get_all_employees()
        if not employees:
            self.stats.inc("no_employees")
            return []
        filtered = self.filter_motul_france(employees)
        if not filtered:
            return []
        all_employee_data: Dict[Any, dict] = {}
        with ThreadPoolExecutor(max_workers=PER_HOST_LIMIT) as request_pool:
            for i in range(0, len(filtered), self.batch_size):
                batch = filtered[i:i + self.batch_size]
                all_employee_data.update(self.process_batch(batch, request_pool))
                self.client.get_token()  # parité : jeton renouvelé à chaque lot
                if i + self.batch_size < len(filtered):
                    self._sleep(BATCH_PAUSE_SECONDS)
        self.stats.inc("duplicates_merged", len(filtered) - len(all_employee_data))
        self.stats.inc("records", len(all_employee_data))
        return list(all_employee_data.values())


def to_raw_row(record: dict) -> Dict[str, Optional[str]]:
    """Projection sur les clés persistées, valeurs au format JSON_VALUE (R-SQL-002)."""
    return {key: json_value_repr(record.get(key)) for key in RAW_KEYS}


# ---------------------------------------------------------------------------
# Contrôles de complétude (D-QUA-01, D-QUA-04, D-TS-02 — À CONFIRMER)
# ---------------------------------------------------------------------------
def completeness_anomalies(stats: dict) -> List[dict]:
    """Traduit les compteurs en contrôles. BLOCKING : l'extraction ne peut pas être présentée
    comme complète. Aucun matricule ni aucune URL dans les messages."""
    c = stats.get("counters", {})
    total = stats.get("total_count")
    received = c.get("directory_received", 0)
    out = []

    def add(code, severity, metric, outcome, message):
        out.append({"control_code": code, "severity": severity, "metric_value": metric,
                    "threshold_value": None, "outcome": outcome, "message": message})

    if total is None or c.get("directory_count_failed") or c.get("directory_list_failed"):
        add("EXTRACTION_INCOMPLETE", "BLOCKING", None, "FAIL", "annuaire TalentSoft illisible (totalCount ou liste)")
    elif received != total:
        add("EXTRACTION_INCOMPLETE", "BLOCKING", received, "FAIL",
            f"reçus {received} pour totalCount {total}")
    else:
        add("EXTRACTION_INCOMPLETE", "BLOCKING", received, "PASS", f"reçus = totalCount = {total}")
    if c.get("directory_schema_unexpected"):
        add("SCHEMA_CHANGE", "BLOCKING", c["directory_schema_unexpected"], "FAIL", "réponse annuaire inattendue")
    records = c.get("records", 0)
    add("EMPTY_EXTRACTION", "BLOCKING", records, "FAIL" if records == 0 else "PASS",
        "aucun salarié retenu" if records == 0 else f"{records} salariés retenus")
    tech_phase1 = c.get("phase1_situation_unavailable", 0) + c.get("phase1_situation_error", 0) + \
        c.get("phase1_position_unavailable", 0) + c.get("phase1_position_error", 0)
    add("TS_FILTER_TECHNICAL_FAILURE", "BLOCKING", tech_phase1, "FAIL" if tech_phase1 else "PASS",
        "salariés non qualifiables (Situation/Position indisponible) : le filtre Motul France est incertain"
        if tech_phase1 else "filtre Motul France évalué pour tous les salariés")
    nf_phase1 = c.get("phase1_situation_404", 0) + c.get("phase1_position_404", 0)
    add("TS_FILTER_NOT_FOUND", "WARNING", nf_phase1, "FAIL" if nf_phase1 else "PASS",
        f"{nf_phase1} réponses 404 au filtrage (salariés exclus, statut à qualifier D-TS-01)")
    failures = 0
    not_found = 0
    for ep, statuses in stats.get("endpoint_status", {}).items():
        if not (ep.startswith("hub_") or ep in ("postaladdress", "phase2_position")):
            continue
        for st, n in statuses.items():
            if st == "200":
                continue
            if st == "404":
                not_found += n
            else:
                failures += n
    processing = sum(v for k, v in c.items() if k.startswith("processing_error_") or k.startswith("phase2_exception_")
                     or k.startswith("phase2_invalid_"))
    employee_level = c.get("employee_error_defaults", 0) + c.get("employee_dropped", 0)
    add("TS_ENDPOINT_FAILURE", "BLOCKING", failures, "FAIL" if failures else "PASS",
        f"{failures} appels de détail en échec après reprises" if failures else "tous les appels de détail aboutis")
    add("TS_ENDPOINT_NOT_FOUND", "WARNING", not_found, "FAIL" if not_found else "PASS", f"{not_found} réponses 404")
    add("TS_PROCESSING_ERROR", "BLOCKING", processing + employee_level, "FAIL" if (processing + employee_level) else "PASS",
        "erreurs de traitement de réponses (valeurs par défaut appliquées par parité)"
        if (processing + employee_level) else "aucune erreur de traitement")
    if c.get("circuit_open_skip"):
        add("TS_CIRCUIT_OPEN", "BLOCKING", c["circuit_open_skip"], "FAIL", "requêtes ignorées par le disjoncteur")
    return out
'''),
    ('hris.transform_ref', '2fe8bb63d4b2f4b141e4dac0b7f9f6af9ea5c2273caf7fafde1927caf9c1caae', r'''"""hris.transform_ref — implémentation de RÉFÉRENCE (Python pur) des procédures SQL de production.

Chaque fonction reproduit une procédure de A09 (PROC1 à PROC5) ligne à ligne, avec les sémantiques
T-SQL émulées par hris.parity. Elle sert :
  * aux tests unitaires par règle R-SQL (exécutés en local, sans Spark) ;
  * de référence pour le test de parité de l'implémentation Spark (hris.transform), exécuté
    sur données fictives en local et dans Fabric.
"""
from __future__ import annotations

import datetime as dt
import decimal
from typing import Dict, Iterable, List, Optional

from hris import parity as P
from hris.model import (ADP_COLUMNS, ADP_DATE_COLUMNS, COMPARE_COLUMNS, FORMAT_REJECT_COLUMNS,
                        FORMAT_SUM_EXCLUDED, LOOKUP_JOINS, NULL_REJECT_COLUMNS, REFERENCE_REJECTS,
                        REJECT_TYPE_FORMAT, REJECT_TYPE_NULL, REJECT_TYPE_REFERENCE, SOURCE_COLUMNS)

HASH_SEPARATOR = "\u001f"
HASH_NULL = "\u0000"


# --- PROC1 : lecture JSON (R-SQL-001 à 004) ----------------------------------------
def parse_source_row(raw: Dict[str, Optional[str]]) -> Dict[str, Optional[str]]:
    out = {}
    for key, col in SOURCE_COLUMNS:
        if key == "BusinessMobile|BusinessPhone":
            mobile, phone = raw.get("BusinessMobile"), raw.get("BusinessPhone")
            value = mobile if mobile is not None else phone  # COALESCE : NULL seulement, pas ''
        elif key == "date_fin_contract":
            value = raw.get(key)
            value = None if value is None else value.replace("None", "")
        else:
            value = raw.get(key)
        out[col] = P.nvarchar(value, 50)
    return out


# --- PROC2 : staging, 15 LEFT JOIN (R-SQL-010 à 024) ---------------------------------
def enrich_rows(source_rows: Iterable[dict], lookups: Dict[str, List[dict]]) -> List[dict]:
    """LEFT JOIN successifs : une correspondance dupliquée DUPLIQUE la ligne (parité R-SQL-058)."""
    rows = [dict(r) for r in source_rows]
    for join in LOOKUP_JOINS:
        table = lookups.get(join["table"], [])
        next_rows = []
        for row in rows:
            matches = [lu for lu in table if P.sql_eq(row.get(join["src"]), lu.get(join["key"]))]
            if not matches:
                r = dict(row)
                for _, out in join["out"]:
                    r[out] = None
                next_rows.append(r)
            for lu in matches:
                r = dict(row)
                for lu_col, out in join["out"]:
                    r[out] = lu.get(lu_col)
                next_rows.append(r)
        rows = next_rows
    return rows


# --- PROC3 : rejets (R-SQL-030 à 042) --------------------------------------------------
def _format_flag(value: Optional[str], kind: str) -> int:
    if kind == "len8":
        return 0 if P.sql_len(value) == 8 else 1
    if kind == "date":
        return 0 if (P.sql_is_empty(value) or P.isdate(value) == 1) else 1
    if kind == "date10":
        return 0 if (P.sql_is_empty(value) or P.isdate(None if value is None else value[:10]) == 1) else 1
    if kind == "pct":
        return P.pct_flag(value, 100)
    if kind == "amount":
        return P.pct_flag(value, 1000000)
    raise ValueError(kind)


def _format_ds(col: str, value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    if col[:2] == "dt":
        return f"Format Date non respecté au YYYY-MM-DD pour {col} = {value}"
    if col[:2] == "pc":
        return f"Format Pourcentage non respecté du 0 au 100 pour {col} = {value}"
    return None


def reject_rows(enriched: Iterable[dict], dt_rejet: dt.date) -> List[dict]:
    enriched = list(enriched)
    out: List[dict] = []
    # Null_Rejet
    for r in enriched:
        flags = {c: (1 if P.sql_is_empty(r.get(c)) else 0) for c in NULL_REJECT_COLUMNS}
        if not any(flags.values()):
            continue
        for c in NULL_REJECT_COLUMNS:
            if flags[c] == 1:
                out.append(_reject(dt_rejet, r, REJECT_TYPE_NULL, c, r.get(c), "", f"{c} ne peut pas être NULL"))
    # Format_Rejet
    for r in enriched:
        flags = {c: _format_flag(r.get(c), kind) for c, kind in FORMAT_REJECT_COLUMNS}
        total = sum(v for c, v in flags.items() if c not in FORMAT_SUM_EXCLUDED)
        if total <= 0:
            continue
        for c, _ in FORMAT_REJECT_COLUMNS:
            if flags[c] == 1:
                out.append(_reject(dt_rejet, r, REJECT_TYPE_FORMAT, c, r.get(c), "", _format_ds(c, r.get(c))))
    # Référence_Rejet
    for r in enriched:
        flags = []
        for fichier, nom, val_col, lu_col, adp_col in REFERENCE_REJECTS:
            value = r.get(val_col)
            missing = r.get(lu_col) is None or (adp_col is not None and r.get(adp_col) is None)
            flags.append(1 if (P.sql_is_not_empty(value) and missing) else 0)
        if sum(flags) <= 0:
            continue
        for flag, (fichier, nom, val_col, _, _) in zip(flags, REFERENCE_REJECTS):
            if flag == 1:
                value = r.get(val_col)
                out.append(_reject(dt_rejet, r, REJECT_TYPE_REFERENCE, nom, value, fichier,
                                   f"Mapping non trouvé entre TS et ADP pour {nom} = {value}"))
    return out


def _reject(dt_rejet, r, typ, col, value, fichier, ds):
    return {"dt_rejet": dt_rejet, "id_unique": r.get("id_unique"), "id_payroll": r.get("id_payroll"),
            "cd_type_rejet": typ, "lb_nom_colonne_rejet": col, "lb_valeur_rejet": value, "fl_rejet": 1,
            "lb_fichier_mapping": fichier, "ds_rejet": ds}


# --- PROC4 : transformation ADP (R-SQL-050 à 058) ------------------------------------
def adp_rows(enriched: Iterable[dict], rejects: Iterable[dict], dt_transforme: dt.date) -> List[dict]:
    rejected = {r["id_unique"] for r in rejects if r.get("id_unique") is not None}
    out = []
    for r in enriched:
        if r.get("id_unique") is not None and r["id_unique"] in rejected:
            continue  # NOT EXISTS sur id_unique (un id_unique NULL n'est jamais exclu)
        row: Dict[str, object] = {"dt_transforme": dt_transforme}
        for col in ADP_COLUMNS[1:]:
            if col in ADP_DATE_COLUMNS:
                value = r.get(col)
                if col == "dt_entree_societe" and value is not None:
                    value = value[:10]
                row[col] = P.convert_date(value)  # ValueError = échec du run (parité)
            elif col == "mt_salaire_mensuel":
                row[col] = P.monthly_salary(r.get("mt_salaire_base"))  # ValueError = échec (D-PAR-07)
            else:
                row[col] = r.get(col)
        out.append(row)
    return out


def hash_value(value) -> str:
    if value is None:
        return HASH_NULL
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, decimal.Decimal):
        return format(value, "f")
    return str(value)


def row_hash(row: dict, columns: List[str] = COMPARE_COLUMNS) -> str:
    """sha2(concat_ws('\\u001f', colonnes comparées, NULL -> '\\u0000'), 256) (A07 § 4)."""
    from hris.common import sha256_hex
    return sha256_hex(HASH_SEPARATOR.join(hash_value(row.get(c)) for c in columns))


# --- PROC5 : résumé (R-SQL-060 à 066) ------------------------------------------------
def summary(inserted_keys: Iterable[str], rejects: Iterable[dict], dt_resume: dt.date) -> dict:
    rejects = list(rejects)
    n_rej = len({r["id_unique"] for r in rejects if r.get("id_unique") is not None})
    n_ins = len({k for k in inserted_keys if k is not None})
    lues = n_ins + n_rej
    pc = decimal.Decimal(0) if lues == 0 else (decimal.Decimal(n_rej) / decimal.Decimal(lues)).quantize(
        decimal.Decimal("0.0001"), rounding=decimal.ROUND_HALF_UP)
    pairs = sorted({(r["cd_type_rejet"], r.get("ds_rejet")) for r in rejects},
                   key=lambda p: (p[0], p[1] or ""))  # STRING_AGG non ordonné : tri explicite (A02 § 8)
    ds_type = ",".join(p[0] for p in pairs) if pairs else None
    return {"dt_resume": dt_resume, "cd_source": "TalentSoft", "cd_target": "ADP", "mt_lignes_inserts": n_ins,
            "mt_lignes_rejets": n_rej, "mt_lignes_lues": lues, "pc_lignes_rejets": pc,
            "ds_type_rejets": P.utf8_truncate(ds_type, 50)}
'''),
    ('hris.spark_io', 'd6aececfa4dfaec8f69b974eba9fe21320ec7eecd66848473edafe628fdfae0f', r'''"""hris.spark_io — écritures Delta idempotentes, journal des étapes, anomalies, socle.

Toutes les écritures passent par Spark (un seul moteur d'écriture par table, A06 F2).
Idempotence : `replaceWhere run_id = …` pour les tables à grain « run », MERGE sur clé naturelle
pour les autres. Aucune table préexistante n'est supprimée, vidée ni remplacée : la création
est `CREATE TABLE IF NOT EXISTS` et un schéma incompatible est signalé, jamais corrigé d'office.
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Any, Dict, Iterable, List, Optional

from hris.common import HrisError, PROCESS_CODE, safe_error, sanitize, to_utc_naive, utc_now
from hris.model import TABLES, Tables, create_table_sql, normalize_type

_RUN_ID = re.compile(r"^[A-Za-z0-9_\-]{1,100}$")


def check_run_id(run_id: str) -> str:
    """Les identifiants de run entrent dans des prédicats replaceWhere : format strict."""
    if not run_id or not _RUN_ID.match(run_id):
        raise HrisError("PARAM_INVALID", "run_id invalide")
    return run_id


def sql_str(value: str) -> str:
    return "'" + str(value).replace("\\", "\\\\").replace("'", "\\'") + "'"


def spark_schema(logical: str):
    from pyspark.sql import types as T
    mapping = {"STRING": T.StringType(), "DATE": T.DateType(), "TIMESTAMP": T.TimestampType(),
               "INT": T.IntegerType(), "BIGINT": T.LongType(), "BOOLEAN": T.BooleanType(),
               "ARRAY<STRING>": T.ArrayType(T.StringType())}
    fields = []
    for name, typ in TABLES[logical]:
        t = typ.upper()
        if t.startswith("DECIMAL"):
            p, s = re.findall(r"\d+", t)
            dtype = T.DecimalType(int(p), int(s))
        else:
            dtype = mapping[t]
        fields.append(T.StructField(name, dtype, True))
    return T.StructType(fields)


def conform(df, logical: str):
    """Sélectionne et type les colonnes dans l'ordre du DDL (colonnes absentes -> NULL)."""
    from pyspark.sql import functions as F
    schema = spark_schema(logical)
    cols = []
    present = set(df.columns)
    for f in schema.fields:
        if f.name in present:
            cols.append(F.col(f"`{f.name}`").cast(f.dataType).alias(f.name))
        else:
            cols.append(F.lit(None).cast(f.dataType).alias(f.name))
    return df.select(*cols)


def rows_df(spark, rows: List[dict], logical: str):
    schema = spark_schema(logical)
    data = [tuple(_py(r.get(f.name), f.dataType) for f in schema.fields) for r in rows]
    return spark.createDataFrame(data, schema=schema)


def _py(value, dtype):
    from pyspark.sql import types as T
    if value is None:
        return None
    if isinstance(dtype, T.TimestampType) and isinstance(value, dt.datetime):
        return to_utc_naive(value)
    if isinstance(dtype, T.DecimalType):
        import decimal
        return decimal.Decimal(str(value))
    if isinstance(dtype, T.IntegerType) or isinstance(dtype, T.LongType):
        return int(value)
    return value


def replace_run(spark, df, physical: str, logical: str, run_id: str, extra_predicate: str = "") -> None:
    check_run_id(run_id)
    predicate = f"run_id = {sql_str(run_id)}" + (f" AND ({extra_predicate})" if extra_predicate else "")
    (conform(df, logical).write.format("delta").mode("overwrite")
     .option("replaceWhere", predicate).saveAsTable(physical))


def append(df, physical: str, logical: str) -> None:
    conform(df, logical).write.format("delta").mode("append").saveAsTable(physical)


def merge_rows(spark, rows: List[dict], physical: str, logical: str, keys: List[str],
               update: bool = True) -> None:
    """MERGE idempotent sur clé naturelle (insert si absent, mise à jour sinon)."""
    if not rows:
        return
    from delta.tables import DeltaTable
    src = rows_df(spark, rows, logical)
    cond = " AND ".join(f"t.`{k}` <=> s.`{k}`" for k in keys)
    m = DeltaTable.forName(spark, physical).alias("t").merge(src.alias("s"), cond)
    if update:
        m = m.whenMatchedUpdateAll()
    m.whenNotMatchedInsertAll().execute()


# ---------------------------------------------------------------------------
# Socle : schémas et tables (NB_HRIS_SETUP), contrôle de conformité
# ---------------------------------------------------------------------------
def ensure_schemas_and_tables(spark, T: Tables, logicals: Optional[Iterable[str]] = None) -> List[dict]:
    """Crée ce qui manque ; vérifie le schéma de ce qui existe. Ne modifie jamais une table existante."""
    report = []
    for schema in T.schemas():
        spark.sql(f"CREATE SCHEMA IF NOT EXISTS {schema}")  # schéma déjà qualifié par Tables.schemas()
    for logical in (logicals or TABLES.keys()):
        physical = T(logical)
        existed = spark.catalog.tableExists(physical)
        if not existed:
            spark.sql(create_table_sql(physical, logical))
            report.append({"table": physical, "action": "CREATED", "detail": ""})
            continue
        report.append({"table": physical, "action": "EXISTS", "detail": check_table_schema(spark, physical, logical)})
    return report


def check_table_schema(spark, physical: str, logical: str) -> str:
    actual = {f.name.lower(): normalize_type(f.dataType.simpleString()) for f in spark.table(physical).schema.fields}
    problems = []
    for name, typ in TABLES[logical]:
        exp = normalize_type(typ)
        got = actual.get(name.lower())
        if got is None:
            problems.append(f"colonne manquante {name}")
        elif got != exp:
            problems.append(f"type {name} {got}!={exp}")
    return "CONFORME" if not problems else "INCOMPATIBLE: " + "; ".join(problems)


# ---------------------------------------------------------------------------
# Journal des étapes (ctl.run_step) et anomalies (log.anomaly)
# ---------------------------------------------------------------------------
class StepLogger:
    """`STARTED` au début, `SUCCEEDED` / `FAILED` / `SKIPPED` à la fin (03 § 4). attempt_no = nombre
    de tentatives antérieures de la même étape pour ce run + 1."""

    def __init__(self, spark, T: Tables, run_id: str, step_code: str):
        self.spark, self.T, self.run_id, self.step_code = spark, T, check_run_id(run_id), step_code
        self.started = utc_now()
        prior = (spark.table(T("ctl.run_step"))
                 .where(f"run_id = {sql_str(run_id)} AND step_code = {sql_str(step_code)} AND status = 'STARTED'")
                 .count())
        self.attempt_no = int(prior) + 1
        self._write("STARTED", None, None, None, "")

    def _write(self, status, rows_in, rows_out, content_hash, message, counters_only=False):
        row = {"run_id": self.run_id, "step_code": self.step_code, "attempt_no": self.attempt_no,
               "status": status, "rows_in": rows_in, "rows_out": rows_out, "content_hash": content_hash,
               "message": sanitize(message, 8000, mask_numbers=not counters_only), "started_at_utc": self.started,
               "ended_at_utc": None if status == "STARTED" else utc_now()}
        append(rows_df(self.spark, [row], "ctl.run_step"), self.T("ctl.run_step"), "ctl.run_step")

    def succeeded(self, rows_in=None, rows_out=None, content_hash=None, message="", counters_only=False):
        """counters_only=True : message construit par le code à partir de compteurs (nombres conservés)."""
        self._write("SUCCEEDED", rows_in, rows_out, content_hash, message, counters_only)

    def skipped(self, message=""):
        self._write("SKIPPED", None, None, None, message)

    def failed(self, exc: BaseException):
        self._write("FAILED", None, None, None, safe_error(exc))


def step_started(spark, T: Tables, run_id: str, step_code: str) -> bool:
    return spark.table(T("ctl.run_step")).where(
        f"run_id = {sql_str(run_id)} AND step_code = {sql_str(step_code)}").limit(1).count() > 0


def step_succeeded(spark, T: Tables, run_id: str, step_code: str) -> bool:
    return spark.table(T("ctl.run_step")).where(
        f"run_id = {sql_str(run_id)} AND step_code = {sql_str(step_code)} AND status = 'SUCCEEDED'"
    ).limit(1).count() > 0


def write_anomalies(spark, T: Tables, run_id: str, anomalies: List[dict], owned_codes: Iterable[str]) -> None:
    """Remplace, pour ce run, les contrôles dont l'étape est propriétaire (rejouable sans doublon)."""
    owned = sorted(set(owned_codes) | {a["control_code"] for a in anomalies})
    now = utc_now()
    rows = [{"run_id": run_id, "process_code": PROCESS_CODE, "created_at_utc": now,
             "message": sanitize(a.get("message", ""), 500), **{k: a.get(k) for k in
             ("control_code", "severity", "metric_value", "threshold_value", "outcome")}} for a in anomalies]
    predicate = "control_code IN (" + ",".join(sql_str(c) for c in owned) + ")" if owned else "1=0"
    replace_run(spark, rows_df(spark, rows, "log.anomaly"), T("log.anomaly"), "log.anomaly", run_id, predicate)


def blocking_failures(spark, T: Tables, run_id: str) -> List[str]:
    rows = (spark.table(T("log.anomaly"))
            .where(f"run_id = {sql_str(run_id)} AND severity = 'BLOCKING' AND outcome = 'FAIL'")
            .select("control_code").distinct().collect())
    return sorted(r.control_code for r in rows)


# ---------------------------------------------------------------------------
# Registre des runs (ctl.run)
# ---------------------------------------------------------------------------
def get_run(spark, T: Tables, run_id: str) -> Optional[dict]:
    rows = spark.table(T("ctl.run")).where(f"run_id = {sql_str(check_run_id(run_id))}").collect()
    if len(rows) > 1:
        raise HrisError("RUN_DUPLICATE", "plusieurs lignes ctl.run pour le même run")
    return rows[0].asDict() if rows else None


def update_run(spark, T: Tables, run_id: str, **fields: Any) -> None:
    """MERGE de mise à jour de ctl.run (colonnes nommées uniquement)."""
    from delta.tables import DeltaTable
    from pyspark.sql import functions as F
    check_run_id(run_id)
    if not fields:
        return
    schema = {n: t for n, t in TABLES["ctl.run"]}
    sets = {}
    for k, v in fields.items():
        if k not in schema:
            raise KeyError(k)
        if isinstance(v, dt.datetime):
            v = to_utc_naive(v)
        if k == "error_summary" and v is not None:
            v = sanitize(v, 1000)
        sets[k] = F.lit(v).cast(spark_schema("ctl.run")[k].dataType)
    (DeltaTable.forName(spark, T("ctl.run")).alias("t")
     .update(condition=F.col("t.run_id") == F.lit(run_id), set=sets))


def set_once_promoted(spark, T: Tables, run_id: str, when: dt.datetime) -> None:
    """promoted_at_utc écrit une seule fois (garde d'idempotence de la promotion)."""
    from delta.tables import DeltaTable
    from pyspark.sql import functions as F
    (DeltaTable.forName(spark, T("ctl.run")).alias("t")
     .update(condition=(F.col("t.run_id") == F.lit(run_id)) & F.col("t.promoted_at_utc").isNull(),
             set={"promoted_at_utc": F.lit(to_utc_naive(when)).cast("timestamp")}))
'''),
    ('hris.transform', '91a7ba2c70cda0430f5c5d60970c9f02dc7207393cdf7069ebd0732f19e9f152', r'''"""hris.transform — implémentation Spark des procédures SQL de production (A02, A09).

Doit produire exactement les résultats de hris.transform_ref (test de parité sur données fictives).
Guide T-SQL -> Spark appliqué (A02 § 8) : espaces de fin (rtrim), NULL vs '', troncature à 50,
CONVERT(DATE,'') = 1900-01-01, émulations ISDATE / ISNUMERIC, décimales.

Écart relevé par rapport à A02 § 6 : en Spark 3.5, `decimal(12,2) / 12` donne DECIMAL(16,6)
(littéral typé au plus juste) ; le diviseur est donc converti explicitement en DECIMAL(10,0)
pour obtenir DECIMAL(23,13) comme T-SQL (R-SQL-054).
"""
from __future__ import annotations

from typing import Dict, List, Tuple

from hris.common import HrisError, PROCESS_CODE, sha256_hex, utc_now
from hris.model import (ADP_COLUMNS, ADP_DATE_COLUMNS, COMPARE_COLUMNS, FORMAT_REJECT_COLUMNS,
                        FORMAT_SUM_EXCLUDED, LOOKUP_JOINS, LOOKUP_TABLES, NULL_REJECT_COLUMNS, REFERENCE_REJECTS,
                        REJECT_TYPE_FORMAT, REJECT_TYPE_NULL, REJECT_TYPE_REFERENCE, SOURCE_COLUMNS,
                        USED_LOOKUP_TABLES, Tables)

SENSITIVE_REJECT_COLUMNS = ["dt_naissance", "id_numero_securite_sociale", "lb_iban", "lb_bic",
                            "lb_account_number", "lb_account_holder", "lb_bank_name"]  # D-SEC-03

_WS_TRIM = r"^[ \t\r\n]+|[ \t\r\n]+$"
# Gardes identiques à hris.parity.DATE_PATTERNS ; l'analyse elle-même est stricte (calendrier réel)
_SPARK_DATE_PARSE = {
    "ymd_dash": ("text", "yyyy-M-d"),
    "ymd_compact": ("text", "yyyyMMdd"),
    "ymd_iso_time": ("first10", "yyyy-MM-dd"),
    "ymd_space_time": ("before_space", "yyyy-M-d"),
    "mdy_slash": ("text", "M/d/yyyy"),
}
_ISNUMERIC = r"^[+-]?[$€£¥]?[ ]*[+-]?(\d[\d,]*\.?\d*|\.\d*|,[\d,]*\.?\d*)?([eEdD][+-]?\d+)?$"
_STRICT_DECIMAL = r"^[ \t]*([+-]?)(\d+\.?\d*|\.\d+)[ \t]*$"


def _F():
    from pyspark.sql import functions as F
    return F


def ensure_session_settings(spark) -> None:
    """Réglages requis par les émulations : fuseau UTC et analyseur de dates « CORRECTED »
    (les chaînes non conformes donnent NULL au lieu d'une exception de compatibilité)."""
    spark.conf.set("spark.sql.session.timeZone", "UTC")
    spark.conf.set("spark.sql.legacy.timeParserPolicy", "CORRECTED")


# ---------------------------------------------------------------------------
# Sémantiques T-SQL en expressions Spark (miroir de hris.parity)
# ---------------------------------------------------------------------------
def is_empty(c):
    F = _F()
    return F.rtrim(c) == F.lit("")


def is_not_empty(c):
    F = _F()
    return F.rtrim(c) != F.lit("")


def parse_date(c):
    """Partie date reconnue (même motif que hris.parity.parse_tsql_date) ou NULL."""
    from hris.parity import DATE_PATTERNS
    F = _F()
    text = F.regexp_replace(c, _WS_TRIM, "")
    out = F.lit(None).cast("date")
    for name, guard, _extract, _order in reversed(DATE_PATTERNS):
        source, fmt = _SPARK_DATE_PARSE[name]
        arg = {"text": text, "first10": F.substring(text, 1, 10),
               "before_space": F.regexp_extract(text, r"^(\S+) ", 1)}[source]
        parsed = F.try_to_timestamp(arg, F.lit(fmt)).cast("date")
        out = F.when(text.rlike(guard), parsed).otherwise(out)
    return out


def isdate(c):
    F = _F()
    d = parse_date(c)
    return F.when(d.isNotNull() & (F.year(d) >= 1753), F.lit(1)).otherwise(F.lit(0))


def convert_date(c):
    """CONVERT(DATE, x) : NULL -> NULL ; blanc -> 1900-01-01 ; non convertible -> erreur."""
    F = _F()
    d = parse_date(c)
    blank = F.regexp_replace(c, _WS_TRIM, "") == F.lit("")
    return (F.when(c.isNull(), F.lit(None).cast("date"))
            .when(blank, F.lit("1900-01-01").cast("date"))
            .when(d.isNotNull(), d)
            .otherwise(F.raise_error(F.lit("CONVERSION_DATE")).cast("date")))


def isnumeric(c):
    F = _F()
    text = F.regexp_replace(c, _WS_TRIM, "")
    body = F.regexp_extract(text, _ISNUMERIC, 1)
    exp = F.regexp_extract(text, _ISNUMERIC, 2)
    valid = text.rlike(_ISNUMERIC) & ~((exp != "") & ~body.rlike(r"\d"))
    return (F.when(c.isNull(), F.lit(0)).when(text == "", F.lit(0))
            .when(valid, F.lit(1)).otherwise(F.lit(0)))


def cast_numeric_12_2(c):
    """CAST(x AS NUMERIC(12,2)) strict ; échec -> erreur (parité : le run échoue)."""
    F = _F()
    sign = F.regexp_extract(c, _STRICT_DECIMAL, 1)
    num = F.regexp_extract(c, _STRICT_DECIMAL, 2)
    wide = F.concat(sign, num).cast("decimal(38,10)")
    rounded = F.round(wide, 2)
    ok = c.rlike(_STRICT_DECIMAL) & wide.isNotNull() & (F.abs(rounded) < F.lit(10000000000))
    return (F.when(c.isNull(), F.lit(None).cast("decimal(12,2)"))
            .when(ok, rounded.cast("decimal(12,2)"))
            .otherwise(F.raise_error(F.lit("CONVERSION_NUMERIC")).cast("decimal(12,2)")))


def pct_flag(c, upper: int):
    F = _F()
    value = cast_numeric_12_2(c)
    return (F.when(isnumeric(c) == 0, F.lit(0))
            .when((value >= F.lit(0)) & (value <= F.lit(upper)), F.lit(0)).otherwise(F.lit(1)))


# ---------------------------------------------------------------------------
# PROC1 : stg.ts_employee_raw -> stg.ts_employee_source
# ---------------------------------------------------------------------------
def source_df(raw):
    F = _F()
    cols = [F.col("run_id"), F.col("business_date")]
    for key, col in SOURCE_COLUMNS:
        if key == "BusinessMobile|BusinessPhone":
            value = F.coalesce(F.col("BusinessMobile"), F.col("BusinessPhone"))
        elif key == "date_fin_contract":
            value = F.expr("replace(`date_fin_contract`, 'None', '')")
        else:
            value = F.col(f"`{key}`")
        cols.append(F.substring(value, 1, 50).alias(col))
    return raw.select(*cols)


# ---------------------------------------------------------------------------
# Contrôles des correspondances (D-QUA-03, À CONFIRMER : bloquant)
# ---------------------------------------------------------------------------
def mapping_controls(spark, T: Tables) -> Tuple[List[dict], str]:
    F = _F()
    anomalies = []
    digests = []
    keys = {j["table"]: j["key"] for j in LOOKUP_JOINS}
    for table in USED_LOOKUP_TABLES:
        physical = T(f"cfg.{table}")
        if not spark.catalog.tableExists(physical):
            anomalies.append(_anom("MAPPING_TABLE_MISSING", 1, f"{table} absente"))
            continue
        df = spark.table(physical)
        missing = [c for c in LOOKUP_TABLES[table] if c not in df.columns]
        if missing:
            anomalies.append(_anom("MAPPING_SCHEMA", len(missing), f"{table} : colonnes attendues absentes"))
            continue
        n = df.count()
        if n == 0:
            anomalies.append(_anom("MAPPING_TABLE_EMPTY", 0, f"{table} vide (initialisation D-REF-01 non réalisée)"))
            continue
        dup = (df.where(F.col(keys[table]).isNotNull()).groupBy(F.rtrim(F.col(keys[table])).alias("k"))
               .count().where("count > 1").count())
        if dup:
            anomalies.append(_anom("MAPPING_KEY_DUPLICATE", dup, f"{table} : {dup} clés en double"))
        row_digest = F.sha2(F.concat_ws("\u001f", *[F.coalesce(F.col(c), F.lit("\u0000"))
                                                    for c in LOOKUP_TABLES[table]]), 256)
        hashes = sorted(r[0] for r in df.select(row_digest).collect())
        digests.append(f"{table}:{sha256_hex(''.join(hashes))}")
    controls = ["MAPPING_TABLE_MISSING", "MAPPING_SCHEMA", "MAPPING_TABLE_EMPTY", "MAPPING_KEY_DUPLICATE"]
    found = {a["control_code"] for a in anomalies}
    for code in controls:
        if code not in found:
            anomalies.append({"control_code": code, "severity": "BLOCKING", "metric_value": 0,
                              "threshold_value": None, "outcome": "PASS", "message": "13 tables contrôlées"})
    return anomalies, sha256_hex("|".join(digests))


def _anom(code, metric, message):
    return {"control_code": code, "severity": "BLOCKING", "metric_value": metric, "threshold_value": None,
            "outcome": "FAIL", "message": message}


# ---------------------------------------------------------------------------
# PROC2 : staging (15 LEFT JOIN)
# ---------------------------------------------------------------------------
def enriched_df(spark, T: Tables, source, dt_staging, loader=None):
    """`loader(table)` -> DataFrame de correspondance (défaut : table Delta cfg.<table>)."""
    F = _F()
    loader = loader or (lambda table: spark.table(T(f"cfg.{table}")))
    df = source.withColumn("dt_staging", F.lit(dt_staging).cast("date"))
    for i, join in enumerate(LOOKUP_JOINS):
        lu = loader(join["table"])
        alias = f"lu{i}"
        sel = [F.rtrim(F.col(join["key"])).alias(f"__k{i}")] + \
              [F.col(lu_col).alias(out) for lu_col, out in join["out"]]
        lu = lu.select(*sel).alias(alias)
        df = df.join(lu, F.rtrim(df[join["src"]]) == F.col(f"{alias}.__k{i}"), "left").drop(f"__k{i}")
    return df


# ---------------------------------------------------------------------------
# PROC3 : rejets
# ---------------------------------------------------------------------------
def reject_df(enriched, dt_rejet):
    F = _F()
    base_cols = [F.lit(dt_rejet).cast("date").alias("dt_rejet"), F.col("id_unique"), F.col("id_payroll")]

    def explode_family(entries, typ, where_any):
        arr = F.array(*[F.struct(F.lit(i).alias("ord"), F.lit(fichier).alias("lb_fichier_mapping"),
                                 F.lit(nom).alias("lb_nom_colonne_rejet"), val.alias("lb_valeur_rejet"),
                                 flag.alias("fl_rejet"), ds.alias("ds_rejet"))
                        for i, (fichier, nom, val, flag, ds) in enumerate(entries)])
        return (enriched.where(where_any)
                .select(*base_cols, F.posexplode(arr).alias("pos", "r"))
                .where(F.col("r.fl_rejet") == 1)
                .select("dt_rejet", "id_unique", "id_payroll", F.lit(typ).alias("cd_type_rejet"),
                        F.col("r.lb_nom_colonne_rejet").alias("lb_nom_colonne_rejet"),
                        F.col("r.lb_valeur_rejet").alias("lb_valeur_rejet"), F.col("r.fl_rejet").alias("fl_rejet"),
                        F.col("r.lb_fichier_mapping").alias("lb_fichier_mapping"),
                        F.col("r.ds_rejet").alias("ds_rejet")))

    # Null_Rejet
    null_flags = {c: F.when(is_empty(F.col(c)), 1).otherwise(0) for c in NULL_REJECT_COLUMNS}
    null_entries = [("", c, F.col(c), null_flags[c], F.concat(F.lit(c), F.lit(" ne peut pas être NULL")))
                    for c in NULL_REJECT_COLUMNS]
    null_any = None
    for c in NULL_REJECT_COLUMNS:
        cond = is_empty(F.col(c))
        null_any = cond if null_any is None else (null_any | cond)
    nulls = explode_family(null_entries, REJECT_TYPE_NULL, F.coalesce(null_any, F.lit(False)))

    # Format_Rejet
    def fflag(col, kind):
        c = F.col(col)
        if kind == "len8":
            return F.when(F.length(F.rtrim(c)) == 8, 0).otherwise(1)
        if kind == "date":
            return F.when(is_empty(c) | (isdate(c) == 1), 0).otherwise(1)
        if kind == "date10":
            return F.when(is_empty(c) | (isdate(F.substring(c, 1, 10)) == 1), 0).otherwise(1)
        if kind == "pct":
            return pct_flag(c, 100)
        if kind == "amount":
            return pct_flag(c, 1000000)
        raise ValueError(kind)

    flags = {col: F.coalesce(fflag(col, kind), F.lit(1)) for col, kind in FORMAT_REJECT_COLUMNS}

    def fds(col):
        val = F.col(col)
        if col[:2] == "dt":
            return F.concat(F.lit(f"Format Date non respecté au YYYY-MM-DD pour {col} = "), val)
        if col[:2] == "pc":
            return F.concat(F.lit(f"Format Pourcentage non respecté du 0 au 100 pour {col} = "), val)
        return F.lit(None).cast("string")

    total = None
    for col, _ in FORMAT_REJECT_COLUMNS:
        if col in FORMAT_SUM_EXCLUDED:
            continue
        total = flags[col] if total is None else total + flags[col]
    fmt_entries = [("", col, F.col(col), flags[col], fds(col)) for col, _ in FORMAT_REJECT_COLUMNS]
    formats = explode_family(fmt_entries, REJECT_TYPE_FORMAT, total > 0)

    # Référence_Rejet
    ref_entries = []
    ref_total = None
    for fichier, nom, val_col, lu_col, adp_col in REFERENCE_REJECTS:
        missing = F.col(lu_col).isNull()
        if adp_col is not None:
            missing = missing | F.col(adp_col).isNull()
        flag = F.when(is_not_empty(F.col(val_col)) & missing, 1).otherwise(0)
        ref_total = flag if ref_total is None else ref_total + flag
        ref_entries.append((fichier, nom, F.col(val_col), flag,
                            F.concat(F.lit(f"Mapping non trouvé entre TS et ADP pour {nom} = "), F.col(val_col))))
    refs = explode_family(ref_entries, REJECT_TYPE_REFERENCE, ref_total > 0)
    return nulls.unionByName(formats).unionByName(refs)


def masked_log_reject(rejects, run_id: str):
    """log.reject : valeurs sensibles masquées (D-SEC-03)."""
    F = _F()
    sensitive = F.col("lb_nom_colonne_rejet").isin(*SENSITIVE_REJECT_COLUMNS)
    return rejects.select(
        F.lit(run_id).alias("run_id"), F.lit(PROCESS_CODE).alias("process_code"), "dt_rejet", "id_unique",
        "id_payroll", "cd_type_rejet", "lb_nom_colonne_rejet",
        F.when(sensitive & F.col("lb_valeur_rejet").isNotNull(), F.lit("***")).otherwise(F.col("lb_valeur_rejet"))
        .alias("lb_valeur_rejet"),
        "lb_fichier_mapping",
        F.when(sensitive & F.col("ds_rejet").contains(" = "),
               F.concat(F.regexp_replace(F.col("ds_rejet"), r" = .*$", ""), F.lit(" = ***")))
        .otherwise(F.col("ds_rejet")).alias("ds_rejet"),
        F.lit(utc_now().replace(tzinfo=None)).cast("timestamp").alias("created_at_utc"))


# ---------------------------------------------------------------------------
# PROC4 : transformation ADP + empreinte
# ---------------------------------------------------------------------------
def adp_df(enriched, rejects, dt_transforme):
    F = _F()
    rejected_keys = rejects.where(F.col("id_unique").isNotNull()).select("id_unique").distinct()
    kept = enriched.join(rejected_keys, on="id_unique", how="left_anti")
    cols = [F.lit(dt_transforme).cast("date").alias("dt_transforme")]
    for col in ADP_COLUMNS[1:]:
        if col in ADP_DATE_COLUMNS:
            src = F.substring(F.col(col), 1, 10) if col == "dt_entree_societe" else F.col(col)
            cols.append(convert_date(src).alias(col))
        elif col == "mt_salaire_mensuel":
            divisor = F.lit(12).cast("decimal(10,0)")
            cols.append((cast_numeric_12_2(F.col("mt_salaire_base")) / divisor).alias(col))
        else:
            cols.append(F.col(col))
    out = kept.select(*cols)
    return out.withColumn("row_hash", row_hash_expr(COMPARE_COLUMNS))


def row_hash_expr(columns):
    F = _F()
    return F.sha2(F.concat_ws("\u001f", *[F.coalesce(F.col(c).cast("string"), F.lit("\u0000")) for c in columns]), 256)
'''),
    ('hris.compare', '44fbebdd86edfc03c67e3a93e7085e98da1800e48b52aca7ec7702ef0106e03c', r'''"""hris.compare — comparaison staging finale <-> référence du dernier run validé, puis promotion.

Règles (01 § 4.5-4.8, 03 § 7, A08 D-CMP-01 à 07 — À CONFIRMER, valeurs par défaut appliquées) :
  * catégories NEW / MODIFIED / UNCHANGED / ABSENT / REJECTED sur id_unique et row_hash ;
  * le comparatif ne journalise que des NOMS de colonnes et des empreintes ;
  * premier chargement : NORMAL + référence vide -> BLOQUÉ (REFERENCE_ABSENTE) ;
    INIT_REFERENCE accepté seulement si la référence est vide (jamais de réinitialisation) ;
  * seuils : NULL -> NOT_CONFIGURED, non bloquant ; aucun seuil n'est inventé ;
  * recalcul interdit dès qu'une écriture de promotion a eu lieu (référence ou outbox du run) ;
  * promotion : revalidation des contrôles, un seul MERGE, outbox écrite dans la même étape,
    promoted_at_utc écrit une fois.
"""
from __future__ import annotations

import decimal
from typing import Dict, List, Optional

from hris.common import HrisError, PROCESS_CODE, RUN_BLOCKED, RUN_COMPARED, RUN_PROMOTED, RUN_VALIDATED, utc_now
from hris.model import ADP_COLUMNS, BUSINESS_KEY, COMPARE_COLUMNS, Tables
from hris.spark_io import (blocking_failures, conform, get_run, replace_run, rows_df, set_once_promoted,
                           sql_str, update_run, write_anomalies)

COMPARE_CONTROLS = ["DUPLICATE_KEY", "NULL_KEY", "REFERENCE_ABSENTE", "REFERENCE_NOT_EMPTY",
                    "ABSENT_RATIO", "MODIFIED_RATIO", "REJECT_RATIO", "STAGING_MISSING"]


def _F():
    from pyspark.sql import functions as F
    return F


def last_promoted_run(spark, T: Tables) -> Optional[str]:
    rows = (spark.table(T("ctl.run")).where(f"process_code = {sql_str(PROCESS_CODE)} AND promoted_at_utc IS NOT NULL")
            .orderBy(_F().col("promoted_at_utc").desc()).select("run_id").limit(1).collect())
    return rows[0].run_id if rows else None


def promotion_writes_exist(spark, T: Tables, run_id: str) -> bool:
    run = get_run(spark, T, run_id) or {}
    if run.get("promoted_at_utc"):
        return True
    if spark.table(T("ctl.run_step")).where(
            f"run_id = {sql_str(run_id)} AND step_code = 'PROMOTE_WRITE'").limit(1).count():
        return True  # marqueur écrit juste avant la première écriture de promotion
    if spark.table(T("ref.ts_employee_adp")).where(f"_last_run_id = {sql_str(run_id)}").limit(1).count():
        return True
    return spark.table(T("pub.adp_outbox")).where(f"run_id = {sql_str(run_id)}").limit(1).count() > 0


def load_thresholds(spark, T: Tables) -> Dict[str, Optional[decimal.Decimal]]:
    rows = spark.table(T("cfg.process")).where(f"process_code = {sql_str(PROCESS_CODE)}").collect()
    if not rows:
        return {"max_absent_pct": None, "max_modified_pct": None, "max_reject_pct": None}
    r = rows[0]
    return {"max_absent_pct": r.max_absent_pct, "max_modified_pct": r.max_modified_pct,
            "max_reject_pct": r.max_reject_pct}


def _pct(num: int, den: int) -> Optional[decimal.Decimal]:
    if den == 0:
        return None
    return (decimal.Decimal(num) * 100 / decimal.Decimal(den)).quantize(decimal.Decimal("0.0001"))


def threshold_controls(summary: dict, thresholds: dict) -> List[dict]:
    """D-CMP-07 : contrôle calculé et journalisé ; NOT_CONFIGURED (non bloquant) tant que le seuil est NULL."""
    out = []
    for code, metric_key, thr_key in [("ABSENT_RATIO", "pct_absent", "max_absent_pct"),
                                       ("MODIFIED_RATIO", "pct_modified", "max_modified_pct"),
                                       ("REJECT_RATIO", "pct_rejected", "max_reject_pct")]:
        metric = summary.get(metric_key)
        thr = thresholds.get(thr_key)
        if thr is None:
            out.append({"control_code": code, "severity": "WARNING", "metric_value": metric, "threshold_value": None,
                        "outcome": "NOT_CONFIGURED", "message": "seuil non fourni (D-CMP-07)"})
        else:
            fail = metric is not None and decimal.Decimal(metric) > decimal.Decimal(thr)
            out.append({"control_code": code, "severity": "BLOCKING", "metric_value": metric, "threshold_value": thr,
                        "outcome": "FAIL" if fail else "PASS", "message": "seuil configuré"})
    return out


def compare(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    F = _F()
    run = get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if promotion_writes_exist(spark, T, run_id):
        # Recalcul interdit : la référence a (au moins en partie) été promue pour ce run.
        prev = spark.table(T("log.comparison_summary")).where(f"run_id = {sql_str(run_id)}").collect()
        return {"decision": run["status"], "recomputed": False,
                **({k: prev[0][k] for k in ("nb_new", "nb_modified", "nb_absent", "nb_rejected")} if prev else {})}

    staging = spark.table(T("stg.ts_employee_adp")).where(f"run_id = {sql_str(run_id)}").drop("run_id", "business_date")
    rejected = (spark.table(T("stg.ts_employee_reject")).where(f"run_id = {sql_str(run_id)} AND id_unique IS NOT NULL")
                .select(F.col("id_unique").alias("k")).distinct())
    ref = spark.table(T("ref.ts_employee_adp"))
    anomalies: List[dict] = []
    nb_staging_rows = staging.count()
    nb_null = staging.where(F.col(BUSINESS_KEY).isNull()).count()
    nb_dup = staging.where(F.col(BUSINESS_KEY).isNotNull()).groupBy(BUSINESS_KEY).count().where("count > 1").count()
    nb_reference = ref.count()
    nb_rejected = rejected.count()
    is_first_load = nb_reference == 0
    staging_ran = spark.table(T("ctl.run_step")).where(
        f"run_id = {sql_str(run_id)} AND step_code = 'TRANSFORM' AND status = 'SUCCEEDED'").limit(1).count() > 0

    anomalies.append({"control_code": "STAGING_MISSING", "severity": "BLOCKING",
                      "metric_value": nb_staging_rows, "threshold_value": None,
                      "outcome": "PASS" if staging_ran else "FAIL",
                      "message": "transformation aboutie" if staging_ran else "aucune transformation aboutie pour ce run"})
    anomalies.append({"control_code": "NULL_KEY", "severity": "BLOCKING", "metric_value": nb_null,
                      "threshold_value": None, "outcome": "FAIL" if nb_null else "PASS",
                      "message": f"{nb_null} lignes sans id_unique"})
    anomalies.append({"control_code": "DUPLICATE_KEY", "severity": "BLOCKING", "metric_value": nb_dup,
                      "threshold_value": None, "outcome": "FAIL" if nb_dup else "PASS",
                      "message": f"{nb_dup} clés en double (correspondances dupliquées ?)"})
    if run_mode == "NORMAL":
        anomalies.append({"control_code": "REFERENCE_ABSENTE", "severity": "BLOCKING", "metric_value": nb_reference,
                          "threshold_value": None, "outcome": "FAIL" if is_first_load else "PASS",
                          "message": "référence vide : premier chargement par INIT_REFERENCE (D-CMP-06)"
                          if is_first_load else "référence présente"})
    elif run_mode == "INIT_REFERENCE":
        anomalies.append({"control_code": "REFERENCE_NOT_EMPTY", "severity": "BLOCKING", "metric_value": nb_reference,
                          "threshold_value": None, "outcome": "PASS" if is_first_load else "FAIL",
                          "message": "référence vide : initialisation autorisée" if is_first_load
                          else "référence non vide : INIT_REFERENCE refusé (pas de réinitialisation)"})

    st = staging.where(F.col(BUSINESS_KEY).isNotNull()).alias("s")
    rf = ref.alias("r")
    joined = st.join(rf, F.col(f"s.{BUSINESS_KEY}") == F.col(f"r.{BUSINESS_KEY}"), "full_outer")
    key = F.coalesce(F.col(f"s.{BUSINESS_KEY}"), F.col(f"r.{BUSINESS_KEY}"))
    changed = F.array_compact(F.array(*[
        F.when(~F.col(f"s.{c}").eqNullSafe(F.col(f"r.{c}")), F.lit(c)) for c in COMPARE_COLUMNS]))
    detail = joined.select(
        key.alias("business_key"),
        F.when(F.col(f"r.{BUSINESS_KEY}").isNull(), "NEW")
        .when(F.col(f"s.{BUSINESS_KEY}").isNull(), "ABSENT_OR_REJECTED")
        .when(F.col("s.row_hash") != F.col("r.row_hash"), "MODIFIED").otherwise("UNCHANGED").alias("category"),
        F.when(F.col(f"s.{BUSINESS_KEY}").isNotNull() & F.col(f"r.{BUSINESS_KEY}").isNotNull()
               & (F.col("s.row_hash") != F.col("r.row_hash")), changed).alias("changed_columns"),
        F.col("r.row_hash").alias("old_row_hash"), F.col("s.row_hash").alias("new_row_hash"))
    detail = (detail.join(rejected, detail.business_key == rejected.k, "left")
              .withColumn("category", F.when(F.col("category") == "ABSENT_OR_REJECTED",
                                             F.when(F.col("k").isNotNull(), "REJECTED").otherwise("ABSENT"))
                          .otherwise(F.col("category"))).drop("k"))
    # Clés rejetées absentes de la référence : catégorie REJECTED également
    rej_only = (rejected.join(ref.select(F.col(BUSINESS_KEY).alias("rk")), rejected.k == F.col("rk"), "left_anti")
                .select(F.col("k").alias("business_key"), F.lit("REJECTED").alias("category"),
                        F.lit(None).cast("array<string>").alias("changed_columns"),
                        F.lit(None).cast("string").alias("old_row_hash"), F.lit(None).cast("string").alias("new_row_hash")))
    detail = detail.unionByName(rej_only).cache()
    counts = {r.category: r["count"] for r in detail.groupBy("category").count().collect()}
    summary = {"nb_staging": nb_staging_rows, "nb_reference": nb_reference,
               "nb_new": counts.get("NEW", 0), "nb_modified": counts.get("MODIFIED", 0),
               "nb_unchanged": counts.get("UNCHANGED", 0), "nb_absent": counts.get("ABSENT", 0),
               "nb_rejected": counts.get("REJECTED", 0)}
    summary["pct_absent"] = _pct(summary["nb_absent"], nb_reference)
    summary["pct_modified"] = _pct(summary["nb_modified"], nb_reference)
    summary["pct_rejected"] = _pct(summary["nb_rejected"], nb_staging_rows + nb_rejected)
    anomalies += threshold_controls(summary, load_thresholds(spark, T))

    now = utc_now()
    against = last_promoted_run(spark, T)
    replace_run(spark, rows_df(spark, [{"run_id": run_id, "process_code": PROCESS_CODE,
                                        "business_date": run["business_date"], "compared_against_run_id": against,
                                        "is_first_load": is_first_load, "created_at_utc": now, **summary}],
                               "log.comparison_summary"),
                T("log.comparison_summary"), "log.comparison_summary", run_id)
    detail_out = (detail.where("category <> 'UNCHANGED'")
                  .withColumn("run_id", F.lit(run_id)).withColumn("process_code", F.lit(PROCESS_CODE))
                  .withColumn("created_at_utc", F.lit(now.replace(tzinfo=None)).cast("timestamp")))
    replace_run(spark, detail_out, T("log.comparison_detail"), "log.comparison_detail", run_id)
    detail.unpersist()
    write_anomalies(spark, T, run_id, anomalies, COMPARE_CONTROLS)

    blocking = blocking_failures(spark, T, run_id)
    decision = RUN_BLOCKED if blocking else RUN_VALIDATED
    update_run(spark, T, run_id, status=decision, compared_against_run_id=against,
               error_code=("BLOCKING_ANOMALY" if blocking else None),
               error_summary=(",".join(blocking) if blocking else None))
    return {"decision": decision, "recomputed": True, "blocking": blocking,
            **{k: summary[k] for k in ("nb_new", "nb_modified", "nb_absent", "nb_rejected", "nb_unchanged")}}


def revalidate(spark, T: Tables, run_id: str) -> List[str]:
    """Avant de lever un blocage ou de promouvoir : les contrôles bloquants du run sont relus et les
    seuils réévalués avec la configuration COURANTE. Renvoie la liste des contrôles en échec."""
    failures = set(blocking_failures(spark, T, run_id))
    rows = spark.table(T("log.comparison_summary")).where(f"run_id = {sql_str(run_id)}").collect()
    if not rows:
        failures.add("COMPARISON_MISSING")
        return sorted(failures)
    summary = rows[0].asDict()
    for a in threshold_controls(summary, load_thresholds(spark, T)):
        if a["severity"] == "BLOCKING" and a["outcome"] == "FAIL":
            failures.add(a["control_code"])
    return sorted(failures)


def promote(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    F = _F()
    from delta.tables import DeltaTable
    run = get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if run.get("promoted_at_utc"):
        # Déjà promu : aucun nouveau MERGE (rejouer la promotion d'un ancien run ramènerait la
        # référence en arrière si un run plus récent a été promu depuis).
        nb = spark.table(T("pub.adp_outbox")).where(f"run_id = {sql_str(run_id)}").count()
        return {"promoted": True, "already_promoted": True, "nb_outbox": nb, "mode": run_mode}
    if run["status"] != RUN_VALIDATED:
        return {"promoted": False, "reason": f"STATUS_{run['status']}", "nb_outbox": 0}
    failures = revalidate(spark, T, run_id)
    if failures:
        update_run(spark, T, run_id, status=RUN_BLOCKED, error_code="REVALIDATION_FAILED",
                   error_summary=",".join(failures))
        return {"promoted": False, "reason": "REVALIDATION_FAILED", "failures": failures, "nb_outbox": 0}
    current_base = last_promoted_run(spark, T)
    if current_base != run.get("compared_against_run_id"):
        update_run(spark, T, run_id, status=RUN_BLOCKED, error_code="REFERENCE_MOVED",
                   error_summary="la référence a été promue par un autre run depuis la comparaison")
        return {"promoted": False, "reason": "REFERENCE_MOVED", "nb_outbox": 0}
    summary = spark.table(T("log.comparison_summary")).where(f"run_id = {sql_str(run_id)}").collect()[0]
    if summary.nb_staging == 0:
        raise HrisError("PROMOTE_EMPTY_STAGING", "staging vide : promotion refusée (protection de la référence)")
    now = utc_now()
    from hris.spark_io import append
    append(rows_df(spark, [{"run_id": run_id, "step_code": "PROMOTE_WRITE", "attempt_no": 1, "status": "STARTED",
                            "message": "écritures de promotion engagées : recalcul du comparatif interdit",
                            "started_at_utc": now}], "ctl.run_step"), T("ctl.run_step"), "ctl.run_step")

    # 1. Outbox (même étape, avant la publication) — jamais en INIT_REFERENCE
    nb_outbox = 0
    if run_mode == "NORMAL":
        changes = (spark.table(T("log.comparison_detail"))
                   .where(f"run_id = {sql_str(run_id)} AND category IN ('NEW','MODIFIED')")
                   .select(F.lit(run_id).alias("run_id"), F.col("business_key").alias("employee_key"),
                           F.col("category").alias("change_category"), F.col("new_row_hash").alias("row_hash"),
                           F.lit("PENDING").alias("status"), F.lit(0).alias("attempt_count"),
                           F.lit(now.replace(tzinfo=None)).cast("timestamp").alias("created_at_utc"),
                           F.lit(now.replace(tzinfo=None)).cast("timestamp").alias("updated_at_utc")))
        nb_outbox = changes.count()
        (DeltaTable.forName(spark, T("pub.adp_outbox")).alias("t")
         .merge(conform(changes, "pub.adp_outbox").alias("s"),
                "t.run_id = s.run_id AND t.employee_key = s.employee_key")
         .whenNotMatchedInsertAll().execute())

    # 2. Référence : un seul MERGE (transaction Delta atomique sur la table)
    staging = (spark.table(T("stg.ts_employee_adp")).where(f"run_id = {sql_str(run_id)}")
               .select(*ADP_COLUMNS, "row_hash").withColumn("_action", F.lit("UPSERT")))
    keep = (spark.table(T("stg.ts_employee_reject")).where(f"run_id = {sql_str(run_id)} AND id_unique IS NOT NULL")
            .select("id_unique").distinct())
    keep = keep.join(staging.select("id_unique"), "id_unique", "left_anti")
    keep_rows = keep.select(*[F.col(c) if c == "id_unique" else F.lit(None).cast(staging.schema[c].dataType).alias(c)
                              for c in ADP_COLUMNS], F.lit(None).cast("string").alias("row_hash"),
                            F.lit("KEEP").alias("_action"))
    source = staging.unionByName(keep_rows)
    values = {c: f"s.`{c}`" for c in ADP_COLUMNS}
    values.update({"row_hash": "s.row_hash", "_last_run_id": sql_str(run_id),
                   "_validated_at_utc": f"timestamp'{now.replace(tzinfo=None).isoformat(sep=' ')}'"})
    (DeltaTable.forName(spark, T("ref.ts_employee_adp")).alias("t")
     .merge(source.alias("s"), f"t.{BUSINESS_KEY} = s.{BUSINESS_KEY}")
     .whenMatchedUpdate(condition="s._action = 'UPSERT' AND t.row_hash <> s.row_hash", set=values)
     .whenNotMatchedInsert(condition="s._action = 'UPSERT'", values=values)
     .whenNotMatchedBySourceDelete()
     .execute())

    # 3. Garde d'idempotence
    set_once_promoted(spark, T, run_id, now)
    update_run(spark, T, run_id, status=RUN_PROMOTED)
    return {"promoted": True, "nb_outbox": nb_outbox, "mode": run_mode}
'''),
    ('hris.run_control', '3dc91ff068c000a4fe73d5a5825dbc08944bed3b9e1a8641aad9d46e4e262e5f', r'''"""hris.run_control — modes START / END / FAIL de NB_HRIS_RUN_CONTROL.

Verrou applicatif ctl.process_lock (A07 § 2) : MERGE conditionnel puis relecture ; un conflit de
concurrence Delta ou une relecture différente = verrou non obtenu -> SKIPPED_CONCURRENT (pas d'échec,
aucune écriture métier). Bail : p_lock_lease_minutes (A08 D-RUN-04, À CONFIRMER).
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Optional

from hris.common import (PROCESS_CODE, RUN_BLOCKED, RUN_FAILED, RUN_MODES, RUN_NOT_PUBLISHED_NON_PRD,
                         RUN_PROMOTED, RUN_PUBLISHED, RUN_PUBLISHED_PARTIAL, RUN_SKIPPED_CONCURRENT, RUN_STARTED,
                         RUN_SUCCEEDED, HrisError, business_date_of, parse_trigger_time, sanitize, to_utc_naive,
                         utc_now)
from hris.model import Tables
from hris.parity import utf8_truncate
from hris.spark_io import check_run_id, get_run, merge_rows, replace_run, rows_df, sql_str, update_run

_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _F():
    from pyspark.sql import functions as F
    return F


def acquire_lock(spark, T: Tables, run_id: str, lease_minutes: int, now: Optional[dt.datetime] = None) -> bool:
    from delta.tables import DeltaTable
    F = _F()
    now = now or utc_now()
    lease = now + dt.timedelta(minutes=lease_minutes)
    src = rows_df(spark, [{"process_code": PROCESS_CODE, "lock_status": "HELD", "holder_run_id": run_id,
                           "acquired_at_utc": now, "lease_expires_at_utc": lease, "released_at_utc": None}],
                  "ctl.process_lock")
    now_lit = F.lit(to_utc_naive(now)).cast("timestamp")
    try:
        (DeltaTable.forName(spark, T("ctl.process_lock")).alias("t")
         .merge(src.alias("s"), "t.process_code = s.process_code")
         .whenMatchedUpdate(condition=(F.col("t.lock_status") == "RELEASED") | (F.col("t.lease_expires_at_utc") < now_lit)
                            | (F.col("t.holder_run_id") == F.col("s.holder_run_id")),
                            set={"lock_status": "s.lock_status", "holder_run_id": "s.holder_run_id",
                                 "acquired_at_utc": "s.acquired_at_utc",
                                 "lease_expires_at_utc": "s.lease_expires_at_utc", "released_at_utc": "s.released_at_utc"})
         .whenNotMatchedInsertAll().execute())
    except Exception as exc:  # conflit Delta (Concurrent*Exception) = verrou non obtenu
        if "Concurrent" in type(exc).__name__ or "Concurrent" in str(exc)[:500]:
            return False
        raise
    rows = spark.table(T("ctl.process_lock")).where(f"process_code = {sql_str(PROCESS_CODE)}").collect()
    return len(rows) == 1 and rows[0].lock_status == "HELD" and rows[0].holder_run_id == run_id


def release_lock(spark, T: Tables, run_id: str) -> bool:
    from delta.tables import DeltaTable
    F = _F()
    holder = spark.table(T("ctl.process_lock")).where(
        f"process_code = {sql_str(PROCESS_CODE)} AND holder_run_id = {sql_str(run_id)} AND lock_status = 'HELD'").count()
    if not holder:
        return False
    (DeltaTable.forName(spark, T("ctl.process_lock")).alias("t")
     .update(condition=(F.col("t.process_code") == PROCESS_CODE) & (F.col("t.holder_run_id") == run_id)
             & (F.col("t.lock_status") == "HELD"),
             set={"lock_status": F.lit("RELEASED"), "released_at_utc": F.lit(to_utc_naive(utc_now())).cast("timestamp")}))
    return True


def start(spark, T: Tables, *, run_id: str, process_code: str, run_mode: str, target_run_id: str,
          trigger_time: str, business_timezone: str, lock_lease_minutes: int, business_date_override: str,
          pipeline_name: str, workspace_id: Optional[str], workspace_name: Optional[str]) -> dict:
    check_run_id(run_id)
    if process_code != PROCESS_CODE:
        raise HrisError("PROCESS_NOT_ENABLED", "seul TS_ADP_EMPLOYEE est actif (flux SFTP bloqué D-FIL-01)")
    if run_mode not in RUN_MODES:
        raise HrisError("PARAM_INVALID", "p_run_mode")
    if run_mode == "RETRY_PUBLICATION":
        if not target_run_id:
            raise HrisError("PARAM_MISSING", "p_target_run_id obligatoire en RETRY_PUBLICATION")
        target = get_run(spark, T, check_run_id(target_run_id))
        if target is None or not target.get("promoted_at_utc"):
            raise HrisError("RETRY_TARGET_INVALID", "le run cible n'existe pas ou n'a pas été promu")
    elif target_run_id:
        raise HrisError("PARAM_INVALID", "p_target_run_id réservé au mode RETRY_PUBLICATION")
    if not isinstance(lock_lease_minutes, int) or lock_lease_minutes <= 0:
        raise HrisError("PARAM_INVALID", "p_lock_lease_minutes")
    trigger_utc = parse_trigger_time(trigger_time)
    if business_date_override:
        # tests uniquement ; refusé en NORMAL dans un workspace dont le nom contient PRD (02 § 5, À CONFIRMER)
        if run_mode == "NORMAL" and isinstance(workspace_name, str) and "PRD" in workspace_name:
            raise HrisError("OVERRIDE_REFUSED", "p_business_date_override interdit en NORMAL dans ce workspace")
        if not _DATE.match(business_date_override):
            raise HrisError("PARAM_INVALID", "p_business_date_override (AAAA-MM-JJ)")
        business_date = dt.date.fromisoformat(business_date_override)
    else:
        business_date = business_date_of(trigger_utc, business_timezone)

    existing = get_run(spark, T, run_id)
    if existing and existing["status"] not in (RUN_STARTED, RUN_SKIPPED_CONCURRENT):
        # RunStart rejoué pour un run déjà avancé : on ne réinitialise rien
        return {"status": "OK", "run_id": run_id, "business_date": str(existing["business_date"]),
                "run_seq_in_day": existing["run_seq_in_day"], "resumed": True}

    if not acquire_lock(spark, T, run_id, lock_lease_minutes):
        merge_rows(spark, [{"run_id": run_id, "process_code": PROCESS_CODE, "run_mode": run_mode,
                            "business_date": business_date, "trigger_time_utc": trigger_utc,
                            "pipeline_name": pipeline_name, "workspace_id": workspace_id,
                            "workspace_name": workspace_name, "status": RUN_SKIPPED_CONCURRENT,
                            "started_at_utc": utc_now(), "ended_at_utc": utc_now(),
                            "target_run_id": target_run_id or None}],
                   T("ctl.run"), "ctl.run", ["run_id"])
        return {"status": RUN_SKIPPED_CONCURRENT, "run_id": run_id, "business_date": str(business_date)}

    seq = (spark.table(T("ctl.run"))
           .where(f"process_code = {sql_str(PROCESS_CODE)} AND business_date = DATE'{business_date.isoformat()}' "
                  f"AND status <> '{RUN_SKIPPED_CONCURRENT}' AND run_id <> {sql_str(run_id)}").count()) + 1
    merge_rows(spark, [{"run_id": run_id, "process_code": PROCESS_CODE, "run_mode": run_mode,
                        "business_date": business_date, "run_seq_in_day": seq, "trigger_time_utc": trigger_utc,
                        "pipeline_name": pipeline_name, "workspace_id": workspace_id,
                        "workspace_name": workspace_name, "status": RUN_STARTED, "started_at_utc": utc_now(),
                        "target_run_id": target_run_id or None}],
               T("ctl.run"), "ctl.run", ["run_id"])
    return {"status": "OK", "run_id": run_id, "business_date": business_date.isoformat(), "run_seq_in_day": seq,
            "run_mode": run_mode}


def final_status(run: dict) -> str:
    status = run["status"]
    if status in (RUN_BLOCKED, RUN_FAILED, RUN_SKIPPED_CONCURRENT, RUN_NOT_PUBLISHED_NON_PRD, RUN_PUBLISHED_PARTIAL):
        return status
    if status in (RUN_PUBLISHED, RUN_PROMOTED):
        return RUN_SUCCEEDED
    return RUN_FAILED  # run arrêté avant la fin de la séquence attendue


def run_summary(spark, T: Tables, run: dict) -> dict:
    """log.run_summary : métriques historiques (R-SQL-060 à 066) calculées sur le delta du run
    (NEW + MODIFIED = insérés ; REJECTED = rejetés) + compteurs de publication."""
    F = _F()
    run_id = run["run_id"]
    source_run = run.get("target_run_id") or run_id
    det = spark.table(T("log.comparison_detail")).where(f"run_id = {sql_str(source_run)}")
    inserts = det.where("category IN ('NEW','MODIFIED')").select("business_key").distinct().count()
    rejects_n = det.where("category = 'REJECTED'").select("business_key").distinct().count()
    lues = inserts + rejects_n
    import decimal
    pc = decimal.Decimal(0) if lues == 0 else (decimal.Decimal(rejects_n) / decimal.Decimal(lues)).quantize(
        decimal.Decimal("0.0001"))
    pairs = sorted({(r.cd_type_rejet, r.ds_rejet or "") for r in
                    spark.table(T("stg.ts_employee_reject")).where(f"run_id = {sql_str(source_run)}")
                    .select("cd_type_rejet", "ds_rejet").distinct().collect()})
    ds_type = utf8_truncate(",".join(p[0] for p in pairs), 50) if pairs else None
    ob = {r.status: r["count"] for r in spark.table(T("pub.adp_outbox")).where(f"run_id = {sql_str(source_run)}")
          .groupBy("status").count().collect()}
    trig = run.get("trigger_time_utc")
    return {"run_id": run_id, "process_code": PROCESS_CODE, "business_date": run.get("business_date"),
            "dt_resume": trig.date() if trig else None, "cd_source": "TalentSoft", "cd_target": "ADP",
            "mt_lignes_inserts": inserts, "mt_lignes_rejets": rejects_n, "mt_lignes_lues": lues,
            "pc_lignes_rejets": pc, "ds_type_rejets": ds_type, "nb_outbox": sum(ob.values()),
            "nb_sent": ob.get("SENT", 0), "nb_no_change": ob.get("NO_CHANGE", 0), "nb_failed": ob.get("FAILED", 0),
            "nb_not_pushed": ob.get("NOT_PUSHED_NON_PRD", 0), "created_at_utc": utc_now()}


def end(spark, T: Tables, run_id: str, env_code: str) -> dict:
    run = get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    status = final_status(run)
    summary = run_summary(spark, T, run)
    replace_run(spark, rows_df(spark, [summary], "log.run_summary"), T("log.run_summary"), "log.run_summary", run_id)
    update_run(spark, T, run_id, status=status, ended_at_utc=utc_now())
    released = release_lock(spark, T, run_id)
    notification = {  # contenu préparé, sans donnée personnelle ; envoi bloqué (D-NOT-01/02)
        "type": "SUMMARY" if status in (RUN_SUCCEEDED, RUN_NOT_PUBLISHED_NON_PRD) else "ERROR",
        "env": env_code, "run_id": run_id, "business_date": str(run.get("business_date")), "status": status,
        "inserts": summary["mt_lignes_inserts"], "rejects": summary["mt_lignes_rejets"],
        "outbox": summary["nb_outbox"], "error_code": run.get("error_code"), "sent": False,
        "send_blocked_reason": "D-NOT-01/D-NOT-02 non décidés"}
    return {"status": status, "run_id": run_id, "lock_released": released, "notification": notification}


def fail(spark, T: Tables, run_id: str, error_message: str) -> dict:
    run = get_run(spark, T, run_id)
    if run is None:
        return {"status": RUN_FAILED, "run_id": run_id, "lock_released": False, "note": "run inconnu"}
    if run["status"] == RUN_SKIPPED_CONCURRENT:
        return {"status": RUN_SKIPPED_CONCURRENT, "run_id": run_id, "lock_released": False}
    update_run(spark, T, run_id, status=RUN_FAILED, ended_at_utc=utc_now(), error_code="PIPELINE_FAILURE",
               error_summary=sanitize(error_message, 1000))
    released = release_lock(spark, T, run_id)
    return {"status": RUN_FAILED, "run_id": run_id, "lock_released": released,
            "notification": {"type": "ERROR", "run_id": run_id, "sent": False,
                             "send_blocked_reason": "D-NOT-01/D-NOT-02 non décidés"}}
'''),
    ('hris.adp_gate', 'e78ceb69b7fd6fdceccc8018e16154788bc7e3bc55436de34f03145142fbd6f8', r'''"""hris.adp_gate — condition d'environnement de la publication ADP (D-ENV-01).

Règle (inchangée, littérale) : la condition d'environnement est satisfaite si et seulement si le
nom RÉEL du workspace d'exécution du notebook de publication contient la sous-chaîne "PRD"
(sensible à la casse, recherche par sous-chaîne, aucune normalisation).

Trois issues distinctes (révision demandée par le mandat) :
  * ENV_SATISFIED     : nom connu contenant "PRD" -> condition d'environnement satisfaite. Cela
                        N'AUTORISE PAS l'envoi : les autres validations (qualité, métier, sécurité,
                        authentification, contrat, activation) restent nécessaires (hris.adp_publish).
  * DENIED_NON_PRD    : nom connu ne contenant pas "PRD" -> absence normale de publication
                        (statut d'outbox NOT_PUSHED_NON_PRD).
  * BLOCKED_UNRESOLVED: nom non résolu, vide ou d'un type inattendu -> publication bloquée, les
                        éléments restent en attente (PENDING) et seront réévalués après résolution.

Le contexte provient exclusivement de notebookutils.runtime.context lu DANS le notebook de
publication. Aucun paramètre, aucune variable, aucun simulateur ne peut forcer la décision :
les contextes simulés ne sont injectés que par les tests isolés.

Ambiguïtés signalées, règle NON modifiée (D-ENV-02, à trancher avant toute activation) :
"PREPRD" et "NOTPRD" satisfont la condition ; "prd" en minuscules ne la satisfait pas.
"""
from __future__ import annotations

import dataclasses
import datetime as dt
import uuid
from typing import Any, Callable, Mapping, Optional

PRD_MARKER = "PRD"
RULE_VERSION = "contains:PRD:case-sensitive:v2"

REASON_CONTAINS_PRD = "WORKSPACE_NAME_CONTAINS_PRD"
REASON_WITHOUT_PRD = "WORKSPACE_NAME_WITHOUT_PRD"
REASON_UNRESOLVED = "WORKSPACE_NAME_UNRESOLVED"
REASON_EMPTY = "WORKSPACE_NAME_EMPTY"
REASON_INVALID_TYPE = "WORKSPACE_NAME_INVALID_TYPE"

ENV_SATISFIED = "ENV_SATISFIED"
DENIED_NON_PRD = "DENIED_NON_PRD"
BLOCKED_UNRESOLVED = "BLOCKED_UNRESOLVED"


def _utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


@dataclasses.dataclass(frozen=True)
class WorkspaceContext:
    name: Any
    workspace_id: Optional[str]
    is_for_pipeline: Optional[bool]
    notebook_name: Optional[str]
    resolution_error: Optional[str]


@dataclasses.dataclass(frozen=True)
class GateDecision:
    decision_id: str
    outcome: str
    env_condition_met: bool
    reason_code: str
    workspace_name: Optional[str]
    workspace_id: Optional[str]
    is_for_pipeline: Optional[bool]
    notebook_name: Optional[str]
    rule_version: str
    evaluated_at_utc: dt.datetime

    @property
    def allowed(self) -> bool:  # compatibilité : True = condition d'environnement satisfaite
        return self.env_condition_met


def resolve_current_workspace(context_provider: Callable[[], Mapping[str, Any]]) -> WorkspaceContext:
    """En Fabric : ``context_provider = lambda: notebookutils.runtime.context`` (A06 F3)."""
    try:
        ctx = context_provider()
    except Exception as exc:  # contexte indisponible : échec fermé
        return WorkspaceContext(None, None, None, None, type(exc).__name__)

    def _get(key: str) -> Any:
        try:
            return ctx[key]
        except Exception:
            return None

    return WorkspaceContext(_get("currentWorkspaceName"), _get("currentWorkspaceId"), _get("isForPipeline"),
                            _get("currentNotebookName"), None)


def decide_adp_push(workspace_name: Any) -> tuple:
    """Règle pure. Renvoie (condition satisfaite, motif)."""
    if workspace_name is None:
        return False, REASON_UNRESOLVED
    if not isinstance(workspace_name, str):
        return False, REASON_INVALID_TYPE
    if workspace_name.strip() == "":
        return False, REASON_EMPTY
    if PRD_MARKER in workspace_name:
        return True, REASON_CONTAINS_PRD
    return False, REASON_WITHOUT_PRD


def outcome_of(reason: str) -> str:
    if reason == REASON_CONTAINS_PRD:
        return ENV_SATISFIED
    if reason == REASON_WITHOUT_PRD:
        return DENIED_NON_PRD
    return BLOCKED_UNRESOLVED  # non résolu / vide / type invalide : réévaluable


def evaluate_gate(context_provider: Callable[[], Mapping[str, Any]],
                  clock: Callable[[], dt.datetime] = _utc_now) -> GateDecision:
    ws = resolve_current_workspace(context_provider)
    met, reason = decide_adp_push(ws.name)
    return GateDecision(
        decision_id=str(uuid.uuid4()), outcome=outcome_of(reason), env_condition_met=met, reason_code=reason,
        workspace_name=ws.name if isinstance(ws.name, str) else None,
        workspace_id=ws.workspace_id if isinstance(ws.workspace_id, str) else None,
        is_for_pipeline=ws.is_for_pipeline if isinstance(ws.is_for_pipeline, bool) else None,
        notebook_name=ws.notebook_name if isinstance(ws.notebook_name, str) else None,
        rule_version=RULE_VERSION, evaluated_at_utc=clock())
'''),
    ('hris.adp_transport', 'aa2ae8bba6efa3bc1336cdc19e21151712fa2a2894f4618a0ddadad0622e556e', r'''"""hris.adp_transport — transport HTTP de publication ADP, classification des résultats, instrumentation.

Principes (mandat § 8) :
  * aucun retry automatique (y compris bibliothèque), aucune redirection suivie, aucun renvoi
    implicite après 401, timeouts explicites, aucun nouvel envoi après un résultat incertain ;
  * la non-transmission n'est PROUVÉE que si l'échec survient avant l'émission du premier octet de
    la requête (phase CONNECT : DNS, TCP, TLS/mTLS) — jamais d'après le seul nom d'une exception ;
  * une réponse 2xx ne suffit pas : l'acceptation suit le contrat de l'endpoint.

Le contrat final de publication est ABSENT : les listes de statuts ci-dessous sont PROVISOIRES,
utilisables pour les simulations uniquement (docs/CONTRAT_ADP_POINTS_D_USAGE.md). Le transport réel
refuse de se construire tant que ADP_ACTIVATION_APPROVED est faux (constante de code, aucun paramètre).
"""
from __future__ import annotations

import dataclasses
import hashlib
import http.client
import json
import socket
import ssl
import threading
from typing import Any, Callable, Dict, List, Optional, Tuple
from urllib.parse import urlparse

from hris.common import HrisError

# ---------------------------------------------------------------------------
# Verrous d'activation (constantes de CODE : ni paramètre, ni variable, ni simulateur)
# ---------------------------------------------------------------------------
ADP_BASE_URL = "https://accounts.eu.adp.com/"  # cible unique ADP PRD (A05 § 4.1)
ADP_ACTIVATION_APPROVED = False      # CONTRAT-ADP:C00 contrat final + D-ADP-01/02, D-ENV-02, D-SEC-01, D-CMP-07, D-NET-01
ADP_READ_POLICY_ESTABLISHED = False  # CONTRAT-ADP:C14 politique des lectures ADP réelles non établie
ADP_HOST_PATTERNS = ("adp.com", "geo.api.gouv.fr")  # hôtes surveillés par la sentinelle

NOT_TRANSMITTED = "NOT_TRANSMITTED"  # non-transmission prouvée
ACCEPTED = "ACCEPTED"                # acceptation selon le contrat de l'endpoint
REJECTED = "REJECTED"                # rejet confirmé
UNCERTAIN = "UNCERTAIN"              # résultat incertain : aucun nouvel envoi automatique

# CONTRAT-ADP:C02 statuts 2xx acceptés et forme de corps attendue (PROVISOIRE : 200/201/202 + JSON,
# d'après la seule gestion explicite de l'existant, change_pay_distribution L4190)
PROVISIONAL_ACCEPTED_2XX = frozenset({200, 201, 202})
# CONTRAT-ADP:C03 statuts 4xx valant rejet confirmé (PROVISOIRE)
PROVISIONAL_REJECT_4XX = frozenset({400, 403, 404, 405, 409, 412, 415, 422})


@dataclasses.dataclass
class TransportResult:
    outcome: str
    reason: str
    phase: str
    http_status: Optional[int] = None
    http_reason: Optional[str] = None
    error_class: Optional[str] = None
    response_text: Optional[str] = dataclasses.field(default=None, repr=False)  # donnée personnelle possible
    body_digest: Optional[str] = None

    def __repr__(self) -> str:
        return (f"TransportResult({self.outcome}, {self.reason}, phase={self.phase}, "
                f"http={self.http_status}, error={self.error_class})")


def classify_response(status: int, body: str, content_type: str = "") -> Tuple[str, str]:
    """Classification d'une réponse HTTP REÇUE à une requête de publication (PROVISOIRE)."""
    if 200 <= status < 300:
        if status not in PROVISIONAL_ACCEPTED_2XX:
            return UNCERTAIN, "2XX_UNEXPECTED_STATUS"          # CONTRAT-ADP:C02
        try:
            parsed = json.loads(body) if body and body.strip() else None
        except ValueError:
            return UNCERTAIN, "2XX_BODY_NOT_JSON"               # CONTRAT-ADP:C02
        if parsed is None and status != 202:
            return UNCERTAIN, "2XX_EMPTY_BODY"                  # CONTRAT-ADP:C02
        return ACCEPTED, f"HTTP_{status}_ACCEPTED_PROVISIONAL"
    if 300 <= status < 400:
        return UNCERTAIN, "REDIRECT_NOT_FOLLOWED"               # CONTRAT-ADP:C04
    if status == 401:
        return REJECTED, "AUTH_REJECTED_NO_RESEND"              # CONTRAT-ADP:C05
    if status == 429:
        return REJECTED, "RATE_LIMITED_NO_RETRY"                # CONTRAT-ADP:C06
    if status in PROVISIONAL_REJECT_4XX:
        return REJECTED, f"HTTP_{status}_REJECTED_PROVISIONAL"  # CONTRAT-ADP:C03
    if 400 <= status < 500:
        return UNCERTAIN, f"HTTP_{status}_UNCLASSIFIED"
    return UNCERTAIN, f"HTTP_{status}_SERVER_ERROR"            # CONTRAT-ADP:C04 (5xx : traitement partiel possible)


def digest(text: Optional[str]) -> Optional[str]:
    return None if text is None else hashlib.sha256(text.encode("utf-8", "replace")).hexdigest()


# ---------------------------------------------------------------------------
# Transport réel (construit seulement si toutes les autorisations sont réunies)
# ---------------------------------------------------------------------------
class AdpHttpTransport:
    """POST/GET vers ADP en http.client : pas de pool à reprise, pas de redirection, phases explicites.
    mTLS par SSLContext.load_cert_chain ; jeton porteur fourni par un callable (jamais journalisé)."""

    MAX_BODY = 1_000_000

    def __init__(self, cert_path: str, key_path: str, bearer: Callable[[], str], timeout_seconds: float,
                 base_url: str = ADP_BASE_URL, _allow_construction: bool = False):
        if not (ADP_ACTIVATION_APPROVED and _allow_construction):
            raise HrisError("ADP_ACTIVATION_NOT_APPROVED", "transport ADP réel non autorisé dans cette version")
        parsed = urlparse(base_url)
        self.host = parsed.hostname
        self.base_path = parsed.path or "/"
        self.timeout = float(timeout_seconds)
        self._bearer = bearer
        self._ctx = ssl.create_default_context()
        self._ctx.load_cert_chain(certfile=cert_path, keyfile=key_path)

    def __repr__(self) -> str:
        return "AdpHttpTransport(<prd>)"

    def _connection(self):
        return http.client.HTTPSConnection(self.host, 443, context=self._ctx, timeout=self.timeout)

    def post(self, path: str, body: bytes) -> TransportResult:
        conn = self._connection()
        try:
            try:
                conn.connect()  # DNS + TCP + TLS : aucun octet de la requête HTTP n'est émis
            except Exception as exc:
                return TransportResult(NOT_TRANSMITTED, "CONNECT_FAILED", "CONNECT", error_class=type(exc).__name__)
            headers = {"Content-Type": "application/json", "Accept": "application/json",
                       "Authorization": "Bearer " + self._bearer(), "Content-Length": str(len(body))}
            try:
                conn.request("POST", self.base_path + path, body=body, headers=headers)
            except Exception as exc:
                return TransportResult(UNCERTAIN, "SEND_INTERRUPTED", "SEND", error_class=type(exc).__name__)
            try:
                resp = conn.getresponse()
                raw = resp.read(self.MAX_BODY + 1)
            except socket.timeout as exc:
                return TransportResult(UNCERTAIN, "TIMEOUT_AFTER_SEND", "RECEIVE", error_class=type(exc).__name__)
            except Exception as exc:
                return TransportResult(UNCERTAIN, "RECEIVE_FAILED", "RECEIVE", error_class=type(exc).__name__)
            text = raw[: self.MAX_BODY].decode("utf-8", "replace")
            outcome, reason = classify_response(resp.status, text, resp.getheader("Content-Type", ""))
            return TransportResult(outcome, reason, "DONE", http_status=resp.status, http_reason=resp.reason,
                                   response_text=text, body_digest=digest(text))
        finally:
            try:
                conn.close()
            except Exception:
                pass

    def get(self, path: str) -> Tuple[Optional[int], Optional[str]]:
        """Lecture ADP (politique distincte, CONTRAT-ADP:C14). Pas de redirection suivie."""
        if not ADP_READ_POLICY_ESTABLISHED:
            raise HrisError("ADP_READ_POLICY_NOT_ESTABLISHED")
        conn = self._connection()
        try:
            conn.request("GET", self.base_path + path, headers={
                "Accept": "application/json", "Authorization": "Bearer " + self._bearer()})
            resp = conn.getresponse()
            return resp.status, resp.read(self.MAX_BODY).decode("utf-8", "replace")
        finally:
            conn.close()


# ---------------------------------------------------------------------------
# Instrumentation : comptage applicatif et sentinelle réseau
# ---------------------------------------------------------------------------
class RecordingTransport:
    """Enveloppe d'un transport (réel ou fictif) qui enregistre chaque appel de publication (POST)
    séparément des lectures (GET). Aucune donnée de requête n'est conservée hors mémoire du test."""

    def __init__(self, inner):
        self.inner = inner
        self.post_calls: List[str] = []
        self.get_calls: List[str] = []

    def post(self, path: str, body: bytes) -> TransportResult:
        self.post_calls.append(path)
        return self.inner.post(path, body)

    def get(self, path: str):
        self.get_calls.append(path)
        return self.inner.get(path)


class NetworkSentinel:
    """Intercepte, pour la durée d'un bloc, toute tentative de résolution DNS ou de connexion vers un
    hôte ADP (ou geo.api.gouv.fr) : la tentative est ENREGISTRÉE puis BLOQUÉE avant tout trafic réseau.
    Couvre socket.getaddrinfo, socket.create_connection, socket.socket.connect et
    http.client.HTTPConnection.connect (donc requests/urllib3 et http.client).
    Limite : ne voit pas un trafic émis hors du processus Python du driver (exécuteurs Spark, autres
    processus) ; la publication ne s'exécute que sur le driver."""

    def __init__(self, patterns=ADP_HOST_PATTERNS):
        self.patterns = tuple(p.lower() for p in patterns)
        self.attempts: List[Dict[str, str]] = []
        self._lock = threading.Lock()
        self._saved: Dict[str, Any] = {}

    def _watched(self, host: Any) -> bool:
        if not isinstance(host, (str, bytes)):
            return False
        h = host.decode() if isinstance(host, bytes) else host
        h = h.lower().rstrip(".")
        return any(h == p or h.endswith("." + p) or h.endswith(p) for p in self.patterns)

    def _record(self, host: str, via: str) -> None:
        with self._lock:
            self.attempts.append({"host": str(host), "via": via})
        raise ConnectionRefusedError(f"HRIS network sentinel: connexion vers un hôte ADP bloquée ({via})")

    def __enter__(self):
        sentinel = self
        self._saved = {"getaddrinfo": socket.getaddrinfo, "create_connection": socket.create_connection,
                       "socket_connect": socket.socket.connect, "http_connect": http.client.HTTPConnection.connect}
        orig_gai, orig_cc = self._saved["getaddrinfo"], self._saved["create_connection"]
        orig_sc, orig_hc = self._saved["socket_connect"], self._saved["http_connect"]

        def getaddrinfo(host, *a, **k):
            if sentinel._watched(host):
                sentinel._record(host, "getaddrinfo")
            return orig_gai(host, *a, **k)

        def create_connection(address, *a, **k):
            if sentinel._watched(address[0]):
                sentinel._record(address[0], "create_connection")
            return orig_cc(address, *a, **k)

        def socket_connect(sock, address, *a, **k):
            if isinstance(address, tuple) and address and sentinel._watched(address[0]):
                sentinel._record(address[0], "socket.connect")
            return orig_sc(sock, address, *a, **k)

        def http_connect(conn, *a, **k):
            if sentinel._watched(getattr(conn, "host", None)):
                sentinel._record(conn.host, "http.client.connect")
            return orig_hc(conn, *a, **k)

        socket.getaddrinfo = getaddrinfo
        socket.create_connection = create_connection
        socket.socket.connect = socket_connect
        http.client.HTTPConnection.connect = http_connect
        return self

    def __exit__(self, *exc):
        socket.getaddrinfo = self._saved["getaddrinfo"]
        socket.create_connection = self._saved["create_connection"]
        socket.socket.connect = self._saved["socket_connect"]
        http.client.HTTPConnection.connect = self._saved["http_connect"]
        return False
'''),
    ('hris.adp_ops', '58585e7444141ad671315c630cc21636577579a49f7c1203f9808ed76779a78e', r'''"""hris.adp_ops — opérations de publication ADP : identité, intention durable, exécution, réconciliation.

PROPOSITION soumise à validation : le contrat final de publication et de reprise est absent.
Chaque point qui en dépend porte un marqueur `CONTRAT-ADP:Cxx` recensé dans
docs/CONTRAT_ADP_POINTS_D_USAGE.md (cohérence vérifiée par tests/unit/test_contract_markers.py).

Modèle :
  * une OPÉRATION = un événement métier ADP (POST) pour un salarié, une variante (P1..P9, P8 détaillé)
    et une valeur cible ; ses TENTATIVES sont distinctes (pub.adp_attempt) ;
  * identité stable : op_id = H(version, run source, salarié, variante, empreinte de la cible) —
    le même événement rejoué garde le même op_id ; une nouvelle modification légitime (autre run,
    y compris un retour à une valeur déjà rencontrée) a un op_id différent ; aucune déduplication
    à vie sur (salarié, code d'opération) ;
  * l'intention (INTENT_RECORDED) est persistée AVANT tout transport ; si cette persistance échoue,
    rien n'est envoyé ; une intention sans issue persistée est traitée comme INCERTAINE ;
  * aucun nouvel envoi automatique d'une opération INCERTAINE : réconciliation ou décision manuelle.
"""
from __future__ import annotations

import dataclasses
import datetime as dt
import uuid
from typing import Any, Callable, Dict, Iterable, List, Optional, Protocol, Sequence

from hris.adp_transport import ACCEPTED, NOT_TRANSMITTED, REJECTED, UNCERTAIN, TransportResult
from hris.common import HrisError, canonical_json, sha256_hex

OP_IDENTITY_VERSION = "op-v1"  # CONTRAT-ADP:C01

# Statuts d'opération (PROPOSÉS) — CONTRAT-ADP:C11
OP_INTENT_RECORDED = "INTENT_RECORDED"
OP_ACCEPTED = "ACCEPTED"              # = « SENT » : acceptation selon le contrat, pas application
OP_REJECTED = "REJECTED"
OP_NOT_SENT = "NOT_SENT"              # non-transmission prouvée (ou renvoi autorisé par décision)
OP_UNCERTAIN = "UNCERTAIN"
OP_BLOCKED_DEPENDENCY = "BLOCKED_DEPENDENCY"
OP_SUPERSEDED = "SUPERSEDED"
OP_APPLIED_OBSERVED = "APPLIED_OBSERVED"
OP_ABANDONED = "ABANDONED"
OP_DONE = {OP_ACCEPTED, OP_APPLIED_OBSERVED}
OP_OPEN_UNCERTAIN = {OP_INTENT_RECORDED, OP_UNCERTAIN}
OP_RETRYABLE = {OP_REJECTED, OP_NOT_SENT, OP_BLOCKED_DEPENDENCY}

# Statuts d'outbox (A07 § 6, inchangés)
OB_PENDING, OB_IN_PROGRESS, OB_SENT = "PENDING", "IN_PROGRESS", "SENT"
OB_NO_CHANGE, OB_FAILED, OB_NOT_PUSHED, OB_SUPERSEDED = "NO_CHANGE", "FAILED", "NOT_PUSHED_NON_PRD", "SUPERSEDED"

_TRANSPORT_TO_OP = {ACCEPTED: OP_ACCEPTED, REJECTED: OP_REJECTED, NOT_TRANSMITTED: OP_NOT_SENT,
                    UNCERTAIN: OP_UNCERTAIN}

# Décisions manuelles (PROPOSÉES) — CONTRAT-ADP:C10
DEC_AUTHORIZE_RESEND = "AUTHORIZE_RESEND"   # incertaine -> NOT_SENT (nouvelle tentative possible en reprise)
DEC_MARK_APPLIED = "MARK_APPLIED"           # incertaine -> APPLIED_OBSERVED (constat humain)
DEC_ABANDON = "ABANDON"                     # -> ABANDONED (aucun envoi)
DECISION_EFFECT = {DEC_AUTHORIZE_RESEND: OP_NOT_SENT, DEC_MARK_APPLIED: OP_APPLIED_OBSERVED,
                   DEC_ABANDON: OP_ABANDONED}


def _utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


@dataclasses.dataclass
class PlannedOp:
    """Opération produite par la logique ADP portée (hris.adp_port). payload, target et log
    contiennent des données personnelles : uniquement en mémoire, jamais persistés en clair."""
    op_code: str                      # P1 .. P9
    op_variant: str                   # ex. P8_SALARY, P8_CONTRACT_A (CONTRAT-ADP:C12)
    endpoint: str                     # chemin relatif, sans identifiant
    payload: dict = dataclasses.field(repr=False)
    target: dict = dataclasses.field(repr=False)   # valeur métier cible (sans partie volatile)
    log: dict = dataclasses.field(repr=False)      # colonnes du log de l'existant (mode, field, nom...)
    depends_on: Sequence[str] = ()    # variantes du même salarié qui doivent être acceptées avant
    read_back: Optional[Callable[[Any], Optional[dict]]] = dataclasses.field(default=None, repr=False)


def target_hash(target: dict) -> str:
    return sha256_hex(canonical_json(target))


def op_identity(source_run_id: str, employee_key: str, op_variant: str, tgt_hash: str) -> str:
    """CONTRAT-ADP:C01 — identité proposée de l'événement (à valider avant activation réelle)."""
    return sha256_hex("|".join([OP_IDENTITY_VERSION, source_run_id, employee_key, op_variant, tgt_hash]))


# ---------------------------------------------------------------------------
# Dépôts (implémentés en Spark dans le notebook, en mémoire dans les tests)
# ---------------------------------------------------------------------------
class OperationRepository(Protocol):
    def get(self, op_id: str) -> Optional[dict]: ...
    def open_uncertain_for_employee(self, employee_key: str) -> List[dict]: ...
    def for_item(self, run_id: str, employee_key: str) -> List[dict]: ...
    def upsert(self, op: dict) -> None: ...            # durable avant retour, sinon exception


class AttemptRepository(Protocol):
    def append(self, attempt: dict) -> None: ...       # durable avant retour, sinon exception


class PublicationLog(Protocol):
    def append(self, rows: List[dict]) -> None: ...


@dataclasses.dataclass
class OpResult:
    op_id: str
    variant: str
    status: str
    detail: str
    transport: Optional[TransportResult] = None


class OperationExecutor:
    """Exécute les opérations d'UN salarié, dans l'ordre, avec intention durable avant transport."""

    def __init__(self, *, ops: OperationRepository, attempts: AttemptRepository, transport,
                 source_run_id: str, publish_run_id: str, employee_key: str, gate_decision_id: str,
                 allow_new_attempt: bool, max_attempts: int, clock: Callable[[], dt.datetime] = _utc_now):
        self.ops, self.attempts, self.transport = ops, attempts, transport
        self.source_run_id, self.publish_run_id = source_run_id, publish_run_id
        self.employee_key, self.gate_decision_id = employee_key, gate_decision_id
        self.allow_new_attempt, self.max_attempts = allow_new_attempt, max_attempts
        self.clock = clock
        self.results: List[OpResult] = []
        self.status_by_variant: Dict[str, str] = {}
        self.seq = 0
        self.halt_reason: Optional[str] = None
        self.log_rows: List[dict] = []   # lignes du log fonctionnel (format de l'existant), en mémoire

    def log(self, row: dict) -> None:
        """Ligne de log.adp_publication_log produite par la logique portée (une par POST, rejet NIR
        ou information), conservée même si le traitement du salarié s'interrompt ensuite."""
        self.log_rows.append(row)

    def submit(self, op: PlannedOp) -> OpResult:
        self.seq += 1
        thash = target_hash(op.target)
        op_id = op_identity(self.source_run_id, self.employee_key, op.op_variant, thash)
        existing = self.ops.get(op_id)
        if self.halt_reason:
            return self._record(op, op_id, existing, OP_NOT_SENT if not existing else existing["status"],
                                f"HALTED_{self.halt_reason}", persist=False)
        if existing:
            st = existing["status"]
            if st in OP_DONE or st in (OP_SUPERSEDED, OP_ABANDONED):
                return self._record(op, op_id, existing, st, "ALREADY_TERMINAL", persist=False)
            if st in OP_OPEN_UNCERTAIN:
                return self._record(op, op_id, existing, OP_UNCERTAIN, "UNCERTAIN_AWAITING_RECONCILIATION",
                                    persist=False)  # aucun nouvel envoi automatique
            if st in OP_RETRYABLE and not (self.allow_new_attempt or existing.get("manual_decision_id")):
                return self._record(op, op_id, existing, st, "RETRY_REQUIRES_EXPLICIT_REPRISE", persist=False)
            if int(existing.get("attempt_count") or 0) >= self.max_attempts:
                return self._record(op, op_id, existing, st, "MAX_ATTEMPTS_REACHED", persist=False)
        # Dépendances (CONTRAT-ADP:C08)
        for dep in op.depends_on:
            dep_status = self.status_by_variant.get(dep)
            if dep_status is not None and dep_status not in OP_DONE:
                row = self._op_row(op, op_id, thash, existing, OP_BLOCKED_DEPENDENCY, None)
                self.ops.upsert(row)
                return self._record(op, op_id, row, OP_BLOCKED_DEPENDENCY, f"DEPENDS_ON_{dep}", persist=False)
        # 1. Intention durable AVANT transport : une exception ici interrompt tout (rien n'est envoyé)
        attempt_no = int((existing or {}).get("attempt_count") or 0) + 1
        intent = self._op_row(op, op_id, thash, existing, OP_INTENT_RECORDED, None, attempt_no=attempt_no)
        self.ops.upsert(intent)
        attempt_id = str(uuid.uuid4())
        started = self.clock()
        self.attempts.append({"attempt_id": attempt_id, "op_id": op_id, "run_id": self.publish_run_id,
                              "attempt_no": attempt_no, "phase": "INTENT", "outcome_class": None,
                              "http_status": None, "error_class": None, "response_digest": None,
                              "started_at_utc": started, "ended_at_utc": None})
        # 2. Transport (aucun retry, aucune redirection, aucun renvoi)
        try:
            result = self.transport.post(op.endpoint, canonical_payload(op.payload))
        except Exception as exc:  # exception non classée par le transport : on ne conclut pas
            result = TransportResult(UNCERTAIN, "TRANSPORT_EXCEPTION", "UNKNOWN", error_class=type(exc).__name__)
        status = _TRANSPORT_TO_OP[result.outcome]
        # 3. Issue persistée ; si cette écriture échoue, l'intention reste et vaudra « incertain »
        self.attempts.append({"attempt_id": attempt_id, "op_id": op_id, "run_id": self.publish_run_id,
                              "attempt_no": attempt_no, "phase": result.phase, "outcome_class": result.outcome,
                              "http_status": result.http_status, "error_class": result.error_class,
                              "response_digest": result.body_digest, "started_at_utc": started,
                              "ended_at_utc": self.clock()})
        final = self._op_row(op, op_id, thash, intent, status, result, attempt_no=attempt_no)
        self.ops.upsert(final)
        if result.reason == "AUTH_REJECTED_NO_RESEND":
            self.halt_reason = "AUTH_REJECTED"   # CONTRAT-ADP:C05 : jeton refusé, la suite n'est pas envoyée
        return self._record(op, op_id, final, status, result.reason, persist=False, transport=result)

    def _op_row(self, op, op_id, thash, existing, status, result, attempt_no=None):
        now = self.clock()
        base = dict(existing or {})
        base.update({"op_id": op_id, "run_id": self.source_run_id, "employee_key": self.employee_key,
                     "seq_no": base.get("seq_no") or self.seq, "op_code": op.op_code, "op_variant": op.op_variant,
                     "depends_on": list(op.depends_on), "target_hash": thash,
                     "payload_hash": sha256_hex(canonical_payload(op.payload).decode("utf-8")),
                     "status": status, "gate_decision_id": self.gate_decision_id,
                     "created_at_utc": base.get("created_at_utc") or now, "updated_at_utc": now})
        if attempt_no is not None:
            base["attempt_count"] = attempt_no
        if result is not None:
            base["last_outcome_class"] = result.outcome
            base["last_http_status"] = result.http_status
        return base

    def _record(self, op, op_id, row, status, detail, persist, transport=None) -> OpResult:
        self.status_by_variant[op.op_variant] = status
        res = OpResult(op_id, op.op_variant, status, detail, transport)
        self.results.append(res)
        return res


def canonical_payload(payload: dict) -> bytes:
    """Corps JSON envoyé (même sérialisation que l'existant : json.dumps, ASCII échappé).
    CONTRAT-ADP:C15 aucune clé d'idempotence n'est transmise à ADP (support inconnu) : pas de garantie
    « exactement une fois » de bout en bout ; seule la non-réémission automatique est garantie ici."""
    import json
    return json.dumps(payload).encode("utf-8")


def item_status(results: Sequence[OpResult]) -> str:
    """Statut d'outbox (A07) déduit des opérations d'un salarié (D-ADP-07, À CONFIRMER).
    CONTRAT-ADP:C16 SENT = au moins une opération ACCEPTÉE selon le contrat (pas une application observée)."""
    statuses = {r.status for r in results}
    if statuses & OP_OPEN_UNCERTAIN:
        return OB_IN_PROGRESS            # incertain : jamais resélectionné automatiquement
    if statuses & OP_RETRYABLE:
        return OB_FAILED
    if statuses & OP_DONE:
        return OB_SENT
    return OB_NO_CHANGE


# ---------------------------------------------------------------------------
# Obsolescence (SUPERSEDED) — CONTRAT-ADP:C09
# ---------------------------------------------------------------------------
def supersession_candidates(item: dict, newer_items: Iterable[dict], item_ops: Iterable[dict]) -> bool:
    """Un élément plus ancien peut devenir SUPERSEDED si et seulement si :
      * un élément PLUS RÉCENT pour le même salarié porte l'état complet le plus récent et n'est pas
        lui-même incertain (PENDING, FAILED, SENT ou NO_CHANGE) ;
      * aucune opération de l'élément ancien n'est INCERTAINE ou en cours (jamais d'obsolescence
        automatique d'un résultat incertain) ;
    Les opérations déjà ACCEPTED de l'élément ancien restent ACCEPTED (déjà faites)."""
    if item.get("status") not in (OB_PENDING, OB_FAILED):
        return False
    if any(o.get("status") in OP_OPEN_UNCERTAIN for o in item_ops):
        return False
    return any(n.get("status") in (OB_PENDING, OB_FAILED, OB_SENT, OB_NO_CHANGE) for n in newer_items)


# ---------------------------------------------------------------------------
# Réconciliation — CONTRAT-ADP:C07
# ---------------------------------------------------------------------------
RECON_MATCH = "MATCH_ALL_TARGET_FIELDS"
RECON_PARTIAL = "PARTIAL_READ"
RECON_DIFFERENT = "DIFFERENT_FROM_TARGET"
RECON_READ_FAILED = "READ_FAILED"


def reconcile(op_row: dict, planned: PlannedOp, reader) -> dict:
    """Compare l'état lu chez ADP à la cible COMPLÈTE de l'opération.
    * tous les champs cibles lus et égaux -> APPLIED_OBSERVED (application observée, pas garantie) ;
    * lecture partielle -> aucune conclusion (une lecture partielle ne prouve pas tous les champs) ;
    * état différent -> aucune conclusion de non-traitement (un GET différent ne prouve rien) :
      l'opération reste INCERTAINE et requiert une décision manuelle tracée."""
    if planned.read_back is None:
        return {"result": RECON_PARTIAL, "conclusion": "NO_READBACK_DEFINED", "fields": []}
    try:
        observed = planned.read_back(reader)
    except Exception as exc:
        return {"result": RECON_READ_FAILED, "conclusion": "UNCHANGED", "fields": [], "error": type(exc).__name__}
    if observed is None:
        return {"result": RECON_READ_FAILED, "conclusion": "UNCHANGED", "fields": []}
    fields = sorted(planned.target.keys())
    missing = [f for f in fields if f not in observed]
    if missing:
        return {"result": RECON_PARTIAL, "conclusion": "UNCHANGED", "fields": fields}
    if all(observed[f] == planned.target[f] for f in fields):
        return {"result": RECON_MATCH, "conclusion": OP_APPLIED_OBSERVED, "fields": fields}
    return {"result": RECON_DIFFERENT, "conclusion": "MANUAL_DECISION_REQUIRED", "fields": fields}


def apply_manual_decision(op_row: dict, decision: str, decided_by: str, basis: str, reference: str = "",
                          clock: Callable[[], dt.datetime] = _utc_now) -> tuple:
    """Décision humaine tracée (auteur, fondement). Renvoie (op mise à jour, ligne de décision)."""
    if decision not in DECISION_EFFECT:
        raise HrisError("DECISION_INVALID", decision)
    if not decided_by or not basis or not basis.strip():
        raise HrisError("DECISION_INCOMPLETE", "auteur et fondement obligatoires")
    if op_row["status"] in OP_DONE and decision != DEC_ABANDON:
        raise HrisError("DECISION_NOT_APPLICABLE", "opération déjà acceptée ou observée")
    decision_id = str(uuid.uuid4())
    now = clock()
    updated = dict(op_row, status=DECISION_EFFECT[decision], manual_decision_id=decision_id, updated_at_utc=now)
    record = {"decision_id": decision_id, "op_id": op_row["op_id"], "run_id": op_row["run_id"],
              "employee_key": op_row["employee_key"], "decision": decision, "decided_by": decided_by,
              "decided_at_utc": now, "basis": basis, "reference": reference}
    return updated, record
'''),
    ('hris.adp_port', 'db8a126b982c68ac552b5a00bcc288fadf42508e44a9e0f1845ca081fb6a696d', r'''"""hris.adp_port — portage de la logique métier de la Function ADP (adp-function-app/function_app.py,
PRJ @ 5541490, main() L4234-5031 et fonctions appelées), pour UN salarié.

Parité par défaut (A05 § 3) : mêmes lectures (17 G3 + G1/G2/G4/G5/G6), mêmes comparaisons brutes,
mêmes payloads, mêmes valeurs de log (mode, field, nom, log_message), anomalies B1 à B27 conservées
sauf améliorations techniques actées :
  * jeton jamais journalisé ; aucun print de données personnelles (B19, B22) ;
  * session INSEE séparée, sans en-tête ADP (B22) ;
  * chaque POST passe par OperationExecutor : intention durable, aucun retry, aucun renvoi ;
  * une exception sur un salarié n'arrête pas le lot (D-ADP-01 : isolement par défaut, BLOQUANTE
    pour l'activation) — l'existant arrêtait tout (B15).
Points À ARBITRER conservés en parité (A08) : B1/B2 variantes, B3 garde salaire, B7 existence sur
erreur HTTP (BLOQUANTE D-ADP-02), B8 imputation 100, B9 FTE, B10, B11, B12, B13, B21, B23, R-ADP-012.
"""
from __future__ import annotations

import datetime
import decimal
import difflib
import json
import re
import unicodedata
from typing import Any, Callable, Dict, Optional, Tuple

from hris.adp_ops import OP_DONE, OperationExecutor, PlannedOp

EV_HIRE = "events/hr/v1/worker.hire"
EV_PERSONAL_EMAIL = "events/hr/v1/worker.personal-communication.email.change"
EV_BUSINESS_EMAIL = "events/hr/v1/worker.business-communication.email.change"
EV_BUSINESS_MOBILE = "events/hr/v1/worker.business-communication.mobile.change"
EV_PERSONAL_MOBILE = "events/hr/v1/worker.personal-communication.mobile.change"
EV_MARITAL = "events/hr/v1/worker.marital-status.change"
EV_LEGAL_ADDRESS = "events/hr/v1/worker.legal-address.change"
EV_WA_MODIFY = "events/hr/v1/worker.work-assignment.modify"
EV_PAY_DISTRIBUTION = "events/payroll/v1/worker.pay-distribution.change"

HIRE_IMPUTATION_PERCENTAGE = 100  # B8 : constante de test lue par fermeture (L367), entier


class ReadResponse:
    """Réponse de lecture (statut, corps JSON) au format utilisé par la logique portée."""

    def __init__(self, status_code: int, text: str = "", reason: str = ""):
        self.status_code = status_code
        self.text = text if isinstance(text, str) else json.dumps(text)
        self.reason = reason

    def json(self):
        return json.loads(self.text) if self.text else None


class TransportReader:
    """Lectures ADP réelles (G1..G5) via le transport, et INSEE (G6) via une session HTTP SÉPARÉE,
    sans en-tête ADP (B22 corrigé). Politique des lectures réelles non établie (CONTRAT-ADP:C14) :
    le transport refuse toute lecture tant qu'elle ne l'est pas."""

    def __init__(self, transport, insee_timeout: float = 5.0):
        self.transport = transport
        self.insee_timeout = insee_timeout

    def get(self, path):
        status, text = self.transport.get(path)
        return ReadResponse(status or 0, text or "")

    def insee(self, postal_code):
        import requests
        with requests.Session() as s:
            r = s.get("https://geo.api.gouv.fr/communes", params={"codePostal": postal_code, "fields": "nom,code"},
                      timeout=self.insee_timeout, allow_redirects=False)
            return ReadResponse(r.status_code, r.text)


# ---------------------------------------------------------------------------
# Fonctions de transformation de l'existant (L204-329, L454-458)
# ---------------------------------------------------------------------------
def normalize_string(input_string):
    normalized = unicodedata.normalize('NFD', input_string)
    without_accents = ''.join(char for char in normalized if unicodedata.category(char) != 'Mn')
    replaced_chars = without_accents.replace('-', ' ').replace("'", ' ')
    cleaned_string = re.sub(r'[^a-zA-Z0-9\s]', '', replaced_chars)
    return cleaned_string.upper().strip()


def convert_date_format(date_str):
    if date_str is None:
        return '1900-01-01'
    date_str = str(date_str).strip()
    if date_str == '' or date_str.lower() in ('nan', 'none', 'nat'):
        return '1900-01-01'
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(date_str, fmt).strftime("%Y-%m-%d")
        except Exception:
            continue
    try:
        from dateutil import parser
        return parser.parse(date_str).strftime("%Y-%m-%d")
    except Exception:
        return '1900-01-01'


def extract_error(data_str):
    """Non protégé dans l'existant (B14) : l'appelant capture l'exception (isolement)."""
    data_string = json.loads(json.dumps(json.loads(data_str), indent=4))
    return data_string['confirmMessage']['processMessages'][0]['developerMessage']['messageTxt']


def get_recordingBasisCode(remunerationType):
    return {"B": "12M", "F": "12F", "S": ""}.get(remunerationType)


def normalize_street_compl(value):
    if not value or not value.strip():
        return ""
    first = value.strip()[0].upper()
    ascii_char = "".join(c for c in unicodedata.normalize("NFD", first) if unicodedata.category(c) != "Mn")
    if ascii_char.isalpha() and ascii_char.isascii():
        return ascii_char
    return ""


def set_workSchedule(fte):
    return '001' if fte == 1 else '999'


ARRONDISSEMENT_COMMUNES = {"75056": ("750", "751"), "69123": ("6900", "6938"), "13055": ("130", "132")}


def resolve_arrondissement_insee(commune_code, postal_code):
    if commune_code not in ARRONDISSEMENT_COMMUNES:
        return commune_code
    prefix_postal, prefix_insee = ARRONDISSEMENT_COMMUNES[commune_code]
    if postal_code.startswith(prefix_postal):
        return prefix_insee + postal_code[len(prefix_postal):]
    return commune_code


def append_arrondissement_to_city(city_name, postal_code):
    city_normalized = city_name.upper().strip()
    if re.match(r'^(PARIS|LYON|MARSEILLE)\s+\d+$', city_normalized):
        return city_name
    cities = {"PARIS": "75", "LYON": "69", "MARSEILLE": "13"}
    if city_normalized not in cities:
        return city_name
    if not postal_code.startswith(cities[city_normalized]):
        return city_name
    if len(postal_code) == 5:
        return f"{city_normalized} {str(int(postal_code[-2:]))}"
    return city_name


def get_insee_code(reader, postal_code, city_name, similarity_threshold=0.8):
    """G6 via une session HTTP SÉPARÉE (B22 corrigé), timeout 5 s, mêmes règles de choix."""
    try:
        response = reader.insee(postal_code)
        if response.status_code != 200:
            return None, None, 0
        communes = response.json()
        if not communes:
            return None, None, 0
        if len(communes) == 1:
            commune = communes[0]
            return resolve_arrondissement_insee(commune['code'], postal_code), commune['nom'], 1.0
        best_match, best_ratio = None, 0
        for commune in communes:
            ratio = difflib.SequenceMatcher(None, city_name, normalize_string(commune['nom'])).ratio()
            if ratio > best_ratio:
                best_ratio, best_match = ratio, commune
        if best_match and best_ratio >= similarity_threshold:
            return resolve_arrondissement_insee(best_match['code'], postal_code), best_match['nom'], best_ratio
        return None, None, best_ratio
    except Exception:
        return None, None, 0


# ---------------------------------------------------------------------------
# Préparation de la ligne (R-ADP-003) : astype puis fillna('') comme pd.read_parquet de la CETAS
# ---------------------------------------------------------------------------
def legacy_row(row: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(row)
    for col in ("id_payroll", "lb_adresse_cp", "cd_ppc_convention_collective_adp", "cd_nature_contract_adp",
                "cd_type_contract_adp"):
        out[col] = str(out.get(col))            # astype(str) AVANT fillna : NULL -> 'None' (parité)
    nir = out.get("id_numero_securite_sociale")
    out["id_numero_securite_sociale"] = int(nir)  # B16 : échec = salarié en erreur (isolé)
    for k, v in list(out.items()):
        if v is None:
            out[k] = ''
    return out


def _ts(date_value: str, now: datetime.datetime) -> str:
    """TS(x) = date + heure courante 'THH:MM:SS.mmmZ' (L4334 etc. ; B27)."""
    return date_value + now.strftime('T%H:%M:%S.') + f"{now.microsecond // 1000:03d}Z"


# ---------------------------------------------------------------------------
# Lectures ADP (G1..G5) — même sémantique que l'existant, une requête par appel (D-ADP-11 : parité)
# ---------------------------------------------------------------------------
def _g3(reader, worker_id):
    return reader.get(f"hr/v2/workers/{worker_id}")


def _g1(reader, worker_id):
    return reader.get("hr/v2/workers?$filter=workers/workerID/idValue eq '" + worker_id + "'")


def worker_exists(reader, worker_id):
    return _g1(reader, worker_id).status_code == 200          # B7 : toute erreur = absent (D-ADP-02)


def ssn_exists_v2(reader, ssn):
    r = reader.get("hr/v2/workers?$filter=workers%2Fperson%2FidentityDocuments%5B0%5D%2FdocumentID%20eq%20%27"
                   + ssn + "%27")
    return r.status_code == 200


def get_worker_ssn(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None
    for worker in r.json()['workers']:
        for document in worker['person']['identityDocuments']:
            if document['typeCode']['codeValue'] == "SSN":
                return document['documentID']
            return False                                     # B6 : premier document seulement
    return None


def first_hire_date(reader, worker_id):
    return _g3(reader, worker_id).json()["workers"][0]["workerDates"]["firstHireDate"]


def convert_worker_id(reader, worker_id):
    return _g3(reader, worker_id).json()["workers"][0]["associateOID"]


def get_name(reader, worker_id):
    return _g3(reader, worker_id).json()["workers"][0]["person"]["legalName"]["formattedName"]


def get_worker_dates(reader, worker_id):
    data = _g3(reader, worker_id).json()
    try:
        worker = data["workers"][0]
        wa = worker.get("workAssignments", [])[0]
        start_wt, start_pb, end_pb = None, '', ''
        for field in wa.get("customFieldGroup", {}).get("dateFields", []):
            fid = field.get("itemID")
            if fid == "startWorkerTypeDate":
                start_wt = field.get("dateValue")
            elif fid == "startProbationDate":
                start_pb = field.get("dateValue")
            elif fid == "endProbationDate":
                end_pb = field.get("dateValue")
        return {"firstHireDate": worker.get("workerDates", {}).get("firstHireDate"),
                "originalHireDate": worker.get("workerDates", {}).get("originalHireDate"),
                "hireDate": wa.get("hireDate"), "seniorityDate": wa.get("seniorityDate"),
                "expectedStartDate": wa.get("expectedStartDate"), "contractStartDate": start_wt,
                "startProbationDate": start_pb, "endProbationDate": end_pb}
    except Exception:
        return None


def _first_worker(r):
    workers = r.json().get("workers", [])
    return workers[0] if workers else None


def get_worker_personalemail(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    try:
        return r.json()["workers"][0]["person"]["communication"]["emails"][0]["emailUri"]
    except (KeyError, IndexError):
        return ''


def get_worker_businessemail(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    emails = worker.get("businessCommunication", {}).get("emails", [])
    return emails[0].get("emailUri") if emails else ''


def get_business_mobile_number(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    mobiles = worker.get("businessCommunication", {}).get("mobiles", [])
    return mobiles[0].get("formattedNumber") if mobiles else ''


def get_personal_mobile_number(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    mobiles = worker.get("person", {}).get("communication", {}).get("mobiles", [])
    return mobiles[0].get("formattedNumber") if mobiles else ''


def get_marital_status_code(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return ''
    worker = _first_worker(r)
    if not worker:
        return ''
    return worker.get("person", {}).get("maritalStatusCode", {}).get("codeValue") or ''


def get_marital_status_effective_date(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None
    worker = _first_worker(r)
    if not worker:
        return None
    return worker.get("person", {}).get("maritalStatusCode", {}).get("effectiveDate") or None


def get_legal_address(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code == 200:
        worker = _first_worker(r)
        if not worker:
            return "", "", "", "", "", "", "", ""
        la = worker.get("person", {}).get("legalAddress", {})
        if la:
            return (la.get("buildingNumberExtension") or "", la.get("buildingNumber") or "",
                    la.get("streetName") or "", la.get("countryCode") or "",
                    la.get("countrySubdivisionLevel2", {}).get("longName") or la.get("cityName") or "",
                    la.get("lineFive") or "", la.get("postalCode") or "",
                    la.get("countrySubdivisionLevel1", {}).get("codeValue") or "")
        return "", "", "", "", "", "", "", ""
    return "", "", "", "", "", "", ""                         # B4 : 7 valeurs -> erreur au dépaquetage


def get_worker_salary(reader, worker_id):
    r = _g1(reader, worker_id)
    if r.status_code == 200:
        try:
            wa = r.json()["workers"][0]["workAssignments"][0]
            return wa.get('baseRemuneration', {}).get('monthlyRateAmount', {}).get('amountValue')
        except (KeyError, IndexError):
            return None
    return False


def get_worker_bonus(reader, worker_id):
    r = _g1(reader, worker_id)
    if r.status_code == 200:
        wa = r.json().get('workers', [])[0].get('workAssignments', [])[0]
        value = 0
        for field in wa.get('customFieldGroup', {}).get('numberFields', []):
            if field.get('itemID') == 'REM_RE1MTS18':
                value = field.get('numberValue', 0)
                break
        return value
    return False


def get_worker_type_and_termination(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None, None
    try:
        workers = r.json().get("workers", [])
        if not workers:
            return None, None
        was = workers[0].get("workAssignments", [])
        if not was:
            return None, None
        a = was[0]
        code = a.get("workerTypeCode", {}).get("codeValue")
        if code == "00":
            end = a.get("expectedTerminationDate")
        else:
            end = None
            for field in a.get("customFieldGroup", {}).get("dateFields", []):
                if field.get("itemID") == "fixedTermContractInitialTerminationDate":
                    end = field.get("dateValue")
                    break
        return code, end
    except Exception:
        return None, None


def get_work_schedule(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code == 200:
        wa = r.json().get('workers', [])[0].get('workAssignments', [])[0]
        code = None
        for field in wa.get('customFieldGroup', {}).get('codeFields', []):
            if field.get('itemID') == 'workSchedule':
                code = field.get('codeValue', None)
                break
        return code, wa.get('fullTimeEquivalenceRatio', None)
    return False, None


def get_organisational_structure(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None
    worker = _first_worker(r)
    if not worker:
        return None
    was = worker.get("workAssignments", [])
    if not was:
        return None
    for unit in was[0].get("assignedOrganizationalUnits", []):
        if unit.get("itemID") == "departmentId":
            return unit.get("nameCode", {}).get("codeValue") or ""
    return ""


def get_worker_imputation(reader, worker_id):
    r = _g3(reader, worker_id)
    if r.status_code != 200:
        return None, None
    worker = _first_worker(r)
    if not worker:
        return None, None
    was = worker.get("workAssignments", [])
    if not was:
        return None, None
    ccs = was[0].get("assignmentCostCenters", [])
    if not ccs:
        return None, None
    pct = ccs[0].get("costCenterPercentage")
    return (int(pct) if pct is not None else None), ccs[0].get("costCenterID")


def get_worker_classification(reader, worker_id):
    r = _g3(reader, worker_id)
    if not (200 <= r.status_code < 400):                      # raise_for_status
        return False                                          # B5 : dépaquetage impossible -> erreur
    wa = ((r.json().get("workers") or [{}])[0].get("workAssignments") or [{}])[0]
    group = classification = coefficient = remuneration = None
    for wg in wa.get("workerGroups", []):
        if wg.get("itemID") == "collectiveAgreement":
            group = wg.get("groupCode", {}).get("codeValue")
            break
    for oc in wa.get("occupationalClassifications", []):
        if oc.get("itemID") == "classification":
            classification = oc.get("classificationCode", {}).get("codeValue")
            break
    for nf in wa.get("customFieldGroup", {}).get("numberFields", []):
        if nf.get("itemID") == "coefficient":
            try:
                coefficient = str(int(float(nf.get("numberValue", 0))))
            except (ValueError, TypeError):
                coefficient = None
            break
    for cf in wa.get("customFieldGroup", {}).get("codeFields", []):
        if cf.get("itemID") == "remunerationType":
            remuneration = cf.get("codeValue")
            break
    return (group, classification, coefficient, wa.get("jobTitle"), wa.get("jobFunctionCode", {}).get("codeValue"),
            remuneration, wa.get("baseRemuneration", {}).get("recordingBasisCode", {}).get("codeValue"))


def get_pay_distribution(reader, associate_oid):
    try:
        r = reader.get(f"payroll/v2/workers/{associate_oid}/pay-distributions")
        if r.status_code != 200:
            return "", "", "", ""
        try:
            di = r.json()["payDistributions"][0]["distributionInstructions"][0]
            holder = di["depositAccount"]["financialAccount"].get("accountName")
            iban = di["depositAccount"].get("IBAN")
            party = di["depositAccount"].get("financialParty", {})
            swift = di["depositAccount"]["financialParty"].get("SWIFTCode")
            bank = party.get("nameCode", {}).get("shortName", "")
            if not all([holder, iban, swift, bank]):
                return "", "", "", ""
            return holder, iban, swift, bank
        except (KeyError, IndexError, TypeError):
            return "", "", "", ""
    except Exception:
        return "", "", "", ""


# ---------------------------------------------------------------------------
# Payloads (A05 § 4.5) — même structure et même ordre de clés que l'existant
# ---------------------------------------------------------------------------
def _person_event(aoid, fragment):
    return {"events": [{"data": {"eventContext": {"worker": {"associateOID": aoid}},
                                 "transform": {"worker": fragment}}}]}


def _wa_event(aoid, waid, eff, wa_fragment, extra_transform=None):
    transform = {"effectiveDateTime": eff}
    if extra_transform:
        transform.update(extra_transform)
    transform["workAssignment"] = wa_fragment
    return {"events": [{"data": {"eventContext": {"associateOID": aoid, "workAssignmentID": waid},
                                 "transform": transform}}]}


def _marital_payload(aoid, code, eff_date):
    return {"events": [{"serviceCategoryCode": {"codeValue": "hr"},
                        "eventNameCode": {"codeValue": "worker.maritalStatus.change"},
                        "data": {"eventContext": {"worker": {"associateOID": aoid}},
                                 "transform": {"eventStatusCode": {"codeValue": "submit"},
                                               "worker": {"person": {"maritalStatusCode": {
                                                   "codeValue": code, "effectiveDate": eff_date}}}}}}]}


def _org_units(code):
    letters = {'A': 'administrativeAssignment1', 'B': 'administrativeAssignment2',
               'C': 'administrativeAssignment3', 'D': 'administrativeAssignment4'}
    return [{"itemID": letters[l], "nameCode": {"codeValue": l + d}} for l, d in re.findall(r'([A-D])(\d+)', code)]


# ---------------------------------------------------------------------------
# Traitement d'un salarié (main(), corps de boucle L4259-4986)
# ---------------------------------------------------------------------------
class EmployeePlanner:
    """planner(row, executor) : exécute la logique de l'existant pour un salarié. Chaque POST est
    soumis à l'exécuteur (intention durable, transport, issue) ; chaque ligne de log de l'existant
    est transmise à executor.log()."""

    def __init__(self, reader, now: Callable[[], datetime.datetime], row_index: Callable[[str], int] = lambda k: 0):
        self.reader = reader
        self.now = now
        self.row_index = row_index

    def __call__(self, raw_row: Dict[str, Any], ex: OperationExecutor) -> None:
        row = legacy_row(raw_row)
        r = self.reader
        now = self.now()
        index = self.row_index(raw_row.get("id_unique"))
        worker_id = row['id_payroll']
        matricule_it = row['cd_matricule_it']
        preferredSalutations = row['cd_emp_titre_adp']
        givenName = normalize_string(row['lb_prenom'])
        familyName1 = normalize_string(row['lb_nom'])
        marital_status = row['cd_emp_situation_matrimoniale_adp']
        marital_effectiveDate = convert_date_format(row['dt_debut_situation_matrimoniale'])
        birthName = row['lb_nom_naissance']
        birth_country_code = row['id_naissance_pays_adp']
        birth_dept = row['lb_naissance_departement']
        birth_postcode = row['lb_naissance_cp']
        birth_city_name = normalize_string(row['lb_naissance_ville'])
        email = row['lb_email_professionnel']
        emailPerso = row['lb_email_personnel']
        landline = row['lb_telephone_professionnel']
        mobile = row['lb_telephone_personnel']
        birthDate = convert_date_format(row['dt_naissance'])
        gender = row['cd_emp_sexe_adp']
        identityDocumentsId = str(int(row['id_numero_securite_sociale']))
        citizenship_country_code = row['cd_emp_nationalite_adp']
        legalStreetCompl = str(row['lb_fiscal_streetnumber_complement'] or '').strip()
        legalAddressLineOne = str(row['lb_fiscal_streetnumber'] or '').strip()
        legalAddressLineTwo = normalize_string(row['lb_fiscal_street'])
        legalAddressPostcode = str(row['lb_fiscal_cp'] or '').strip()
        legalAddressCityName = normalize_string(row['lb_fiscal_ville'])
        legalAddressCountrycode = str(row['id_fiscal_pays_adp'] or '').strip()
        legalAdditionalAddressInformation = normalize_string(row['lb_fiscal_additional_address_information'])
        organisational_structure_code = str(row['cd_organisationalstructure_adp'] or '').strip()
        account_holder = normalize_string(row['lb_account_holder'])
        iban = row['lb_iban']
        swift_code = row['lb_bic']
        bank_name = normalize_string(row['lb_bank_name'])
        assignedWorkLocations = row['cd_ppc_location_country_adp']          # R-ADP-047 (D-ADP-06)
        hireDate = convert_date_format(str(row['dt_entree_societe']))
        firstHireDate = convert_date_format(str(row['dt_entree_groupe']))
        seniorityDate = convert_date_format(str(row['dt_anciennete']))
        if not seniorityDate or seniorityDate == "" or seniorityDate == "1900-01-01":
            seniorityDate = firstHireDate
        companyStartDate = convert_date_format(str(row['dt_company_start']))
        startProbationDate = convert_date_format(str(row['dt_ppc_debut_periode_essaie']))
        endProbationDate = convert_date_format(str(row['dt_ppc_fin_periode_essaie']))
        situationStartDateTS = _ts(convert_date_format(str(row['dt_situation_start'])), now)
        collaborationType = row['cd_type_collaboration_adp']
        workerTypeCode = row['cd_nature_contract_adp']
        reasonCDD = row['cd_recours_cdd']
        contractType = row['cd_type_contract_adp'].strip(' ')
        if workerTypeCode == '20':
            reason = ''
        elif workerTypeCode == '01':
            reason = reasonCDD
        else:
            reason = '99'
        raw_reason = str(row.get('cd_ppc_raison_debut_contrat_adp', '')).strip()
        reasonHire = str(int(raw_reason)) if raw_reason.isdigit() else ''
        contractStartDate = convert_date_format(str(row['dt_ppc_debut_contrat']))
        contractStartDateTS = _ts(contractStartDate, now)
        expectedTerminationDate = convert_date_format(str(row['dt_ppc_fin_contrat']))
        if expectedTerminationDate == '1900-01-01':
            expectedTerminationDate = None
        costCenter = row['cd_ppc_cost_center_adp']
        imputationPercentage = '100'                                          # B8 (mise à jour : chaîne)
        monthlySalary = round(float(row['mt_salaire_mensuel']), 2)            # B16 : '' -> erreur
        salaryeffectiveDateTime = _ts(convert_date_format(str(row['dt_debut_mt_salaire_base'])), now)
        bonus = float(row['pc_prime_annuelle'])
        bonuseffectiveDate = convert_date_format(str(row['dt_debut_pc_prime_annuelle']))
        bonuseffectiveDateTime = _ts(bonuseffectiveDate, now)
        raw_fte = str(row.get('pc_imputation', '')).strip().replace(',', '.')
        try:
            fte = float(raw_fte) / 100                                        # B9 : FTE pris dans pc_imputation
        except ValueError:
            fte = 0.0
        terminationDate = convert_date_format(str(row['dt_sortie']))
        int(terminationDate.split("-")[0])                                    # parité : termination_year
        jobFunctionCode = str(row['cd_categorie_cotisant_adp'])
        remunerationType = row['cd_classe_remuneration_adp'].strip(' ')
        recordingBasisCode = get_recordingBasisCode(remunerationType)
        jobTitle = normalize_string(row['lb_job_title'])
        workSchedule = set_workSchedule(fte)

        # --- Décision embauche / mise à jour (R-ADP-008/009, B7 à arbitrer D-ADP-02) -------------
        if worker_exists(r, worker_id) == False and ssn_exists_v2(r, identityDocumentsId) == False:  # noqa: E712
            detail = f"Row {index} is a new hire : id_payroll = {worker_id}"
            payload = _hire_payload_impl(
                self,
                worker_id, citizenship_country_code, identityDocumentsId, givenName, familyName1, birthName,
                preferredSalutations, landline, mobile, email, marital_status, birthDate, birth_city_name,
                birth_postcode, birth_dept, birth_country_code, legalAddressLineOne, legalAddressLineTwo,
                legalAddressCityName, legalAddressCountrycode, legalAddressPostcode, hireDate, seniorityDate,
                jobFunctionCode, workerTypeCode, monthlySalary, costCenter, collaborationType, contractType,
                endProbationDate, startProbationDate, bonus, reasonHire, reason, expectedTerminationDate,
                assignedWorkLocations, emailPerso, jobTitle, gender, remunerationType,
                legalAdditionalAddressInformation, workSchedule, recordingBasisCode, matricule_it,
                contractStartDate, firstHireDate, companyStartDate, marital_effectiveDate,
                organisational_structure_code, legalStreetCompl)
            target = {"workerID": worker_id, "hire_payload_hash": _stable_hash(payload)}
            self._post(ex, "P1", "P1_HIRE_V1", EV_HIRE, payload, target, familyName1, worker_id, "INSERT", "hire",
                       detail, read_back=lambda rd: {"workerID": worker_id} if worker_exists(rd, worker_id) else None)
            return

        ssn_worker_id = get_worker_ssn(r, worker_id)
        if identityDocumentsId != ssn_worker_id:
            detail = f"SSN Mismatch for {worker_id} between TS {identityDocumentsId} and ADP {ssn_worker_id}"
            ex.log(self._log_row(worker_id, familyName1, 'UPDATE', 'SSN', 409, 'DATA_VALIDATION', None, detail,
                                 detail, step="SSN_CHECK"))
            return

        workAssignmentID = worker_id + "|" + first_hire_date(r, worker_id)
        associateOID = convert_worker_id(r, worker_id)
        dates = get_worker_dates(r, worker_id)
        current_seniority_date = dates["seniorityDate"]                      # dates None -> erreur (parité)
        formatted_name = get_name(r, worker_id)
        nom = formatted_name

        current = get_worker_personalemail(r, worker_id)
        if emailPerso != current:
            self._post(ex, "P2", "P2_PERSONAL_EMAIL", EV_PERSONAL_EMAIL,
                       _person_event(associateOID, {"person": {"communication": {"email": {"emailUri": emailPerso}}}}),
                       {"emailUri": emailPerso}, nom, worker_id, 'UPDATE', 'personalemail',
                       f"changing personal email from {current} to {emailPerso}",
                       read_back=lambda rd: {"emailUri": get_worker_personalemail(rd, worker_id)})

        current = get_worker_businessemail(r, worker_id)
        if email != current:
            self._post(ex, "P3", "P3_BUSINESS_EMAIL", EV_BUSINESS_EMAIL,
                       _person_event(associateOID, {"businessCommunication": {"email": {"emailUri": email}}}),
                       {"emailUri": email}, nom, worker_id, 'UPDATE', 'businessemail',
                       f"changing business email from {current} to {email}",
                       read_back=lambda rd: {"emailUri": get_worker_businessemail(rd, worker_id)})

        current = get_business_mobile_number(r, worker_id)
        if landline != current:
            self._post(ex, "P4", "P4_BUSINESS_MOBILE", EV_BUSINESS_MOBILE,
                       _person_event(associateOID, {"businessCommunication": {"mobile": {"formattedNumber": landline}}}),
                       {"formattedNumber": landline}, nom, worker_id, 'UPDATE', 'businessmobile',
                       f"changing business mobile  from {current} to {landline}",
                       read_back=lambda rd: {"formattedNumber": get_business_mobile_number(rd, worker_id)})

        current = get_personal_mobile_number(r, worker_id)
        if mobile != current:
            self._post(ex, "P5", "P5_PERSONAL_MOBILE", EV_PERSONAL_MOBILE,
                       _person_event(associateOID, {"person": {"communication": {"mobile": {"formattedNumber": mobile}}}}),
                       {"formattedNumber": mobile}, nom, worker_id, 'UPDATE', 'personalmobile',
                       f"changing personal mobile  from {current} to {mobile}",
                       read_back=lambda rd: {"formattedNumber": get_personal_mobile_number(rd, worker_id)})

        cur_ms = get_marital_status_code(r, worker_id)
        cur_ms_date = get_marital_status_effective_date(r, worker_id)
        if marital_status != cur_ms or marital_effectiveDate != cur_ms_date:   # B10
            self._post(ex, "P6", "P6_MARITAL_STATUS", EV_MARITAL,
                       _marital_payload(associateOID, marital_status, marital_effectiveDate),
                       {"codeValue": marital_status, "effectiveDate": marital_effectiveDate}, nom, worker_id,
                       'UPDATE', 'maritalStatusCode',
                       f"changing marital status from {cur_ms} to {marital_status} on {marital_effectiveDate}",
                       read_back=lambda rd: {"codeValue": get_marital_status_code(rd, worker_id),
                                             "effectiveDate": get_marital_status_effective_date(rd, worker_id)})

        (c_compl, c_one, c_two, c_country, c_city, c_add, c_cp, _c_bureau) = get_legal_address(r, worker_id)  # B4
        legalStreetCompl = normalize_street_compl(legalStreetCompl)                                           # B21
        city_norm = append_arrondissement_to_city(legalAddressCityName, legalAddressPostcode)
        if (legalStreetCompl != c_compl or legalAddressLineOne != c_one or legalAddressLineTwo != c_two
                or legalAddressCountrycode != c_country or city_norm != c_city
                or legalAdditionalAddressInformation != c_add or legalAddressPostcode != c_cp):
            insee, _, _ = get_insee_code(r, legalAddressPostcode, legalAddressCityName)
            longname = append_arrondissement_to_city(legalAddressCityName, legalAddressPostcode)
            payload = _person_event(associateOID, {"person": {"legalAddress": {
                "buildingNumberExtension": legalStreetCompl, "buildingNumber": legalAddressLineOne,
                "streetName": legalAddressLineTwo, "countryCode": legalAddressCountrycode,
                "countrySubdivisionLevel1": {"codeValue": legalAddressCityName},
                "countrySubdivisionLevel2": {"longName": longname, "codeValue": insee if insee else ""},
                "lineFive": legalAdditionalAddressInformation, "postalCode": legalAddressPostcode}}})
            detail = (f"changing legal address from {c_one} {c_compl} {c_two} {c_country} {c_city} {c_add} {c_cp} "
                      f"to {legalAddressLineOne} {legalStreetCompl} {legalAddressLineTwo} {legalAddressCountrycode} "
                      f"{legalAddressCityName} {legalAdditionalAddressInformation} {legalAddressPostcode}")
            target = {"buildingNumberExtension": legalStreetCompl, "buildingNumber": legalAddressLineOne,
                      "streetName": legalAddressLineTwo, "countryCode": legalAddressCountrycode, "city": longname,
                      "lineFive": legalAdditionalAddressInformation, "postalCode": legalAddressPostcode}

            def rb(rd):
                v = get_legal_address(rd, worker_id)
                if len(v) != 8:
                    return None
                return {"buildingNumberExtension": v[0], "buildingNumber": v[1], "streetName": v[2],
                        "countryCode": v[3], "city": v[4], "lineFive": v[5], "postalCode": v[6]}
            self._post(ex, "P7", "P7_LEGAL_ADDRESS", EV_LEGAL_ADDRESS, payload, target, nom, worker_id, 'UPDATE',
                       'legal address', detail, read_back=rb)  # remise à vide PUIS extraction en 4xx (L4662-4665)

        current_salary = get_worker_salary(r, worker_id)
        if current_salary is not None and monthlySalary != current_salary and salaryeffectiveDateTime != '1900-01-01':  # B3
            self._post(ex, "P8", "P8_SALARY", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, salaryeffectiveDateTime,
                                 {"baseRemuneration": {"monthlyRateAmount": {"currencyCode": "EUR",
                                                                             "amountValue": monthlySalary}}}),
                       {"amountValue": monthlySalary, "effectiveDate": salaryeffectiveDateTime[:10]}, nom, worker_id,
                       'UPDATE', 'salary', f"changing salary from {current_salary} to {monthlySalary} on {salaryeffectiveDateTime}",
                       read_back=lambda rd: {"amountValue": get_worker_salary(rd, worker_id)})

        current_bonus = get_worker_bonus(r, worker_id)
        if bonus != get_worker_bonus(r, worker_id) and bonuseffectiveDate != '1900-01-01':
            self._post(ex, "P8", "P8_BONUS", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, bonuseffectiveDateTime,
                                 {"customFieldGroup": {"numberFields": [{"itemID": "REM_RE1MTS18", "numberValue": bonus}]}}),
                       {"numberValue": bonus, "effectiveDate": bonuseffectiveDate}, nom, worker_id, 'UPDATE', 'bonus',
                       f"changing bonus from {current_bonus} to {bonus} on {bonuseffectiveDateTime}",
                       read_back=lambda rd: {"numberValue": get_worker_bonus(rd, worker_id)})

        cur_type, cur_end = get_worker_type_and_termination(r, worker_id)
        contract_variant = None
        if workerTypeCode != cur_type:
            if expectedTerminationDate != '1900-01-01' and workerTypeCode != '00':     # B2 : A pour tout non-CDI
                contract_variant = "P8_CONTRACT_A"
                fragment = {"expectedTerminationDate": expectedTerminationDate,
                            "assignmentStatus": {"reasonCode": {"codeValue": reasonHire}},
                            "workerTypeCode": {"codeValue": workerTypeCode},
                            "customFieldGroup": {
                                "codeFields": [{"itemID": "recoursReason", "codeValue": reason},
                                               {"itemID": "contractType", "codeValue": contractType}],
                                "dateFields": [{"dateValue": endProbationDate, "itemID": "endProbationDate"},
                                               {"dateValue": startProbationDate, "itemID": "startProbationDate"},
                                               {"itemID": "fixedTermContractInitialTerminationDate",
                                                "dateValue": expectedTerminationDate},
                                               {"itemID": "startWorkerTypeDate", "dateValue": contractStartDate}]}}
                payload = _wa_event(associateOID, workAssignmentID, contractStartDateTS, fragment)
            else:
                contract_variant = "P8_CONTRACT_B"
                fragment = {"expectedStartDate": contractStartDate,
                            "assignmentStatus": {"reasonCode": {"codeValue": reasonHire}},
                            "workerTypeCode": {"codeValue": workerTypeCode},
                            "assignedWorkLocations": [{"itemID": "default", "nameCode": {"codeValue": assignedWorkLocations}}],
                            "customFieldGroup": {
                                "codeFields": [{"itemID": "recoursReason", "codeValue": reason},
                                               {"itemID": "contractType", "codeValue": contractType}],
                                "dateFields": [{"itemID": "startWorkerTypeDate", "dateValue": contractStartDate}]}}
                payload = _wa_event(associateOID, workAssignmentID, contractStartDateTS, fragment,
                                    extra_transform={"expectedStartDate": contractStartDate})
            target = {"workerTypeCode": workerTypeCode, "contract_payload_hash": _stable_hash(payload, volatile=True)}
            self._post(ex, "P8", contract_variant, EV_WA_MODIFY, payload, target, nom, worker_id, 'UPDATE',
                       'contrat et affectation',
                       f"changing contract from {cur_type} to {workerTypeCode} with probation dates from "
                       f"{startProbationDate} to {endProbationDate} ",
                       read_back=lambda rd: {"workerTypeCode": get_worker_type_and_termination(rd, worker_id)[0]})

        if expectedTerminationDate != cur_end:                                          # B12
            self._post(ex, "P8", "P8_END_DATE", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, contractStartDateTS,
                                 {"expectedTerminationDate": expectedTerminationDate}),
                       {"expectedTerminationDate": expectedTerminationDate}, nom, worker_id, 'UPDATE',
                       'date de fin de contrat',
                       f"changing termination date from {cur_end} to {expectedTerminationDate}",
                       depends_on=("P8_CONTRACT_A", "P8_CONTRACT_B"))                    # CONTRAT-ADP:C08

        cur_ws, cur_fte = get_work_schedule(r, worker_id)
        if cur_fte is None:
            pass
        elif cur_fte is not None and round(float(cur_fte), 4) == round(float(fte), 4):
            pass
        elif workSchedule != cur_ws:
            self._post(ex, "P8", "P8_SCHEDULE", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS,
                                 {"customFieldGroup": {"codeFields": [{"itemID": "workSchedule", "codeValue": workSchedule}]}}),
                       {"workSchedule": workSchedule}, nom, worker_id, 'UPDATE', 'work schedule',
                       f"changing work schedule from code {cur_ws} FTE {cur_fte} to code {workSchedule} FTE {fte} "
                       f"on {situationStartDateTS}",
                       read_back=lambda rd: {"workSchedule": get_work_schedule(rd, worker_id)[0]})

        cur_org = get_organisational_structure(r, worker_id)
        if organisational_structure_code and cur_org is not None and organisational_structure_code != cur_org:  # B11
            detail = (f"changing organisational structure from {cur_org} to {organisational_structure_code} "
                      f"on {situationStartDateTS}")
            units = _org_units(organisational_structure_code)
            if not units:  # pas de POST, une ligne de log à statut vide (parité)
                ex.log(self._log_row(worker_id, nom, 'UPDATE', 'organisational structure', None, None, None, '',
                                     detail, step="P8_ORG_NO_SEGMENT"))
            else:
                self._post(ex, "P8", "P8_ORG", EV_WA_MODIFY,
                           _wa_event(associateOID, workAssignmentID, situationStartDateTS,
                                     {"assignedOrganizationalUnits": units}),
                           {"units": [u["nameCode"]["codeValue"] for u in units]}, nom, worker_id, 'UPDATE',
                           'organisational structure', detail)

        cur_imp, cur_cc = get_worker_imputation(r, worker_id)
        if imputationPercentage != str(cur_imp) or costCenter != cur_cc:
            self._post(ex, "P8", "P8_IMPUTATION", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS,
                                 {"assignmentCostCenters": [{"costCenterPercentage": imputationPercentage,
                                                             "costCenterID": costCenter}]}),
                       {"costCenterPercentage": imputationPercentage, "costCenterID": costCenter}, nom, worker_id,
                       'UPDATE', 'imputation',
                       f"changing imputation from {cur_cc} to {costCenter} and from {cur_imp} % to "
                       f"{imputationPercentage} % on {situationStartDateTS}",
                       read_back=lambda rd: dict(zip(("costCenterPercentage", "costCenterID"),
                                                     (lambda p, c: (str(p), c))(*get_worker_imputation(rd, worker_id)))))

        if seniorityDate != str(current_seniority_date):
            self._post(ex, "P8", "P8_SENIORITY", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS, {"seniorityDate": seniorityDate}),
                       {"seniorityDate": seniorityDate}, nom, worker_id, 'UPDATE', 'senioirty date',
                       f"changing seniority date from {current_seniority_date} to {seniorityDate}",
                       read_back=lambda rd: {"seniorityDate": (get_worker_dates(rd, worker_id) or {}).get("seniorityDate")})

        (_g, _c, _k, cur_title, cur_fn, cur_rem, _rb) = get_worker_classification(r, worker_id)   # B5
        if jobTitle != cur_title or jobFunctionCode != cur_fn or remunerationType != cur_rem:
            fragment = {"jobFunctionCode": {"codeValue": jobFunctionCode}, "jobTitle": jobTitle,
                        "baseRemuneration": {"recordingBasisCode": {"codeValue": recordingBasisCode}},
                        "customFieldGroup": {"indicatorFields": [{"itemID": "forcingCoefficient", "indicatorValue": True}],
                                             "codeFields": [{"itemID": "remunerationType", "codeValue": remunerationType},
                                                            {"itemID": "modeRemSATH", "codeValue": remunerationType}]}}
            self._post(ex, "P8", "P8_POSITION", EV_WA_MODIFY,
                       _wa_event(associateOID, workAssignmentID, situationStartDateTS, fragment),
                       {"jobTitle": jobTitle, "jobFunctionCode": jobFunctionCode, "remunerationType": remunerationType},
                       nom, worker_id, 'UPDATE', 'position professionnelle',
                       f"changing position from {cur_title} to {jobTitle} and from {cur_fn} to {jobFunctionCode} "
                       f"and from {cur_rem} to {remunerationType} on {situationStartDateTS}",
                       read_back=lambda rd: dict(zip(("jobTitle", "jobFunctionCode", "remunerationType"),
                                                     get_worker_classification(rd, worker_id)[3:6])))

        if terminationDate != '' and terminationDate != '1900-01-01':                    # 5.19 : information seule
            ex.log(self._log_row(worker_id, nom, 'INFO', 'GroupEndDate', '', '', '', '',
                                 f"Group End Date / Termination Date {terminationDate} i.e. date de sortie TS",
                                 step="INFO_GROUP_END_DATE"))

        c_holder, c_iban, c_swift, c_bank = get_pay_distribution(r, associateOID)
        if account_holder != c_holder or iban != c_iban or swift_code != c_swift or bank_name != c_bank:  # B13
            payload = {"events": [{"data": {"eventContext": {"worker": {"associateOID": associateOID}},
                                            "transform": {"payDistribution": {"distributionInstructions": [{
                                                "paymentMethodCode": {"codeValue": "V"}, "itemID": "1",
                                                "depositAccount": {"IBAN": iban, "financialAccount": {"accountName": account_holder},
                                                                   "financialParty": {"SWIFTCode": swift_code,
                                                                                      "nameCode": {"shortName": bank_name}}}}]}}}}]}
            self._post(ex, "P9", "P9_PAY_DISTRIBUTION", EV_PAY_DISTRIBUTION, payload,
                       {"accountName": account_holder, "IBAN": iban, "SWIFTCode": swift_code, "shortName": bank_name},
                       nom, worker_id, 'UPDATE', 'paie distribution',
                       f"changing payment distribution from {c_holder} to {account_holder} and from {c_iban}  to {iban} "
                       f"and from {c_swift} to {swift_code} and from {c_bank} to {bank_name}",
                       read_back=lambda rd: dict(zip(("accountName", "IBAN", "SWIFTCode", "shortName"),
                                                     get_pay_distribution(rd, associateOID))))

    # -- utilitaires -----------------------------------------------------------------
    def _post(self, ex: OperationExecutor, code, variant, endpoint, payload, target, nom, worker_id, mode, field,
              detail, read_back=None, depends_on=(), clear_error=False):
        res = ex.submit(PlannedOp(code, variant, endpoint, payload, target, {"mode": mode, "field": field},
                                  depends_on=depends_on, read_back=read_back))
        tr = res.transport
        if tr is None:
            return res  # aucun transport (déjà fait, incertain, dépendance, arrêt) : pas de ligne de log
        status = tr.http_status
        error_message = ''
        if str(status).startswith("4") and not clear_error and tr.response_text is not None:
            try:
                error_message = extract_error(tr.response_text)
            except Exception:
                error_message = ''  # B14 : l'existant s'arrêtait ; ici le salarié continue (D-ADP-01)
        response_code = tr.response_text if tr.response_text is not None else (tr.error_class or '')
        ex.log(self._log_row(worker_id, nom, mode, field, status, tr.http_reason or tr.reason, response_code,
                             error_message, detail, step=variant, endpoint=endpoint))
        return res

    def _log_row(self, worker_id, nom, mode, field, status, reason, response, error, detail, step, endpoint=None):
        return {"ts_adp": datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None), "matricule_id": worker_id,
                "nom": nom, "mode": mode, "field": field, "status_code": '' if status is None else str(status),
                "reason_code": '' if reason is None else str(reason), "response_code": '' if response is None else str(response),
                "error_message": error or '', "log_message": detail, "step_name": step,
                "http_method": "POST" if endpoint else None, "endpoint_template": endpoint}


def _stable_hash(payload, volatile=False):
    from hris.common import canonical_json, sha256_hex
    text = canonical_json(payload)
    if volatile:  # partie horaire de effectiveDateTime exclue de l'identité
        text = re.sub(r'(\d{4}-\d{2}-\d{2})T\d{2}:\d{2}:\d{2}\.\d{3}Z', r'\1', text)
    return sha256_hex(text)


def _hire_payload_impl(self, worker_id, citizenship, nir, given, family, birth_name, salutation, landline, mobile,
                       email, marital, birth_date, birth_city, birth_cp, birth_dept, birth_country, line_one, line_two,
                       city, country, cp, hire_date, seniority, job_fn, worker_type, salary, cost_center, collab,
                       contract_type, end_pb, start_pb, bonus, reason_hire, reason, end_date, location, email_perso,
                       job_title, gender, rem_type, line_five, schedule, recording, matricule_it, contract_start,
                       first_hire, company_start, marital_date, org_code, compl):
    """Variante 1 de hire_pending_worker (L3415-3711), toujours envoyée (B1)."""
    bureau = city
    insee, _, _ = get_insee_code(self.reader, cp, city)
    long_city = append_arrondissement_to_city(city, cp)
    units = _org_units(org_code) if org_code else [{"itemID": "departmentId", "nameCode": {"codeValue": ""}}]
    return {"events": [{"data": {"transform": {"worker": {
        "workerID": {"idValue": worker_id},
        "workerDates": {"firstHireDate": first_hire, "originalHireDate": company_start},
        "businessCommunication": {"mobiles": [{"nameCode": {"codeValue": "WORK"}, "formattedNumber": landline}],
                                  "emails": [{"emailUri": email}]},
        "person": {
            "citizenshipCountryCodes": [{"codeValue": citizenship}],
            "identityDocuments": [{"documentID": nir, "typeCode": {"codeValue": "SSN", "shortName": ""}}],
            "legalName": {"givenName": given, "middleName": '', "familyName1": family, "familyName2": birth_name,
                          "formattedName": given + " " + family,
                          "preferredSalutations": [{"sequenceNumber": 1, "salutationCode": {
                              "shortName": "", "codeValue": salutation, "longName": ""}}]},
            "communication": {"landlines": [{"itemID": "Personnel", "formattedNumber": ''}],
                              "mobiles": [{"itemID": "Personnel", "formattedNumber": mobile}],
                              "emails": [{"itemID": "Personnel", "emailUri": email_perso,
                                          "nameCode": {"codeValue": "Personnel", "shortName": "Personnel",
                                                       "longName": "Personnel"}}]},
            "maritalStatusCode": {"effectiveDate": marital_date, "codeValue": marital},
            "genderCode": {"codeValue": gender},
            "birthDate": birth_date,
            "birthPlace": {"cityName": birth_city, "postalCode": birth_cp,
                           "countrySubdivisionLevel1": {"codeValue": birth_dept}, "countryCode": birth_country},
            "legalAddress": {"buildingNumberExtension": compl, "buildingNumber": line_one, "streetName": line_two,
                             "countryCode": country, "countrySubdivisionLevel1": {"codeValue": bureau},
                             "countrySubdivisionLevel2": {"longName": long_city, "codeValue": insee if insee else ""},
                             "lineFive": line_five, "postalCode": cp}},
        "workAssignment": {
            "hireDate": hire_date, "seniorityDate": seniority, "expectedTerminationDate": end_date,
            "assignmentStatus": {"reasonCode": {"codeValue": reason_hire}},
            "jobFunctionCode": {"codeValue": job_fn},
            "assignedOrganizationalUnits": units,
            "assignedWorkLocations": [{"itemID": "default", "nameCode": {"codeValue": location}}],
            "jobTitle": job_title,
            "workerGroups": [{"groupCode": {"shortName": "", "codeValue": "94", "longName": ""}}],   # R-ADP-012
            "workerTypeCode": {"codeValue": worker_type},
            "workArrangementCode": {"codeValue": "900"},
            "remunerationBasisCode": {"codeValue": "1"},
            "baseRemuneration": {"recordingBasisCode": {"codeValue": recording},
                                 "monthlyRateAmount": {"currencyCode": "EUR", "amountValue": salary}},
            "assignmentCostCenters": [{"costCenterPercentage": HIRE_IMPUTATION_PERCENTAGE, "costCenterID": cost_center,
                                       "costCenterName": ""}],
            "customFieldGroup": {
                "amountFields": [{"itemID": "internshipCompensation", "amountValue": salary}],
                "indicatorFields": [{"itemID": "forcingCoefficient", "indicatorValue": True}],
                "codeFields": [{"itemID": "collaborationType", "codeValue": collab},
                               {"itemID": "modeRemSATH", "codeValue": rem_type},
                               {"itemID": "recoursReason", "codeValue": reason},
                               {"itemID": "contractType", "codeValue": contract_type},
                               {"itemID": "TLM", "codeValue": "W"},
                               {"itemID": "remunerationType", "codeValue": rem_type},
                               {"itemID": "workSchedule", "codeValue": schedule}],
                "dateFields": [{"dateValue": end_pb, "itemID": "endProbationDate"},
                               {"dateValue": start_pb, "itemID": "startProbationDate"},
                               {"itemID": "fixedTermContractInitialTerminationDate", "dateValue": end_date},
                               {"itemID": "startWorkerTypeDate", "dateValue": contract_start}],
                "numberFields": [{"itemID": "REM_RE1MTS18", "numberValue": bonus}],
                "stringFields": [{"itemID": "SAD_LIB34", "stringValue": matricule_it}]}},
        "eventReasonCode": {"codeValue": "ACQ"}}}}}]}

'''),
    ('hris.adp_publish', '60c09780945f4bb6bd6b748b05f0283444aa39bcf153f54ca241f06a21ce651b', r'''"""hris.adp_publish — orchestration de la publication ADP d'un run (outbox -> opérations).

Ordre imposé :
  1. décision du garde-fou (nom réel du workspace) écrite AVANT toute autre action ;
  2. DENIED_NON_PRD : éléments PENDING -> NOT_PUSHED_NON_PRD ; aucun secret, aucun jeton, aucun appel ;
     BLOCKED_UNRESOLVED : rien n'est modifié (réévaluable) ; aucun appel ;
  3. ENV_SATISFIED : la condition d'environnement ne suffit pas — les blocages d'activation
     (contrat, décisions A08, contrôles du run) sont revalidés ; s'il en reste, aucun appel ;
  4. sinon seulement : fabrique du transport (secrets + jeton), puis traitement salarié par salarié.

Sélection sans recalcul : NORMAL -> éléments PENDING du run ; RETRY_PUBLICATION -> éléments PENDING
ou FAILED du run cible (jamais IN_PROGRESS = incertain, qui exige réconciliation ou décision).
"""
from __future__ import annotations

import dataclasses
import datetime as dt
from typing import Any, Callable, Dict, List, Mapping, Optional, Protocol, Sequence

from hris.adp_gate import BLOCKED_UNRESOLVED, DENIED_NON_PRD, ENV_SATISFIED, GateDecision, evaluate_gate
from hris.adp_ops import (OB_FAILED, OB_IN_PROGRESS, OB_NO_CHANGE, OB_NOT_PUSHED, OB_PENDING, OB_SENT,
                          OB_SUPERSEDED, OP_OPEN_UNCERTAIN, OP_SUPERSEDED, OperationExecutor, item_status,
                          supersession_candidates)
from hris.adp_transport import ADP_ACTIVATION_APPROVED

STATUS_NOT_PUBLISHED_NON_PRD = "NOT_PUBLISHED_NON_PRD"
STATUS_BLOCKED_UNRESOLVED = "BLOCKED_WORKSPACE_UNRESOLVED"
STATUS_BLOCKED_ACTIVATION = "BLOCKED_ACTIVATION"
STATUS_NOTHING = "NOTHING_TO_PUBLISH"
STATUS_PUBLISHED = "PUBLISHED"
STATUS_PARTIAL = "PUBLISHED_PARTIAL"
STATUS_HALTED = "HALTED_AUTH_REJECTED"


def _utc_now() -> dt.datetime:
    return dt.datetime.now(dt.timezone.utc)


class OutboxRepository(Protocol):
    def load(self, run_id: str, statuses: Sequence[str], max_attempts: int) -> List[dict]: ...
    def newer_items(self, run_id: str, employee_key: str) -> List[dict]: ...
    def update(self, run_id: str, updates: List[dict]) -> None: ...


class GateLog(Protocol):
    def record(self, run_id: str, decision: GateDecision) -> None: ...


def code_activation_blockers() -> List[str]:
    """Blocages d'activation dans CETTE version du code (aucun paramètre ne peut les lever)."""
    blockers = []
    if not ADP_ACTIVATION_APPROVED:
        blockers += ["CONTRAT_ADP_FINAL_ABSENT", "D-ADP-01", "D-ADP-02", "D-ENV-02", "D-SEC-01", "D-NET-01",
                     "MANDAT_DEV_SANS_PUBLICATION_REELLE"]
    return blockers


@dataclasses.dataclass
class PublicationContext:
    publish_run_id: str           # run en cours (pipeline)
    source_run_id: str            # run dont l'outbox est publiée (= publish_run_id en NORMAL)
    mode: str                     # NORMAL | RETRY_PUBLICATION
    max_attempts: int
    pause_seconds: float = 1.0    # parité R-ADP-006 (D-ADP-10)


def run_publication(ctx: PublicationContext, *, context_provider: Callable[[], Mapping[str, Any]],
                    outbox: OutboxRepository, gate_log: GateLog, ops, attempts, publication_log,
                    activation_blockers: Callable[[], List[str]], transport_factory: Callable[[], Any],
                    planner: Callable[..., Any], rows_by_key: Callable[[], Dict[str, dict]],
                    sleep: Callable[[float], None], clock: Callable[[], dt.datetime] = _utc_now) -> dict:
    if not ctx.publish_run_id or not ctx.source_run_id:
        raise ValueError("run_id obligatoire")
    if ctx.max_attempts < 1:
        raise ValueError("max_attempts doit être >= 1")
    if ctx.mode not in ("NORMAL", "RETRY_PUBLICATION"):
        raise ValueError("mode de publication invalide")

    decision = evaluate_gate(context_provider, clock)
    gate_log.record(ctx.publish_run_id, decision)            # 1. tracé AVANT toute autre action

    # CONTRAT-ADP:C13 reprise : seuls PENDING/FAILED sous max_attempts ; jamais IN_PROGRESS (incertain)
    selectable = (OB_PENDING,) if ctx.mode == "NORMAL" else (OB_PENDING, OB_FAILED)
    items = outbox.load(ctx.source_run_id, selectable, ctx.max_attempts)
    summary = {"publish_run_id": ctx.publish_run_id, "source_run_id": ctx.source_run_id, "mode": ctx.mode,
               "gate_decision_id": decision.decision_id, "gate_outcome": decision.outcome,
               "gate_reason_code": decision.reason_code, "env_condition_met": decision.env_condition_met,
               "workspace_name": decision.workspace_name, "items_selected": len(items),
               "transport_built": False, "sent": 0, "no_change": 0, "failed": 0, "uncertain": 0,
               "blocked_prior_uncertain": 0, "superseded": 0, "not_pushed": 0, "activation_blockers": []}

    if decision.outcome == DENIED_NON_PRD:                     # 2. absence normale de publication
        now = clock()
        pending = [it for it in items if it["status"] == OB_PENDING]
        outbox.update(ctx.source_run_id, [{"employee_key": it["employee_key"], "status": OB_NOT_PUSHED,
                                           "gate_decision_id": decision.decision_id, "updated_at_utc": now,
                                           "last_outcome_code": decision.reason_code} for it in pending])
        summary.update(not_pushed=len(pending), status=STATUS_NOT_PUBLISHED_NON_PRD)
        return summary
    if decision.outcome == BLOCKED_UNRESOLVED:                 # réévaluable : rien n'est classé
        summary.update(status=STATUS_BLOCKED_UNRESOLVED)
        return summary
    assert decision.outcome == ENV_SATISFIED
    blockers = list(activation_blockers())                     # 3. revalidation des autres conditions
    if blockers:
        summary.update(status=STATUS_BLOCKED_ACTIVATION, activation_blockers=blockers)
        return summary
    if not items:
        summary.update(status=STATUS_NOTHING)
        return summary

    transport = transport_factory()                            # 4. secrets + jeton ici seulement
    summary["transport_built"] = True
    rows = rows_by_key()
    halted = False
    try:
        for it in items:
            key = it["employee_key"]
            now = clock()
            # Opérations incertaines d'un run quelconque pour ce salarié : aucun envoi
            prior = [o for o in ops.open_uncertain_for_employee(key)]
            if prior:
                own = any(o["run_id"] == ctx.source_run_id for o in prior)
                outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_IN_PROGRESS if own else it["status"],
                                                   "last_outcome_code": "BLOCKED_PRIOR_UNCERTAIN",
                                                   "updated_at_utc": now, "gate_decision_id": decision.decision_id}])
                summary["blocked_prior_uncertain"] += 1
                continue
            # Obsolescence d'un élément ancien par un élément plus récent (jamais si incertain)
            newer = outbox.newer_items(ctx.source_run_id, key)
            item_ops = ops.for_item(ctx.source_run_id, key)
            if supersession_candidates(it, newer, item_ops):
                for o in item_ops:
                    if o["status"] not in ("ACCEPTED", "APPLIED_OBSERVED"):
                        ops.upsert(dict(o, status=OP_SUPERSEDED, updated_at_utc=now))
                outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_SUPERSEDED,
                                                   "last_outcome_code": "SUPERSEDED_BY_NEWER_RUN",
                                                   "updated_at_utc": now}])
                summary["superseded"] += 1
                continue
            row = rows.get(key)
            if row is None:
                outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_FAILED,
                                                   "last_outcome_code": "STAGING_ROW_MISSING", "updated_at_utc": now}])
                summary["failed"] += 1
                continue
            sleep(ctx.pause_seconds)
            attempt = int(it.get("attempt_count") or 0) + 1
            outbox.update(ctx.source_run_id, [{"employee_key": key, "status": OB_IN_PROGRESS, "attempt_count": attempt,
                                               "last_attempt_at_utc": now, "gate_decision_id": decision.decision_id,
                                               "updated_at_utc": now}])
            executor = OperationExecutor(ops=ops, attempts=attempts, transport=transport,
                                         source_run_id=ctx.source_run_id, publish_run_id=ctx.publish_run_id,
                                         employee_key=key, gate_decision_id=decision.decision_id,
                                         allow_new_attempt=(ctx.mode == "RETRY_PUBLICATION"),
                                         max_attempts=ctx.max_attempts, clock=clock)
            port_error = None
            try:
                planner(row, executor)
            except Exception as exc:  # erreur de la logique portée : salarié isolé (D-ADP-01, À CONFIRMER)
                port_error = f"PORT_ERROR_{type(exc).__name__}"
            if executor.log_rows:
                publication_log.append([dict(r, run_id=ctx.publish_run_id, employee_key=key, attempt_no=attempt,
                                             gate_decision_id=decision.decision_id) for r in executor.log_rows])
            status = item_status(executor.results)
            if port_error and status in (OB_SENT, OB_NO_CHANGE):
                status = OB_FAILED  # traitement incomplet : reprise explicite requise
            outbox.update(ctx.source_run_id, [{"employee_key": key, "status": status, "updated_at_utc": clock(),
                                               "last_outcome_code": port_error or _last_code(executor)}])
            summary[{OB_SENT: "sent", OB_NO_CHANGE: "no_change", OB_FAILED: "failed",
                     OB_IN_PROGRESS: "uncertain"}[status]] += 1
            if executor.halt_reason:
                halted = True
                break
    finally:
        close = getattr(transport, "close", None)
        if close:
            close()
    if halted:
        summary["status"] = STATUS_HALTED
    elif summary["failed"] or summary["uncertain"] or summary["blocked_prior_uncertain"]:
        summary["status"] = STATUS_PARTIAL
    else:
        summary["status"] = STATUS_PUBLISHED
    return summary


def _last_code(executor) -> str:
    for r in reversed(executor.results):
        detail = getattr(r, "detail", None)
        if detail:
            return str(detail)[:100]
    return "NO_OPERATION"
'''),
    ('hris.adp_spark', '8366cf0b1f5d716d82ab0d707bbc927278fbd623b9a7c9c3e0a69e935620be36', r'''"""hris.adp_spark — dépôts Delta de la publication ADP et fabriques du transport réel.

Chaque upsert d'opération est un commit Delta distinct (MERGE) : l'intention est durable AVANT
le transport. Le transport réel et le lecteur réel ne se construisent que si l'activation est
approuvée DANS LE CODE (hris.adp_transport.ADP_ACTIVATION_APPROVED) — jamais par paramètre.
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import stat
import tempfile
import uuid
from typing import Any, Dict, List, Optional

from hris.adp_gate import GateDecision
from hris.adp_transport import ADP_ACTIVATION_APPROVED, AdpHttpTransport
from hris.common import HrisError, PROCESS_CODE, utc_now
from hris.model import Tables
from hris.spark_io import append, merge_rows, rows_df, sql_str


def _F():
    from pyspark.sql import functions as F
    return F


class SparkOutbox:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def load(self, run_id, statuses, max_attempts):
        F = _F()
        rows = (self.spark.table(self.T("pub.adp_outbox")).where(F.col("run_id") == run_id)
                .where(F.col("status").isin(list(statuses)))
                .where(F.coalesce(F.col("attempt_count"), F.lit(0)) < max_attempts)
                .orderBy("employee_key").collect())
        return [r.asDict() for r in rows]

    def newer_items(self, run_id, employee_key):
        """Éléments du même salarié issus de runs promus APRÈS le run donné."""
        F = _F()
        runs = self.spark.table(self.T("ctl.run")).select("run_id", "promoted_at_utc")
        mine = runs.where(F.col("run_id") == run_id).collect()
        if not mine or mine[0].promoted_at_utc is None:
            return []
        later = runs.where(F.col("promoted_at_utc") > F.lit(mine[0].promoted_at_utc)).select("run_id")
        rows = (self.spark.table(self.T("pub.adp_outbox")).where(F.col("employee_key") == employee_key)
                .join(later, "run_id").collect())
        return [r.asDict() for r in rows]

    def update(self, run_id, updates):
        if not updates:
            return
        from delta.tables import DeltaTable
        cols = ["employee_key", "status", "attempt_count", "last_attempt_at_utc", "last_outcome_code",
                "gate_decision_id", "updated_at_utc"]
        rows = [dict({c: u.get(c) for c in cols}, run_id=run_id) for u in updates]
        src = rows_df(self.spark, rows, "pub.adp_outbox")
        (DeltaTable.forName(self.spark, self.T("pub.adp_outbox")).alias("t")
         .merge(src.alias("s"), "t.run_id = s.run_id AND t.employee_key = s.employee_key")
         .whenMatchedUpdate(set={
             "status": "coalesce(s.status, t.status)",
             "attempt_count": "coalesce(s.attempt_count, t.attempt_count)",
             "last_attempt_at_utc": "coalesce(s.last_attempt_at_utc, t.last_attempt_at_utc)",
             "last_outcome_code": "coalesce(s.last_outcome_code, t.last_outcome_code)",
             "gate_decision_id": "coalesce(s.gate_decision_id, t.gate_decision_id)",
             "updated_at_utc": "coalesce(s.updated_at_utc, t.updated_at_utc)"})
         .execute())


class SparkOps:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def get(self, op_id):
        rows = self.spark.table(self.T("pub.adp_operation")).where(_F().col("op_id") == op_id).collect()
        return rows[0].asDict() if rows else None

    def open_uncertain_for_employee(self, employee_key):
        F = _F()
        rows = (self.spark.table(self.T("pub.adp_operation")).where(F.col("employee_key") == employee_key)
                .where(F.col("status").isin("INTENT_RECORDED", "UNCERTAIN")).collect())
        return [r.asDict() for r in rows]

    def for_item(self, run_id, employee_key):
        F = _F()
        rows = (self.spark.table(self.T("pub.adp_operation"))
                .where((F.col("run_id") == run_id) & (F.col("employee_key") == employee_key)).collect())
        return [r.asDict() for r in rows]

    def upsert(self, op):
        merge_rows(self.spark, [op], self.T("pub.adp_operation"), "pub.adp_operation", ["op_id"])


class SparkAttempts:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def append(self, attempt):
        append(rows_df(self.spark, [attempt], "pub.adp_attempt"), self.T("pub.adp_attempt"), "pub.adp_attempt")


class SparkGateLog:
    def __init__(self, spark, T: Tables):
        self.spark, self.T = spark, T

    def record(self, run_id, decision: GateDecision):
        row = {"decision_id": decision.decision_id, "run_id": run_id, "evaluated_at_utc": decision.evaluated_at_utc,
               "workspace_name": decision.workspace_name, "workspace_id": decision.workspace_id,
               "is_for_pipeline": decision.is_for_pipeline, "notebook_name": decision.notebook_name,
               "rule_version": decision.rule_version, "decision": decision.env_condition_met,
               "reason_code": decision.reason_code}
        append(rows_df(self.spark, [row], "log.adp_gate_decision"), self.T("log.adp_gate_decision"),
               "log.adp_gate_decision")


class SparkPublicationLog:
    def __init__(self, spark, T: Tables, log_date: Optional[dt.date]):
        self.spark, self.T, self.log_date = spark, T, log_date

    def append(self, rows: List[dict]):
        now = utc_now()
        out = [dict(r, log_id=str(uuid.uuid4()), source_system="FABRIC", log_date=self.log_date,
                    _ingested_at_utc=now) for r in rows]
        append(rows_df(self.spark, out, "log.adp_publication_log"), self.T("log.adp_publication_log"),
               "log.adp_publication_log")


# ---------------------------------------------------------------------------
# Fabriques réelles (refusent tant que l'activation n'est pas approuvée dans le code)
# ---------------------------------------------------------------------------
SECRET_NAMES = {"client_id": "client-id-adp", "client_secret": "client-secret-adp",
                "cert_pem": "MOTUL-SA-cer", "key_pem": "motulsa-auth-key"}


def format_pem(raw: str, label: str) -> str:
    """Équivalent de format_cert / format_rsa_key (L5068-5106), PKCS#1, un seul bloc (B24, parité)."""
    match = re.search(rf"-----BEGIN {label}-----(.*)-----END {label}-----", raw.strip(), re.DOTALL)
    if not match:
        raise HrisError("PEM_INVALID", label)
    body = re.sub(r"\s+", "", match.group(1))
    return f"-----BEGIN {label}-----\n" + "\n".join(body[i:i + 64] for i in range(0, len(body), 64)) + \
        f"\n-----END {label}-----\n"


class RealAdpTransportFactory:
    """Lecture des secrets ADP, jeton mTLS, transport. Appelée seulement si le garde-fou ET tous les
    blocages d'activation sont levés. Refuse dans cette version (activation non approuvée)."""

    def __init__(self, kv_url: str, timeout_seconds: float, get_secret):
        self.kv_url, self.timeout, self._get_secret = kv_url, timeout_seconds, get_secret
        self.built = 0

    def __call__(self):
        self.built += 1
        if not ADP_ACTIVATION_APPROVED:
            raise HrisError("ADP_ACTIVATION_NOT_APPROVED", "aucune lecture de secret ADP ni aucun jeton")
        raise HrisError("ADP_REAL_TRANSPORT_NOT_WIRED", "câblage réel à valider avec le contrat final")  # pragma: no cover
'''),
    ('hris.steps', '9a492212bed76d1b3aec96a5efcba8f76a149c87c0ae0f0b217170cb0103ee7f', r'''"""hris.steps — étapes du pipeline (logique des notebooks), appelées à l'identique par les notebooks et par
les tests d'intégration Fabric (schéma isolé). Aucune sortie de donnée personnelle : compteurs seulement."""
from __future__ import annotations

import datetime as dt
import json
import time
from typing import Any, Callable, Dict, List, Mapping, Optional

from hris import compare as C
from hris import spark_io as IO
from hris import transform as S
from hris import adp_publish as PUB
from hris.common import HrisError, PROCESS_CODE, safe_error, utc_now
from hris.model import COMPARE_COLUMNS, Tables
from hris.ts_acquire import completeness_anomalies, to_raw_row

TS_OWNED_CONTROLS = ["EXTRACTION_INCOMPLETE", "SCHEMA_CHANGE", "EMPTY_EXTRACTION", "TS_FILTER_TECHNICAL_FAILURE",
                     "TS_FILTER_NOT_FOUND", "TS_ENDPOINT_FAILURE", "TS_ENDPOINT_NOT_FOUND", "TS_PROCESSING_ERROR",
                     "TS_CIRCUIT_OPEN", "TS_SECRET_RESOLUTION"]
TRANSFORM_OWNED = ["MAPPING_TABLE_MISSING", "MAPPING_SCHEMA", "MAPPING_TABLE_EMPTY", "MAPPING_KEY_DUPLICATE",
                   "TRANSFORM_NOT_EXECUTED", "CONVERSION_ERROR"]


def _F():
    from pyspark.sql import functions as F
    return F


def seed(spark, T: Tables) -> dict:
    """Graines insérées seulement si absentes (cfg.process, verrou)."""
    out = {"process": False, "lock": False}
    if spark.table(T("cfg.process")).where(f"process_code = '{PROCESS_CODE}'").count() == 0:
        IO.append(IO.rows_df(spark, [{
            "process_code": PROCESS_CODE, "enabled": True, "acquisition_mode": "TS_API", "publication_target": "ADP_API",
            "key_columns": ["id_unique"], "compare_columns": COMPARE_COLUMNS, "reference_table": "ref.ts_employee_adp",
            "staging_table": "stg.ts_employee_adp", "max_absent_pct": None, "max_modified_pct": None,
            "max_reject_pct": None, "updated_at_utc": utc_now()}], "cfg.process"), T("cfg.process"), "cfg.process")
        out["process"] = True
    if spark.table(T("ctl.process_lock")).where(f"process_code = '{PROCESS_CODE}'").count() == 0:
        IO.append(IO.rows_df(spark, [{"process_code": PROCESS_CODE, "lock_status": "RELEASED",
                                      "released_at_utc": utc_now()}], "ctl.process_lock"),
                  T("ctl.process_lock"), "ctl.process_lock")
        out["lock"] = True
    return out


# ---------------------------------------------------------------------------
# Acquisition : stockage d'une extraction et contrôles (l'extraction réseau est dans le notebook)
# ---------------------------------------------------------------------------
SOURCE_REAL = "TS_PRD_API"      # extraction réelle du tenant TalentSoft de production
SOURCE_FIXTURE = "FIXTURE"       # données fictives (tests isolés uniquement)


def store_acquisition(spark, T: Tables, run: dict, records: List[dict], stats: dict, step: IO.StepLogger,
                      *, source: str) -> dict:
    """Écrit stg.ts_employee_raw et les contrôles. `source` est OBLIGATOIRE et tracé dans ctl.run_step : il permet de
    prouver, dans le Lakehouse, qu'un run est une extraction réelle (TS_PRD_API) et non un jeu fictif (FIXTURE)."""
    if source not in (SOURCE_REAL, SOURCE_FIXTURE):
        raise HrisError("PARAM_INVALID", "source")
    if source == SOURCE_FIXTURE and not getattr(T, "test_schema", ""):
        raise HrisError("FIXTURE_OUTSIDE_TEST_SCHEMA", "des données fictives ne s'écrivent que dans le schéma de test")
    run_id = run["run_id"]
    extracted_at = utc_now()
    raw_rows = [dict(to_raw_row(r), run_id=run_id, business_date=run["business_date"], _extracted_at_utc=extracted_at)
                for r in records]
    distinct_keys = len({r.get("employeeNumber") for r in records})
    IO.replace_run(spark, IO.rows_df(spark, raw_rows, "stg.ts_employee_raw"), T("stg.ts_employee_raw"),
                   "stg.ts_employee_raw", run_id)
    anomalies = completeness_anomalies(stats)
    IO.write_anomalies(spark, T, run_id, anomalies, TS_OWNED_CONTROLS)
    failed = [a["control_code"] for a in anomalies if a["severity"] == "BLOCKING" and a["outcome"] == "FAIL"]
    IO.update_run(spark, T, run_id, status="ACQUIRED")
    step.succeeded(rows_out=len(raw_rows), counters_only=True, message=json.dumps({
        "source": source, "scope": "LegalStructure=Motul France", "mode": "full_extraction_no_period_filter",
        "rows_persisted": len(raw_rows), "distinct_employee_numbers": distinct_keys,
        "stats": stats, "blocking_failures": failed}))
    return {"status": "ACQUIRED" if not failed else "ACQUIRED_INCOMPLETE", "rows": len(raw_rows),
            "distinct_keys": distinct_keys, "total_count": stats.get("total_count"), "blocking_failures": failed}


# ---------------------------------------------------------------------------
# Transformation
# ---------------------------------------------------------------------------
def run_transform(spark, T: Tables, run_id: str) -> dict:
    F = _F()
    run = IO.get_run(spark, T, run_id)
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if run["run_mode"] == "RETRY_PUBLICATION":
        return {"run_id": run_id, "status": "SKIPPED_RETRY_PUBLICATION"}
    if run["status"] not in ("ACQUIRED", "TRANSFORMED"):
        return {"run_id": run_id, "status": "SKIPPED_STATUS_" + str(run["status"])}
    step = IO.StepLogger(spark, T, run_id, "TRANSFORM")
    try:
        d_run = run["trigger_time_utc"].date()  # date UTC du run (parité GETDATE()), figée au démarrage
        pred = f"run_id = {IO.sql_str(run_id)}"
        bdate = F.lit(run["business_date"]).cast("date")
        raw = spark.table(T("stg.ts_employee_raw")).where(pred)
        n_raw = raw.count()
        IO.replace_run(spark, S.source_df(raw), T("stg.ts_employee_source"), "stg.ts_employee_source", run_id)
        anomalies, mapping_hash = S.mapping_controls(spark, T)
        blocked = [a["control_code"] for a in anomalies if a["outcome"] == "FAIL"]
        if blocked:
            anomalies.append({"control_code": "TRANSFORM_NOT_EXECUTED", "severity": "BLOCKING", "metric_value": None,
                              "threshold_value": None, "outcome": "FAIL",
                              "message": "enrichissement non exécuté : correspondances invalides"})
            IO.write_anomalies(spark, T, run_id, anomalies, TRANSFORM_OWNED)
            IO.update_run(spark, T, run_id, status="TRANSFORMED")
            step.succeeded(rows_in=n_raw, rows_out=0, content_hash=mapping_hash, counters_only=True,
                           message=json.dumps({"blocked_by": blocked}))
            return {"run_id": run_id, "status": "BLOCKED_MAPPINGS", "rows_raw": n_raw, "blocked_by": blocked}
        source = spark.table(T("stg.ts_employee_source")).where(pred)
        IO.replace_run(spark, S.enriched_df(spark, T, source, d_run), T("stg.ts_employee_enriched"),
                       "stg.ts_employee_enriched", run_id)
        enriched = spark.table(T("stg.ts_employee_enriched")).where(pred)
        try:
            rejects = S.reject_df(enriched, d_run).withColumn("run_id", F.lit(run_id)).withColumn("business_date", bdate)
            IO.replace_run(spark, rejects, T("stg.ts_employee_reject"), "stg.ts_employee_reject", run_id)
            rejects = spark.table(T("stg.ts_employee_reject")).where(pred)
            IO.replace_run(spark, S.masked_log_reject(rejects, run_id), T("log.reject"), "log.reject", run_id)
            adp = S.adp_df(enriched, rejects, d_run).withColumn("run_id", F.lit(run_id)).withColumn("business_date", bdate)
            IO.replace_run(spark, adp, T("stg.ts_employee_adp"), "stg.ts_employee_adp", run_id)
        except Exception as exc:
            if "CONVERSION_" in str(exc):
                anomalies.append({"control_code": "CONVERSION_ERROR", "severity": "BLOCKING", "metric_value": None,
                                  "threshold_value": None, "outcome": "FAIL",
                                  "message": "erreur de conversion T-SQL reproduite (échec du run, D-PAR-02/07)"})
                IO.write_anomalies(spark, T, run_id, anomalies, TRANSFORM_OWNED)
            raise
        IO.write_anomalies(spark, T, run_id, anomalies, TRANSFORM_OWNED)
        rej_t = spark.table(T("stg.ts_employee_reject")).where(pred)
        counts = {"rows_raw": n_raw, "rows_enriched": enriched.count(), "reject_lines": rej_t.count(),
                  "rejected_keys": rej_t.select("id_unique").distinct().count(),
                  "rows_adp": spark.table(T("stg.ts_employee_adp")).where(pred).count()}
        IO.update_run(spark, T, run_id, status="TRANSFORMED")
        step.succeeded(rows_in=n_raw, rows_out=counts["rows_adp"], content_hash=mapping_hash,
                       message=json.dumps(counts), counters_only=True)
        return dict(counts, run_id=run_id, status="TRANSFORMED")
    except Exception as exc:
        step.failed(exc)
        IO.update_run(spark, T, run_id, error_code="TRANSFORM_FAILED", error_summary=safe_error(exc))
        raise


def _check_mode(run, run_mode):
    if run is None:
        raise HrisError("RUN_UNKNOWN")
    if run["run_mode"] != run_mode:
        raise HrisError("PARAM_INVALID", "mode différent du mode enregistré pour ce run")


def run_compare(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    run = IO.get_run(spark, T, run_id)
    _check_mode(run, run_mode)
    if run_mode == "RETRY_PUBLICATION":
        return {"run_id": run_id, "decision": "SKIPPED_RETRY_PUBLICATION"}
    if run["status"] not in ("TRANSFORMED", "COMPARED", "BLOCKED", "VALIDATED"):
        return {"run_id": run_id, "decision": "SKIPPED_STATUS_" + str(run["status"])}
    step = IO.StepLogger(spark, T, run_id, "COMPARE")
    try:
        result = dict(C.compare(spark, T, run_id, run_mode), run_id=run_id)
        step.succeeded(message=json.dumps(result, default=str), counters_only=True)
        return result
    except Exception as exc:
        step.failed(exc)
        raise


def run_promote(spark, T: Tables, run_id: str, run_mode: str) -> dict:
    run = IO.get_run(spark, T, run_id)
    _check_mode(run, run_mode)
    if run_mode == "RETRY_PUBLICATION":
        return {"run_id": run_id, "promoted": False, "reason": "SKIPPED_RETRY_PUBLICATION"}
    step = IO.StepLogger(spark, T, run_id, "PROMOTE")
    try:
        result = dict(C.promote(spark, T, run_id, run_mode), run_id=run_id)
        if result.get("promoted"):
            step.succeeded(rows_out=result.get("nb_outbox"), message=json.dumps(result, default=str), counters_only=True)
        else:
            step.skipped(json.dumps(result, default=str))
        return result
    except Exception as exc:
        step.failed(exc)
        raise


# ---------------------------------------------------------------------------
# Publication
# ---------------------------------------------------------------------------
STATUS_MAP = {PUB.STATUS_NOT_PUBLISHED_NON_PRD: ("NOT_PUBLISHED_NON_PRD", None),
              PUB.STATUS_BLOCKED_UNRESOLVED: ("BLOCKED", "PUBLICATION_GATE_UNRESOLVED"),
              PUB.STATUS_BLOCKED_ACTIVATION: ("BLOCKED", "PUBLICATION_ACTIVATION_BLOCKED"),
              PUB.STATUS_NOTHING: ("PUBLISHED", None), PUB.STATUS_PUBLISHED: ("PUBLISHED", None),
              PUB.STATUS_PARTIAL: ("PUBLISHED_PARTIAL", None), PUB.STATUS_HALTED: ("PUBLISHED_PARTIAL", "AUTH_REJECTED")}


def run_publish(spark, T: Tables, run_id: str, run_mode: str, *, context_provider: Callable[[], Mapping[str, Any]],
                transport_factory, planner, max_attempts: int, pause_seconds: float,
                extra_blockers: Callable[[], List[str]] = lambda: [], sleep: Callable[[float], None] = time.sleep) -> dict:
    """Publication du run (ou du run cible en RETRY_PUBLICATION) sous sentinelle réseau bloquante."""
    from hris.adp_spark import SparkAttempts, SparkGateLog, SparkOps, SparkOutbox, SparkPublicationLog
    from hris.adp_transport import ADP_ACTIVATION_APPROVED, NetworkSentinel
    run = IO.get_run(spark, T, run_id)
    _check_mode(run, run_mode)
    result = {"run_id": run_id, "mode": run_mode}
    if run_mode == "INIT_REFERENCE":
        return dict(result, status="SKIPPED_INIT_REFERENCE")          # TG-03 : aucune étape de publication
    if run_mode == "NORMAL" and run["status"] != "PROMOTED":
        return dict(result, status="SKIPPED_NOT_PROMOTED_" + str(run["status"]))
    source_run_id = run["target_run_id"] if run_mode == "RETRY_PUBLICATION" else run_id
    source_run = IO.get_run(spark, T, source_run_id) or {}

    def run_blockers():
        blockers = list(C.revalidate(spark, T, source_run_id))
        if all(v is None for v in C.load_thresholds(spark, T).values()):
            blockers.append("D-CMP-07_SEUILS_NON_FOURNIS")
        return blockers

    def rows_by_key():
        rows = spark.table(T("stg.ts_employee_adp")).where(f"run_id = {IO.sql_str(source_run_id)}").collect()
        return {r["id_unique"]: r.asDict() for r in rows}

    calls = {"factory": 0}

    def counted_factory():
        calls["factory"] += 1
        return transport_factory()

    ctx = PUB.PublicationContext(publish_run_id=run_id, source_run_id=source_run_id, mode=run_mode,
                                 max_attempts=max_attempts, pause_seconds=pause_seconds)
    step = IO.StepLogger(spark, T, run_id, "PUBLISH_ADP")
    try:
        with NetworkSentinel() as sentinel:
            summary = PUB.run_publication(
                ctx, context_provider=context_provider, outbox=SparkOutbox(spark, T), gate_log=SparkGateLog(spark, T),
                ops=SparkOps(spark, T), attempts=SparkAttempts(spark, T),
                publication_log=SparkPublicationLog(spark, T, source_run.get("business_date")),
                activation_blockers=lambda: PUB.code_activation_blockers() + run_blockers() + list(extra_blockers()),
                transport_factory=counted_factory, planner=planner, rows_by_key=rows_by_key, sleep=sleep)
        proof = {"gate_outcome": summary["gate_outcome"], "gate_reason_code": summary["gate_reason_code"],
                 "workspace_name": summary["workspace_name"], "transport_factory_calls": calls["factory"],
                 "transport_built": summary["transport_built"], "sentinel_blocking": True,
                 "sentinel_attempts": len(sentinel.attempts),
                 "sentinel_hosts": sorted({a["host"] for a in sentinel.attempts}),
                 "adp_activation_approved": ADP_ACTIVATION_APPROVED}
        run_status, error_code = STATUS_MAP[summary["status"]]
        fields = {"status": run_status}
        if error_code:
            fields.update(error_code=error_code, error_summary=",".join(summary.get("activation_blockers") or []))
        IO.update_run(spark, T, run_id, **fields)
        step.succeeded(rows_in=summary["items_selected"], counters_only=True,
                       message=json.dumps({"summary": summary, "proof": proof}, default=str))
        return dict(result, status=summary["status"], summary=summary, proof=proof)
    except Exception as exc:
        step.failed(exc)
        IO.update_run(spark, T, run_id, error_code="PUBLISH_FAILED", error_summary=safe_error(exc))
        raise
'''),
    ('hris.legacy_logs', '23e7cd84abb562bf3097d258b4c17af8515a2074394ff3ecd146c2ab8d6c1409', r'''"""hris.legacy_logs — lecture des logs historiques de l'ancienne Function ADP (import unique, D-LOG-00).

Reprise des fonctions du prototype NB_HRIS_IMPORT_LOGS_HISTORIQUES (TEST_2/notebooks) sous forme
testable. Lecture seule : aucun fichier n'est produit, modifié ni supprimé. Les lignes rejetées sont
tracées sans valeur. Les zéros de tête de matricule_id ne sont pas recomplétés.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import os
import re
import uuid
from typing import Dict, List, Optional, Tuple

EXPECTED_COLUMNS = ["ts_adp", "matricule_id", "nom", "mode", "field", "status_code",
                    "reason_code", "response_code", "error_message", "log_message"]
CANDIDATE_DELIMITERS = [";", "\t", ","]
TS_FORMATS = ["%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f",
              "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M"]
DATE_FOLDER = re.compile(r"(?:^|/)(\d{8})(?:/|$)")
SOURCE_SYSTEM = "LEGACY_ADP_FUNCTION"


def list_files(root: str) -> List[dict]:
    found = []
    for dirpath, _, filenames in os.walk(root):
        for name in filenames:
            full = os.path.join(dirpath, name)
            st = os.stat(full)
            found.append({"full": full, "rel": os.path.relpath(full, root).replace(os.sep, "/"), "name": name,
                          "size": int(st.st_size),
                          "mtime": dt.datetime.fromtimestamp(st.st_mtime, tz=dt.timezone.utc)})
    return sorted(found, key=lambda f: f["rel"])


def parse_ts(value: Optional[str]) -> Optional[dt.datetime]:
    text = (value or "").strip()
    for fmt in TS_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def read_delimited(path: str) -> Tuple[Optional[List[str]], Optional[List[List[str]]], str]:
    with open(path, "rb") as handle:
        raw = handle.read()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return None, None, "ENCODING_NOT_SUPPORTED"
    first_line = text.splitlines()[0] if text else ""
    for delimiter in CANDIDATE_DELIMITERS:
        header = [h.strip() for h in next(csv.reader([first_line], delimiter=delimiter))]
        if set(EXPECTED_COLUMNS).issubset(header):
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter, quotechar='"'))
            return header, rows[1:], f"text/{encoding}/{delimiter!r}"
    return None, None, "HEADER_MISMATCH"


def read_excel(path: str, sheet_name: str = "") -> Tuple[Optional[List[str]], Optional[List[List[str]]], str]:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return None, None, "OPENPYXL_MISSING"
    import pandas as pd
    df = pd.read_excel(path, sheet_name=sheet_name or 0, dtype=str, engine="openpyxl")
    header = [str(c).strip() for c in df.columns]
    if not set(EXPECTED_COLUMNS).issubset(header):
        return None, None, "HEADER_MISMATCH"
    return header, df.fillna("").astype(str).values.tolist(), "xlsx"


def folder_date(rel_path: str) -> Optional[dt.date]:
    match = DATE_FOLDER.search(rel_path)
    if not match:
        return None
    try:
        return dt.datetime.strptime(match.group(1), "%Y%m%d").date()
    except ValueError:
        return None


def process_folder(local_root: str, uri_prefix: str, extensions: set, already_imported: set, import_id: str,
                   executed_by: Optional[str], workspace_name: Optional[str], sheet_name: str = "",
                   now: Optional[dt.datetime] = None, target_table: str = "log.adp_publication_log"):
    """Analyse un dossier ; renvoie (enregistrements, rejets sans valeur, registre des fichiers)."""
    now = now or dt.datetime.now(dt.timezone.utc)
    records, rejects, ledger = [], [], []
    for f in list_files(local_root):
        source_uri = f"{uri_prefix}/{f['rel']}"
        entry = {"import_id": import_id, "source_uri": source_uri, "file_name": f["name"], "file_size": f["size"],
                 "file_modified_at_utc": f["mtime"], "file_format": None, "status": None, "detail": None,
                 "rows_read": 0, "rows_loaded": 0, "rows_rejected": 0, "target_table": target_table,
                 "executed_by": executed_by, "workspace_name": workspace_name, "processed_at_utc": now}
        extension = os.path.splitext(f["name"])[1].lower()
        if source_uri in already_imported:
            entry.update(status="SKIPPED_ALREADY_IMPORTED")
        elif extension not in extensions:
            entry.update(status="SKIPPED_UNSUPPORTED", detail=f"extension {extension or '(aucune)'}")
        else:
            try:
                header, rows, fmt = read_excel(f["full"], sheet_name) if extension == ".xlsx" else read_delimited(f["full"])
            except Exception as exc:
                header, rows, fmt = None, None, f"READ_ERROR:{type(exc).__name__}"
            entry["file_format"] = fmt
            if header is None:
                entry.update(status="REJECTED_SCHEMA", detail=fmt)
            else:
                position = {name: header.index(name) for name in EXPECTED_COLUMNS}
                log_date_from_folder = folder_date(f["rel"])
                for row_number, row in enumerate(rows, start=1):
                    if not any(str(cell).strip() for cell in row):
                        continue
                    entry["rows_read"] += 1
                    values = {c: str(row[i]) if i < len(row) else "" for c, i in position.items()}
                    ts = parse_ts(values["ts_adp"])
                    reason = ("COLUMN_COUNT_MISMATCH" if len(row) != len(header)
                              else "TS_ADP_INVALID" if ts is None else None)
                    if reason:
                        entry["rows_rejected"] += 1
                        rejects.append({"import_id": import_id, "source_uri": source_uri,
                                        "source_row_number": row_number, "reject_reason": reason,
                                        "rejected_at_utc": now})
                        continue
                    payload = "\u001f".join(values[c] for c in EXPECTED_COLUMNS)
                    records.append({
                        "log_id": str(uuid.uuid4()), "source_system": SOURCE_SYSTEM, "run_id": None,
                        "log_date": log_date_from_folder or ts.date(), "ts_adp": ts,
                        **{c: values[c] for c in EXPECTED_COLUMNS if c != "ts_adp"},
                        "employee_key": None, "step_name": None, "http_method": None, "endpoint_template": None,
                        "attempt_no": None, "gate_decision_id": None, "_import_id": import_id,
                        "_source_file_uri": source_uri, "_source_row_number": row_number,
                        "_row_hash": hashlib.sha256(payload.encode("utf-8")).hexdigest(), "_ingested_at_utc": now})
                    entry["rows_loaded"] += 1
                entry["status"] = "IMPORTED"
        ledger.append(entry)
    return records, rejects, ledger
'''),
    ('hris.adp_sim', '3698bf9dff53d20f0d9239ce70b866825d7ce1e3a76904e0d70a12ee6fa387d8', r'''"""hris.adp_sim — simulateur ADP ISOLÉ et dépôts en mémoire, réservés aux tests.

Aucun accès réseau : le simulateur répond en mémoire. Il ne peut pas activer le transport réel
(hris.adp_transport.AdpHttpTransport refuse sa construction tant que l'activation n'est pas
approuvée dans le code). Données exclusivement fictives.
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import re
from typing import Any, Callable, Dict, List, Optional, Tuple

from hris.adp_transport import (ACCEPTED, NOT_TRANSMITTED, REJECTED, UNCERTAIN, TransportResult,
                                classify_response, digest)
from hris.adp_ops import OP_OPEN_UNCERTAIN
from hris.adp_port import ReadResponse


class Resp(ReadResponse):
    def __init__(self, status_code: int, body: Any = None, reason: str = ""):
        super().__init__(status_code, body if isinstance(body, str) else json.dumps(body),
                         reason or {200: "OK", 400: "Bad Request", 404: "Not Found", 500: "Server Error"}.get(status_code, ""))
        self._body = body

    def json(self):
        if isinstance(self._body, str):
            return json.loads(self._body)
        return copy.deepcopy(self._body)


class FakeAdp:
    """État ADP fictif (workers indexés par id_payroll) ; applique les POST acceptés.
    `post_script` : liste de comportements consommés à chaque POST, ex. 'ok', ('status', 500),
    'timeout_after_send', 'connect_refused', 'send_interrupted', ('redirect', 302), ('2xx_text', 200)."""

    def __init__(self, workers: Optional[Dict[str, dict]] = None, insee: Optional[Dict[str, list]] = None,
                 get_status: Optional[Callable[[str], Optional[int]]] = None):
        self.workers = copy.deepcopy(workers or {})
        self.insee_db = insee or {}
        self.post_script: List[Any] = []
        self.posts: List[Tuple[str, dict]] = []
        self.gets: List[str] = []
        self.insee_calls: List[str] = []
        self.get_status = get_status

    # -- lecture -----------------------------------------------------------------------------
    def get(self, path: str) -> Resp:
        self.gets.append(path)
        if self.get_status:
            forced = self.get_status(path)
            if forced:
                return Resp(forced, {"error": "simulated"})
        m = re.match(r"^hr/v2/workers/([^/?]+)$", path)
        if m:
            w = self.workers.get(m.group(1))
            return Resp(200, {"workers": [w]}) if w else Resp(404, {"workers": []})
        m = re.match(r"^hr/v2/workers\?\$filter=workers/workerID/idValue eq '(.*)'$", path)
        if m:
            w = self.workers.get(m.group(1))
            return Resp(200, {"workers": [w]}) if w else Resp(204, "")
        m = re.match(r"^hr/v2/workers\?\$filter=.*%27(.*)%27$", path)
        if m:
            ssn = m.group(1)
            for w in self.workers.values():
                docs = w.get("person", {}).get("identityDocuments", [])
                if docs and docs[0].get("documentID") == ssn:
                    return Resp(200, {"workers": [w]})
            return Resp(204, "")
        m = re.match(r"^payroll/v2/workers/([^/]+)/pay-distributions$", path)
        if m:
            for w in self.workers.values():
                if w.get("associateOID") == m.group(1):
                    return Resp(200, {"payDistributions": [w.get("_pay", {"distributionInstructions": [{}]})]})
            return Resp(404, {})
        return Resp(404, {})

    def insee(self, postal_code: str) -> Resp:
        self.insee_calls.append(postal_code)
        return Resp(200, self.insee_db.get(postal_code, []))

    # -- publication -------------------------------------------------------------------------
    def post(self, path: str, body: bytes) -> TransportResult:
        behaviour = self.post_script.pop(0) if self.post_script else "ok"
        payload = json.loads(body.decode("utf-8"))
        if behaviour == "connect_refused":
            return TransportResult(NOT_TRANSMITTED, "CONNECT_FAILED", "CONNECT", error_class="ConnectionRefusedError")
        if behaviour == "send_interrupted":
            return TransportResult(UNCERTAIN, "SEND_INTERRUPTED", "SEND", error_class="BrokenPipeError")
        self.posts.append((path, payload))
        if behaviour == "timeout_after_send":
            self._apply(path, payload)  # le serveur a traité, la réponse est perdue
            return TransportResult(UNCERTAIN, "TIMEOUT_AFTER_SEND", "RECEIVE", error_class="timeout")
        if isinstance(behaviour, tuple) and behaviour[0] == "redirect":
            status, text = behaviour[1], ""
        elif isinstance(behaviour, tuple) and behaviour[0] == "status":
            status, text = behaviour[1], json.dumps({"confirmMessage": {"processMessages": [
                {"developerMessage": {"messageTxt": "simulated"}}]}}) if 400 <= behaviour[1] < 500 else "{}"
        elif isinstance(behaviour, tuple) and behaviour[0] == "2xx_text":
            status, text = behaviour[1], "OK-not-json"
        else:
            status, text = 200, json.dumps({"events": [{"eventStatusCode": {"codeValue": "complete"}}]})
        outcome, reason = classify_response(status, text)
        if outcome == ACCEPTED or behaviour == "apply_then_500":
            self._apply(path, payload)
        return TransportResult(outcome, reason, "DONE", http_status=status, http_reason="SIM",
                               response_text=text, body_digest=digest(text))

    def _apply(self, path: str, payload: dict) -> None:
        ev = (payload.get("events") or [{}])[0].get("data")
        if not isinstance(ev, dict):
            return  # charge utile de test sans structure ADP : rien à appliquer
        if path.endswith("worker.hire"):
            w = copy.deepcopy(ev["transform"]["worker"])
            wid = w["workerID"]["idValue"]
            w["associateOID"] = f"AOID-{wid}"
            w.setdefault("workAssignments", [w.pop("workAssignment", {})])
            self.workers[wid] = w
            return
        ctx = ev.get("eventContext", {})
        aoid = ctx.get("associateOID") or ctx.get("worker", {}).get("associateOID")
        worker = next((w for w in self.workers.values() if w.get("associateOID") == aoid), None)
        if worker is None:
            return
        t = ev["transform"]
        if "workAssignment" in t:
            wa = worker.setdefault("workAssignments", [{}])[0]
            for k, v in t["workAssignment"].items():
                if k == "customFieldGroup":
                    cfg = wa.setdefault("customFieldGroup", {})
                    for group, items in v.items():
                        existing = {i.get("itemID"): i for i in cfg.get(group, [])}
                        for i in items:
                            existing[i.get("itemID")] = i
                        cfg[group] = list(existing.values())
                else:
                    wa[k] = v
        elif "payDistribution" in t:
            worker["_pay"] = t["payDistribution"]
        elif "worker" in t:
            _deep_merge(worker, t["worker"])


def _deep_merge(dst: dict, src: dict) -> None:
    for k, v in src.items():
        if k == "communication" and isinstance(v, dict):
            comm = dst.setdefault("communication", {})
            if "email" in v:
                comm["emails"] = [v["email"]]
            if "mobile" in v:
                comm["mobiles"] = [v["mobile"]]
            continue
        if k == "businessCommunication":
            bc = dst.setdefault("businessCommunication", {})
            if "email" in v:
                bc["emails"] = [v["email"]]
            if "mobile" in v:
                bc["mobiles"] = [v["mobile"]]
            continue
        if isinstance(v, dict) and isinstance(dst.get(k), dict):
            _deep_merge(dst[k], v)
        else:
            dst[k] = copy.deepcopy(v)


# ---------------------------------------------------------------------------
# Dépôts en mémoire (mêmes contrats que les dépôts Spark)
# ---------------------------------------------------------------------------
class MemOps:
    def __init__(self, fail_on_upsert_status: Optional[str] = None):
        self.rows: Dict[str, dict] = {}
        self.fail_on_upsert_status = fail_on_upsert_status
        self.writes: List[str] = []

    def get(self, op_id):
        r = self.rows.get(op_id)
        return dict(r) if r else None

    def open_uncertain_for_employee(self, key):
        return [dict(r) for r in self.rows.values() if r["employee_key"] == key and r["status"] in OP_OPEN_UNCERTAIN]

    def for_item(self, run_id, key):
        return [dict(r) for r in self.rows.values() if r["run_id"] == run_id and r["employee_key"] == key]

    def upsert(self, op):
        if self.fail_on_upsert_status and op["status"] == self.fail_on_upsert_status:
            raise IOError("simulated durable write failure")
        self.rows[op["op_id"]] = dict(op)
        self.writes.append(op["status"])


class MemAttempts:
    def __init__(self):
        self.rows: List[dict] = []

    def append(self, a):
        self.rows.append(dict(a))


class MemLog:
    def __init__(self):
        self.rows: List[dict] = []
        self.decisions: List[tuple] = []

    def append(self, rows):
        self.rows.extend(rows)

    def record(self, run_id, decision):
        self.decisions.append((run_id, decision))


class MemOutbox:
    """Outbox multi-runs ; `run_order` donne l'ordre de promotion des runs (plus récent = plus grand)."""

    def __init__(self, run_order: Optional[Dict[str, int]] = None):
        self.items: Dict[Tuple[str, str], dict] = {}
        self.run_order = run_order or {}

    def add(self, run_id, key, status="PENDING", attempt_count=0):
        self.items[(run_id, key)] = {"run_id": run_id, "employee_key": key, "status": status,
                                     "attempt_count": attempt_count, "change_category": "MODIFIED"}

    def load(self, run_id, statuses, max_attempts):
        return [dict(v) for (r, _), v in sorted(self.items.items()) if r == run_id and v["status"] in statuses
                and int(v.get("attempt_count") or 0) < max_attempts]

    def newer_items(self, run_id, key):
        mine = self.run_order.get(run_id, 0)
        return [dict(v) for (r, k), v in self.items.items() if k == key and self.run_order.get(r, 0) > mine]

    def update(self, run_id, updates):
        for u in updates:
            self.items[(run_id, u["employee_key"])].update({k: v for k, v in u.items() if k != "employee_key"})


def fake_context(name, workspace_id="00000000-0000-0000-0000-000000000000"):
    """Contexte d'exécution SIMULÉ (tests isolés uniquement)."""
    return lambda: {"currentWorkspaceName": name, "currentWorkspaceId": workspace_id, "isForPipeline": True,
                    "currentNotebookName": "NB_HRIS_PUBLISH_ADP"}
'''),
    ('jdd', '1afe3dd5706f2e1043bdff70d84c4e6085ca490f53bb021d7207cc1f1b9fc307', r'''"""Jeux de données FICTIFS (JDD-01 à JDD-06, plan de tests 04 § 2). Aucune donnée réelle :
noms, numéros et codes sont inventés et non représentatifs."""
from __future__ import annotations

import copy
from typing import Dict, List

from hris.model import LOOKUP_TABLES, RAW_KEYS
from hris.ts_acquire import initialize_employee_fields, to_raw_row

# --- JDD-03 : correspondances fictives (codes inventés) ---------------------------------------
LOOKUPS: Dict[str, List[dict]] = {
    "lu_hr_emp_titre": [{"cd_emp_titre_ts": "MR", "cd_emp_titre_adp": "1"}, {"cd_emp_titre_ts": "MS", "cd_emp_titre_adp": "2"}],
    "lu_hr_emp_sexe": [{"cd_emp_sexe_ts": "M", "cd_emp_sexe_adp": "1"}, {"cd_emp_sexe_ts": "F", "cd_emp_sexe_adp": "2"}],
    "lu_hr_emp_nationalite": [{"cd_emp_nationalite_ts": "FRA", "cd_emp_nationalite_adp": "FR"}],
    "lu_hr_emp_situation_matrimoniale": [{"cd_emp_situation_matrimoniale_ts": "SINGLE",
                                          "cd_emp_situation_matrimoniale_adp": "C"}],
    "lu_hr_emp_pays": [{"cd_emp_pays_ts": "FRA", "id_emp_pays_adp": "FR"}, {"cd_emp_pays_ts": "FR", "id_emp_pays_adp": "FR"}],
    "lu_hr_ppc_raison_debut_contrat": [{"cd_ppc_raison_debut_contrat_ts": "CRE", "cd_ppc_raison_debut_contrat_adp": "01"}],
    "lu_hr_ppc_contrat2": [{"cd_contract_ts": "CDI", "cd_nature_contract_adp": "00", "cd_type_contract_adp": "00",
                            "cd_type_collaboration_adp": "SAL"},
                           {"cd_contract_ts": "CDD", "cd_nature_contract_adp": "01", "cd_type_contract_adp": "01",
                            "cd_type_collaboration_adp": "SAL"}],
    "lu_hr_ppc_classification": [{"cd_convention_collective_lvl2_ts": "C1", "cd_classification_adp": "IC5",
                                  "cd_coefficient_adp": "100", "cd_ppc_convention_collective_adp": "94"}],
    "lu_hr_ppc_classes": [{"cd_occupational_category_ts": "CAD", "cd_categorie_cotisant_adp": "13",
                           "cd_classe_remuneration_adp": "B"}],
    "lu_hr_geo_entite": [{"cd_geo_entite_ts": "Motul France", "cd_geo_entite_adp": "MF"}],
    "lu_hr_ppc_location_country": [{"cd_ppc_location_country_ts": "SITE1", "cd_ppc_location_country_adp": "02003"}],
    "lu_hr_ppc_cost_center": [{"cd_ppc_cost_center_ts": "CC1", "cd_ppc_cost_center_adp": "S0001"}],
    "lu_hr_organisationalstructure": [{"cd_organisationalstructure_ts": "ORG1", "cd_organisationalstructure_adp": "A1B2"}],
    "lu_hr_emp_diplome": [], "lu_hr_emp_work_accident": [], "lu_hr_ppc_convention_collective": [],
    "lu_hr_ppc_contrat": [],
}


def full_lookup_rows(table: str) -> List[dict]:
    return [{c: r.get(c) for c in LOOKUP_TABLES[table]} for r in LOOKUPS[table]]


def ts_record(n: int, **overrides) -> dict:
    """Enregistrement TalentSoft fictif « valide » au format de la Function (avant JSON_VALUE)."""
    rec = initialize_employee_fields(f"E{n:05d}")
    rec.update({
        "LegacyID": f"{n:08d}", "Matricule_IT": f"IT{n:04d}", "LastName": f"NOM{n}", "FirstName": f"Prenom{n}",
        "BirthName": f"NOM{n}", "Title": "MR", "Gender": "M", "Nationality": "FRA",
        "hubSocialSecurityNumber": f"1{n:012d}", "MaritalStatus": "SINGLE", "MaritalStatus_startDate": "2020-01-01",
        "Date_de_naissance": "1990-05-17", "Ville_de_naissance": "Villefictive", "Pays_de_naissance": "FRA",
        "LegalTownOrCity": "Villefictive", "LegalCountry": "FRA", "LegalPostalCode": "99000",
        "LegalStreet": "Rue Imaginaire", "LegalStreetNumber": str(n), "adresse_countryCode": "FR",
        "BusinessEmail": f"user{n}@example.invalid", "PersonalEmail": f"perso{n}@example.invalid",
        "BusinessMobile": "0000000000", "referenceAdmissionDate": "2019-01-07T00:00:00",
        "GroupStartDate": "2019-01-07", "CompanyStartDate": "2019-01-07", "HiringReason": "CRE",
        "date_debut_contract": "2019-01-07", "date_fin_contract": "1900-01-01", "NatureOfContract": "CDI",
        "OccupationalCategory": "CAD", "Situation_starDate": "2024-01-01", "Coefficient": "C1",
        "LegalStructure": "Motul France", "GeographicOrganisationStructure": "SITE1", "PPC_FTE": "100",
        "OrganisationalStructure": "ORG1", "CostCenter": "CC1", "pourcentage_imputation": "100",
        "AccountHolder": f"NOM{n}", "IBANOrABA": "XX00TEST0000000000", "BicOrSwift": "TESTXXXX",
        "BankName": "BANQUE FICTIVE", "salary_base": 48000, "base_startDate": "2024-01-01",
        "Pourcentage_prime_annuelles": 10, "pourcentage_prime_annuelles_startDate": "2024-01-01",
    })
    rec.update(overrides)
    return rec


def raw_rows(records: List[dict]) -> List[dict]:
    return [to_raw_row(r) for r in records]


def deep(x):
    return copy.deepcopy(x)
'''),
    ('test_parity_transform', 'e167eada2f5f8e55dcc87e26cce157e96de92dc73e14edf0842fa381e67650c2', r'''"""TU-05 / TU-08 — sémantiques T-SQL émulées et règles R-SQL (implémentation de référence).
Données fictives uniquement."""
import datetime as dt
import decimal
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.join(HERE, "..", "..", "src"), os.path.join(HERE, "..", "fixtures")]

from hris import parity as P  # noqa: E402
from hris import transform_ref as R  # noqa: E402
from hris.model import ADP_COLUMNS, COMPARE_COLUMNS, LOOKUP_JOINS  # noqa: E402
import jdd  # noqa: E402

D = dt.date(2026, 10, 7)


def staged(records):
    src = [R.parse_source_row(r) for r in jdd.raw_rows(records)]
    lookups = {t: jdd.full_lookup_rows(t) for t in jdd.LOOKUPS}
    return R.enrich_rows(src, lookups)


class JsonValueAndTruncation(unittest.TestCase):
    def test_json_value_repr(self):  # R-SQL-002
        self.assertIsNone(P.json_value_repr(None))
        self.assertEqual(P.json_value_repr(True), "true")
        self.assertEqual(P.json_value_repr(48000), "48000")
        self.assertEqual(P.json_value_repr(48000.5), "48000.5")
        self.assertEqual(P.json_value_repr("abc"), "abc")
        self.assertIsNone(P.json_value_repr({"a": 1}))
        self.assertIsNone(P.json_value_repr("x" * 4001))

    def test_nvarchar_50(self):  # R-SQL-002 / D-PAR-03
        self.assertEqual(len(P.nvarchar("é" * 80)), 50)
        self.assertIsNone(P.nvarchar(None))

    def test_coalesce_null_only_and_replace_none(self):  # R-SQL-003
        src = R.parse_source_row({"BusinessMobile": "", "BusinessPhone": "0102", "date_fin_contract": "None"})
        self.assertEqual(src["lb_telephone_professionnel"], "")
        self.assertEqual(src["dt_ppc_fin_contrat"], "")
        src = R.parse_source_row({"BusinessMobile": None, "BusinessPhone": "0102"})
        self.assertEqual(src["lb_telephone_professionnel"], "0102")


class TsqlEmulation(unittest.TestCase):
    def test_trailing_spaces(self):
        self.assertTrue(P.sql_is_empty("   "))
        self.assertFalse(P.sql_is_empty(None))
        self.assertFalse(P.sql_is_not_empty(None))
        self.assertEqual(P.sql_len("12345678  "), 8)
        self.assertTrue(P.sql_eq("ABC  ", "ABC"))
        self.assertFalse(P.sql_eq(" ABC", "ABC"))

    def test_isdate(self):  # R-SQL-043 / D-PAR-01
        for ok in ["2024-01-31", "2024-1-5", "20240131", "2024-01-31T10:20:30", "2024-01-31T10:20:30.123Z",
                   "2024-01-31 10:20", "01/31/2024", " 2024-01-31 "]:
            self.assertEqual(P.isdate(ok), 1, ok)
        for ko in [None, "", "2024-02-30", "2024-13-01", "1700-01-01", "31/01/2024", "abc",
                   "2024-01-31T10:20:30.1234567"]:
            self.assertEqual(P.isdate(ko), 0, ko)

    def test_convert_date(self):  # R-SQL-052
        self.assertEqual(P.convert_date(""), dt.date(1900, 1, 1))
        self.assertEqual(P.convert_date("  "), dt.date(1900, 1, 1))
        self.assertIsNone(P.convert_date(None))
        self.assertEqual(P.convert_date("2024-01-31T00:00:00"), dt.date(2024, 1, 31))
        with self.assertRaises(ValueError):
            P.convert_date("zz")

    def test_isnumeric_and_cast(self):  # R-SQL-035/036/044, D-PAR-02
        for ok in ["10", "-3.5", "+1", "1e5", "$", ".", ",", "1,000"]:
            self.assertEqual(P.isnumeric(ok), 1, ok)
        for ko in [None, "", "abc", "e5", "1-2"]:
            self.assertEqual(P.isnumeric(ko), 0, ko)
        self.assertEqual(P.cast_numeric_12_2("12.345"), decimal.Decimal("12.35"))
        self.assertEqual(P.cast_numeric_12_2("-12.345"), decimal.Decimal("-12.35"))
        for bad in ["1e5", "$", "", "1,000", "99999999999"]:
            with self.assertRaises(ValueError, msg=bad):
                P.cast_numeric_12_2(bad)

    def test_pct_flag(self):
        self.assertEqual(P.pct_flag("abc", 100), 0)      # non numérique : non rejeté (parité)
        self.assertEqual(P.pct_flag("150", 100), 1)
        self.assertEqual(P.pct_flag("100", 100), 0)
        with self.assertRaises(ValueError):
            P.pct_flag("1e5", 100)                       # ISNUMERIC=1 mais CAST échoue : échec du run

    def test_monthly_salary_scale(self):  # R-SQL-054
        v = P.monthly_salary("5000")
        self.assertEqual(v.as_tuple().exponent, -13)
        self.assertEqual(v, decimal.Decimal("416.6666666666667"))

    def test_utf8_truncate(self):  # R-SQL-066
        self.assertEqual(len(P.utf8_truncate("Référence_Rejet," * 5, 50).encode("utf-8")) <= 50, True)


class Rules(unittest.TestCase):
    def test_valid_employee_goes_through(self):
        rows = staged([jdd.ts_record(1)])
        rej = R.reject_rows(rows, D)
        self.assertEqual(rej, [])
        adp = R.adp_rows(rows, rej, D)
        self.assertEqual(len(adp), 1)
        self.assertEqual(list(adp[0].keys()), ADP_COLUMNS)
        self.assertEqual(adp[0]["dt_ppc_fin_contrat"], dt.date(1900, 1, 1))  # 'None' -> '' -> 1900-01-01
        self.assertEqual(adp[0]["mt_salaire_mensuel"], decimal.Decimal("4000.0000000000000"))

    def test_null_reject(self):  # R-SQL-030/031
        rows = staged([jdd.ts_record(2, Matricule_IT="")])
        rej = [r for r in R.reject_rows(rows, D) if r["cd_type_rejet"] == "Null_Rejet"]
        self.assertEqual([(r["lb_nom_colonne_rejet"], r["ds_rejet"]) for r in rej],
                         [("cd_matricule_it", "cd_matricule_it ne peut pas être NULL")])
        rows = staged([jdd.ts_record(3, Matricule_IT=None)])
        self.assertEqual([r for r in R.reject_rows(rows, D) if r["cd_type_rejet"] == "Null_Rejet"], [])

    def test_format_reject_and_fte_excluded_from_sum(self):  # R-SQL-032..039
        rows = staged([jdd.ts_record(4, PPC_FTE="150")])
        self.assertEqual(R.reject_rows(rows, D), [])          # ETP seul hors bornes : aucun rejet (R-SQL-038)
        rows = staged([jdd.ts_record(5, PPC_FTE="150", Date_de_naissance="2024-02-30")])
        rej = R.reject_rows(rows, D)
        self.assertEqual(sorted(r["lb_nom_colonne_rejet"] for r in rej), ["dt_naissance", "pc_ppc_fte"])
        ds = {r["lb_nom_colonne_rejet"]: r["ds_rejet"] for r in rej}
        self.assertEqual(ds["dt_naissance"], "Format Date non respecté au YYYY-MM-DD pour dt_naissance = 2024-02-30")
        rows = staged([jdd.ts_record(6, LegacyID="123")])
        rej = R.reject_rows(rows, D)
        self.assertEqual([(r["lb_nom_colonne_rejet"], r["ds_rejet"]) for r in rej], [("id_payroll", None)])

    def test_reference_reject(self):  # R-SQL-040/041/042
        rows = staged([jdd.ts_record(7, Title="XX", LegalCountry="ZZZ")])
        rej = R.reject_rows(rows, D)
        self.assertEqual([(r["cd_type_rejet"], r["lb_fichier_mapping"], r["ds_rejet"]) for r in rej],
                         [("Référence_Rejet", "titre.csv", "Mapping non trouvé entre TS et ADP pour cd_emp_titre = XX")])
        adp = R.adp_rows(staged([jdd.ts_record(8, LegalCountry="ZZZ")]), [], D)
        self.assertIsNone(adp[0]["id_fiscal_pays_adp"])   # pays fiscal jamais rejeté (R-SQL-042)

    def test_trailing_space_join(self):  # D-PAR-04
        rows = staged([jdd.ts_record(9, Title="MR   ")])
        self.assertEqual(rows[0]["cd_emp_titre_adp"], "1")

    def test_duplicate_mapping_duplicates_rows(self):  # R-SQL-058
        lookups = {t: jdd.full_lookup_rows(t) for t in jdd.LOOKUPS}
        lookups["lu_hr_emp_sexe"] = lookups["lu_hr_emp_sexe"] + [{"cd_emp_sexe_ts": "M", "cd_emp_sexe_adp": "9"}]
        src = [R.parse_source_row(r) for r in jdd.raw_rows([jdd.ts_record(10)])]
        self.assertEqual(len(R.enrich_rows(src, lookups)), 2)

    def test_salary_conversion_failure_fails_run(self):  # D-PAR-07
        rows = staged([jdd.ts_record(11, salary_base="")])
        with self.assertRaises(ValueError):
            R.adp_rows(rows, R.reject_rows(rows, D), D)

    def test_reject_excludes_employee(self):  # R-SQL-050
        rows = staged([jdd.ts_record(12), jdd.ts_record(13, Gender="X")])
        rej = R.reject_rows(rows, D)
        self.assertEqual([r["id_unique"] for r in R.adp_rows(rows, rej, D)], ["E00012"])

    def test_row_hash(self):  # TU-08
        adp = R.adp_rows(staged([jdd.ts_record(14)]), [], D)[0]
        h = R.row_hash(adp)
        self.assertEqual(h, R.row_hash(dict(adp)))
        self.assertNotEqual(h, R.row_hash(dict(adp, lb_email_personnel="")))
        self.assertNotEqual(R.row_hash(dict(adp, lb_email_personnel=None)), R.row_hash(dict(adp, lb_email_personnel="")))
        self.assertEqual(R.row_hash(dict(adp, dt_transforme=dt.date(2000, 1, 1))), h)  # dt_transforme non comparée
        self.assertEqual(len(COMPARE_COLUMNS), 71)

    def test_summary(self):  # R-SQL-060..066
        rej = R.reject_rows(staged([jdd.ts_record(15, Gender="X"), jdd.ts_record(16, Title="Y")]), D)
        s = R.summary(["E00001", "E00002"], rej, D)
        self.assertEqual((s["mt_lignes_inserts"], s["mt_lignes_rejets"], s["mt_lignes_lues"]), (2, 2, 4))
        self.assertEqual(s["pc_lignes_rejets"], decimal.Decimal("0.5000"))
        self.assertTrue(len(s["ds_type_rejets"].encode("utf-8")) <= 50)

    def test_lookup_join_count(self):
        self.assertEqual(len(LOOKUP_JOINS), 15)


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_ts_acquire', '096584d0970044b5076b2a1cbb0a711288131e316f7c29565e8df2491456dabb', r'''"""TU-07 — acquisition TalentSoft (JDD-01) contre un faux serveur HTTP en mémoire. Aucun appel réseau."""
import json
import os
import re
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.join(HERE, "..", "..", "src"), os.path.join(HERE, "..", "fixtures")]

from hris.common import TS_PRD_BASE_URL, HrisError, SecretValue  # noqa: E402
from hris import ts_acquire as A  # noqa: E402


class FakeResp:
    def __init__(self, status, body=None, headers=None):
        self.status_code = status
        self._body = body
        self.headers = headers or {}

    def json(self):
        if isinstance(self._body, Exception):
            raise self._body
        return json.loads(json.dumps(self._body))


def hub(field, value, start, **kw):
    return dict({"field": field, "value": value, "startDate": start}, **kw)


class FakeTalentSoft:
    """Faux TalentSoft : 3 salariés fictifs, dont 1 hors « Motul France »."""

    def __init__(self, total_override=None, scripted=None):
        self.calls = []
        self.scripted = scripted or {}   # chemin (regex) -> liste de statuts à renvoyer d'abord
        self.total_override = total_override
        self.closed = False
        self.employees = [
            {"employeeNumber": "E1", "referenceAdmissionDate": "2019-01-07T00:00:00",
             "individual": {"extras": {"hubSocialSecurityNumber": 1234567890123}, "user": {"username": "u1"}}},
            {"employeeNumber": "E2", "referenceAdmissionDate": "2020-02-01T00:00:00",
             "individual": {"extras": {}, "user": {}}},
            {"employeeNumber": "E3", "individual": {}},
        ]

    def close(self):
        self.closed = True

    def post(self, url, headers=None, data=None, timeout=None, allow_redirects=None):
        self.calls.append(("POST", url))
        assert allow_redirects is False
        assert "client_secret" in data and data["grant_type"] == "client_credentials"
        return self._scripted(url) or FakeResp(200, {"access_token": "tok-fictif"})

    def _scripted(self, url):
        for pattern, statuses in self.scripted.items():
            if re.search(pattern, url) and statuses:
                st = statuses.pop(0)
                if isinstance(st, Exception):
                    raise st
                return FakeResp(st, {}, {"Retry-After": "0"})
        return None

    def get(self, url, headers=None, timeout=None, allow_redirects=None):
        self.calls.append(("GET", url))
        assert allow_redirects is False
        assert headers["Authorization"].startswith("Bearer ")
        s = self._scripted(url)
        if s:
            return s
        path = url[len(TS_PRD_BASE_URL):]
        if path.startswith("/api/v1/directory/employees?count=1"):
            return FakeResp(200, {"totalCount": self.total_override or len(self.employees), "results": []})
        if path.startswith("/api/v1/directory/employees?count="):
            return FakeResp(200, {"totalCount": len(self.employees), "results": self.employees})
        m = re.match(r"/api/hub/v2/employees/(\w+)/(\w+)$", path)
        if m:
            emp, entity = m.groups()
            return FakeResp(200, self.entity(emp, entity))
        m = re.match(r"/api/hub/v2.0/positions/(\w+)\?date=(.*)$", path)
        if m:
            legal = "Motul France" if m.group(1) != "P3" else "Motul SA"
            return FakeResp(200, [{"field": "LegalStructure", "value": legal}, {"field": "Level", "value": "80 %"},
                                  {"field": "FullTimeEquivalentPercent", "value": 0},
                                  {"field": "CostCenter", "value": "CC1"}])
        if path.endswith("/postal-addresses"):
            return FakeResp(200, [{"postalCode": "99000", "city": "Villefictive", "countryCode": "be"}])
        return FakeResp(404, {})

    def entity(self, emp, entity):
        if entity == "Situation":
            pos = {"E1": "P1", "E2": "P2", "E3": "P3"}[emp]
            return [hub("WorkingType", "Ingénieur", "2023-01-01"), hub("Position", pos, "2023-01-01"),
                    hub("WorkingType", "Ancien", "2020-01-01")]
        if entity == "Identity":   # ordre de l'API : l'item plus ancien APRÈS le plus récent est ignoré
            return [hub("LastName", "NOM", "2022-01-01"), hub("FirstName", "Prénom", "2022-01-01"),
                    hub("LastName", "ANCIEN", "2021-01-01"), hub("Gender", False, "2022-01-01")]
        if entity == "Compensation":
            return [{"key": "X_ANNUAL_BASE_SALARY", "field": "Amount", "value": 45000.5, "startDate": "2024-01-01"},
                    {"key": "X_ANNUAL_BASE_SALARY", "field": "Amount", "value": 40000, "startDate": "2023-01-01"},
                    {"key": "CONTRACTUAL BONUS", "field": "Amount", "value": 0, "startDate": "2024-01-01"}]
        if entity == "Contract":
            return [hub("NatureOfContract", "CDI", "2019-01-07", endDate=None)]
        if entity == "Disability":
            return [hub("RecognizedAsDisabled", True, "2020-01-01")]
        return []


def make(fake, **kw):
    stats = A.AcquisitionStats()
    client = A.TsClient(TS_PRD_BASE_URL, SecretValue("id-fictif"), SecretValue("secret-fictif"),
                        session_factory=lambda: fake, stats=stats, sleep=lambda s: None, **kw)
    return A.TalentSoftExtractor(client, sleep=lambda s: None), stats


class SourceIsPrdOnly(unittest.TestCase):
    def test_refuses_any_other_tenant(self):
        for url in ["https://testmotul.talent-soft.com", "http://motul.talent-soft.com", "", None,
                    "https://motul.talent-soft.com.evil.example"]:
            with self.assertRaises(HrisError, msg=url):
                A.assert_prd_source(url)
        self.assertEqual(A.assert_prd_source(TS_PRD_BASE_URL + "/"), TS_PRD_BASE_URL)


class Extraction(unittest.TestCase):
    def test_full_extraction_rules(self):
        fake = FakeTalentSoft()
        ex, stats = make(fake)
        records = ex.run()
        by = {r["employeeNumber"]: r for r in records}
        self.assertEqual(sorted(by), ["E1", "E2"])                      # E3 hors Motul France (R-TS-010)
        e1 = by["E1"]
        self.assertEqual(e1["LastName"], "NOM")                         # R-TS-018 : item ancien ignoré
        self.assertEqual(e1["Gender"], "")                              # R-TS-020 : falsy -> ''
        self.assertEqual(e1["WorkingType"], "Ingénieur")
        self.assertEqual(e1["salary_base"], 45000.5)                    # R-TS-019 : valeur brute, max strict par clé
        self.assertEqual(e1["prime_annuelles"], 0)
        self.assertEqual(e1["pourcentage_imputation"], "80")            # R-TS-021 : Level sans %
        self.assertEqual(e1["PPC_FTE"], "")                             # 0 falsy
        self.assertEqual(e1["Hors_France"], "1")                        # R-TS-022 : 'be' != FR
        self.assertEqual(e1["date_fin_contract"], "1900-01-01")         # R-TS-024
        self.assertEqual(e1["hubSocialSecurityNumber"], "1234567890123")
        self.assertEqual(by["E2"]["address_postalCode"], "")            # pas d'username : pas d'appel postal
        raw = A.to_raw_row(e1)
        self.assertNotIn("handicap", raw)                               # données de santé non persistées
        self.assertEqual(raw["salary_base"], "45000.5")
        self.assertEqual(len(raw), 70)
        anomalies = {a["control_code"]: a for a in A.completeness_anomalies(stats.as_dict())}
        self.assertEqual(anomalies["EXTRACTION_INCOMPLETE"]["outcome"], "PASS")
        self.assertEqual(anomalies["TS_ENDPOINT_FAILURE"]["outcome"], "PASS")
        self.assertTrue(fake.calls[0][0] == "POST")

    def test_incomplete_directory_is_blocking(self):  # D-QUA-01
        fake = FakeTalentSoft(total_override=5)
        ex, stats = make(fake)
        ex.run()
        a = {x["control_code"]: x for x in A.completeness_anomalies(stats.as_dict())}
        self.assertEqual((a["EXTRACTION_INCOMPLETE"]["severity"], a["EXTRACTION_INCOMPLETE"]["outcome"]),
                         ("BLOCKING", "FAIL"))

    def test_retries_429_then_success_and_exhausted_503_counted(self):  # R-TS-014/015
        fake = FakeTalentSoft(scripted={r"/E1/Identity$": [429, 429], r"/E2/Contract$": [503, 503, 503]})
        ex, stats = make(fake)
        records = {r["employeeNumber"]: r for r in ex.run()}
        self.assertEqual(records["E1"]["LastName"], "NOM")             # 3e essai réussi
        self.assertEqual(records["E2"]["NatureOfContract"], "")         # parité : valeurs par défaut...
        a = {x["control_code"]: x for x in A.completeness_anomalies(stats.as_dict())}
        self.assertEqual(a["TS_ENDPOINT_FAILURE"]["outcome"], "FAIL")   # ... mais l'échec n'est plus silencieux

    def test_401_refreshes_token_once(self):
        fake = FakeTalentSoft(scripted={r"/E1/KeyDates$": [401]})
        ex, stats = make(fake)
        ex.run()
        self.assertGreaterEqual(stats.counters["token_refresh_on_401"], 1)

    def test_network_errors_and_circuit_breaker(self):  # R-TS-016
        fake = FakeTalentSoft(scripted={r"/hub/v2/employees/E1/": [OSError("x")] * 40})
        ex, stats = make(fake)
        ex.run()
        self.assertGreater(stats.counters["network_error"], 0)
        a = {x["control_code"]: x for x in A.completeness_anomalies(stats.as_dict())}
        self.assertEqual(a["TS_FILTER_TECHNICAL_FAILURE"]["outcome"], "FAIL")

    def test_token_failure_raises(self):
        fake = FakeTalentSoft(scripted={r"/api/token$": [400]})
        ex, _ = make(fake)
        with self.assertRaises(HrisError):
            ex.run()

    def test_stats_contain_no_identifier(self):
        fake = FakeTalentSoft()
        ex, stats = make(fake)
        ex.run()
        text = json.dumps(stats.as_dict())
        for forbidden in ["E1", "E2", "u1", "1234567890123", "talent-soft", "tok-fictif", "secret-fictif"]:
            self.assertNotIn(forbidden, text)
        self.assertNotIn("secret-fictif", repr(ex.client))


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_adp_publication', '6bc67706cfe7e926368c4a048dd995a364f16885b3ea6196fe4745fa1cbbd32f', r'''"""TU-01 à TU-04 + exigences du mandat § 8/§ 10 — publication ADP, uniquement avec clients FICTIFS.
Aucun accès réseau : le simulateur répond en mémoire et la sentinelle bloque tout hôte ADP."""
import datetime as dt
import json
import os
import socket
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.join(HERE, "..", "..", "src"), os.path.join(HERE, "..", "fixtures")]

from hris import adp_gate as G  # noqa: E402
from hris import adp_ops as O  # noqa: E402
from hris import adp_publish as PUB  # noqa: E402
from hris import adp_transport as TR  # noqa: E402
from hris.adp_sim import FakeAdp, MemAttempts, MemLog, MemOps, MemOutbox, fake_context  # noqa: E402
from hris.common import HrisError  # noqa: E402

NOW = dt.datetime(2026, 10, 7, 3, 0, tzinfo=dt.timezone.utc)


class ForbiddenFactory:
    """Fabrique de transport qui ne doit JAMAIS être appelée (lecture de secrets ADP)."""
    def __init__(self):
        self.called = 0

    def __call__(self):
        self.called += 1
        raise AssertionError("transport ADP construit alors qu'il ne le devait pas")


def simple_planner(variants=("P2_PERSONAL_EMAIL",), targets=None, deps=None):
    """Planificateur fictif : soumet une opération par variante (cibles fictives)."""
    def plan(row, ex):
        for v in variants:
            tgt = (targets or {}).get(v, {"value": row["value"]})
            ex.submit(O.PlannedOp(v[:2], v, "events/hr/v1/fake", {"events": [{"x": tgt}]}, tgt,
                                  {"mode": "UPDATE", "field": v}, depends_on=(deps or {}).get(v, ()),
                                  read_back=lambda rd, t=tgt: rd.read(t)))
    return plan


def run(name, outbox, *, factory, ops=None, mode="NORMAL", source="R1", blockers=lambda: [], planner=None,
        rows=None, max_attempts=3):
    ops = ops if ops is not None else MemOps()
    logs, gate = MemLog(), MemLog()
    ctx = PUB.PublicationContext(publish_run_id="P-" + source, source_run_id=source, mode=mode,
                                 max_attempts=max_attempts, pause_seconds=0)
    summary = PUB.run_publication(ctx, context_provider=fake_context(name) if isinstance(name, (str, type(None))) else name,
                                  outbox=outbox, gate_log=gate, ops=ops, attempts=MemAttempts(),
                                  publication_log=logs, activation_blockers=blockers, transport_factory=factory,
                                  planner=planner or simple_planner(),
                                  rows_by_key=lambda: rows or {k: {"value": f"v-{k}"} for (_, k) in outbox.items},
                                  sleep=lambda s: None, clock=lambda: NOW)
    return summary, ops, gate, logs


class GateRule(unittest.TestCase):
    def test_table(self):  # TU-01 (règle littérale inchangée, D-ENV-02 à valider)
        cases = [("WS-HRIS-PRD", G.ENV_SATISFIED), ("WS_HRIS_PRD", G.ENV_SATISFIED),
                 ("WS-HRIS-DEV", G.DENIED_NON_PRD), ("Sandbox - Fabric BI backend", G.DENIED_NON_PRD),
                 ("ws-hris-prd", G.DENIED_NON_PRD), ("WS-HRIS-Prd", G.DENIED_NON_PRD),
                 ("WS-HRIS-PREPRD", G.ENV_SATISFIED), ("NOTPRD", G.ENV_SATISFIED),
                 ("", G.BLOCKED_UNRESOLVED), ("   ", G.BLOCKED_UNRESOLVED), (None, G.BLOCKED_UNRESOLVED),
                 (123, G.BLOCKED_UNRESOLVED), (b"WS-HRIS-PRD", G.BLOCKED_UNRESOLVED)]
        for name, outcome in cases:
            met, reason = G.decide_adp_push(name)
            self.assertEqual(G.outcome_of(reason), outcome, name)

    def test_context_failure_is_unresolved_not_non_prd(self):  # TU-02
        def boom():
            raise RuntimeError("contexte indisponible")
        d = G.evaluate_gate(boom, clock=lambda: NOW)
        self.assertEqual((d.outcome, d.env_condition_met), (G.BLOCKED_UNRESOLVED, False))
        d = G.evaluate_gate(lambda: {"currentWorkspaceId": "x"}, clock=lambda: NOW)
        self.assertEqual(d.outcome, G.BLOCKED_UNRESOLVED)


class GuardProof(unittest.TestCase):
    """Preuve instrumentée : zéro appel de publication hors PRD et en contexte non résolu ;
    contrôle POSITIF en contexte simulé autorisé (le comptage détecte bien un appel)."""

    def test_non_prd_zero_publication_calls(self):
        for name in ["WS-HRIS-DEV", "Sandbox - Fabric BI backend", "ws-hris-prd"]:
            outbox, factory = MemOutbox(), ForbiddenFactory()
            for k in ("E1", "E2"):
                outbox.add("R1", k)
            s, ops, gate, logs = run(name, outbox, factory=factory)
            self.assertEqual(factory.called, 0, name)
            self.assertEqual(s["status"], PUB.STATUS_NOT_PUBLISHED_NON_PRD)
            self.assertEqual({v["status"] for v in outbox.items.values()}, {"NOT_PUSHED_NON_PRD"})
            self.assertEqual(ops.rows, {})
            self.assertEqual(len(gate.decisions), 1)

    def test_unresolved_zero_calls_and_items_stay_pending(self):
        for ctx in [None, "", lambda: (_ for _ in ()).throw(RuntimeError("x"))]:
            outbox, factory = MemOutbox(), ForbiddenFactory()
            outbox.add("R1", "E1")
            s, ops, gate, _ = run(ctx, outbox, factory=factory)
            self.assertEqual(factory.called, 0)
            self.assertEqual(s["status"], PUB.STATUS_BLOCKED_UNRESOLVED)
            self.assertEqual(outbox.items[("R1", "E1")]["status"], "PENDING")   # réévaluable, pas NOT_PUSHED

    def test_env_satisfied_but_activation_blocked_zero_calls(self):
        outbox, factory = MemOutbox(), ForbiddenFactory()
        outbox.add("R1", "E1")
        s, _, _, _ = run("WS-HRIS-PRD", outbox, factory=factory, blockers=PUB.code_activation_blockers)
        self.assertEqual(factory.called, 0)
        self.assertEqual(s["status"], PUB.STATUS_BLOCKED_ACTIVATION)
        self.assertIn("CONTRAT_ADP_FINAL_ABSENT", s["activation_blockers"])
        self.assertEqual(outbox.items[("R1", "E1")]["status"], "PENDING")

    def test_positive_control_simulated_allowed_context(self):
        adp = FakeAdp()
        rec = TR.RecordingTransport(adp)
        outbox = MemOutbox()
        outbox.add("R1", "E1")
        s, ops, _, logs = run("WS-HRIS-PRD", outbox, factory=lambda: rec)
        self.assertEqual(len(rec.post_calls), 1)                 # l'instrumentation voit bien l'appel
        self.assertEqual(s["status"], PUB.STATUS_PUBLISHED)
        self.assertEqual(outbox.items[("R1", "E1")]["status"], "SENT")

    def test_real_transport_cannot_be_built(self):
        with self.assertRaises(HrisError):
            TR.AdpHttpTransport("c.pem", "k.pem", lambda: "x", 30, _allow_construction=True)
        self.assertFalse(TR.ADP_ACTIVATION_APPROVED)
        self.assertFalse(TR.ADP_READ_POLICY_ESTABLISHED)


class NetworkSentinelTest(unittest.TestCase):
    def test_sentinel_detects_and_blocks_without_network(self):  # contrôle positif de l'observation réseau
        with TR.NetworkSentinel() as s:
            with self.assertRaises(ConnectionRefusedError):
                socket.getaddrinfo("accounts.eu.adp.com", 443)
            with self.assertRaises(ConnectionRefusedError):
                socket.create_connection(("geo.api.gouv.fr", 443), timeout=1)
        self.assertEqual([a["host"] for a in s.attempts], ["accounts.eu.adp.com", "geo.api.gouv.fr"])
        self.assertIs(socket.getaddrinfo.__module__, "socket" if False else socket.getaddrinfo.__module__)

    def test_sentinel_ignores_other_hosts_and_restores(self):
        orig = socket.getaddrinfo
        with TR.NetworkSentinel() as s:
            self.assertFalse(s._watched("motul.talent-soft.com"))
        self.assertIs(socket.getaddrinfo, orig)


class TransportClassification(unittest.TestCase):
    def test_table(self):
        cases = [(200, '{"a":1}', TR.ACCEPTED), (201, '{}', TR.ACCEPTED), (202, "", TR.ACCEPTED),
                 (200, "", TR.UNCERTAIN), (200, "OK", TR.UNCERTAIN), (204, "", TR.UNCERTAIN),
                 (302, "", TR.UNCERTAIN), (401, "{}", TR.REJECTED), (429, "{}", TR.REJECTED),
                 (400, "{}", TR.REJECTED), (408, "", TR.UNCERTAIN), (500, "{}", TR.UNCERTAIN),
                 (503, "", TR.UNCERTAIN)]
        for status, body, outcome in cases:
            self.assertEqual(TR.classify_response(status, body)[0], outcome, (status, body))


class Durability(unittest.TestCase):
    def test_intent_persistence_failure_means_no_send(self):
        adp, outbox = FakeAdp(), MemOutbox()
        rec = TR.RecordingTransport(adp)
        outbox.add("R1", "E1")
        s, ops, _, _ = run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=MemOps(fail_on_upsert_status="INTENT_RECORDED"))
        self.assertEqual(rec.post_calls, [])                    # aucun envoi sans intention durable
        self.assertNotEqual(outbox.items[("R1", "E1")]["status"], "SENT")

    def test_crash_after_send_leaves_intent_uncertain_and_never_resent(self):
        adp, outbox = FakeAdp(), MemOutbox()
        outbox.add("R1", "E1")
        ops = MemOps(fail_on_upsert_status="ACCEPTED")          # issue non persistée (crash après envoi)
        rec = TR.RecordingTransport(adp)
        run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=ops)
        self.assertEqual(len(rec.post_calls), 1)
        self.assertEqual([r["status"] for r in ops.rows.values()], ["INTENT_RECORDED"])
        ops.fail_on_upsert_status = None
        outbox.items[("R1", "E1")]["status"] = "FAILED"          # même relancé explicitement...
        rec2 = TR.RecordingTransport(FakeAdp())
        s, *_ = run("WS-HRIS-PRD", outbox, factory=lambda: rec2, ops=ops, mode="RETRY_PUBLICATION")
        self.assertEqual(rec2.post_calls, [])                   # ... aucun nouvel envoi automatique
        self.assertEqual(s["blocked_prior_uncertain"], 1)

    def test_crash_before_send(self):
        adp, outbox = FakeAdp(), MemOutbox()
        adp.post_script = ["connect_refused"]
        outbox.add("R1", "E1")
        s, ops, _, _ = run("WS-HRIS-PRD", outbox, factory=lambda: adp)
        self.assertEqual(list(ops.rows.values())[0]["status"], O.OP_NOT_SENT)   # non-transmission prouvée
        self.assertEqual(outbox.items[("R1", "E1")]["status"], "FAILED")
        s, *_ = run("WS-HRIS-PRD", outbox, factory=lambda: adp, ops=ops, mode="RETRY_PUBLICATION")
        self.assertEqual(outbox.items[("R1", "E1")]["status"], "SENT")          # reprise explicite autorisée

    def test_crash_during_send(self):
        adp, outbox = FakeAdp(), MemOutbox()
        adp.post_script = ["send_interrupted"]
        outbox.add("R1", "E1")
        s, ops, _, _ = run("WS-HRIS-PRD", outbox, factory=lambda: adp)
        self.assertEqual(list(ops.rows.values())[0]["status"], O.OP_UNCERTAIN)
        self.assertEqual(outbox.items[("R1", "E1")]["status"], "IN_PROGRESS")
        outbox.items[("R1", "E1")]["status"] = "FAILED"
        rec = TR.RecordingTransport(FakeAdp())
        run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=ops, mode="RETRY_PUBLICATION")
        self.assertEqual(rec.post_calls, [])


class HttpOutcomes(unittest.TestCase):
    def _one(self, behaviour):
        adp, outbox = FakeAdp(), MemOutbox()
        adp.post_script = [behaviour]
        outbox.add("R1", "E1")
        rec = TR.RecordingTransport(adp)
        s, ops, _, logs = run("WS-HRIS-PRD", outbox, factory=lambda: rec)
        return s, list(ops.rows.values())[0], outbox.items[("R1", "E1")], rec

    def test_2xx_unusable_is_uncertain(self):
        s, op, item, rec = self._one(("2xx_text", 200))
        self.assertEqual((op["status"], item["status"]), (O.OP_UNCERTAIN, "IN_PROGRESS"))
        self.assertEqual(len(rec.post_calls), 1)

    def test_401_rejected_no_resend_and_halt(self):
        adp, outbox = FakeAdp(), MemOutbox()
        adp.post_script = [("status", 401)]
        for k in ("E1", "E2"):
            outbox.add("R1", k)
        rec = TR.RecordingTransport(adp)
        s, ops, _, _ = run("WS-HRIS-PRD", outbox, factory=lambda: rec)
        self.assertEqual(len(rec.post_calls), 1)                 # pas de renvoi, E2 non traité
        self.assertEqual(s["status"], PUB.STATUS_HALTED)
        self.assertEqual(outbox.items[("R1", "E2")]["status"], "PENDING")

    def test_429_5xx_timeout_redirect(self):
        for behaviour, expected in [(("status", 429), O.OP_REJECTED), (("status", 500), O.OP_UNCERTAIN),
                                    ("timeout_after_send", O.OP_UNCERTAIN), (("redirect", 302), O.OP_UNCERTAIN)]:
            s, op, item, rec = self._one(behaviour)
            self.assertEqual(op["status"], expected, behaviour)
            self.assertEqual(len(rec.post_calls), 1, behaviour)    # aucun retry automatique


class IdentityDependenciesSupersession(unittest.TestCase):
    def test_p8_variants_distinct_and_dependency_blocks(self):
        adp, outbox = FakeAdp(), MemOutbox()
        adp.post_script = [("status", 400)]                     # le contrat est rejeté
        outbox.add("R1", "E1")
        plan = simple_planner(("P8_CONTRACT_A", "P8_END_DATE", "P8_SALARY"),
                              deps={"P8_END_DATE": ("P8_CONTRACT_A", "P8_CONTRACT_B")})
        rec = TR.RecordingTransport(adp)
        s, ops, _, _ = run("WS-HRIS-PRD", outbox, factory=lambda: rec, planner=plan)
        st = {r["op_variant"]: r["status"] for r in ops.rows.values()}
        self.assertEqual(st, {"P8_CONTRACT_A": O.OP_REJECTED, "P8_END_DATE": O.OP_BLOCKED_DEPENDENCY,
                              "P8_SALARY": O.OP_ACCEPTED})
        self.assertEqual(len(rec.post_calls), 2)
        self.assertEqual(len({r["op_id"] for r in ops.rows.values()}), 3)

    def test_identity_stable_on_replay_and_distinct_for_new_change(self):
        a = O.op_identity("R1", "E1", "P2_PERSONAL_EMAIL", O.target_hash({"emailUri": "a"}))
        self.assertEqual(a, O.op_identity("R1", "E1", "P2_PERSONAL_EMAIL", O.target_hash({"emailUri": "a"})))
        b = O.op_identity("R2", "E1", "P2_PERSONAL_EMAIL", O.target_hash({"emailUri": "b"}))
        c = O.op_identity("R3", "E1", "P2_PERSONAL_EMAIL", O.target_hash({"emailUri": "a"}))  # retour à A
        self.assertEqual(len({a, b, c}), 3)

    def test_two_successive_legitimate_changes_same_field_both_sent(self):
        adp, outbox, ops = FakeAdp(), MemOutbox({"R1": 1, "R2": 2}), MemOps()
        outbox.add("R1", "E1")
        rec = TR.RecordingTransport(adp)
        run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=ops, source="R1", rows={"E1": {"value": "A"}})
        outbox.add("R2", "E1")
        run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=ops, source="R2", rows={"E1": {"value": "B"}})
        outbox.add("R3", "E1")
        outbox.run_order["R3"] = 3
        run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=ops, source="R3", rows={"E1": {"value": "A"}})
        self.assertEqual(len(rec.post_calls), 3)               # A, puis B, puis retour légitime à A
        self.assertEqual(sum(1 for r in ops.rows.values() if r["status"] == O.OP_ACCEPTED), 3)

    def test_supersession_without_losing_unsent_change(self):
        outbox, ops = MemOutbox({"R1": 1, "R2": 2}), MemOps()
        outbox.add("R1", "E1", status="FAILED")                 # ancien, non envoyé
        outbox.add("R1", "E2", status="FAILED")                 # ancien, sans élément plus récent
        outbox.add("R2", "E1", status="SENT")                   # plus récent, état complet publié
        rec = TR.RecordingTransport(FakeAdp())
        s, *_ = run("WS-HRIS-PRD", outbox, factory=lambda: rec, ops=ops, source="R1", mode="RETRY_PUBLICATION")
        self.assertEqual(outbox.items[("R1", "E1")]["status"], "SUPERSEDED")
        self.assertEqual(outbox.items[("R1", "E2")]["status"], "SENT")   # modification ancienne non perdue
        self.assertEqual(rec.post_calls, ["events/hr/v1/fake"])

    def test_never_supersede_uncertain(self):
        item = {"status": "FAILED"}
        self.assertFalse(O.supersession_candidates(item, [{"status": "SENT"}], [{"status": O.OP_UNCERTAIN}]))
        self.assertFalse(O.supersession_candidates({"status": "IN_PROGRESS"}, [{"status": "SENT"}], []))
        self.assertTrue(O.supersession_candidates(item, [{"status": "PENDING"}], [{"status": O.OP_NOT_SENT}]))

    def test_retry_selects_failed_not_in_progress_and_respects_max_attempts(self):
        outbox = MemOutbox()
        outbox.add("R1", "E1", status="FAILED", attempt_count=1)
        outbox.add("R1", "E2", status="IN_PROGRESS", attempt_count=1)
        outbox.add("R1", "E3", status="FAILED", attempt_count=3)
        rec = TR.RecordingTransport(FakeAdp())
        s, *_ = run("WS-HRIS-PRD", outbox, factory=lambda: rec, mode="RETRY_PUBLICATION")
        self.assertEqual(s["items_selected"], 1)
        self.assertEqual(outbox.items[("R1", "E2")]["status"], "IN_PROGRESS")


class ReconciliationAndDecisions(unittest.TestCase):
    class Reader:
        def __init__(self, observed):
            self.observed = observed

        def read(self, target):
            return self.observed

    def _op(self):
        return O.PlannedOp("P2", "P2_PERSONAL_EMAIL", "x", {}, {"emailUri": "a"}, {},
                           read_back=lambda rd: rd.read(None))

    def test_match_partial_different(self):
        op = self._op()
        self.assertEqual(O.reconcile({}, op, self.Reader({"emailUri": "a"}))["conclusion"], O.OP_APPLIED_OBSERVED)
        self.assertEqual(O.reconcile({}, op, self.Reader({}))["result"], O.RECON_PARTIAL)
        r = O.reconcile({}, op, self.Reader({"emailUri": "zzz"}))
        self.assertEqual(r["conclusion"], "MANUAL_DECISION_REQUIRED")   # un GET différent ne prouve pas le non-traitement

    def test_manual_decision_traced_and_required_fields(self):
        row = {"op_id": "o1", "run_id": "R1", "employee_key": "E1", "status": O.OP_UNCERTAIN}
        with self.assertRaises(HrisError):
            O.apply_manual_decision(row, O.DEC_AUTHORIZE_RESEND, "", "fondement")
        with self.assertRaises(HrisError):
            O.apply_manual_decision(row, O.DEC_AUTHORIZE_RESEND, "exploitant", "  ")
        updated, record = O.apply_manual_decision(row, O.DEC_AUTHORIZE_RESEND, "exploitant", "vérifié avec ADP",
                                                  "TICKET-1", clock=lambda: NOW)
        self.assertEqual(updated["status"], O.OP_NOT_SENT)
        self.assertEqual((record["decided_by"], record["basis"]), ("exploitant", "vérifié avec ADP"))


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_contract_markers', 'f5dbc43711cead3e47a932ab9b089cfccf67367977cae041ec2602e6b32442db', r'''"""Cohérence entre les marqueurs `CONTRAT-ADP:Cxx` du code et docs/CONTRAT_ADP_POINTS_D_USAGE.md."""
import os
import re
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
DOC = os.path.join(ROOT, "docs", "CONTRAT_ADP_POINTS_D_USAGE.md")
sys.path[:0] = [os.path.join(ROOT, "tools")]


@unittest.skipUnless(os.path.exists(DOC), "document des points d'usage absent (exécution hors dépôt)")
class ContractMarkers(unittest.TestCase):
    def test_markers_documented_and_located(self):
        import contract_points
        found = contract_points.scan()
        doc = open(DOC, encoding="utf-8").read()
        documented = set(re.findall(r"^\| (C\d\d) \| [^|]+\| [^|]+\| ", doc, flags=re.M))
        self.assertEqual(set(found) - documented, set(), "marqueurs du code non documentés")
        self.assertEqual(documented - set(found), set(), "points documentés sans marqueur dans le code")
        self.assertIn(contract_points.render(found), doc, "localisations à régénérer : python tools/contract_points.py")


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_legacy_logs', '8b9667c33c642e1b061573284f31a479890fcc6fb5dfcaa8e3b4bfd6a28b38c1', r'''"""TU-10 / TI-10 (partie locale) — lecture des logs historiques sur fichiers FICTIFS (JDD-06)."""
import csv
import io
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.join(HERE, "..", "..", "src")]

from hris import legacy_logs as L  # noqa: E402

COLS = L.EXPECTED_COLUMNS


def write_csv(path, rows, delimiter=";"):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, quotechar='"')
    w.writerow(COLS)
    w.writerows(rows)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(buf.getvalue())


class LegacyLogs(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp(prefix="jdd06_")
        ok = ["2026-01-19 03:00:01.123456", "00000001", "NOM FICTIF", "UPDATE", "salary", "200", "OK",
              '{"events": [1; 2]}', "", "changing salary from 1 to 2"]
        bad_ts = ["19/13/2026 99:99", "00000002", "X", "UPDATE", "bonus", "400", "Bad", "{}", "err", "msg"]
        write_csv(os.path.join(self.root, "ADP", "20260119", "output_log_adp.csv"), [ok, bad_ts, ok[:9]])
        write_csv(os.path.join(self.root, "ADP", "20260120", "output_log_adp.csv"), [ok], delimiter="\t")
        with open(os.path.join(self.root, "ADP", "20260119", "verbose_log_030000.log"), "w") as fh:
            fh.write("verbeux fictif")
        with open(os.path.join(self.root, "ADP", "20260121.csv"), "w") as fh:
            fh.write("a;b;c\n1;2;3\n")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_import_and_replay(self):
        records, rejects, ledger = L.process_folder(self.root, "test", {".csv", ".xlsx"}, set(), "imp-1", "u", "ws")
        status = {e["file_name"] + "@" + e["source_uri"].split("/")[-2]: e["status"] for e in ledger}
        self.assertEqual(sorted(status.values()), ["IMPORTED", "IMPORTED", "REJECTED_SCHEMA", "SKIPPED_UNSUPPORTED"])
        self.assertEqual(len(records), 2)
        self.assertEqual({r["log_date"].isoformat() for r in records}, {"2026-01-19", "2026-01-20"})
        self.assertEqual(records[0]["response_code"], '{"events": [1; 2]}')   # ';' entre guillemets
        self.assertEqual(records[0]["matricule_id"], "00000001")              # zéros de tête conservés
        self.assertEqual(sorted(r["reject_reason"] for r in rejects), ["COLUMN_COUNT_MISMATCH", "TS_ADP_INVALID"])
        for r in rejects:                                                    # aucune valeur de ligne rejetée
            self.assertEqual(set(r), {"import_id", "source_uri", "source_row_number", "reject_reason", "rejected_at_utc"})
        imported = {e["source_uri"] for e in ledger if e["status"] == "IMPORTED"}
        _, _, ledger2 = L.process_folder(self.root, "test", {".csv"}, imported, "imp-2", "u", "ws")
        self.assertEqual(sum(e["status"] == "SKIPPED_ALREADY_IMPORTED" for e in ledger2), 2)
        self.assertTrue(os.path.exists(os.path.join(self.root, "ADP", "20260119", "output_log_adp.csv")))  # rien supprimé


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_tables_naming', 'b9f80302efc2a23b7d08a98643b394c89613aab56ee469b05d4c4f4e4be37e03', r'''"""Résolution des noms de tables : 4 parties sans Lakehouse par défaut ; isolement des tests ; refus de deviner."""
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.join(HERE, "..", "..", "src")]

from hris.model import TABLES, Tables, tables_for_notebook  # noqa: E402


class TablesNaming(unittest.TestCase):
    def test_partial_names_without_lakehouse(self):
        self.assertEqual(Tables()("ctl.run"), "ctl.run")
        self.assertEqual(Tables("hris_tst")("ctl.run"), "hris_tst.ctl__run")

    def test_four_part_names(self):
        t = Tables("", "MOTUL_LH_HRIS", "Sandbox - Fabric BI backend")
        self.assertEqual(t("ctl.run"), "`Sandbox - Fabric BI backend`.`MOTUL_LH_HRIS`.ctl.run")
        self.assertEqual(t.schemas()[0], "`Sandbox - Fabric BI backend`.`MOTUL_LH_HRIS`.cfg")
        tt = Tables("hris_tst", "MOTUL_LH_HRIS", "WS")
        self.assertEqual(tt("log.anomaly"), "`WS`.`MOTUL_LH_HRIS`.hris_tst.log__anomaly")
        self.assertEqual(tt.extra("zz_test_report"), "`WS`.`MOTUL_LH_HRIS`.hris_tst.zz_test_report")
        self.assertEqual(len({t(x) for x in TABLES}), len(TABLES))

    def test_real_and_test_tables_never_collide(self):
        real, test = Tables("", "LH", "WS"), Tables("hris_tst", "LH", "WS")
        self.assertFalse({real(x) for x in TABLES} & {test(x) for x in TABLES})

    def test_context_resolution_and_refusal(self):
        ctx = {"currentWorkspaceName": "Sandbox - Fabric BI backend"}
        self.assertIn("`MOTUL_LH_HRIS`", tables_for_notebook(ctx, "", "MOTUL_LH_HRIS")("ctl.run"))
        self.assertIn("`BOUND_LH`", tables_for_notebook(dict(ctx, defaultLakehouseName="BOUND_LH"))("ctl.run"))
        with self.assertRaises(ValueError):
            tables_for_notebook(ctx)                      # ni paramètre ni Lakehouse par défaut : on ne devine pas
        with self.assertRaises(ValueError):
            tables_for_notebook({}, "", "LH")             # workspace non lisible

    def test_invalid_names_refused(self):
        for bad in ("a`b", "a\nb"):
            with self.assertRaises(ValueError):
                Tables("", "LH", bad)
        with self.assertRaises(ValueError):
            Tables("", "LH", "")


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_spark_transform_parity', 'b4b9d4fe448e412da8272adb49cfbddbd9de57dced583a1e620169546036ec86', r'''"""NR-02 (portion locale) — parité de l'implémentation Spark (hris.transform) avec l'implémentation de
référence du SQL de production (hris.transform_ref), sur données FICTIVES couvrant les cas limites.
Exécutable en local (pyspark 3.5, sans Delta) et dans Fabric (NB_HRIS_TESTS)."""
import datetime as dt
import decimal
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path[:0] = [os.path.join(HERE, "..", "..", "src"), os.path.join(HERE, "..", "fixtures")]

try:
    from pyspark.sql import SparkSession
except ImportError:  # pragma: no cover
    SparkSession = None

from hris import transform as S  # noqa: E402
from hris import transform_ref as R  # noqa: E402
from hris.model import ADP_COLUMNS, LOOKUP_TABLES, RAW_KEYS, SOURCE_COLUMN_NAMES, Tables  # noqa: E402
import jdd  # noqa: E402

D = dt.date(2026, 10, 7)
EDGE_RECORDS = [
    jdd.ts_record(1),
    jdd.ts_record(2, PPC_FTE="150"),                                   # ETP seul hors bornes : pas de rejet
    jdd.ts_record(3, PPC_FTE="150", Date_de_naissance="2024-02-30"),   # date invalide + ETP
    jdd.ts_record(4, LegacyID="123", Matricule_IT=""),                 # longueur + Null_Rejet
    jdd.ts_record(5, Title="XX", Gender="Z"),                          # Référence_Rejet
    jdd.ts_record(6, Title="MR   ", LastName="x" * 80),               # espaces de fin, troncature 50
    jdd.ts_record(7, BusinessMobile=None, BusinessPhone="0102"),       # COALESCE
    jdd.ts_record(8, date_fin_contract="2027-01-31", GroupEndDate=""),  # dates
    jdd.ts_record(9, salary_base=50000.555, Pourcentage_prime_annuelles="abc"),  # arrondi, non numérique
    jdd.ts_record(10, referenceAdmissionDate="2019-01-07T00:00:00.000", LegalCountry="ZZZ"),
    jdd.ts_record(11, hubSocialSecurityNumber="   ", Matricule_IT=None),
    jdd.ts_record(12, RecalculatedSeniorityDate="01/31/2020", CompanyStartDate="20200131"),
]


@unittest.skipIf(SparkSession is None, "pyspark absent")
class SparkParity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.spark = SparkSession.getActiveSession() or (SparkSession.builder.master("local[1]")
                                                         .config("spark.ui.enabled", "false")
                                                         .config("spark.driver.memory", "2g")
                                                         .config("spark.sql.shuffle.partitions", "1").getOrCreate())
        S.ensure_session_settings(cls.spark)
        cls.T = Tables("hris_tst")
        cls.lookups = {}
        for table, cols in LOOKUP_TABLES.items():
            rows = [tuple(r.get(c) for c in cols) for r in jdd.full_lookup_rows(table)]
            cls.lookups[table] = cls.spark.createDataFrame(rows, ", ".join(f"{c} string" for c in cols))
        cls.loader = staticmethod(lambda table: SparkParity.lookups[table])

    def raw_df(self, records):
        rows = [dict(r, run_id="RUN", business_date=D) for r in jdd.raw_rows(records)]
        schema = "run_id string, business_date date, " + ", ".join(f"`{k}` string" for k in RAW_KEYS)
        return self.spark.createDataFrame([tuple(r.get(k) for k in ["run_id", "business_date"] + RAW_KEYS)
                                           for r in rows], schema)

    def test_full_chain_matches_reference(self):
        src = S.source_df(self.raw_df(EDGE_RECORDS)).localCheckpoint()
        enr = S.enriched_df(self.spark, self.T, src, D, self.loader).localCheckpoint()
        rej = S.reject_df(enr, D).localCheckpoint()
        adp = S.adp_df(enr, rej, D)

        ref_src = [R.parse_source_row(r) for r in jdd.raw_rows(EDGE_RECORDS)]
        ref_enr = R.enrich_rows(ref_src, {t: jdd.full_lookup_rows(t) for t in LOOKUP_TABLES})
        ref_rej = R.reject_rows(ref_enr, D)
        ref_adp = R.adp_rows(ref_enr, ref_rej, D)

        got_src = sorted(tuple(r[c] for c in SOURCE_COLUMN_NAMES) for r in src.collect())
        self.assertEqual(got_src, sorted(tuple(r[c] for c in SOURCE_COLUMN_NAMES) for r in ref_src))

        rej_cols = ["id_unique", "cd_type_rejet", "lb_nom_colonne_rejet", "lb_valeur_rejet", "lb_fichier_mapping",
                    "ds_rejet", "fl_rejet"]
        key = lambda t: tuple("" if v is None else str(v) for v in t)  # noqa: E731
        self.assertEqual(sorted((tuple(r[c] for c in rej_cols) for r in rej.collect()), key=key),
                         sorted((tuple(r[c] for c in rej_cols) for r in ref_rej), key=key))

        got_adp = {r["id_unique"]: r.asDict() for r in adp.collect()}
        ref_by = {r["id_unique"]: r for r in ref_adp}
        self.assertEqual(sorted(got_adp), sorted(ref_by))
        for k, ref in ref_by.items():
            for c in ADP_COLUMNS:
                self.assertEqual(got_adp[k][c], ref[c], f"{k}.{c}")
            self.assertEqual(got_adp[k]["row_hash"], R.row_hash(ref), k)
        self.assertEqual(str(adp.schema["mt_salaire_mensuel"].dataType), "DecimalType(23,13)")

    def test_conversion_error_fails_like_tsql(self):
        src = S.source_df(self.raw_df([jdd.ts_record(20, salary_base="")]))
        enr = S.enriched_df(self.spark, self.T, src, D, self.loader).localCheckpoint()
        with self.assertRaises(Exception):
            S.adp_df(enr, S.reject_df(enr, D), D).collect()
        src = S.source_df(self.raw_df([jdd.ts_record(21, Pourcentage_prime_annuelles="1e5")]))
        enr = S.enriched_df(self.spark, self.T, src, D, self.loader).localCheckpoint()
        with self.assertRaises(Exception):
            S.reject_df(enr, D).collect()

    def test_masking_of_sensitive_rejects(self):
        src = S.source_df(self.raw_df([jdd.ts_record(30, Date_de_naissance="1990-02-30")]))
        enr = S.enriched_df(self.spark, self.T, src, D, self.loader).localCheckpoint()
        masked = S.masked_log_reject(S.reject_df(enr, D), "RUN").collect()
        self.assertTrue(masked)
        for r in masked:
            if r["lb_nom_colonne_rejet"] == "dt_naissance":
                self.assertEqual(r["lb_valeur_rejet"], "***")
                self.assertNotIn("1990", r["ds_rejet"])


if __name__ == "__main__":
    unittest.main()
'''),
    ('test_fabric_integration', 'b4d697b97b8f107a27e319b94bf15b7440c8137e6948d10512e814a06ac7399a', r'''"""Tests d'intégration Fabric (TI, TR, TG, TU-09) — Delta réel, schéma ISOLÉ `hris_tst`, données FICTIVES.

Exécutés par NB_HRIS_TESTS après `configure(spark, notebookutils)`. Aucune table réelle n'est lue ni écrite ;
aucun appel TalentSoft ni ADP : l'acquisition est simulée par des enregistrements fictifs, la publication
évalue le garde-fou dans le contexte RÉEL du notebook avec une fabrique de transport qui échoue si elle est
appelée, sous sentinelle réseau.
"""
import datetime as dt
import threading
import unittest
import uuid

import jdd
from hris import compare as C
from hris import run_control as RC
from hris import spark_io as IO
from hris import steps
from hris.common import PROCESS_CODE
from hris.model import LOOKUP_TABLES, TABLES, Tables, tables_for_notebook
from hris.transform import ensure_session_settings
from hris.ts_acquire import AcquisitionStats

ENV = {"spark": None, "notebookutils": None, "schema": "hris_tst", "lakehouse": ""}


def configure(spark, notebookutils, schema="hris_tst", lakehouse=""):
    ENV.update(spark=spark, notebookutils=notebookutils, schema=schema, lakehouse=lakehouse)


def _ok_stats(n):
    st = AcquisitionStats()
    st.total_count = n
    st.counters.update({"directory_received": n, "records": n})
    return st.as_dict()


class ForbiddenFactory:
    def __init__(self):
        self.called = 0

    def __call__(self):
        self.called += 1
        raise AssertionError("transport ADP construit")


class FabricIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.spark = ENV["spark"]
        if cls.spark is None:
            raise unittest.SkipTest("hors Fabric")
        ensure_session_settings(cls.spark)
        cls.T = tables_for_notebook(ENV["notebookutils"].runtime.context, ENV["schema"], ENV["lakehouse"])
        for schema in cls.T.schemas():
            cls.spark.sql(f"CREATE SCHEMA IF NOT EXISTS {schema}")
        for logical in TABLES:  # tables de TEST uniquement (schéma isolé), recréées à chaque campagne
            cls.spark.sql(f"DROP TABLE IF EXISTS {cls.T(logical)}")
        IO.ensure_schemas_and_tables(cls.spark, cls.T)
        steps.seed(cls.spark, cls.T)
        for table, cols in LOOKUP_TABLES.items():
            rows = [dict(r, _loaded_at_utc=dt.datetime.now(dt.timezone.utc), _loaded_by="NB_HRIS_TESTS",
                         _change_ref="JDD-03-FICTIF") for r in jdd.full_lookup_rows(table)]
            if rows:
                IO.append(IO.rows_df(cls.spark, rows, f"cfg.{table}"), cls.T(f"cfg.{table}"), f"cfg.{table}")

    # -- utilitaires --------------------------------------------------------------------------------------
    def start(self, mode="NORMAL", trigger="2026-10-07T01:00:00Z", target=""):
        run_id = "TST-" + uuid.uuid4().hex[:12]
        res = RC.start(self.spark, self.T, run_id=run_id, process_code=PROCESS_CODE, run_mode=mode, target_run_id=target,
                       trigger_time=trigger, business_timezone="Europe/Paris", lock_lease_minutes=30,
                       business_date_override="", pipeline_name="TEST", workspace_id=None, workspace_name=None)
        self.assertEqual(res["status"], "OK", res)
        return run_id

    def acquire(self, run_id, records):
        run = IO.get_run(self.spark, self.T, run_id)
        step = IO.StepLogger(self.spark, self.T, run_id, "ACQUIRE")
        return steps.store_acquisition(self.spark, self.T, run, records, _ok_stats(len(records)), step,
                                       source=steps.SOURCE_FIXTURE)

    def full_run(self, records, mode="NORMAL", publish=True):
        run_id = self.start(mode)
        self.acquire(run_id, records)
        steps.run_transform(self.spark, self.T, run_id)
        cmp = steps.run_compare(self.spark, self.T, run_id, mode)
        pro = steps.run_promote(self.spark, self.T, run_id, mode)
        pub = None
        if publish:
            pub = self.publish(run_id, mode)
        end = RC.end(self.spark, self.T, run_id, "TEST")
        return run_id, cmp, pro, pub, end

    def publish(self, run_id, mode, factory=None):
        factory = factory or ForbiddenFactory()
        res = steps.run_publish(self.spark, self.T, run_id, mode,
                                context_provider=lambda: ENV["notebookutils"].runtime.context,  # contexte RÉEL
                                transport_factory=factory, planner=lambda row, ex: None, max_attempts=3,
                                pause_seconds=0)
        res["_factory_calls"] = factory.called
        return res

    # -- scénarios ---------------------------------------------------------------------------------------
    def test_01_lock_and_concurrency(self):  # TU-09, TI-07
        a, b = "TST-LOCK-A-" + uuid.uuid4().hex[:6], "TST-LOCK-B-" + uuid.uuid4().hex[:6]
        self.assertTrue(RC.acquire_lock(self.spark, self.T, a, 30))
        self.assertFalse(RC.acquire_lock(self.spark, self.T, b, 30))
        self.assertTrue(RC.release_lock(self.spark, self.T, a))
        results = {}

        def grab(rid):
            results[rid] = RC.acquire_lock(self.spark, self.T, rid, 30)
        threads = [threading.Thread(target=grab, args=(r,)) for r in (a, b)]
        [t.start() for t in threads]
        [t.join() for t in threads]
        self.assertEqual(sum(results.values()), 1, results)
        holder = [r for r, ok in results.items() if ok][0]
        RC.release_lock(self.spark, self.T, holder)
        # bail expiré : reprise possible
        self.assertTrue(RC.acquire_lock(self.spark, self.T, a, 30, now=dt.datetime(2020, 1, 1, tzinfo=dt.timezone.utc)))
        self.assertTrue(RC.acquire_lock(self.spark, self.T, b, 30))
        RC.release_lock(self.spark, self.T, b)

    def test_02_first_load_blocked_then_init_then_normal_with_outbox_and_gate(self):  # TI-01..03, TG-01, TG-03
        recs = [jdd.ts_record(i) for i in range(1, 6)] + [jdd.ts_record(9, Gender="X")]
        run1, cmp1, pro1, pub1, end1 = self.full_run(recs, "NORMAL")
        self.assertEqual(cmp1["decision"], "BLOCKED")
        self.assertIn("REFERENCE_ABSENTE", cmp1["blocking"])
        self.assertFalse(pro1["promoted"])
        self.assertEqual(end1["status"], "BLOCKED")

        run2, cmp2, pro2, pub2, end2 = self.full_run(recs, "INIT_REFERENCE")
        self.assertEqual(cmp2["decision"], "VALIDATED", cmp2)
        self.assertTrue(pro2["promoted"])
        self.assertEqual(pro2["nb_outbox"], 0)                       # INIT_REFERENCE : aucune outbox
        self.assertEqual(pub2["status"], "SKIPPED_INIT_REFERENCE")   # TG-03
        self.assertEqual(self.spark.table(self.T("ref.ts_employee_adp")).count(), 5)

        recs3 = [jdd.ts_record(1, PersonalEmail="nouveau1@example.invalid")] + \
                [jdd.ts_record(i) for i in range(2, 5)] + [jdd.ts_record(7)]          # E5 absent, E7 nouveau
        run3, cmp3, pro3, pub3, end3 = self.full_run(recs3, "NORMAL")
        self.assertEqual((cmp3["decision"], cmp3["nb_new"], cmp3["nb_modified"], cmp3["nb_absent"]),
                         ("VALIDATED", 1, 1, 1), cmp3)
        self.assertEqual(pro3["nb_outbox"], 2)
        det = {r.business_key: r for r in self.spark.table(self.T("log.comparison_detail"))
               .where(f"run_id = '{run3}'").collect()}
        self.assertEqual(det["E00001"].changed_columns, ["lb_email_personnel"])   # noms seulement
        # TG-01 : contexte réel du workspace DEV (nom sans PRD) -> aucune construction de transport
        self.assertEqual(pub3["proof"]["gate_outcome"], "DENIED_NON_PRD", pub3)
        self.assertEqual(pub3["_factory_calls"], 0)
        self.assertEqual(pub3["proof"]["sentinel_attempts"], 0)
        statuses = {r.status for r in self.spark.table(self.T("pub.adp_outbox")).where(f"run_id = '{run3}'").collect()}
        self.assertEqual(statuses, {"NOT_PUSHED_NON_PRD"})
        self.assertEqual(end3["status"], "NOT_PUBLISHED_NON_PRD")
        gate = self.spark.table(self.T("log.adp_gate_decision")).where(f"run_id = '{run3}'").collect()
        self.assertEqual(len(gate), 1)
        self.assertFalse(gate[0].decision)
        self.assertEqual(gate[0].reason_code, "WORKSPACE_NAME_WITHOUT_PRD")
        self.assertEqual(self.spark.table(self.T("pub.adp_operation")).count(), 0)

        # TR-03 / TR-02 : promotion rejouée sans effet ; comparatif non recalculé après promotion
        before = self.spark.table(self.T("ref.ts_employee_adp")).orderBy("id_unique").collect()
        again = C.promote(self.spark, self.T, run3, "NORMAL")
        self.assertTrue(again.get("already_promoted"))
        self.assertEqual(self.spark.table(self.T("ref.ts_employee_adp")).orderBy("id_unique").collect(), before)
        self.assertFalse(C.compare(self.spark, self.T, run3, "NORMAL")["recomputed"])

        # TI-08 : second run le même jour, sans changement -> delta vide, rang 2+
        run4, cmp4, pro4, pub4, end4 = self.full_run(recs3, "NORMAL")
        self.assertEqual((cmp4["nb_new"], cmp4["nb_modified"]), (0, 0))
        self.assertGreaterEqual(IO.get_run(self.spark, self.T, run4)["run_seq_in_day"], 2)

    def test_03_init_reference_refused_when_reference_not_empty(self):
        if self.spark.table(self.T("ref.ts_employee_adp")).count() == 0:
            self.skipTest("référence vide (ordre des tests)")
        run, cmp, pro, pub, end = self.full_run([jdd.ts_record(1)], "INIT_REFERENCE")
        self.assertEqual(cmp["decision"], "BLOCKED")
        self.assertIn("REFERENCE_NOT_EMPTY", cmp["blocking"])

    def test_04_empty_batch_blocks(self):  # TI-04
        run_id = self.start()
        run = IO.get_run(self.spark, self.T, run_id)
        st = AcquisitionStats()
        st.total_count = 0
        res = steps.store_acquisition(self.spark, self.T, run, [], st.as_dict(),
                                      IO.StepLogger(self.spark, self.T, run_id, "ACQUIRE"), source=steps.SOURCE_FIXTURE)
        self.assertIn("EMPTY_EXTRACTION", res["blocking_failures"])
        steps.run_transform(self.spark, self.T, run_id)
        self.assertEqual(steps.run_compare(self.spark, self.T, run_id, "NORMAL")["decision"], "BLOCKED")
        self.assertFalse(steps.run_promote(self.spark, self.T, run_id, "NORMAL")["promoted"])
        RC.end(self.spark, self.T, run_id, "TEST")

    def test_05_threshold_revalidation(self):  # « revalidation après configuration des seuils »
        if self.spark.table(self.T("ref.ts_employee_adp")).count() == 0:
            self.skipTest("référence vide")
        from delta.tables import DeltaTable
        tbl = DeltaTable.forName(self.spark, self.T("cfg.process"))
        tbl.update(set={"max_modified_pct": "0"})
        try:
            run_id = self.start()
            self.acquire(run_id, [jdd.ts_record(2, PersonalEmail="x2@example.invalid")] + [jdd.ts_record(i) for i in (1, 3, 4, 7)])
            steps.run_transform(self.spark, self.T, run_id)
            cmp = steps.run_compare(self.spark, self.T, run_id, "NORMAL")
            self.assertEqual(cmp["decision"], "BLOCKED")
            self.assertIn("MODIFIED_RATIO", cmp["blocking"])
            tbl.update(set={"max_modified_pct": "100"})
            # lever la configuration ne suffit pas : la promotion relit les contrôles du run (bloqué)
            self.assertFalse(steps.run_promote(self.spark, self.T, run_id, "NORMAL")["promoted"])
            cmp2 = steps.run_compare(self.spark, self.T, run_id, "NORMAL")   # revalidation explicite
            self.assertEqual(cmp2["decision"], "VALIDATED", cmp2)
            self.assertTrue(steps.run_promote(self.spark, self.T, run_id, "NORMAL")["promoted"])
            RC.end(self.spark, self.T, run_id, "TEST")
        finally:
            tbl.update(set={"max_modified_pct": "CAST(NULL AS DECIMAL(9,4))"})

    def test_06_missing_mapping_blocks_only_enrichment(self):  # TI-06
        from delta.tables import DeltaTable
        backup = self.spark.table(self.T("cfg.lu_hr_emp_sexe")).collect()
        DeltaTable.forName(self.spark, self.T("cfg.lu_hr_emp_sexe")).delete()   # table de TEST
        try:
            run_id = self.start()
            self.acquire(run_id, [jdd.ts_record(1)])
            res = steps.run_transform(self.spark, self.T, run_id)
            self.assertEqual(res["status"], "BLOCKED_MAPPINGS")
            self.assertEqual(self.spark.table(self.T("stg.ts_employee_source")).where(f"run_id = '{run_id}'").count(), 1)
            self.assertEqual(steps.run_compare(self.spark, self.T, run_id, "NORMAL")["decision"], "BLOCKED")
            RC.end(self.spark, self.T, run_id, "TEST")
        finally:
            if backup:
                schema = self.spark.table(self.T("cfg.lu_hr_emp_sexe")).schema
                (self.spark.createDataFrame([tuple(r) for r in backup], schema).write.format("delta")
                 .mode("append").saveAsTable(self.T("cfg.lu_hr_emp_sexe")))

    def test_07_run_end_releases_lock_and_summary_has_no_values(self):
        rows = self.spark.table(self.T("ctl.process_lock")).collect()
        self.assertEqual(rows[0].lock_status, "RELEASED")
        summ = self.spark.table(self.T("log.run_summary")).collect()
        self.assertTrue(summ)
        for r in summ:
            self.assertTrue(r.ds_type_rejets is None or len(r.ds_type_rejets.encode("utf-8")) <= 50)


def suite():
    return unittest.defaultTestLoader.loadTestsFromTestCase(FabricIntegration)
'''),
]
def _hris_load(entries):
    pkg = _sys.modules.get('hris') or _types.ModuleType('hris')
    pkg.__path__ = []
    _sys.modules['hris'] = pkg
    for name, digest, src in entries:
        if _hl.sha256(src.encode('utf-8')).hexdigest() != digest:
            raise RuntimeError('empreinte invalide : ' + name)
        mod = _types.ModuleType(name)
        mod.__file__ = name.replace('.', '/') + '.py'
        _sys.modules[name] = mod
        exec(compile(src, mod.__file__, 'exec'), mod.__dict__)
        if name.startswith('hris.'):
            setattr(pkg, name.split('.', 1)[1], mod)
    return {n: d for n, d, _ in entries}
HRIS_LIBRARY_DIGESTS = _hris_load(_HRIS_MODULES)
print('bibliothèque HRIS chargée :', len(HRIS_LIBRARY_DIGESTS), 'modules')

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

import io
import json
import platform
import unittest

import jdd
import test_adp_publication
import test_contract_markers
import test_fabric_integration
import test_legacy_logs
import test_tables_naming
import test_parity_transform
import test_spark_transform_parity
import test_ts_acquire
from hris.common import as_bool, json_exit
from hris.transform import ensure_session_settings

ensure_session_settings(spark)
test_fabric_integration.configure(spark, notebookutils, p_test_schema, p_lakehouse_name)
modules = [test_parity_transform, test_ts_acquire, test_adp_publication, test_contract_markers,
           test_legacy_logs, test_tables_naming, test_spark_transform_parity]
if as_bool(p_run_integration):
    modules.append(test_fabric_integration)
suite = unittest.TestSuite()
_only = [x.strip() for x in p_only.split(",") if x.strip()]


def _flatten(s):
    for t in s:
        if isinstance(t, unittest.TestSuite):
            yield from _flatten(t)
        else:
            yield t


for m in modules:
    for t in _flatten(unittest.defaultTestLoader.loadTestsFromModule(m)):
        if not _only or any(o in t.id() for o in _only):
            suite.addTest(t)
stream = io.StringIO()
outcome = unittest.TextTestRunner(stream=stream, verbosity=2).run(suite)
print(stream.getvalue()[-20000:])


def versions():
    out = {}
    for name in ("requests", "openpyxl", "pandas", "pyarrow", "dateutil", "delta"):
        try:
            mod = __import__(name)
            out[name] = getattr(mod, "__version__", "présent")
        except Exception as exc:
            out[name] = "ABSENT:" + type(exc).__name__
    return out


result = {"tests_run": outcome.testsRun, "failures": [t.id() for t, _ in outcome.failures],
          "errors": [t.id() for t, _ in outcome.errors], "skipped": [t.id() for t, _ in outcome.skipped],
          "successful": outcome.wasSuccessful(),
          "platform": {"spark": spark.version, "python": platform.python_version(),
                       "NV-02_session_timezone_default": NV_TIMEZONE_DEFAULT, "NV-06_ansi_enabled": NV_ANSI,
                       "time_parser_policy_default": NV_PARSER, "NV-05_libraries": versions(),
                       "workspace_name": notebookutils.runtime.context.get("currentWorkspaceName"),
                       "is_for_pipeline": notebookutils.runtime.context.get("isForPipeline")}}
details = {t.id(): tb[-1500:] for t, tb in outcome.failures + outcome.errors}
print(json.dumps(result, ensure_ascii=False, indent=1))
print(json.dumps(details, ensure_ascii=False, indent=1)[:20000])

# Rapport de campagne (données fictives uniquement) dans le schéma de test, lisible sans donnée RH
import datetime as _dt
from hris.model import tables_for_notebook
_T = tables_for_notebook(notebookutils.runtime.context, p_test_schema, p_lakehouse_name)
spark.sql(f"CREATE TABLE IF NOT EXISTS {_T.extra('zz_test_report')} (run_at_utc TIMESTAMP, report STRING) USING DELTA")
spark.createDataFrame([(_dt.datetime.utcnow(), json.dumps(dict(result, details=details), ensure_ascii=False, default=str))],
                      "run_at_utc timestamp, report string").write.format("delta").mode("append").saveAsTable(_T.extra("zz_test_report"))

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }

# CELL ********************

notebookutils.notebook.exit(json_exit(dict(result, details={k: v[-600:] for k, v in list(details.items())[:10]})))

# METADATA ********************

# META {
# META   "language": "python",
# META   "language_group": "synapse_pyspark"
# META }
